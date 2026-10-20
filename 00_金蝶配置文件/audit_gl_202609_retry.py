# -*- coding: utf-8 -*-
"""
稽核补拉：audit_gl_202609.py 跑完后，392 张凭证因金蝶限流（1000次/分）失败。
本脚本：
  1) 重新探针取全量凭证清单
  2) 对比 output/gl_voucher_audit_202609/vouchers_raw.json，找出缺失凭证
  3) 2 线程 + 全局限速（≤12 次/秒，撞限流睡 3 秒重试）补拉
  4) 合并后重跑稽核，输出最终 Excel + 汇总日志
"""

import os, sys, io, json, time, threading
from collections import defaultdict
from queue import Queue

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
# 注意：audit_gl_202609 导入时会自行包 stdout（勿重复包装，否则旧包装被 GC 会关掉共享缓冲）
from audit_gl_202609 import (KdClient, audit_book, FIELDS, fnum,
                             OUT_DIR, DATE_FROM, DATE_TO, FOCUS_KEYWORD,
                             CASH_PREFIXES, log)

from openpyxl import Workbook

THREADS = 2
CALL_INTERVAL = 0.085          # 全局限速：两次调用最小间隔（≈12次/秒 < 限流16.6次/秒）
RATE_LIMIT_SLEEP = 3.2         # 撞限流后休眠
RATE_LIMIT_RETRY = 8

_rate_lock = threading.Lock()
_last_call = [0.0]
def pace():
    with _rate_lock:
        wait = CALL_INTERVAL - (time.time() - _last_call[0])
        if wait > 0: time.sleep(wait)
        _last_call[0] = time.time()

def is_rate_limit(msg):
    return "限流" in msg or "超过最大值" in msg

def voucher_entries_safe(client, bno, zi, no):
    flt = f"FAccountBookID.FNumber='{bno}' and FVOUCHERGROUPNO={no}"
    if zi and zi != "?":
        flt += f" and FVOUCHERGROUPID.FName='{zi}'"
    keys = ",".join(f for f, _ in FIELDS)
    rows, start = [], 0
    while True:
        for attempt in range(RATE_LIMIT_RETRY):
            pace()
            try:
                data = client.query(keys, flt, "FVOUCHERGROUPNO", start=start)
                break
            except Exception as e:
                if is_rate_limit(str(e)) and attempt < RATE_LIMIT_RETRY - 1:
                    time.sleep(RATE_LIMIT_SLEEP); continue
                raise
        rows.extend(data)
        if len(data) < 2000:
            break
        start += 2000
    return rows


def main():
    t0 = time.time()
    log("登录金蝶云星空 ...")
    client = KdClient()
    log("  登录成功")

    # 1) 全量凭证清单
    log(f"\n[1/5] 探针 {DATE_FROM}~{DATE_TO} 凭证清单 ...")
    keys = "FAccountBookID.FNumber,FAccountBookID.FName,FVOUCHERGROUPID.FName,FVOUCHERGROUPNO"
    probe = client.query(keys, f"FDATE>='{DATE_FROM}' and FDATE<='{DATE_TO}'",
                         "FAccountBookID.FNumber,FVOUCHERGROUPID.FName,FVOUCHERGROUPNO")
    book_name, vouchers, seen = {}, [], set()
    for r in probe:
        bno, bname, zi, no = r[0], r[1], r[2] or "?", r[3]
        book_name[bno] = bname
        try: no_i = int(float(no))
        except (TypeError, ValueError): continue
        if (bno, zi, no_i) in seen: continue
        seen.add((bno, zi, no_i))
        vouchers.append((bno, zi, no_i))
    log(f"  共 {len(vouchers)} 张")

    # 2) 已有数据 → 缺失清单
    raw_path = os.path.join(OUT_DIR, "vouchers_raw.json")
    with open(raw_path, encoding="utf-8") as fp:
        raw = json.load(fp)
    headers = raw["headers"]
    idx = {cn: i for i, cn in enumerate(headers)}
    have_rows = defaultdict(list)
    for r in raw["rows"]:
        have_rows[(r[idx["账簿编码"]], r[idx["凭证字"]] or "?", int(float(r[idx["凭证号"]])))].append(r)
    missing = [v for v in vouchers if v not in have_rows]
    log(f"\n[2/5] 已有 {len(have_rows)} 张，缺失 {len(missing)} 张，开始补拉（{THREADS} 线程 + 限速）...")

    # 3) 补拉
    q = Queue()
    for v in missing: q.put(v)
    result, errors, slow = {}, [], []
    done = [0]
    def worker():
        c = KdClient()
        while True:
            try: bno, zi, no = q.get_nowait()
            except Exception: return
            t1 = time.time()
            try:
                rows = voucher_entries_safe(c, bno, zi, no)
                result[(bno, zi, no)] = rows
                dt = time.time() - t1
                if dt > 10 or len(rows) >= 2000:
                    slow.append((dt, bno, zi, no, len(rows)))
            except Exception as e:
                errors.append((bno, zi, no, str(e)))
            with threading.Lock():
                done[0] += 1
                if done[0] % 25 == 0:
                    log(f"    补拉进度 {done[0]}/{len(missing)}（{time.time()-t0:.0f}s，失败 {len(errors)}）")
    threads = [threading.Thread(target=worker, daemon=True) for _ in range(THREADS)]
    for t in threads: t.start()
    for t in threads: t.join()
    log(f"  补拉完成 {done[0]}/{len(missing)}，失败 {len(errors)}，用时 {time.time()-t0:.0f}s")
    for e in errors[:15]:
        log(f"    ⚠ 仍失败：{e}")
    for dt, bno, zi, no, n in sorted(slow, reverse=True)[:10]:
        log(f"    大凭证：{zi}字{no}号 [{bno}] {n} 条，{dt:.0f}s")

    # 4) 合并
    book_rows = defaultdict(list)
    for k, rows in have_rows.items():
        book_rows[k[0]].extend(rows)
    for (bno, zi, no), rows in result.items():
        book_rows[bno].extend(rows)
    total_rows = sum(len(r) for r in book_rows.values())
    got_v = sum(len(set((r[idx["凭证字"]], r[idx["凭证号"]]) for r in rows)) for rows in book_rows.values())
    log(f"\n[3/5] 合并后：凭证 {got_v} 张 / 分录 {total_rows} 条")

    # 5) 重跑稽核 + 落盘
    log(f"\n[4/5] 稽核 ...")
    all_v, all_e, summaries = [], [], []
    for bno in sorted(book_rows):
        v_iss, e_iss, summ = audit_book(book_name[bno], book_rows[bno], idx)
        focus = "  ← 重点关注" if FOCUS_KEYWORD in book_name[bno] else ""
        log(f"\n■ {bno} {book_name[bno]}{focus}")
        log(f"  凭证 {summ['凭证张数']} 张 / 分录 {summ['分录条数']} 条 / 借方合计 {summ['借方合计']:,.2f}")
        log(f"  未审核 {summ['未审核']} / 未过账 {summ['未过账']} / 凭证级问题 {summ['凭证级问题']} / 分录级问题 {summ['分录级问题']}")
        if summ["怪兽凭证"]:
            mstr = "、".join(f"{zi}字{no}号({n:,}条)" for zi, no, n in sorted(summ["怪兽凭证"], key=lambda x: -x[2])[:8])
            log(f"  超大凭证TOP：{mstr}")
        shown = 0
        for lv, zi, no, msg in v_iss:
            if shown < 20:
                log(f"    [{lv}] {zi}字 {no}号：{msg}"); shown += 1
        if len(v_iss) > 20:
            log(f"    ... 其余 {len(v_iss)-20} 项见 Excel")
        for it in e_iss[:5]:
            log(f"    [分录] {it[0]} {it[1]}字 {it[2]}号 {it[3]} {it[4]}")
        for lv, zi, no, msg in v_iss:
            all_v.append([bno, book_name[bno], lv, zi, no, msg])
        for it in e_iss:
            all_e.append([bno, book_name[bno], *it])
        summaries.append([bno, book_name[bno], summ['凭证张数'], summ['分录条数'],
                          round(summ['借方合计'], 2), summ['未审核'], summ['未过账'],
                          summ['凭证级问题'], summ['分录级问题'], len(summ["怪兽凭证"])])

    log(f"\n[5/5] 保存结果 ...")
    wb = Workbook()
    ws = wb.active; ws.title = "汇总"
    ws.append(["账簿编码","账簿","凭证张数","分录条数","借方合计","未审核","未过账","凭证级问题数","分录级问题数","超大凭证数"])
    for row in summaries: ws.append(row)
    ws1 = wb.create_sheet("凭证级问题")
    ws1.append(["账簿编码","账簿","级别","凭证字","凭证号","问题"])
    for row in all_v: ws1.append(row)
    ws2 = wb.create_sheet("分录级问题")
    ws2.append(["账簿编码","账簿","类型","凭证字","凭证号","科目编码","科目名称","备注"])
    for row in all_e: ws2.append(row)
    # 全部分录按账簿分 sheet；超大账簿截断（Excel 单表 104 万行上限）
    for bno in sorted(book_rows):
        rows = book_rows[bno]
        if len(rows) > 900000:
            rows = rows[:900000]
            note = True
        else:
            note = False
        wsn = wb.create_sheet(f"分录_{bno}")
        wsn.append(headers + (["⚠ 已截断"] if note else []))
        for r in rows: wsn.append(list(r))
    xlsx = os.path.join(OUT_DIR, "凭证稽核_全主体_202609.xlsx")
    wb.save(xlsx)
    merged = [r for bno in sorted(book_rows) for r in book_rows[bno]]
    with open(raw_path, "w", encoding="utf-8") as fp:
        json.dump({"headers": headers, "rows": merged}, fp, ensure_ascii=False)
    log(f"  Excel：{xlsx}")
    log(f"  原始数据：{raw_path}")
    log(f"完成。总用时 {time.time()-t0:.0f}s")


if __name__ == "__main__":
    main()
