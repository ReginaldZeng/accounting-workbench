# -*- coding: utf-8 -*-
"""
金蝶云星空 2026年9月 总账凭证稽核（全主体，重点输出孝感星期九）

取数策略（血泪经验）：
  - ExecuteBillQuery 限 2000 行/页；凭证含分录字段时一行=一条分录
  - 单张凭证可能超 2000 条分录（月末计提结转怪兽凭证，实测一张 4.9 万行），
    按账簿/日期/号区间查询都会触顶且深分页大查询会卡死服务器
  - 故：逐张凭证等值过滤（FVOUCHERGROUPNO=N）+ StartRow 分页，范围小、速度快
  - 4 线程并发，每线程独立登录会话

稽核项：
  凭证级：借贷平衡 / 断号 / 未审核 / 未过账 / 制审同人 / 怪兽凭证提示
  分录级：摘要空白 / 同分录借贷同录；资金类科目银行账号维度单独核查
"""

import os, sys, io, json, configparser, threading, time
from collections import defaultdict
from queue import Queue

import requests
from openpyxl import Workbook

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8")

FORM_ID = "GL_VOUCHER"
DATE_FROM = "2026-09-01"
DATE_TO   = "2026-09-30"
FOCUS_KEYWORD = "孝感"
THREADS = 4
PAGE = 2000

CONF_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "conf.ini")
CONF_NODE = "config"
LOGIN_SVC = "Kingdee.BOS.WebApi.ServicesStub.AuthService.LoginByAppSecret.common.kdsvc"
QUERY_SVC = "Kingdee.BOS.WebApi.ServicesStub.DynamicFormService.ExecuteBillQuery.common.kdsvc"

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR  = os.path.join(BASE_DIR, "output", "gl_voucher_audit_202609")

FIELDS = [
    ("FAccountBookID.FNumber",  "账簿编码"),
    ("FAccountBookID.FName",    "账簿"),
    ("FDate",                   "日期"),
    ("FYEAR",                   "会计年度"),
    ("FPERIOD",                 "期间"),
    ("FVOUCHERGROUPID.FName",   "凭证字"),
    ("FVOUCHERGROUPNO",         "凭证号"),
    ("FCreatorId.FName",        "制单"),
    ("FCHECKERID.FName",        "审核"),
    ("FPOSTERID.FName",         "过账"),
    ("FCASHIERID.FName",        "出纳"),
    ("FACCOUNTID.FNumber",      "科目编码"),
    ("FACCOUNTID.FName",        "科目名称"),
    ("FCURRENCYID.FName",       "币别"),
    ("FEXCHANGERATE",           "汇率"),
    ("FAMOUNTFOR",              "原币金额"),
    ("FDEBIT",                  "借方金额"),
    ("FCREDIT",                 "贷方金额"),
    ("FEXPLANATION",            "摘要"),
]
CASH_PREFIXES = ("1001", "1002", "1012")

_lock = threading.RLock()
def log(msg):
    with _lock:
        print(msg, flush=True)

def load_conf():
    cfg = configparser.ConfigParser(); cfg.read(CONF_PATH, encoding="utf-8"); c = cfg[CONF_NODE]
    return {
        "acct_id": c.get("X-KDApi-AcctID","").strip(), "username": c.get("X-KDApi-UserName","").strip(),
        "app_id": c.get("X-KDApi-AppID","").strip(), "app_secret": c.get("X-KDApi-AppSec","").strip(),
        "server_url": c.get("X-KDApi-ServerUrl","").strip().rstrip("/"),
        "lcid": int(c.get("X-KDApi-LCID","2052") or "2052"),
    }

CONF = load_conf()

class KdClient:
    def __init__(self):
        self.s = requests.Session()
        r = self.s.post(f"{CONF['server_url']}/{LOGIN_SVC}",
                        data=json.dumps({"parameters":[CONF["acct_id"], CONF["username"], CONF["app_id"], CONF["app_secret"], CONF["lcid"]]}, ensure_ascii=False).encode("utf-8"),
                        headers={"Content-Type": "application/json;charset=utf-8"}, timeout=90)
        r.raise_for_status()
        res = r.json()
        if not (isinstance(res, dict) and res.get("LoginResultType") == 1):
            raise RuntimeError(f"登录失败：{res}")

    def query(self, field_keys, flt, order, start=0, limit=PAGE, timeout=60, retries=2):
        q = {"FormId": FORM_ID, "FieldKeys": field_keys, "FilterString": flt,
             "OrderString": order, "TopRowCount": 0, "StartRow": start, "Limit": limit}
        body = json.dumps(q, ensure_ascii=False)
        for attempt in range(retries):
            try:
                r = self.s.post(f"{CONF['server_url']}/{QUERY_SVC}",
                                data=json.dumps({"parameters":[body]}, ensure_ascii=False).encode("utf-8"),
                                headers={"Content-Type": "application/json;charset=utf-8"}, timeout=timeout)
                r.raise_for_status()
                data = r.json()
            except Exception as e:
                if attempt == retries-1: raise
                time.sleep(1.5*(attempt+1)); continue
            if isinstance(data, dict):
                raise RuntimeError(f"查询报错：{json.dumps(data, ensure_ascii=False)[:300]}")
            if data:
                first = data[0]
                cell = first[0] if isinstance(first, list) and first else first
                if isinstance(cell, dict):
                    try:
                        errs = cell["Result"]["ResponseStatus"]["Errors"]
                        msg = "；".join(e.get("Message","") for e in errs)
                    except Exception:
                        msg = json.dumps(cell, ensure_ascii=False)
                    raise RuntimeError(f"查询报错：{msg}")
            return data
        return []

    def voucher_entries(self, bno, zi, no):
        """逐张凭证拉全部分录（自动分页）。"""
        flt = f"FAccountBookID.FNumber='{bno}' and FVOUCHERGROUPNO={no}"
        if zi and zi != "?":
            flt += f" and FVOUCHERGROUPID.FName='{zi}'"
        keys = ",".join(f for f, _ in FIELDS)
        rows, start = [], 0
        while True:
            data = self.query(keys, flt, "FVOUCHERGROUPNO", start=start)
            rows.extend(data)
            if len(data) < PAGE:
                break
            start += PAGE
        return rows


def fnum(v):
    try: return float(v) if v not in (None, "") else 0.0
    except (TypeError, ValueError): return 0.0


def audit_book(book_name, rows, idx):
    def col(r, cn): return r[idx[cn]]

    vouchers = defaultdict(list)
    for r in rows:
        vouchers[(col(r,"凭证字") or "", col(r,"凭证号"))].append(r)

    v_issues, e_issues = [], []
    unchecked = unposted = 0
    monsters = []
    for (zi, no), items in sorted(vouchers.items(), key=lambda kv: (kv[0][0], fnum(kv[0][1]))):
        debit  = sum(fnum(col(r,"借方金额")) for r in items)
        credit = sum(fnum(col(r,"贷方金额")) for r in items)
        checker = (col(items[0],"审核") or "").strip()
        poster  = (col(items[0],"过账") or "").strip()
        creator = (col(items[0],"制单") or "").strip()
        if len(items) >= 2000:
            monsters.append((zi, no, len(items)))
        if abs(debit - credit) > 0.005:
            v_issues.append(("错误", zi, no, f"借贷不平衡：借 {debit:,.2f} / 贷 {credit:,.2f}"))
        if not checker:
            unchecked += 1
            v_issues.append(("提示", zi, no, "未审核"))
        if not poster:
            unposted += 1
        if checker and creator and checker == creator:
            v_issues.append(("风险", zi, no, f"制单人与审核人同为 {creator}（制审未分离）"))

    by_zi = defaultdict(list)
    for (zi, no) in vouchers:
        try: by_zi[zi].append(int(float(no)))
        except (TypeError, ValueError): pass
    for zi, nums in by_zi.items():
        ns = sorted(set(nums))
        if not ns: continue
        missing = sorted(set(range(ns[0], ns[-1]+1)) - set(ns))
        if missing:
            tail = "..." if len(missing) > 20 else ""
            v_issues.append(("提示", zi, "-", f"凭证号断号：缺 {missing[:20]}{tail}（区间 {ns[0]}~{ns[-1]}）"))

    for (zi, no), items in vouchers.items():
        for r in items:
            exp  = (col(r,"摘要") or "").strip()
            d, c = fnum(col(r,"借方金额")), fnum(col(r,"贷方金额"))
            if not exp:
                e_issues.append(("摘要空白", zi, no, col(r,"科目编码"), col(r,"科目名称"), ""))
            if d > 0 and c > 0:
                e_issues.append(("借贷同录", zi, no, col(r,"科目编码"), col(r,"科目名称"), f"借{d:,.2f}/贷{c:,.2f}"))

    summary = {
        "凭证张数": len(vouchers), "分录条数": len(rows),
        "借方合计": sum(fnum(col(r,"借方金额")) for r in rows),
        "未审核": unchecked, "未过账": unposted,
        "凭证级问题": len(v_issues), "分录级问题": len(e_issues),
        "怪兽凭证": monsters,
    }
    return v_issues, e_issues, summary


def main():
    log("登录金蝶云星空 ...")
    probe_client = KdClient()
    log("  登录成功")

    # 0) 凭证清单（头字段 → 一行一张凭证）
    log(f"\n[0/4] 拉取 {DATE_FROM}~{DATE_TO} 凭证清单 ...")
    keys = "FAccountBookID.FNumber,FAccountBookID.FName,FVOUCHERGROUPID.FName,FVOUCHERGROUPNO,FDate,FCreatorId.FName,FCHECKERID.FName,FPOSTERID.FName"
    flt = f"FDATE>='{DATE_FROM}' and FDATE<='{DATE_TO}'"
    probe_rows = probe_client.query(keys, flt, "FAccountBookID.FNumber,FVOUCHERGROUPID.FName,FVOUCHERGROUPNO")
    book_name = {}
    vouchers = []   # (bno, zi, no)
    seen = set()
    for r in probe_rows:
        bno, bname, zi, no = r[0], r[1], r[2] or "?", r[3]
        book_name[bno] = bname
        try: no_i = int(float(no))
        except (TypeError, ValueError): continue
        if (bno, zi, no_i) in seen: continue
        seen.add((bno, zi, no_i))
        vouchers.append((bno, zi, no_i))
    log(f"  共 {len(vouchers)} 张凭证，{len(book_name)} 个账簿")

    # 1) 逐张凭证拉分录（多线程）
    log(f"\n[1/4] 逐张凭证拉取分录（{THREADS} 线程）...")
    q = Queue()
    for v in vouchers: q.put(v)
    result = {}
    errors = []
    slow = []
    done = [0]
    def worker():
        client = KdClient()
        while True:
            try: bno, zi, no = q.get_nowait()
            except Exception: return
            t1 = time.time()
            try:
                rows = client.voucher_entries(bno, zi, no)
                result[(bno, zi, no)] = rows
                dt = time.time() - t1
                if dt > 10 or len(rows) >= PAGE:
                    slow.append((dt, bno, zi, no, len(rows)))
            except Exception as e:
                errors.append((bno, zi, no, str(e)))
            with _lock:
                done[0] += 1
                if done[0] % 25 == 0:
                    log(f"    进度 {done[0]}/{len(vouchers)}（已用 {time.time()-t0:.0f}s，失败 {len(errors)}）")
    threads = [threading.Thread(target=worker, daemon=True) for _ in range(THREADS)]
    t0 = time.time()
    for t in threads: t.start()
    for t in threads: t.join()
    log(f"  完成 {done[0]}/{len(vouchers)} 张，用时 {time.time()-t0:.0f}s，失败 {len(errors)} 张")
    for e in errors[:10]:
        log(f"    ⚠ 拉取失败：{e}")
    for dt, bno, zi, no, n in sorted(slow, reverse=True)[:15]:
        log(f"    慢/大凭证：{zi}字{no}号 [{bno}] {n} 条，{dt:.0f}s")

    headers = [cn for _, cn in FIELDS]
    idx = {cn: i for i, cn in enumerate(headers)}
    book_rows = defaultdict(list)
    for (bno, zi, no), rows in result.items():
        book_rows[bno].extend(rows)

    # 2) 资金类科目维度核查
    log(f"\n[2/4] 检查资金类科目银行账号维度 ...")
    cash_no_dim = []
    for bno in sorted(book_name):
        pre = " or ".join(f"FACCOUNTID.FNumber like '{p}%'" for p in CASH_PREFIXES)
        flt = (f"({pre}) and FDATE>='{DATE_FROM}' and FDATE<='{DATE_TO}' "
               f"and FAccountBookID.FNumber='{bno}'")
        dim_keys = ("FAccountBookID.FNumber,FVOUCHERGROUPID.FName,FVOUCHERGROUPNO,"
                    "FACCOUNTID.FNumber,FDEBIT,FCREDIT,FEXPLANATION,"
                    "FDetailID.FF100002.FNumber,FDetailID.FF100002.FName")
        try:
            rows = probe_client.query(dim_keys, flt, "FVOUCHERGROUPID.FName,FVOUCHERGROUPNO")
        except Exception as e:
            log(f"    ⚠ [{bno}] 维度查询失败：{e}（跳过）"); continue
        miss = [r for r in rows if not (r[7] or "").strip()]
        for r in miss:
            cash_no_dim.append([bno, book_name[bno], r[1], r[2], r[3], r[4], r[5]])
        log(f"    [{bno}] {book_name[bno]}：资金类分录 {len(rows)} 条，缺维度 {len(miss)} 条")

    # 3) 稽核
    log(f"\n[3/4] 稽核 ...")
    all_v, all_e, summaries = [], [], []
    for bno in sorted(book_rows):
        v_iss, e_iss, summ = audit_book(book_name[bno], book_rows[bno], idx)
        focus = "  ← 重点关注" if FOCUS_KEYWORD in book_name[bno] else ""
        log(f"\n■ {bno} {book_name[bno]}{focus}")
        log(f"  凭证 {summ['凭证张数']} 张 / 分录 {summ['分录条数']} 条 / 借方合计 {summ['借方合计']:,.2f}")
        log(f"  未审核 {summ['未审核']} / 未过账 {summ['未过账']} / 凭证级问题 {summ['凭证级问题']} / 分录级问题 {summ['分录级问题']}")
        if summ["怪兽凭证"]:
            mstr = "、".join(f"{zi}字{no}号({n}条)" for zi, no, n in summ["怪兽凭证"][:8])
            log(f"  超大凭证（≥2000条分录）：{mstr}{' ...' if len(summ['怪兽凭证'])>8 else ''}")
        shown = 0
        for lv, zi, no, msg in v_iss:
            if shown < 15:
                log(f"    [{lv}] {zi}字 {no}号：{msg}"); shown += 1
        if len(v_iss) > 15:
            log(f"    ... 其余 {len(v_iss)-15} 项见 Excel")
        for it in e_iss[:5]:
            log(f"    [分录] {it[0]} {it[1]}字 {it[2]}号 {it[3]} {it[4]}")
        if len(e_iss) > 5:
            log(f"    ... 其余 {len(e_iss)-5} 条见 Excel")
        if not v_iss and not e_iss:
            log("  无问题")
        for lv, zi, no, msg in v_iss:
            all_v.append([bno, book_name[bno], lv, zi, no, msg])
        for it in e_iss:
            all_e.append([bno, book_name[bno], *it])
        summaries.append([bno, book_name[bno], summ['凭证张数'], summ['分录条数'],
                          round(summ['借方合计'], 2), summ['未审核'], summ['未过账'],
                          summ['凭证级问题'], summ['分录级问题'], len(summ["怪兽凭证"])])
    for row in cash_no_dim:
        all_e.append([row[0], row[1], "缺银行账号维度", row[2], row[3], row[4], "", f"借{row[5]}/贷{row[6]}"])

    # 4) 落盘
    log(f"\n[4/4] 保存结果 ...")
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = Workbook()
    ws = wb.active; ws.title = "汇总"
    ws.append(["账簿编码","账簿","凭证张数","分录条数","借方合计","未审核","未过账","凭证级问题数","分录级问题数","超大凭证数"])
    for row in summaries: ws.append(row)
    ws1 = wb.create_sheet("凭证级问题")
    ws1.append(["账簿编码","账簿","级别","凭证字","凭证号","问题"])
    for row in all_v: ws1.append(row)
    ws2 = wb.create_sheet("分录级问题")
    ws2.append(["账簿编码","账簿","类型","凭证字","凭证号","科目编码","科目名称","备注"])
    for row in all_e: ws2.append(row)
    ws3 = wb.create_sheet("全部分录")
    ws3.append(headers)
    for bno in sorted(book_rows):
        for r in book_rows[bno]: ws3.append(list(r))
    xlsx = os.path.join(OUT_DIR, "凭证稽核_全主体_202609.xlsx")
    wb.save(xlsx)
    all_rows = [r for bno in sorted(book_rows) for r in book_rows[bno]]
    with open(os.path.join(OUT_DIR, "vouchers_raw.json"), "w", encoding="utf-8") as fp:
        json.dump({"headers": headers, "rows": all_rows}, fp, ensure_ascii=False)
    log(f"  Excel：{xlsx}")
    log("完成。")


if __name__ == "__main__":
    main()
