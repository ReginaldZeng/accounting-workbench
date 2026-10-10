"""抖音月结接口层：真的路由 + 一次性临时库；金蝶整个打桩，不连任何真实服务。"""
import gzip, hashlib, io, json, os, sys, tempfile, time, types, unittest, zipfile
from pathlib import Path
from unittest.mock import patch
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / '01_Current_Deliverables/app/backend'))
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import insert, select
from kernels import ec_settle, ec_preparation, ec_douyin as m

SHOP, TMALL = '星期零STARFIELD 抖音官旗店', ec_preparation.TARGET
P, Q, Z = '6917926768823643751', '6917926768823643752', '6917926768823643753'
PLATFORM_HEAD = '主订单编号,子订单编号,选购商品,商品规格,商品数量,商家编码,订单应付金额,运费,收件人,收件人手机号,详细地址,订单提交时间,商家备注,订单完成时间,订单状态,售后状态,取消原因,发货时间,平台实际承担优惠金额,商家实际承担优惠金额,达人实际承担优惠金额'
SETTLE_HEAD = '动账时间,动帐流水号,动账方向,动账金额,动账账户,动账场景,计费类型,子订单号,订单号,订单实付应结,运费实付,实际平台补贴,订单退款,平台服务费,佣金,招商服务费,站外推广费,备注'
NEW, OLD = '1791430000_new.csv', '1791213782_old.csv'            # 文件名开头是导出时刻：NEW 比 OLD 晚两天半


def platform_csv(q_status='已发货', q_done='-', q_after='-', memo='-'):
    return ('﻿' + PLATFORM_HEAD + '\n' + '\n'.join([
        f'"\t{P}","\t{P}","\t辣丝丝","\t1袋",1,"\t300400008",28.90,0.00,##密文,$$密文,##密文,2026-09-22 00:54:10,{memo},2026-09-24 13:18:29,已完成,退款成功,,2026-09-22 08:53:45,21.00,25.00,0.00',
        f'"\t{Q}","\t{Q}","\t辣丝丝","\t1袋",1,"\t300400008",20.00,0.00,##密文,$$密文,##密文,2026-09-25 10:00:00,-,{q_done},{q_status},{q_after},,2026-09-25 18:00:00,5.00,0.00,8.00',
        f'"\t{Z}","\t{Z}","\t辣丝丝","\t1袋",1,"\t300400008",15.90,0.00,##密文,$$密文,##密文,2026-10-02 10:00:00,-,-,待发货,-,,-,0.00,0.00,0.00']) + '\n').encode('utf-8')


def settle_csv(*lines):
    return ('﻿' + SETTLE_HEAD + '\n' + '\n'.join(lines) + '\n').encode('utf-8')


FIRST_HALF = f"2026-09-03 08:00:00,'T0,入账,9.00,聚合账户,货款结算入账,小店自卖,'6917926768823643799,'6917926768823643799,10,0,0,0,-1.00,0,0,0,订单结算"
SECOND_HALF = f"2026-09-28 07:54:38,'T1,入账,1.84,聚合账户,货款结算入账,巨量千川,'{P},'{P},28.9,0,0.81,-27.78,-0.09,0,0,0,订单结算"


class DouyinApiTests(unittest.TestCase):
    MODULES = ('db', 'core', 'kingdee_client', 'routers.ec', 'routers.ec_workbench', 'routers.ec_douyin')

    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix='ec-douyin-api-')
        cls.previous_env = os.environ.get('DB_URL')
        os.environ['DB_URL'] = 'sqlite:///' + (Path(cls.tmp.name) / 'test.sqlite').as_posix()
        import routers
        cls.routers = routers
        cls.modules = {k: sys.modules.get(k) for k in cls.MODULES}                 # 用完原样放回去，别串到其它测试文件
        cls.attrs = {k: getattr(routers, k, None) for k in ('ec', 'ec_workbench', 'ec_douyin')}
        for k in cls.MODULES: sys.modules.pop(k, None)
        for k in cls.attrs:                                                        # 包上要是挂着别的测试留下的真模块，下面的 import 会直接拿到旧的
            if hasattr(routers, k): delattr(routers, k)
        import db
        cls.db = db
        cls.role = ['admin']
        def require(request, permission):
            if permission == 'ec_settle_upload' and cls.role[0] not in ('admin', 'writer'): raise HTTPException(403, '测试权限拦截')   # writer＝有上传 / 跑批权限的普通账号
            return {'name': '测试会计', 'role': cls.role[0]}
        core = types.ModuleType('core'); core.db = db; core._require_perm = require; core.pull_token_ok = lambda request: False
        sys.modules['core'] = core
        cls.kd = {'calls': [], 'rows': [], 'err': None, 'boom': ''}
        kc = types.ModuleType('kingdee_client')
        def login():
            if cls.kd['boom']: raise RuntimeError(cls.kd['boom'])
            return object(), {}
        def query_raw(s, conf, form, keys, flt, start):
            cls.kd['calls'].append(flt); return cls.kd['rows'], cls.kd['err']
        kc.login, kc._query_raw = login, query_raw
        sys.modules['kingdee_client'] = kc
        cls.now = ['2026-10-06 09:00:00']
        fake_ec = types.SimpleNamespace(fulfillment=types.SimpleNamespace(TARGET_SHOP=TMALL), es=ec_settle, _now=lambda: cls.now[0], _KD_REFRESH={},
                                        _kd_cache_meta=lambda p: None, EC_UPLOAD_DIR=str(Path(cls.tmp.name) / 'ec_uploads'), _tmall_latest_rows=lambda p: {})
        sys.modules['routers.ec'] = fake_ec; routers.ec = fake_ec
        from routers import ec_workbench as wbr, ec_douyin as dy
        cls.wbr, cls.dy = wbr, dy
        app = FastAPI(); app.include_router(wbr.router); app.include_router(dy.router)
        cls.client = TestClient(app)

    @classmethod
    def tearDownClass(cls):
        cls.client.close(); cls.db._engine.dispose()
        for k, v in cls.modules.items():
            if v is None: sys.modules.pop(k, None)
            else: sys.modules[k] = v
        for k, v in cls.attrs.items():
            if v is None:
                if hasattr(cls.routers, k): delattr(cls.routers, k)
            else: setattr(cls.routers, k, v)
        if cls.previous_env is None: os.environ.pop('DB_URL', None)
        else: os.environ['DB_URL'] = cls.previous_env
        cls.tmp.cleanup()

    def setUp(self):                                                                # 每条用例一张白纸
        with self.db._engine.begin() as cx: cx.execute(self.wbr.TABLE.delete())
        self.db.set_setting('ec_douyin_ar_meta', {}); self.db.set_setting('ec_preparation_rules', {})
        self.dy._cache.clear(); self.dy._bill_cache.clear(); self.dy._jobs.clear()
        self.kd.update(calls=[], rows=[], err=None, boom=''); self.role[0] = 'admin'

    # ---- 小工具 ----
    def up(self, *files):
        r = self.client.post('/api/ec/douyin/upload', data={'shop': SHOP}, files=[('files', (n, b, 'application/octet-stream')) for n, b in files])
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()['results']

    def latest(self, period, kind):
        T = self.wbr.TABLE
        with self.db._engine.connect() as cx:
            return cx.execute(select(T).where(T.c.period == period, T.c.shop == SHOP, T.c.kind == kind).order_by(T.c.id.desc())).fetchall()

    def rows(self, period, kind):
        return json.loads(gzip.decompress(self.latest(period, kind)[0].payload))['rows']

    def cards(self, period='2026-09', shop=SHOP):
        r = self.client.get('/api/ec/workbench/sources', params={'period': period, 'shop': shop})
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def sync_ar(self, bills, period='2026-09'):
        """假装同步过金蝶应收：写快照文件和同步记录。"""
        path = self.dy._ar_path(period, SHOP); os.makedirs(os.path.dirname(path), exist_ok=True)
        with gzip.open(path, 'wt', encoding='utf-8') as f: json.dump({'bills': bills, 'book': {}}, f)
        self.db.set_setting('ec_douyin_ar_meta', {'%s|%s' % (period, SHOP): {'ts': self.now[0], 'operator': '测试会计', 'bills': len(bills), 'since': '2026-06-01'}})

    @staticmethod
    def bill(no, date, amount, order, written=0, ws='A'):
        return dict(no=no, date=date, amount=amount, written=written, open=round(amount - written, 2), ws=ws, ds='C', order=order, ship='')

    # ---- 上传：判重按内容、旧导出不盖新状态 ----
    def test_platform_upload_dedupes_by_content_and_keeps_newer_status(self):
        first = self.up((OLD, platform_csv()))[0]
        self.assertEqual((first['ok'], first['kind'], first['duplicate'], first['rows']), (True, 'dy_platform', False, 3))
        self.assertEqual((len(self.latest('2026-09', 'dy_platform')), len(self.latest('2026-10', 'dy_platform'))), (1, 1))    # 按下单月份各存一份
        self.assertEqual(self.up((OLD, platform_csv()))[0]['duplicate'], True)                                             # 原样再传：重复
        self.assertEqual(len(self.latest('2026-09', 'dy_platform')), 1)
        row = self.rows('2026-09', 'dy_platform')[0]
        self.assertEqual(row['x'], '2026-10-05 23:23')                                                                     # 状态是哪次导出时的：取自文件名
        self.assertFalse(any('密文' in str(v) for r in self.rows('2026-09', 'dy_platform') for v in r.values()))
        newer = self.up((NEW, platform_csv(q_status='已完成', q_done='2026-10-02 09:00:00')))[0]
        self.assertEqual(newer['duplicate'], False)                                                                        # 状态变了：存成新版本
        self.assertEqual(len(self.latest('2026-10', 'dy_platform')), 1)                                                    # 10 月那份没变：不多存
        back = self.up((OLD, platform_csv()))[0]                                                                           # 手滑又传了一遍旧导出
        self.assertEqual(back['duplicate'], True); self.assertIn('这份导出比系统里的旧，2 行没有覆盖', back['warnings'])   # 9 月两行；10 月那行内容没变过，谈不上新旧
        q = next(r for r in self.rows('2026-09', 'dy_platform') if r['order'] == Q)
        self.assertEqual((q['status'], q['done']), ('已完成', '2026-10-02 09:00:00'))                                        # 库里还是新的状态
        only_memo = self.up(('1791500000_memo.csv', platform_csv(q_status='已完成', q_done='2026-10-02 09:00:00', memo='补发一包')))[0]
        self.assertEqual(only_memo['duplicate'], False)                                                                    # 只有备注变了：也算新内容
        self.assertEqual(next(r for r in self.rows('2026-09', 'dy_platform') if r['order'] == P)['memo'], '补发一包')

    def test_memo_with_contact_details_never_reaches_the_snapshot(self):
        self.up((OLD, platform_csv(memo='改地址 广东省深圳市南山区科技路1号 13800138000')))
        memo = next(r for r in self.rows('2026-09', 'dy_platform') if r['order'] == P)['memo']
        self.assertEqual(memo, m.MEMO_HIDDEN)
        packed = gzip.decompress(self.latest('2026-09', 'dy_platform')[0].payload).decode('utf-8')
        self.assertNotIn('13800138000', packed); self.assertNotIn('科技路', packed)

    def test_bad_files_get_a_plain_reason_each_and_do_not_sink_the_batch(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z: z.writestr('订单.csv', platform_csv())
        locked = bytearray(buf.getvalue())
        for sig, at in ((b'PK\x03\x04', 6), (b'PK\x01\x02', 8)):
            i = locked.find(sig); locked[i + at] |= 1
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z: z.writestr('订单.csv', platform_csv() * 50)
        broken = bytearray(buf.getvalue()); broken[60:70] = b'\xff' * 10                      # 压缩数据坏了
        slashed = settle_csv("2026/9/10 8:00,'R2,入账,48.90,聚合账户,货款结算入账,巨量千川,'%s,'%s,49.9,0,0,0,-1.00,0,0,0,订单结算" % (Q, Q))
        got = self.up(('带密码.zip', bytes(locked)), ('坏了.zip', bytes(broken)), ('另存过.csv', slashed), ('不相干.csv', '甲,乙\n1,2\n'.encode('utf-8')), (OLD, platform_csv()))
        self.assertEqual([r['ok'] for r in got], [False, False, False, False, True])           # 前四个各有各的原因，最后一个照常入库
        self.assertIn('带密码', got[0]['error']); self.assertIn('压缩包损坏', got[1]['error']); self.assertIn('被改过格式', got[2]['error']); self.assertIn('表头不是', got[3]['error'])
        T = self.wbr.TABLE
        with self.db._engine.connect() as cx:
            periods = {r.period for r in cx.execute(select(T.c.period))}
        self.assertEqual(periods, {'2026-09', '2026-10'})                                      # 没有存进“2026/9/”这种不存在的月份
        with patch.object(self.dy, 'MAX_UPLOAD', 100):
            self.assertIn('超过 60 MB', self.up(('太大.csv', platform_csv()))[0]['error'])
        self.role[0] = 'reader'
        self.assertEqual(self.client.post('/api/ec/douyin/upload', data={'shop': SHOP}, files=[('files', (OLD, platform_csv()))]).status_code, 403)

    def test_old_format_settle_is_still_ready_and_reupload_in_two_files_clears_the_note(self):
        # 造一份系统升级前存的动账明细：行里没有买家实付 / 补贴，摘要里没有 extra，来源是上半月、下半月两个文件
        strip = lambda rows: [{k: v for k, v in r.items() if k not in ('paid', 'subsidy', 'ot', 'goods')} for r in rows]
        old_rows = strip(m.parse(settle_csv(FIRST_HALF, SECOND_HALF), 'x.csv')[0]['rows'])
        T = self.wbr.TABLE
        with self.db._engine.begin() as cx:
            cx.execute(insert(T).values(period='2026-09', shop=SHOP, kind='dy_settle', digest=hashlib.sha256(b'old').hexdigest(), filenames=json.dumps(['上半月.csv', '下半月.csv']),
                                        payload=gzip.compress(json.dumps({'rows': old_rows, 'status': 'ready'}).encode()), operator='老会计', ts='2026-10-05 15:48:00',
                                        summary=json.dumps({'rows': 2, 'valid': True, 'warnings': []})))
        j = self.cards(); card = next(c for c in j['sources'] if c['kind'] == 'dy_settle')
        self.assertEqual((card['state'], card['available'], card['by'], card['span']), ('ready', True, '老会计', ['2026-09-03', '2026-09-28']))   # 资料是好的：不算缺
        self.assertIn('系统升级前传的', card['warnings'][0]); self.assertNotIn('解析', card['warnings'][0])
        self.assertIn('extra', json.loads(self.latest('2026-09', 'dy_settle')[0].summary))                                     # 摘要补写回库，下次不用再解压
        self.assertEqual(self.up(('上半月.csv', settle_csv(FIRST_HALF)))[0]['duplicate'], False)
        self.assertTrue(next(c for c in self.cards()['sources'] if c['kind'] == 'dy_settle')['warnings'])                       # 只重传了一半：提醒还在
        self.assertEqual(self.up(('下半月.csv', settle_csv(SECOND_HALF)))[0]['duplicate'], False)                                 # 另一半不会被当成重复
        card = next(c for c in self.cards()['sources'] if c['kind'] == 'dy_settle')
        self.assertEqual((card['state'], card['warnings']), ('ready', []))
        self.assertTrue(all('paid' in r for r in self.rows('2026-09', 'dy_settle')))

    def test_same_rows_in_another_order_is_a_duplicate_and_renamed_export_does_not_block_newer_one(self):
        a = "2026-09-03 08:00:00,'S1,入账,9.00,聚合账户,货款结算入账,小店自卖,'%s,'%s,10,0,0,0,-1.00,0,0,0,订单结算" % (P, P)
        b = "2026-09-03 08:00:00,'S2,入账,8.00,聚合账户,货款结算入账,小店自卖,'%s,'%s,9,0,0,0,-1.00,0,0,0,订单结算" % (Q, Q)
        self.assertEqual(self.up(('DL.csv', settle_csv(a, b)))[0]['duplicate'], False)
        self.assertEqual(self.up(('DL.csv', settle_csv(b, a)))[0]['duplicate'], True)                    # 同一秒的两行换了先后：内容没变，算重复
        self.assertEqual(len(self.latest('2026-09', 'dy_settle')), 1)
        # 改过名的订单导出：认不出导出时刻就不记；之后传一份文件名正常、状态更新的，照样盖得进去
        self.now[0] = '2026-10-06 09:00:00'
        self.up(('9月订单导出.csv', platform_csv()))
        self.assertEqual({r['x'] for r in self.rows('2026-09', 'dy_platform')}, {''})
        newer = self.up((OLD, platform_csv(q_status='已完成', q_done='2026-10-02 09:00:00')))[0]           # 导出时刻 10-05，早于刚才的上传时刻
        self.assertEqual(newer['duplicate'], False); self.assertFalse(any('没有覆盖' in w for w in newer['warnings']))
        self.assertEqual(next(r for r in self.rows('2026-09', 'dy_platform') if r['order'] == Q)['status'], '已完成')
        back = self.up(('又改了名.csv', platform_csv()))[0]                                                # 没有导出时刻的旧状态再传：靠“状态倒退”拦住
        self.assertIn('这份导出比系统里的旧，1 行没有覆盖', back['warnings'])

    def test_wdt_orders_with_reformatted_dates_are_refused_whole_and_old_snapshot_gets_a_note(self):
        from openpyxl import Workbook
        def xlsx(rows):
            wb = Workbook(); ws = wb.active
            ws.append(['订单编号', '店铺', '子单原始单号', '订单状态', '订单退款状态', '交易时间', '付款时间', '发货时间', '收件人', '买家实付', '应收金额', '货品名称', '实发数量', '分摊后总价'])
            for r in rows: ws.append(r)
            buf = io.BytesIO(); wb.save(buf); return buf.getvalue()
        good = ['JY1', SHOP, P, '已完成', '', '2026-09-03 10:00:00', '2026-09-03 10:00:05', '2026-09-03 18:00:00', '张三', 15.9, 15.9, '豆腐', 1, 15.9]
        bad = ['JY2', SHOP, Q, '已完成', '', '2026/10/3 10:00', '2026/10/3 10:00', '2026/10/3 18:00', '张三', 19.8, 19.8, '年糕', 1, 19.8]
        got = self.up(('订单明细.xlsx', xlsx([good, bad])))[0]
        self.assertEqual(got['ok'], False); self.assertIn('日期被改过格式', got['error']); self.assertNotIn('YYYY', got['error'])
        self.assertEqual(self.latest('2026-09', 'dy_orders'), [])                                          # 没有存进去一半
        # 系统升级前存的旺店通订单明细（摘要里没有 format）：清单上提醒重传；重传后提醒消失
        T = self.wbr.TABLE
        old_rows = [dict(r, paid=1.0) for r in m.parse(xlsx([good]), '订单明细.xlsx', SHOP)[0]['rows']]
        with self.db._engine.begin() as cx:
            cx.execute(insert(T).values(period='2026-09', shop=SHOP, kind='dy_orders', digest=hashlib.sha256(b'old-wdt').hexdigest(), filenames=json.dumps(['订单明细.xlsx']),
                                        payload=gzip.compress(json.dumps({'rows': old_rows, 'status': 'ready'}).encode()), operator='老会计', ts='2026-10-05 22:45:00',
                                        summary=json.dumps({'rows': 1, 'valid': True, 'warnings': []})))
        card = next(c for c in self.cards()['sources'] if c['kind'] == 'dy_orders')
        self.assertEqual(card['state'], 'ready'); self.assertIn('买家实付可能偏小', card['warnings'][0])
        self.assertEqual(self.up(('订单明细.xlsx', xlsx([good])))[0]['duplicate'], False)
        card = next(c for c in self.cards()['sources'] if c['kind'] == 'dy_orders')
        self.assertEqual((card['warnings'], self.rows('2026-09', 'dy_orders')[0]['paid']), ([], 15.9))

    # ---- 数据准备清单 ----
    def test_checklist_always_has_four_douyin_kinds_and_reports_sync_state(self):
        self.db.set_setting('ec_preparation_rules', {SHOP: []})                                # 基础资料页那一行被动过：抖音店照样四类
        j = self.cards()
        self.assertEqual([(c['kind'], bool(c.get('optional'))) for c in j['sources']], [('dy_settle', False), ('dy_ledger', False), ('dy_orders', False), ('dy_platform', False), ('dy_returns', True)])
        self.assertTrue(all(c['purpose'] for c in j['sources']))
        self.assertEqual((j['douyin_ar'], j['douyin_job']), (None, {'running': False, 'error': ''}))
        self.sync_ar([self.bill('AR1', '2026-09-22', 49.9, P)])
        self.dy._jobs[('2026-09', SHOP)] = {'running': True, 'error': ''}
        j = self.cards()
        self.assertEqual((j['douyin_ar']['bills'], j['douyin_ar']['by'], j['douyin_job']['running']), (1, '测试会计', True))
        self.dy._jobs[('2026-09', SHOP)] = {'running': False, 'error': '金蝶同步失败（ConnectionError），请检查金蝶连接后重试'}
        self.assertIn('金蝶同步失败', self.cards()['douyin_job']['error'])
        # 真走一次同步、让金蝶连不上：留给页面的那句话是白话，不带程序里的英文名
        self.kd['boom'] = 'HTTPSConnectionPool(host=kd.example, port=8090): Max retries exceeded'
        self.assertEqual(self.client.post('/api/ec/douyin/ar-refresh', data={'period': '2026-09', 'shop': SHOP}).status_code, 200)
        for _ in range(200):
            if not self.dy._jobs[('2026-09', SHOP)]['running']: break
            time.sleep(0.02)
        said = self.cards()['douyin_job']['error']
        self.assertIn('没有同步成', said); self.assertNotRegex(said, r'[A-Za-z]{4,}')
        tm = self.cards(shop=TMALL)
        self.assertEqual((tm['douyin_ar'], tm['douyin_job']), (None, None))
        self.assertTrue(all('by' in c and 'span' in c for c in tm['sources']))

    # ---- 对账结果：缓存对得上就不再解压；平台状态只给还挂着的蓝字 ----
    def test_result_cache_skips_decompress_and_pstate_skips_red_bills(self):
        self.up(('DL.csv', settle_csv(FIRST_HALF)), (OLD, platform_csv()))
        self.sync_ar([self.bill('ARQ', '2026-09-25', 33.0, Q), self.bill('ARR', '2026-09-26', -5.0, Q)])
        args = {'period': '2026-09', 'shop': SHOP}
        with patch.object(self.dy, '_sources', wraps=self.dy._sources) as loaded:
            first = self.client.get('/api/ec/douyin/bills', params=args).json()
            self.client.get('/api/ec/douyin/bills', params=args); self.client.get('/api/ec/douyin/settle', params=args)
            detail = self.client.get('/api/ec/douyin/order', params=dict(args, order=Q)).json()
            self.assertEqual(loaded.call_count, 1)                                              # 四类资料只解开一次
        got = {b['no']: b['pstate'] for b in first['rows']}
        self.assertEqual(got['ARQ'], '买家还没确认收货（按 10-05 的订单导出）'); self.assertEqual(got['ARR'], '')       # 红字不说“钱来不来”
        self.assertEqual((detail['platform_gross'], detail['platform_unshipped'], detail['platform_closed']), (33.0, 0.0, False))
        self.assertEqual(self.client.get('/api/ec/douyin/settle', params=args).json()['sources'], {'dy_settle': 1, 'dy_ledger': 0, 'dy_orders': 0, 'dy_platform': 2, 'dy_insure': 0, 'dy_returns': 0})
        self.up((NEW, platform_csv(q_status='已关闭', q_after='退款成功')))                                             # 资料有了新版本：缓存作废，重算
        again = {b['no']: b['pstate'] for b in self.client.get('/api/ec/douyin/bills', params=args).json()['rows']}
        self.assertEqual(again['ARQ'], '平台订单已关闭（退款成功），这笔钱不会结算了（按 10-08 的订单导出）')

    def test_flows_endpoint_lists_premiums_by_order_and_ledger_rows(self):
        ledger_head = '动账流水号,关联订单号,关联子订单号,动账时间,账户方向,动账金额(元),动账场景,账户余额(元),备注'
        ledger = ('\ufeff' + ledger_head + '\n' + '\n'.join([
            f"T0,,,2026-09-03 08:00:00,收入,9.00,货款结算入账,9.00,结算",
            f"X1,TRA2026,,2026-09-12 10:00:00,支出,-0.40,退换货运费险,8.60,保费扣除（2笔保单）",
            f"X2,{P},,2026-09-13 10:00:00,支出,-5.00,退款-结算后退款-退用户,3.60,退款"]) + '\n').encode('utf-8')
        insure = ('\ufeff保险单号,动账流水号,关联子订单号,摘要描述,动账时间,金额(元)\n' + f"B1,X1,{P},退换货运费险,2026-09-12 10:00:01,0.2\nB2,X1,{Q},退换货运费险,2026-09-12 10:00:02,0.2\n").encode('utf-8')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z: z.writestr('动账明细 - 聚合账户.csv', ledger); z.writestr('保费支出 - 聚合账户.csv', insure)
        got = self.up(('DL.csv', settle_csv(FIRST_HALF, SECOND_HALF)), ('FL.zip', buf.getvalue()))
        self.assertEqual([(r['kind'], r['rows']) for r in got], [('dy_settle', 2), ('dy_ledger', 3), ('dy_insure', 2)])
        self.sync_ar([self.bill('AR1', '2026-09-22', 49.9, P)])
        ask = lambda scene, **kw: self.client.get('/api/ec/douyin/flows', params=dict({'period': '2026-09', 'shop': SHOP, 'scene': scene}, **kw)).json()
        j = ask('退换货运费险')
        self.assertEqual((j['kind'], j['total'], j['amount'], j['insure_missing']), ('insure', 2, -0.4, False))
        self.assertEqual([(r['id'], r['flow'], r['order'], r['known']) for r in j['rows']], [('B1', 'X1', P, True), ('B2', 'X1', Q, False)])   # P 有应收，点得开；Q 没有
        self.assertEqual(ask('退换货运费险', q=Q)['total'], 1)
        j = ask('退款-结算后退款-退用户')
        self.assertEqual((j['kind'], [(r['id'], r['order'], r['amt'], r['known']) for r in j['rows']]), ('ledger', [('X2', P, -5.0, True)]))
        self.assertEqual(ask('没有这种')['total'], 0)
        self.assertEqual(self.client.get('/api/ec/douyin/settle', params={'period': '2026-09', 'shop': SHOP}).json()['sources']['dy_insure'], 2)

    def test_matched_order_with_a_red_bill_says_the_system_will_not_push_it(self):
        self.up(('DL.csv', settle_csv(FIRST_HALF, SECOND_HALF)))
        self.sync_ar([self.bill('AR1', '2026-09-22', 29.71, P), self.bill('AR1R', '2026-09-26', -27.78, P)])       # 蓝 − 红 ＝ 1.93 ＝ 流水应冲
        get = lambda path, **kw: self.client.get('/api/ec/douyin/' + path, params=dict({'period': '2026-09', 'shop': SHOP}, **kw)).json()
        self.assertEqual(sorted((b['no'], b['cat'], b['hold']) for b in get('bills', q=P)['rows']), [('AR1', 'ok', 'red'), ('AR1R', 'ok', 'red')])   # 对得上，但清单上要标出系统不推
        self.assertIn('红字', get('order', order=P)['hold'])                                                        # 抽屉里也说
        self.assertEqual(get('order', order='6917926768823643799')['hold'], '')

    def test_wdt_returns_upload_is_optional_and_annotates_list_and_drawer(self):
        head = '退换单号,类型,退款阶段,退换原因,处理状态,平台退款状态,货品名称,货品编号,登记数量,入库数量,登记时间,店铺,原始单号,原始子订单号,分摊退款金额,退款总额,退款成功时间'
        body = f'TK1,退款不退货,售后,与商家协商一致退款,已完成,退款成功,辣丝丝,="A1",1,0,="2026-09-27 10:00:00",{SHOP},="{P}",="{P}",27.7800,27.7800,="2026-09-27 10:00:05"'
        before = self.cards()
        got = self.up(('DL.csv', settle_csv(FIRST_HALF, SECOND_HALF)), ('退换单.csv', (head + '\n' + body + '\n合计:,NA,NA,NA,NA,NA,NA,NA,1,0,NA,NA,NA,NA,27.78,27.78,NA\n').encode('gb18030')))
        self.assertEqual([(r['kind'], r['rows']) for r in got], [('dy_settle', 2), ('dy_returns', 1)])
        after = self.cards()
        card = next(c for c in after['sources'] if c['kind'] == 'dy_returns')
        self.assertEqual((card['optional'], card['available'], card['rows'], bool(card['purpose'])), (True, True, 1, True))
        self.assertEqual([c['kind'] for c in after['sources'] if not c.get('optional')], ['dy_settle', 'dy_ledger', 'dy_orders', 'dy_platform'])
        self.assertEqual(after['progress']['required'], before['progress']['required'])                               # 不算进齐套
        self.sync_ar([self.bill('AR1', '2026-09-22', 29.71, P)])                                                          # 抖音结算时退了 27.78，金蝶只有蓝字
        get = lambda path, **kw: self.client.get('/api/ec/douyin/' + path, params=dict({'period': '2026-09', 'shop': SHOP}, **kw)).json()
        row = get('bills', q=P)['rows'][0]
        self.assertEqual(row['cat'], 'mismatch'); self.assertIn('登记了退款不退货 27.78', row['rnote']); self.assertIn('要手工补红字 27.78', row['rnote'])
        d = get('order', order=P)
        self.assertEqual(([(r['tk'], r['amt'], r['back']) for r in d['returns']], d['rnote']), ([('TK1', 27.78, 0.0)], row['rnote']))
        self.assertEqual(get('order', order='6917926768823643799')['returns'], [])

    def test_push_settings_open_to_anyone_who_may_push_and_every_change_is_recorded(self):
        post = lambda **kw: self.client.post('/api/ec/douyin/push/mode', data=kw)
        self.role[0] = 'viewer'
        try: self.assertEqual(post(size=10).status_code, 403)                                     # 只能看的账号改不了
        finally: self.role[0] = 'writer'
        try:
            r = post(mode='dry', size=10)                                                         # 有上传 / 跑批权限的普通账号（不是管理员）可以改
            self.assertEqual((r.status_code, r.json()['conf']['mode'], r.json()['conf']['size'], r.json()['conf']['by']), (200, 'dry', 10, '测试会计'))
            at = r.json()['conf']['at']; self.assertTrue(at)
            self.assertEqual(post(size=10001).status_code, 400)                                   # 上限还是金蝶一张收款单能挂的行数
            self.assertEqual(post(mode='dry', size=10).json()['conf']['at'], at)                  # 没改动不算一次改动
        finally: self.role[0] = 'admin'
        self.assertEqual(self.db.get_setting('ec_douyin_push')['by'], '测试会计')

    def test_pushed_list_can_be_downloaded_per_batch_and_skips_undone_batches(self):
        from openpyxl import load_workbook
        self.up(('DL.csv', settle_csv(FIRST_HALF, SECOND_HALF)))
        self.sync_ar([self.bill('AR1', '2026-09-22', 29.71, P), self.bill('AR2', '2026-09-01', 10.0, '6917926768823643799')])
        line = lambda amount: [dict(kind='cash', amount=amount, memo='这批订单结算到账的钱')]
        self.dy._push_log_save('2026-09', SHOP, [
            dict(fid=11, at='2026-10-10 15:00:00', by='测试会计', count=1, total=10.0, first='AR2', last='AR2', bills=['AR2'], lines=line(10.0)),
            dict(fid=12, at='2026-10-10 15:05:00', by='测试会计', count=1, total=29.71, first='AR1', last='AR1', bills=['AR1'], lines=line(29.71), deleted=True)])
        get = lambda **kw: self.client.get('/api/ec/douyin/push/export', params=dict({'period': '2026-09', 'shop': SHOP}, **kw))
        r = get()
        self.assertEqual(r.status_code, 200, r.text)
        rows = [list(x)[:3] for x in load_workbook(io.BytesIO(r.content))['下推清单'].iter_rows(values_only=True)]
        self.assertEqual(rows[1:], [['AR2', 10.0, '6917926768823643799'], ['合计 1 张', 10.0, None]])       # 撤回的那批不在清单里
        self.assertEqual(get(fid='11').status_code, 200)
        self.assertEqual(get(fid='12').status_code, 404)                                              # 已撤回的批没有清单可下
        self.role[0] = 'viewer'
        try: self.assertEqual(get().status_code, 200)                                                 # 只能看的账号也能下载（只读）
        finally: self.role[0] = 'admin'

    # ---- 抽屉里现查金蝶 ----
    def test_order_bills_only_asks_for_known_bills_caches_and_never_leaks_errors(self):
        self.up(('DL.csv', settle_csv(FIRST_HALF)))
        self.sync_ar([self.bill('AR1', '2026-09-22', 49.9, P), self.bill('AR1R', '2026-09-27', -49.9, P)])
        ask = lambda nos, shop=SHOP: self.client.get('/api/ec/douyin/order-bills', params={'period': '2026-09', 'shop': shop, 'nos': nos})
        self.kd['rows'] = [['AR1', '标准应收单', 'CK1', 'JY1', '开发', '2026-09-22T09:11:20.353', '黄春艳', '2026-09-24T15:06:26.51', '300400008', '辣丝丝', '袋', 1.0, 49.9, 13.0, 44.1593, 5.7407, 49.9, 'XQLCK1', 'SAL_OUTSTOCK']]
        r = ask("AR1, AR1R,AR1,不认识的,AR9') or 1=1 --")
        self.assertEqual((r.status_code, list(r.json()['bills'])), (200, ['AR1']))
        self.assertEqual(self.kd['calls'], ["FBillNo in ('AR1','AR1R')"])                       # 只认这家店对账结果里有的单号
        ask("AR1, AR1R"); self.assertEqual(len(self.kd['calls']), 1)                            # 几分钟内同一批单不重复查金蝶
        self.assertEqual(ask('不认识的').json(), {'ok': True, 'bills': {}}); self.assertEqual(len(self.kd['calls']), 1)
        self.dy._bill_cache.clear(); self.kd.update(rows=None, err='会话已过期 SessionId=abcdef0123456789 http://kd.example:8090/k3cloud')
        j = ask('AR1').json()
        self.assertEqual(j['bills'], {}); self.assertIn('金蝶这会儿没连上', j['error'])
        self.assertNotIn('SessionId', j['error']); self.assertNotIn('kd.example', j['error'])
        self.kd.update(err=None, boom='HTTPSConnectionPool(host=kd.example, port=8090): Max retries exceeded')
        j = ask('AR1').json()
        self.assertIn('金蝶这会儿没连上', j['error']); self.assertNotIn('kd.example', j['error'])
        self.assertEqual(ask('AR1', TMALL).status_code, 400)                                    # 不是抖音店
        # 应收重新同步以后，短缓存不再给同步前查的那份
        self.kd.update(boom='', rows=[['AR1'] + [''] * 18]); self.dy._bill_cache.clear()
        ask('AR1'); n = len(self.kd['calls']); ask('AR1'); self.assertEqual(len(self.kd['calls']), n)
        self.now[0] = '2026-10-06 10:00:00'
        self.sync_ar([self.bill('AR1', '2026-09-22', 49.9, P), self.bill('AR1R', '2026-09-27', -49.9, P)])
        ask('AR1'); self.assertEqual(len(self.kd['calls']), n + 1)


if __name__ == '__main__':
    unittest.main()
