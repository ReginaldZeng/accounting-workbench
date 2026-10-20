import sys
import unittest
from pathlib import Path
from io import BytesIO
from copy import deepcopy
from openpyxl import load_workbook
sys.path.insert(0,str(Path('02_Revision_History/20260703_V1.2_接真数据/app/backend').resolve()))
from kernels.temp_attendance_export import build_data, export_xlsx, person_id, HEADERS

def fixture():
 p={'姓名':'测试甲','归属':'服务甲','部门':'车间','岗位':'普工','班型':'夜班','上报总工时':10.5,
    '夜班工资单价':19,'夜班管理费单价':3,'应付工资':199.5,'应付管理费':31.5}
 r={**{k:p[k] for k in ('姓名','归属','部门','岗位','班型')},'日':31,'手机尾号':'0123',
    '上班打卡':'19:55','下班打卡':'次日07:00','无效卡':['23:58'],'跨度':11.0833,
    '规则扣减':.5,'重算工时':10.5,'上报工时':10.5,'差异':0,'档':'ok','判定':'正常'}
 p2={**p,'姓名':'测试乙','归属':'服务乙'}
 r2={**r,'姓名':'测试乙','归属':'服务乙','档':'hard','判定':'缺卡'}
 return {'people':[p,p2],'rows':[r,r2],'params':{}}

class ExportTests(unittest.TestCase):
 def test_scope_and_calendar(self):
  res=fixture();data,summary=build_data(res,'2026-08','agency',['服务甲'])
  self.assertEqual(len(data),33);self.assertEqual(len(summary),1)
  self.assertTrue(all(len(r)==21 for r,_,_ in data))
  self.assertEqual(str(data[0][0][3]),'2026-07-31');self.assertEqual(str(data[-1][0][3]),'2026-09-01')
  self.assertEqual(summary[0][10],231);self.assertEqual(summary[0][14],'一致')
  self.assertEqual(data[-2][0][9],11.0833)
  self.assertEqual(data[-2][0][1],'尾号 0123')
 def test_person_and_invalid_scope(self):
  res=fixture();d,s=build_data(res,'2026-08','person',people=[person_id(res['people'][1])])
  self.assertEqual(s[0][0],'测试乙')
  for mode in ('agency','person','bad'):
   with self.assertRaises(ValueError):build_data(res,'2026-08',mode)
 def test_missing_price_never_zero(self):
  res=fixture();res['people'][0]['合同缺档']='缺合同'
  d,s=build_data(res,'2026-08','agency',['服务甲'])
  self.assertIsNone(d[-2][0][18]);self.assertIsNone(s[0][10])
  self.assertEqual(s[0][14],'金额未完整计算')
 def test_identity_duplicates_rejected(self):
  res=fixture();res['people'].append(deepcopy(res['people'][0]))
  with self.assertRaises(ValueError):build_data(res,'2026-08')
 def test_workbook_features_and_text(self):
  res=fixture();res['people'][0]['姓名']='=1+1';res['rows'][0]['姓名']='=1+1'
  wb=load_workbook(BytesIO(export_xlsx(res,'2026-08')))
  self.assertEqual(wb.sheetnames,['服务方汇总','逐人汇总','逐日核对'])
  ws=wb['逐日核对'];self.assertEqual([c.value for c in ws[4]],HEADERS)
  cell=next(row[0] for row in ws.iter_rows(min_row=5) if row[0].value=='=1+1')
  self.assertEqual(cell.data_type,'s')
  self.assertEqual(ws.freeze_panes,'D5');self.assertTrue(ws.auto_filter.ref)
  self.assertEqual(ws['B5'].number_format,'@')
  self.assertEqual(ws['A5'].fill.fgColor.rgb,'00EEF0F3')
 def test_mixed_missing_shift_not_guessed(self):
  res=fixture();res['people'][0]['班型']='白夜混合';res['rows'][0].update(班型='白夜混合',跨度=None)
  d,s=build_data(res,'2026-08','agency',['服务甲'])
  self.assertIsNone(d[-2][0][18])
 def test_boundary_punch_and_ack(self):
  res=fixture();res['rows'][0]['已认定']={'认定人':'复核员','时间':'2026-09-01','理由':'核实'}
  raw={'by_key':{'测试甲':[{'标识':'0123','bnd':{0:[1200],32:[420]},'days':{}}]}}
  d,s=build_data(res,'2026-08','agency',['服务甲'],punch=raw)
  self.assertEqual(d[-1][0][6],'07:00');self.assertIn('复核员',d[-2][0][19])
  self.assertEqual(d[-2][1],'acked');self.assertIsNone(d[-1][0][18])

if __name__=='__main__':unittest.main()
