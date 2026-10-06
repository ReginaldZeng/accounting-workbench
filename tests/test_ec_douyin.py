import io
import unittest
import zipfile
from kernels import ec_douyin as m

SETTLE_HEAD = '动账时间,动帐流水号,动账方向,动账金额,动账账户,动账场景,计费类型,子订单号,订单号,订单实付应结,运费实付,实际平台补贴,订单退款,平台服务费,佣金,招商服务费,站外推广费,备注'
LEDGER_HEAD = '动账流水号,关联订单号,关联子订单号,动账时间,账户方向,动账金额(元),动账场景,账户余额(元),备注'


def settle_csv(*lines):
    return ('﻿' + SETTLE_HEAD + '\n' + '\n'.join(lines) + '\n').encode('utf-8')


def ledger_csv(*lines):
    return ('﻿' + LEDGER_HEAD + '\n' + '\n'.join(lines) + '\n').encode('utf-8')


def bill(no, date, amount, order, written=0, ws='A', ds='C'):
    return dict(no=no, date=date, amount=amount, written=written, open=round(amount - written, 2), ws=ws, ds=ds, order=order)


O1, O2, O3, O4, O5, O6 = ('69301627455895%05d' % i for i in range(1, 7))
SETTLE = settle_csv(
    f"2026-09-10 08:00:00,'S1,入账,48.50,聚合账户,货款结算入账,巨量千川,'{O1},'{O1},40,0,10,0,-1.50,0,0,0,订单结算",
    f"2026-09-10 08:00:00,'S2,入账,26.00,聚合账户,货款结算入账,精选联盟,'{O2},'{O2},30,0,0,0,-1.00,-2.00,-0.50,-0.50,订单结算",
    f"2026-09-11 09:00:00,'S3,入账,19.00,聚合账户,货款结算入账,小店自卖,'{O3},'{O3},20,0,0,0,-1.00,0,0,0,订单结算",
    f"2026-09-12 09:00:00,'S6,入账,9.00,聚合账户,货款结算入账,小店自卖,'{O6},'{O6},10,0,0,0,-1.00,0,0,0,订单结算",
    "2026-09-12 10:00:00,'X1,出账,5.00,聚合账户,退换货运费险,,,,0,0,0,0,0,0,0,0,保费扣除",
    f"2026-08-20 10:00:00,'S0,入账,9.50,聚合账户,货款结算入账,小店自卖,'{O5},'{O5},10,0,0,0,-0.50,0,0,0,订单结算")
LEDGER = ledger_csv(
    f"S0,`{O5},`{O5},2026-08-20 10:00:00,收入,9.50,货款结算入账,109.50,商家货款入账",
    f"S2,`{O2},`{O2},2026-09-10 08:00:00,收入,26.00,货款结算入账,184.00,商家货款入账",      # 同一秒两笔，文件里顺序是反的
    f"S1,`{O1},`{O1},2026-09-10 08:00:00,收入,48.50,货款结算入账,158.00,商家货款入账",
    f"S3,`{O3},`{O3},2026-09-11 09:00:00,收入,19.00,货款结算入账,203.00,商家货款入账",
    f"S6,`{O6},`{O6},2026-09-12 09:00:00,收入,9.00,货款结算入账,212.00,商家货款入账",
    "X1,,,2026-09-12 10:00:00,支出,-5.00,退换货运费险,207.00,保费扣除")
BILLS = [bill('AR1', '2026-09-05', 50, O1), bill('AR2', '2026-09-05', 30, O2), bill('AR3', '2026-09-06', 25, O3),
         bill('AR4', '2026-09-08', 60, O4), bill('AR4R', '2026-09-09', -60, O4),
         bill('AR5', '2026-08-18', 10, O5, written=10, ws='C'),
         bill('AR7', '2026-09-15', 70, '6930162745589500007'), bill('AR8', '2026-09-28', 80, '6930162745589500008'),
         bill('AR9', '2026-09-29', 5, ''), bill('ARX', '2026-10-02', 99, '6930162745589500009')]


class DouyinTests(unittest.TestCase):
    def rows(self):
        s = m.parse(SETTLE, '明细.csv')[0]; l = m.parse(LEDGER, '流水.csv')[0]
        self.assertEqual((s['kind'], l['kind']), ('dy_settle', 'dy_ledger'))
        return s['rows'], l['rows']

    def test_parse_recognises_by_header_inside_zip_and_skips_insurance(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z:
            z.writestr('a.csv', LEDGER); z.writestr('b.csv', '﻿保险单号,动账流水号,关联子订单号,摘要描述,动账时间,金额(元)\n1,X1,,退换货运费险,2026-09-12 10:00:00,5\n'.encode())
        parts = m.parse(buf.getvalue(), '包.zip')
        self.assertEqual([p['kind'] for p in parts], ['dy_ledger', 'dy_insure'])                       # 保费支出：列齐了就逐单接入
        self.assertEqual(parts[1]['rows'][0], {'id': '1', 'flow': 'X1', 'order': '', 't': '2026-09-12 10:00:00', 'amt': -5.0, 'memo': '退换货运费险', 'n': 0})
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z:                                                           # 列不齐的老样子、或者是空的：跳过，不连累同包的账户流水
            z.writestr('a.csv', LEDGER); z.writestr('b.csv', '\ufeff保险单号,动账流水号,金额(元)\n1,X1,5\n'.encode('utf-8'))
            z.writestr('c.csv', '\ufeff保险单号,动账流水号,关联子订单号,摘要描述,动账时间,金额(元)\n'.encode('utf-8'))
        skipped = m.parse(buf.getvalue(), 'x.zip')
        self.assertEqual([p['kind'] for p in skipped], ['dy_ledger', '', '']); self.assertIn('不单独接入', skipped[1]['skip'])
        with self.assertRaises(ValueError): m.parse('a,b\n1,2\n'.encode(), 'x.csv')
        s, _ = self.rows()
        self.assertEqual((s[1]['id'], s[1]['order'], s[1]['amt'], s[1]['fees']), ('S2', O2, 26.0, {'平台服务费': 1.0, '佣金': 2.0, '招商服务费': 0.5, '站外推广费': 0.5}))
        self.assertEqual(s[4]['amt'], -5.0)

    def test_merge_dedupes_and_months_split(self):
        s, _ = self.rows()
        self.assertEqual(len(m.merge(s, s)), 6)
        self.assertEqual({k: len(v) for k, v in m.by_month(s).items()}, {'2026-09': 5, '2026-08': 1})

    def test_chain_orders_same_second_rows_by_balance(self):
        _, l = self.rows()
        ordered, breaks = m.chain(m.merge([], l))
        self.assertEqual(([r['id'] for r in ordered], breaks), (['S0', 'S1', 'S2', 'S3', 'S6', 'X1'], 0))
        self.assertEqual(m.chain([dict(l[0]), dict(l[3])])[1], 1)

    def test_reconcile_september(self):
        s, l = self.rows()
        r = m.reconcile('2026-09', m.merge([], s), m.merge([], l), BILLS)
        cat = {b['no']: b['cat'] for b in r['bills']}
        self.assertEqual(cat, {'AR1': 'ok', 'AR2': 'ok', 'AR3': 'mismatch', 'AR4': 'pair', 'AR4R': 'pair', 'AR7': 'overdue', 'AR8': 'transit', 'AR9': 'no_order'})
        self.assertEqual(r['categories']['ok'], {'count': 2, 'orders': 2, 'amount': 80.0})
        self.assertEqual(next(b for b in r['bills'] if b['no'] == 'AR3')['diff'], 5.0)
        d = r['draft']
        self.assertEqual((d['net'], d['deduction_total'], d['total'], d['settled_gross'], d['from_ledger']), (97.5, 12.5, 110.0, 110.0, True))
        self.assertEqual({x['name']: x['amount'] for x in d['deductions']}, {'退换货运费险': 5.0, '平台服务费': 4.5, '佣金': 2.0, '招商服务费': 0.5, '站外推广费': 0.5})
        b = r['buckets']
        self.assertEqual((b['ok'], b['mismatch'], b['missing']), ({'orders': 2, 'flow': 80.0, 'ar': 80.0}, {'orders': 1, 'flow': 20.0, 'ar': 25.0}, {'orders': 1, 'flow': 10.0, 'ar': 0.0}))
        self.assertEqual(round(sum(x['flow'] for x in b.values()), 2), d['total'])            # 四个桶把收款明细合计分完
        self.assertEqual(r['missing'], [{'order': O6, 'flow': 10.0, 'bucket': 'missing'}])
        self.assertEqual({k: r['balance'][k] for k in ('open', 'close', 'net', 'breaks', 'rows')}, {'open': 109.5, 'close': 207.0, 'net': 97.5, 'breaks': 0, 'rows': 5})
        self.assertIn('2026年09月结算单扣款项 佣金2.00元', m.memo('2026-09', '佣金', 2.0))

    def test_august_already_written_off_and_no_ledger_fallback(self):
        s, _ = self.rows()
        r = m.reconcile('2026-08', m.merge([], s), [], BILLS)
        self.assertEqual(r['buckets']['written'], {'orders': 1, 'flow': 10.0, 'ar': 0.0})
        self.assertEqual((r['draft']['net'], r['draft']['from_ledger'], r['balance'], r['bills']), (9.5, False, None, []))

    def test_page_select_export(self):
        s, l = self.rows()
        r = m.reconcile('2026-09', m.merge([], s), m.merge([], l), BILLS)
        self.assertEqual([b['no'] for b in m.select(r, 'pair')], ['AR4R', 'AR4']); self.assertEqual(len(m.select(r, q=O1[-6:])), 1)
        self.assertEqual(m.page(m.select(r, 'ok'))['amount'], 80.0)
        ok = next(b for b in r['bills'] if b['no'] == 'AR2')
        self.assertEqual((ok['cash'], ok['fee'], ok['flow'], ok['settled_at']), (26.0, 4.0, 30.0, '2026-09-10 08:00:00'))
        d = m.order_detail('2026-09', m.merge([], s), m.merge([], l), BILLS, O2)
        self.assertEqual((len(d['flows']), d['flows'][0]['id'], d['flows'][0]['bal'], d['flows'][0]['gross'], d['flows'][0]['fees']['佣金']), (1, 'S2', 184.0, 30.0, 2.0))
        self.assertEqual(([b['no'] for b in d['bills']], d['flow_total'], d['cash_total'], d['fee_total'], d['open_total']), (['AR2'], 30.0, 26.0, 4.0, 30.0))
        self.assertEqual(m.order_detail('2026-09', m.merge([], s), [], BILLS, O4)['flows'], [])
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(m.export(r, '抖音店', {'account': '抖音177', 'open': 109.5, 'close': 109.5})))
        self.assertEqual(wb.sheetnames, ['本月账户进出汇总', '交人工处理', '可核销(下推源单)', '金额对不上', '红蓝互冲', '该结没结', '在途', '无订单号', '流水有·应收对不上号'])
        first = wb['本月账户进出汇总']
        self.assertEqual([[c.value for c in first[i]][:2] for i in (4, 5, 6)], [['货款结算到账', 102.5], ['　退换货运费险', -5], ['本月账户净变动', 97.5]])
        self.assertEqual((r['draft']['settle_cash'], r['draft']['fee_total'], r['draft']['other_total']), (102.5, 7.5, 5.0))
        self.assertEqual(wb['可核销(下推源单)'].max_row, 4); self.assertEqual(wb['可核销(下推源单)']['I3'].value + wb['可核销(下推源单)']['J3'].value, wb['可核销(下推源单)']['K3'].value)

    def orders_xlsx(self, rows):
        from openpyxl import Workbook
        wb = Workbook(); ws = wb.active
        ws.append(['订单编号', '店铺', '子单原始单号', '订单状态', '订单退款状态', '交易时间', '付款时间', '发货时间', '收件人', '买家实付', '应收金额', '货品名称', '实发数量', '分摊后总价'])
        for r in rows: ws.append(r)
        buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    def test_wangdiantong_orders_group_merged_shipments(self):
        A, B = '6930162745589599991', '6930162745589599992'
        blob = self.orders_xlsx([
            ['JY1', '抖音店', A, '已完成', '', '2026-09-05 10:00:00', '2026-09-05 10:00:05', '2026-09-06 09:00:00', '张三', 19.8, 35.7, '豆腐 & 辣条', 1, 19.8],
            ['JY1', '抖音店', B + '-1', '已完成', '', '2026-09-05 10:01:00', '', '2026-09-06 09:00:00', '张三', 19.8, 35.7, '辣丝丝', 1, 15.9],
            ['JY2', '抖音店', '6930162745589599993', '已取消', '全部退款', '2026-09-07 10:00:00', '', '', '李四', 0, 0, '辣丝丝', 0, 0],
            ['JY3', '天猫店', '5127801625403010443A', '已完成', '', '2026-09-07 10:00:00', '', '', '王五', 9, 9, '辣丝丝', 1, 9],
            ['JY4', '抖音店', 'AD202610020054', '已完成', '', '2026-09-07 10:00:00', '', '', '王五', 9, 9, '补发', 1, 9]])
        part = m.parse(blob, '订单明细.xlsx', '抖音店')[0]
        self.assertEqual((part['kind'], sorted(r['id'] for r in part['rows'])), ('dy_orders', ['JY1|' + A, 'JY1|' + B, 'JY2|6930162745589599993']))
        first = next(r for r in part['rows'] if r['order'] == A)
        self.assertEqual((first['recv'], first['amount'], first['ship'], first['goods'], first['t']), (35.7, 19.8, '2026-09-06 09:00:00', ['豆腐 & 辣条'], '2026-09-05 10:00:00'))
        self.assertNotIn('张三', str(part['rows']))                                   # 收件人不留
        with self.assertRaises(ValueError): m.parse(blob, '订单明细.xlsx', '别的店')
        # 金蝶照合并后的旺店通订单开一张应收，只记了 A；抖音 A、B 各结各的
        settle = m.parse(settle_csv(
            f"2026-09-10 08:00:00,'M1,入账,19.40,聚合账户,货款结算入账,小店自卖,'{A},'{A},19.8,0,0,0,-0.40,0,0,0,订单结算",
            f"2026-09-11 08:00:00,'M2,入账,15.58,聚合账户,货款结算入账,小店自卖,'{B},'{B},15.9,0,0,0,-0.32,0,0,0,订单结算"), 'x.csv')[0]['rows']
        bills = [bill('ARM', '2026-09-06', 35.7, A)]
        alone = m.reconcile('2026-09', settle, [], bills)
        self.assertEqual((alone['bills'][0]['cat'], alone['buckets']['missing']['orders']), ('mismatch', 1))
        r = m.reconcile('2026-09', settle, [], bills, part['rows'])
        self.assertEqual((r['bills'][0]['cat'], r['bills'][0]['flow'], r['bills'][0]['merged'], r['buckets']['ok'], r['coverage']['merged_groups']), ('ok', 35.7, 2, {'orders': 2, 'flow': 35.7, 'ar': 35.7}, 1))
        half = m.reconcile('2026-09', settle[:1], [], bills, part['rows'])
        self.assertEqual((half['bills'][0]['cat'], half['bills'][0]['reason']), ('transit', '合单发货的 2 个订单里还有没结算的'))
        d = m.order_detail('2026-09', settle, [], bills, B, part['rows'])
        self.assertEqual((d['members'], [f['order'] for f in d['flows']], [b['no'] for b in d['bills']], d['flow_total'], d['open_total'], d['wdt'][0]['jy'], d['wdt'][0]['orders']), ([A, B], [A, B], ['ARM'], 35.7, 35.7, 'JY1', [A, B]))

    def test_mismatch_reasons_and_settle_then_refund(self):
        X, Y, Z = '6930162745589588881', '6930162745589588882', '6930162745589588883'
        settle = m.parse(settle_csv(
            f"2026-09-10 08:00:00,'R1,入账,48.90,聚合账户,货款结算入账,巨量千川,'{X},'{X},49.9,0,0,0,-1.00,0,0,0,订单结算",
            f"2026-09-12 08:00:00,'R2,出账,49.90,聚合账户,退款-结算后退款-退用户,,'{X},'{X},0,0,0,0,0,0,0,0,退款",
            f"2026-09-12 08:00:01,'R3,入账,1.00,聚合账户,退款-订单退款触发-退分账,,'{X},'{X},0,0,0,0,0,0,0,0,服务费返还",
            f"2026-09-10 09:00:00,'R4,入账,48.90,聚合账户,货款结算入账,巨量千川,'{Y},'{Y},49.9,0,0,0,-1.00,0,0,0,订单结算",
            f"2026-09-10 10:00:00,'R5,入账,44.00,聚合账户,货款结算入账,巨量千川,'{Z},'{Z},49.9,0,0,-5,-0.90,0,0,0,订单结算"), 'x.csv')[0]['rows']
        bills = [bill('BX', '2026-09-05', 49.9, X), bill('BXR', '2026-09-13', -49.9, X), bill('BY', '2026-09-05', 49.9, Y), bill('BYR', '2026-09-08', -49.9, Y), bill('BZ', '2026-09-05', 49.9, Z)]
        r = m.reconcile('2026-09', settle, [], bills)
        got = {b['no']: (b['cat'], b['reason']) for b in r['bills']}
        self.assertEqual(got['BX'][0], 'pair'); self.assertIn('先结算、后来又把钱退回去', got['BX'][1])
        self.assertEqual(got['BY'], ('mismatch', '金蝶红字蓝字已全冲掉，抖音没退款、全额结了 49.90'))
        self.assertEqual(got['BZ'], ('mismatch', '抖音结算时退了 5.00，金蝶没开红字'))
        self.assertEqual((r['buckets']['refunded'], r['buckets']['mismatch']['orders']), ({'orders': 1, 'flow': 49.9, 'ar': 0.0}, 2))
        self.assertEqual(round(sum(x['flow'] for x in r['buckets'].values()), 2), r['draft']['total'])
        from openpyxl import load_workbook
        todo = load_workbook(io.BytesIO(m.export(r, '抖音店')))['交人工处理']
        self.assertEqual((todo.max_row, todo['A3'].value, todo['O3'].value), (5, '钱已到账·金额对不上', '抖音结算时退了 5.00，金蝶没开红字'))

    def test_refund_before_settle_takes_subsidy_back(self):
        # 49.9 的一袋 26 小包：买家实付 28.9 + 平台补贴 21；吃了一包退 25 包 → 退买家 27.78，补贴同比例收回、只剩 0.81
        P, Q = '6930162745589588891', '6930162745589588892'
        settle = m.parse(settle_csv(
            f"2026-09-28 07:54:38,'T1,入账,1.84,聚合账户,货款结算入账,巨量千川,'{P},'{P},28.9,0,0.81,-27.78,-0.09,0,0,0,订单结算",
            f"2026-09-28 08:00:00,'T2,入账,44.00,聚合账户,货款结算入账,巨量千川,'{Q},'{Q},49.9,0,0,-5,-0.90,0,0,0,订单结算"), 'x.csv')[0]['rows']
        bills = [bill('BP', '2026-09-22', 49.9, P), bill('BPR', '2026-09-27', -49.9, P), bill('BQ', '2026-09-22', 49.9, Q)]
        d = m.order_detail('2026-09', settle, [], bills, P)
        self.assertEqual(d['subsidy_back'], {'before': 21.0, 'back': 20.19, 'blue': 49.9, 'red': -49.9, 'keep': 1.93, 'red_should': -47.97})
        self.assertEqual((d['flow_total'], d['open_total']), (1.93, 0.0))
        self.assertIsNone(m.order_detail('2026-09', settle, [], bills, Q)['subsidy_back'])      # 没有补贴被收回：照旧一条算式

    def test_platform_orders_parse_index_and_state(self):
        head = '主订单编号,子订单编号,选购商品,商品规格,商品数量,商家编码,订单应付金额,运费,收件人,收件人手机号,详细地址,订单提交时间,商家备注,订单完成时间,订单状态,售后状态,取消原因,发货时间,平台实际承担优惠金额,商家实际承担优惠金额,达人实际承担优惠金额'
        P, Q, Z = '6917926768823643751', '6917926768823643752', '6917926768823643753'
        blob = ('\ufeff' + head + '\n' + '\n'.join([
            f'"\t{P}","\t{P}","\t辣丝丝","\t1袋",1,"\t300400008",28.90,0.00,##密文,$$密文,##密文,2026-09-22 00:54:10,吃了一包,2026-09-24 13:18:29,已完成,退款成功,,2026-09-22 08:53:45,21.00,25.00,0.00',
            f'"\t{Q}","\t{Q}","\t辣丝丝","\t1袋",1,"\t300400008",20.00,0.00,##密文,$$密文,##密文,2026-09-25 10:00:00,-,2026-10-02 09:00:00,已完成,-,,2026-09-25 18:00:00,5.00,0.00,8.00',
            f'"\t{Q}","\t{Q}9","\t赠品","\t默认",1,"\t300400009",0.00,0.00,##密文,$$密文,##密文,2026-09-25 10:00:00,-,-,已关闭,-,主品订单售后完成,-,0.00,0.00,0.00',
            f'"\t{Z}","\t{Z}","\t辣丝丝","\t1袋",1,"\t300400008",15.90,0.00,##密文,$$密文,##密文,2026-09-26 10:00:00,-,-,已关闭,退款成功,,2026-09-26 18:00:00,0.00,0.00,0.00']) + '\n').encode('utf-8')
        part = m.parse(blob, '订单导出.csv')[0]
        self.assertEqual((part['kind'], len(part['rows'])), ('dy_platform', 4))
        first = part['rows'][0]
        self.assertEqual((first['id'], first['order'], first['pay'], first['plat'], first['shop'], first['status'], first['after'], first['done'], first['code']),
                         (P, P, 28.9, 21.0, 25.0, '已完成', '退款成功', '2026-09-24 13:18:29', '300400008'))
        self.assertFalse(any('密文' in str(v) for r in part['rows'] for v in r.values()))             # 收件人、电话、地址不留
        self.assertEqual(set(first), {'id', 'order', 't', 'goods', 'spec', 'code', 'qty', 'pay', 'plat', 'kol', 'shop', 'freight', 'status', 'after', 'ship', 'done', 'cancel', 'memo', 'n'})
        index = m.platform_index(part['rows'])
        # Q 的赠品子订单没发货就关闭：不算进订单状态，金额记在 unshipped（这里是 0 元赠品）
        self.assertEqual((index[P]['gross'], index[Q]['gross'], index[Q]['pay'], index[Q]['status'], index[Q]['done'], index[Q]['unshipped']), (49.9, 33.0, 20.0, ['已完成'], '2026-10-02 09:00:00', 0.0))
        self.assertEqual(m.platform_state(index[Q], '2026-09-30'), '买家 10-02 确认收货（月底时还没确认），等平台结算')
        self.assertEqual(m.platform_state(index[Z], '2026-09-30'), '平台订单已关闭（退款成功），这笔钱不会结算了')
        self.assertEqual(m.platform_state(index[P], '2026-09-30'), '平台上有退款成功的售后，会少结或不结')
        self.assertEqual(m.platform_state(None, '2026-09-30'), '')
        settle = m.parse(settle_csv(f"2026-09-28 07:54:38,'T1,入账,1.84,聚合账户,货款结算入账,巨量千川,'{P},'{P},28.9,0,0.81,-27.78,-0.09,0,0,0,订单结算"), 'x.csv')[0]['rows']
        bills = [bill('BP', '2026-09-22', 49.9, P), bill('BPR', '2026-09-27', -49.9, P), bill('BQ', '2026-09-25', 33.0, Q)]
        d = m.order_detail('2026-09', settle, [], bills, P, None, part['rows'])
        self.assertEqual((d['platform_gross'], d['platform_pay'], d['blue_total'], len(d['platform']), d['platform_state'], d['platform_at']), (49.9, 28.9, 49.9, 1, '', '2026-09-22 00:54:10'))
        q = m.order_detail('2026-09', settle, [], bills, Q, None, part['rows'])
        self.assertEqual((q['platform_gross'], len(q['platform']), q['platform_state']), (33.0, 2, '买家 10-02 确认收货（月底时还没确认），等平台结算'))
        self.assertEqual(m.order_detail('2026-09', settle, [], bills, P)['platform_gross'], None)

    def test_memo_drops_anything_that_looks_like_contact_details(self):
        # 号码、地址都是编的
        self.assertEqual(m._memo('顾客要改地址：广东省深圳市南山区科技路1号 张三 13800138000 已登记'), m.MEMO_HIDDEN)
        self.assertEqual(m._memo('已电话沟通，补偿 2 元'), m.MEMO_HIDDEN)                              # 有“电话”字样就不留，宁可多挡
        self.assertEqual(m._memo('好' * 75 + '13800138000'), m.MEMO_HIDDEN)                           # 先判断再截断：号码跨在 80 字的截断线上也认得出
        self.assertEqual(m._memo('补发(JY202609220185) 运单 SF1234567890123'), '补发(JY202609220185) 运单 SF1234567890123')
        self.assertEqual(m._memo('吃了一包反馈不喜欢要退货退款  计算金额为27.78元'), '吃了一包反馈不喜欢要退货退款  计算金额为27.78元')
        self.assertEqual((m._memo('-'), m._memo(None), len(m._memo('好' * 200))), ('', '', 80))
        for hidden in ('改寄 张三 vx13800138000', 'tel13800138000', '138.0013.8000 张三', '１３８００１３８０００', '座机 02012345678', '新收货：张三 A3-1201',
                       '转寄 某某大厦8楼 李四', '加微信聊', '身份证 440301199001011234', '发到公司前台'):
            self.assertEqual(m._memo(hidden), m.MEMO_HIDDEN, hidden)
        for kept in ('补发 圆通 YT8812345678901', '顺丰 SF1234567890123 已拦截', '退款 21.82 元已登记【客服如皓2 09-23 18:38】', '缺货，15号补发', '14位单号 12345678901234'):
            self.assertEqual(m._memo(kept), kept, kept)
        import time
        started = time.time(); m._memo('广东' * 2500 + '市'); self.assertLess(time.time() - started, 0.5)          # 很长的备注也不会卡住

    def test_platform_state_covers_every_after_sale_status(self):
        def order(status, after='', done='', ship='2026-09-20 10:00:00', freight=0.0, x=''):
            rows = [{'id': '1', 'order': 'O', 't': '2026-09-20 09:00:00', 'pay': 10.0, 'plat': 0.0, 'kol': 0.0, 'freight': freight, 'status': status, 'after': after, 'ship': ship, 'done': done}]
            if x: rows = [dict(r, x=x) for r in rows]
            return m.platform_index(rows)['O']
        say = lambda *a, **k: m.platform_state(order(*a, **k), '2026-09-30')
        for after in ('售后关闭', '售后已拒绝', '换货成功', '补寄成功'):                                    # 售后办完了、不退钱：照正常流程说
            self.assertEqual(say('已完成', after, '2026-09-25 08:00:00'), '买家 09-25 确认收货，等平台结算')
            self.assertEqual(say('已发货', after), '买家还没确认收货')
        for after in ('待收退货', '待退货', '售后待处理'):
            self.assertEqual(say('已发货', after), '平台上售后还在处理（%s）' % after)
        self.assertEqual(say('已完成', '退款成功', '2026-09-25 08:00:00'), '平台上有退款成功的售后，会少结或不结')
        self.assertEqual(say('已发货', '以后新出的状态'), '平台上有售后记录（以后新出的状态）')                     # 不认识的不猜
        self.assertEqual(say('已关闭', '退款成功'), '平台订单已关闭（退款成功），这笔钱不会结算了')
        self.assertEqual(say('已关闭', '退款成功', freight=26.0), '平台订单已关闭（退款成功），货款不会结算，运费 26.00 元可能照结')
        self.assertEqual(say('已关闭', '退款成功', ship='', freight=26.0), '平台订单已关闭（退款成功），这笔钱不会结算了')   # 没发货：运费也不会结
        self.assertEqual(say('已发货', x='2026-10-05 23:23'), '买家还没确认收货（按 10-05 的订单导出）')
        self.assertTrue(m.platform_closed(order('已关闭'))); self.assertFalse(m.platform_closed(order('已完成'))); self.assertFalse(m.platform_closed(None))

    def test_unshipped_closed_suborder_is_not_this_bills_business(self):
        # 一个订单两个子订单：一个 12.90 没发货就退款关闭（金蝶不开应收），一个 37.20 正常发货、已确认收货
        O = '6917926768823643761'
        rows = [{'id': O, 'order': O, 't': '2026-09-20 09:00:00', 'pay': 30.0, 'plat': 7.2, 'kol': 0.0, 'freight': 0.0, 'status': '已完成', 'after': '', 'ship': '2026-09-20 18:00:00', 'done': '2026-09-24 08:00:00'},
                {'id': O + '9', 'order': O, 't': '2026-09-20 09:00:00', 'pay': 10.0, 'plat': 2.9, 'kol': 0.0, 'freight': 0.0, 'status': '已关闭', 'after': '退款成功', 'ship': '', 'done': ''}]
        o = m.platform_index(rows)[O]
        self.assertEqual((o['gross'], o['unshipped'], o['status'], o['after']), (50.1, 12.9, ['已完成'], []))
        self.assertEqual(m.platform_state(o, '2026-09-30'), '买家 09-24 确认收货，等平台结算')              # 不再说“会少结或不结”
        d = m.order_detail('2026-09', [], [], [bill('B1', '2026-09-20', 37.2, O)], O, None, rows)
        self.assertEqual((d['platform_gross'], d['platform_unshipped'], d['blue_total'], d['platform_closed']), (50.1, 12.9, 37.2, False))
        # 带 0 元赠品的订单：货款那个子订单已经退款关闭，赠品还是“已完成”——照样认成整单已关闭
        G = '6917926768823643781'
        gift = [{'id': G, 'order': G, 't': '2026-09-20 09:00:00', 'pay': 49.9, 'plat': 0.0, 'kol': 0.0, 'freight': 0.0, 'status': '已关闭', 'after': '退款成功', 'ship': '2026-09-20 18:00:00', 'done': ''},
                {'id': G + '9', 'order': G, 't': '2026-09-20 09:00:00', 'pay': 0.0, 'plat': 0.0, 'kol': 0.0, 'freight': 0.0, 'status': '已完成', 'after': '', 'ship': '2026-09-20 18:00:00', 'done': '2026-09-24 08:00:00'}]
        g = m.platform_index(gift)[G]
        self.assertEqual((g['status'], m.platform_closed(g), m.platform_state(g, '2026-09-30')), (['已关闭'], True, '平台订单已关闭（退款成功），这笔钱不会结算了'))
        # 月底发货、下月初退货：红字记在下个月，期末看不到——单独带出来，卡片上好说明
        later = m.order_detail('2026-09', [], [], [bill('GB', '2026-09-30', 49.9, G), bill('GR', '2026-10-05', -49.9, G)], G, None, gift)
        self.assertEqual((later['open_total'], later['platform_closed'], later['later_red'], later['later_red_at'], [b['no'] for b in later['bills']]), (49.9, True, -49.9, '2026-10-05', ['GB']))
        self.assertEqual((d['later_red'], d['later_red_at']), (0, ''))
        # 整单都没发货就关闭：看全部子订单，才说钱不会来
        gone = m.platform_index([dict(rows[1], id='X', order='X')])['X']
        self.assertEqual((gone['status'], m.platform_closed(gone)), (['已关闭'], True))

    def test_platform_state_only_for_unsettled_orders_with_open_blue(self):
        A, B = '6930162745589588801', '6930162745589588802'
        orders = m.parse(self.orders_xlsx([
            ['JY1', '抖音店', A, '已完成', '', '2026-09-03 10:00:00', '2026-09-03 10:00:05', '2026-09-03 18:00:00', '张三', 15.9, 35.7, '豆腐', 1, 15.9],
            ['JY1', '抖音店', B, '已完成', '', '2026-09-03 10:00:00', '2026-09-03 10:00:05', '2026-09-03 18:00:00', '张三', 19.8, 35.7, '年糕', 1, 19.8]]), '订单明细.xlsx', '抖音店')[0]['rows']
        plat = [{'id': A, 'order': A, 't': '2026-09-03 10:00:00', 'pay': 15.9, 'plat': 0.0, 'kol': 0.0, 'freight': 0.0, 'status': '已完成', 'after': '', 'ship': '2026-09-03 18:00:00', 'done': '2026-09-08 08:00:00'},
                {'id': B, 'order': B, 't': '2026-09-03 10:00:00', 'pay': 19.8, 'plat': 0.0, 'kol': 0.0, 'freight': 0.0, 'status': '已关闭', 'after': '退款成功', 'ship': '2026-09-03 18:00:00', 'done': ''}]
        settle = m.parse(settle_csv(f"2026-09-12 08:00:00,'M1,入账,15.00,聚合账户,货款结算入账,小店自卖,'{A},'{A},15.9,0,0,0,-0.90,0,0,0,订单结算"), 'x.csv')[0]['rows']
        bills = [bill('ARM', '2026-09-03', 35.7, A)]                                                     # 金蝶应收只记了已经结算的那个订单号
        r = m.reconcile('2026-09', settle, [], bills, orders)
        self.assertEqual((r['bills'][0]['cat'], r['bills'][0]['pending']), ('transit', [B]))               # 合单：带出还没结算的那个订单
        index = m.platform_index(plat)
        self.assertEqual(m.platform_states(index, r['bills'][0]['pending'], r['end']), '平台订单已关闭（退款成功），这笔钱不会结算了')
        self.assertEqual(m.platform_states(index, r['bills'][0]['pending'], r['end'], 2), '尾号 8802：平台订单已关闭（退款成功），这笔钱不会结算了')   # 知道是合单里的一部分：带尾号
        d = m.order_detail('2026-09', settle, [], bills, A, orders, plat)
        self.assertEqual((d['platform_state'], d['platform_closed']), ('尾号 8802：平台订单已关闭（退款成功），这笔钱不会结算了', True))   # 说的是没结算的 B，不是已结算的 A
        self.assertEqual(m.platform_states(index, [A, B], '2026-09-30'), '尾号 8801：买家 09-08 确认收货，等平台结算；尾号 8802：平台订单已关闭（退款成功），这笔钱不会结算了')
        # 只剩红字挂着（蓝字早已核销）：不说“等平台结算”
        red_only = [bill('BL', '2026-07-01', 15.9, A, written=15.9, ws='C'), bill('RD', '2026-07-31', -15.9, A)]
        self.assertEqual(m.order_detail('2026-09', [], [], red_only, A, None, plat)['platform_state'], '')
        self.assertEqual(r['bills'][0]['merged'], 2)
        single = m.reconcile('2026-09', settle, [], [bill('S1', '2026-09-03', 15.9, A)])
        self.assertEqual(single['bills'][0]['pending'], [])                                              # 不是合单：不带
        # 合单里钱已到账、和未核销应收一分不差（另一个订单没发货就关闭，金蝶本来就没给它开应收）：可核销，不再说“这笔钱不会结算了”
        even = m.order_detail('2026-09', settle, [], [bill('ARE', '2026-09-03', 15.9, A)], A, orders, plat)
        self.assertEqual((even['flow_total'], even['open_total'], even['platform_state'], even['platform_closed']), (15.9, 15.9, '', False))
        # 0 元的赠品单不会有结算：合单里不说它；几句话一样时尾号并在一起
        gift = dict(plat[0], id='G', order='G', pay=0.0, plat=0.0)
        idx = m.platform_index(plat + [gift, dict(plat[0], id='H', order='H' * 4 + '9901')])
        self.assertEqual(m.platform_states(idx, ['G', B], '2026-09-30', 3), '尾号 8802：平台订单已关闭（退款成功），这笔钱不会结算了')
        self.assertEqual(m.platform_states(idx, [A, 'HHHH9901'], '2026-09-30', 3), '尾号 8801、9901：买家 09-08 确认收货，等平台结算')
        self.assertEqual(m.platform_states(idx, ['G'], '2026-09-30'), '买家 09-08 确认收货，等平台结算')    # 单独一个 0 元订单（不是合单）照常说

    def test_blue_must_fit_platform_amount_before_saying_blue_is_right(self):
        # 退款接近全额时，「按比例」那条判断几乎什么蓝字都放行；有平台订单明细时先要金额对得上
        P = '6917926768823643771'
        settle = m.parse(settle_csv(f"2026-09-28 07:54:38,'K1,入账,1.60,聚合账户,货款结算入账,巨量千川,'{P},'{P},30,0,1.0,-28.5,-0.90,0,0,0,订单结算"), 'x.csv')[0]['rows']
        plat = [{'id': P, 'order': P, 't': '2026-09-22 00:54:10', 'pay': 30.0, 'plat': 20.0, 'kol': 0.0, 'freight': 0.0, 'status': '已完成', 'after': '退款成功', 'ship': '2026-09-22 08:00:00', 'done': '2026-09-24 08:00:00'}]
        wrong, right = [bill('W', '2026-09-22', 50.4, P)], [bill('G', '2026-09-22', 50.0, P)]
        self.assertIsNotNone(m.order_detail('2026-09', settle, [], wrong, P)['subsidy_back'])               # 没有平台明细：照旧（V2.827 的行为）
        self.assertIsNone(m.order_detail('2026-09', settle, [], wrong, P, None, plat)['subsidy_back'])      # 有平台明细：蓝字 50.40 ≠ 平台 50.00，不再说“蓝字没错”
        self.assertEqual(m.order_detail('2026-09', settle, [], right, P, None, plat)['subsidy_back']['before'], 20.0)

    def test_files_resaved_by_excel_are_refused_whole(self):
        ok = f"2026-09-10 08:00:00,'R1,入账,48.90,聚合账户,货款结算入账,巨量千川,'{O1},'{O1},49.9,0,0,0,-1.00,0,0,0,订单结算"
        slashed = f"2026/9/10 8:00,'R2,入账,48.90,聚合账户,货款结算入账,巨量千川,'{O2},'{O2},49.9,0,0,0,-1.00,0,0,0,订单结算"
        with self.assertRaisesRegex(ValueError, '有 1 行的日期或单号被改过格式'): m.parse(settle_csv(ok, slashed), 'x.csv')
        head = '主订单编号,子订单编号,订单应付金额,订单提交时间,订单状态,售后状态,平台实际承担优惠金额'
        sci = ('\ufeff' + head + '\n6.91793E+18,6.91793E+18,10,2026-09-10 08:00:00,已完成,-,0\n').encode('utf-8')
        with self.assertRaisesRegex(ValueError, '被改过格式'): m.parse(sci, '订单导出.csv')
        with self.assertRaisesRegex(ValueError, '老版 Excel'): m.parse(b'\xd0\xcf\x11\xe0' + b'\x00' * 64, '旧.xls')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z: z.writestr('[Content_Types].xml', '<Types/>'); z.writestr('word/document.xml', '<w/>')
        with self.assertRaisesRegex(ValueError, '没有工作表'): m.parse(buf.getvalue(), '误传的.docx')

    def test_older_export_does_not_overwrite_newer_status(self):
        new = {'id': 'S1', 'order': 'O1', 't': '2026-09-20 09:00:00', 'status': '已完成', 'done': '2026-09-25 08:00:00', 'x': '2026-10-08 10:00'}
        old = dict(new, status='已发货', done='', x='2026-10-05 23:23')
        skipped = []
        self.assertEqual(m.merge([new], [old], 'dy_platform', skipped), [new]); self.assertEqual(skipped, ['S1'])
        self.assertEqual(m.merge([old], [new], 'dy_platform'), [new])                                    # 正常方向：新的盖旧的
        nox = lambda r: {k: v for k, v in r.items() if k != 'x'}
        self.assertTrue(m.older('dy_platform', nox(new), nox(old)))                                      # 不知道哪天导出的：靠状态倒退认
        self.assertFalse(m.older('dy_platform', nox(old), nox(new)))
        self.assertFalse(m.older('dy_platform', dict(new, status='已完成'), dict(new, status='已关闭', x='2026-10-09 10:00')))
        self.assertTrue(m.older('dy_orders', {'status': '已取消'}, {'status': '已发货'}))                     # 旺店通：取消了的订单不会被旧导出“复活”进合单组
        self.assertFalse(m.older('dy_orders', {'status': '已发货'}, {'status': '已取消'}))
        self.assertEqual(m.merge([{'id': 'a', 't': '1'}], [{'id': 'a', 't': '1', 'amt': 2}]), [{'id': 'a', 't': '1', 'amt': 2}])   # 流水类：照旧后传的盖先传的
        self.assertEqual((m.export_time('1791213782_3edeb47b86bdb5b9.csv'), m.export_time('订单导出.csv'), m.export_time('dir/1791213782_x.csv')), ('2026-10-05 23:23', '', '2026-10-05 23:23'))

    def test_wdt_buyer_paid_adds_up_item_lines(self):
        A = '6930162745589588811'
        part = m.parse(self.orders_xlsx([
            ['JY9', '抖音店', A, '已完成', '', '2026-09-03 10:00:00', '2026-09-03 10:00:05', '2026-09-03 18:00:00', '张三', 9.59, 79.9, '套餐物料一', 1, 30.0],
            ['JY9', '抖音店', A, '已完成', '', '2026-09-03 10:00:00', '2026-09-03 10:00:05', '2026-09-03 18:00:00', '张三', 38.35, 79.9, '套餐物料二', 1, 49.9]]), '订单明细.xlsx', '抖音店')[0]
        self.assertEqual((part['rows'][0]['paid'], part['rows'][0]['recv'], part['rows'][0]['amount']), (47.94, 79.9, 79.9))   # 买家实付逐行加；应收金额是整单数，不加
        B = '6930162745589588812'
        line = lambda: ['JY8', '抖音店', B, '已完成', '', '2026-09-03 10:00:00', '2026-09-03 10:00:05', '2026-09-03 18:00:00', '张三', 9.588, 79.9, '套餐物料', 1, 15.98]
        five = m.parse(self.orders_xlsx([line() for _ in range(5)]), '订单明细.xlsx', '抖音店')[0]['rows'][0]
        self.assertEqual((five['paid'], five['amount']), (47.94, 79.9))                                   # 9.588×5：每加一行取一次整会得 47.95

    def test_gbk_file_reads_and_zip_with_too_many_files_is_refused(self):
        line = f"2026-09-10 08:00:00,'R1,入账,48.90,聚合账户,货款结算入账,巨量千川,'{O1},'{O1},49.9,0,0,0,-1.00,0,0,0,订单结算"
        gbk = (SETTLE_HEAD + '\n' + line + '\n').encode('gb18030')
        self.assertEqual((m.parse(gbk, 'x.csv')[0]['kind'], m.parse(gbk, 'x.csv')[0]['rows'][0]['memo']), ('dy_settle', '订单结算'))
        with self.assertRaisesRegex(ValueError, '编码'): m.parse(b'\xff\xfe' + '动账时间'.encode('utf-16-le'), 'x.csv')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z:
            for i in range(m.MAX_ZIP_FILES + 1): z.writestr('%d.csv' % i, settle_csv(line))
        with self.assertRaisesRegex(ValueError, '超过 20 个'): m.parse(buf.getvalue(), '很多.zip')

    def test_flow_rows_list_what_makes_up_each_account_movement(self):
        O = '6930162745589588821'
        ledger = [{'id': 'L1', 't': '2026-09-12 10:00:00', 'amt': -0.4, 'scene': '退换货运费险', 'order': 'TRA2026', 'bal': 100.0, 'memo': '保费扣除（2笔保单）'},
                  {'id': 'L2', 't': '2026-09-13 10:00:00', 'amt': -49.9, 'scene': '退款-结算后退款-退用户', 'order': O, 'bal': 50.1, 'memo': '退款'},
                  {'id': 'L3', 't': '2026-09-14 10:00:00', 'amt': 30.0, 'scene': m.SETTLE, 'order': O, 'bal': 80.1, 'memo': '结算'},
                  {'id': 'L4', 't': '2026-08-14 10:00:00', 'amt': -9.0, 'scene': '退款-结算后退款-退用户', 'order': O, 'bal': 1.0, 'memo': '上月的'}]
        insure = [{'id': 'P2', 'flow': 'L1', 'order': O + '9', 't': '2026-09-12 10:00:00', 'amt': -0.2, 'memo': '退换货运费险'},
                  {'id': 'P1', 'flow': 'L1', 'order': O, 't': '2026-09-12 10:00:00', 'amt': -0.2, 'memo': '退换货运费险'}]
        kind, rows = m.flow_rows('2026-09', ledger, [], insure, m.INSURE_SCENE, {O}, {O + '9': O})
        self.assertEqual((kind, [(r['id'], r['flow'], r['order'], r['amt'], r['known']) for r in rows]), ('insure', [('P1', 'L1', O, -0.2, True), ('P2', 'L1', O, -0.2, True)]))   # 逐单；子订单号换成主订单号
        kind, rows = m.flow_rows('2026-09', ledger, [], [], m.INSURE_SCENE, {O})
        self.assertEqual((kind, [(r['id'], r['order'], r['known']) for r in rows]), ('ledger', [('L1', 'TRA2026', False)]))    # 没有保费明细：退回按账户流水一批一笔
        kind, rows = m.flow_rows('2026-09', ledger, [], insure, '退款-结算后退款-退用户', {O})
        self.assertEqual([(r['id'], r['order'], r['amt'], r['bal'], r['known']) for r in rows], [('L2', O, -49.9, 50.1, True)])   # 只要本期的，不含结算，不含上月
        settle = [{'id': 'S9', 't': '2026-09-13 10:00:00', 'amt': -49.9, 'scene': '退款-结算后退款-退用户', 'order': O, 'memo': '退款'}]
        self.assertEqual([r['id'] for r in m.flow_rows('2026-09', [], settle, [], '退款-结算后退款-退用户', {O})[1]], ['S9'])   # 没有账户流水：用订单维度动账明细
        # 保单跟着它所在的那笔流水走：平台改场景名之前叫「权益保险」，保单摘要却写运费险——归到权益保险那一行，运费险那一行不多算
        ledger.append({'id': 'L5', 't': '2026-09-01 09:00:00', 'amt': -0.01, 'scene': '权益保险', 'order': 'TRA2025', 'bal': 0.0, 'memo': '保费扣除（1笔保单）'})
        insure.append({'id': 'P0', 'flow': 'L5', 'order': O, 't': '2026-09-01 09:00:12', 'amt': -0.01, 'memo': '退换货运费险'})
        self.assertEqual([r['id'] for r in m.flow_rows('2026-09', ledger, [], insure, m.INSURE_SCENE, {O})[1]], ['P1', 'P2'])
        self.assertEqual((lambda k, rows: (k, [r['id'] for r in rows]))(*m.flow_rows('2026-09', ledger, [], insure, '权益保险', {O})), ('insure', ['P0']))
        # 有一笔流水没配上保单（保费明细不全）：逐单列出来会比汇总少，整项退回按流水列
        ledger.append({'id': 'L6', 't': '2026-09-20 09:00:00', 'amt': -0.3, 'scene': '退换货运费险', 'order': 'TRA2027', 'bal': 0.0, 'memo': '保费扣除（1笔保单）'})
        self.assertEqual((lambda k, rows: (k, [r['id'] for r in rows]))(*m.flow_rows('2026-09', ledger, [], insure, m.INSURE_SCENE, {O})), ('ledger', ['L1', 'L6']))

    def test_password_zip_is_refused_with_plain_words(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as z: z.writestr('a.csv', 'x')
        raw = bytearray(buf.getvalue())
        for sig, at in ((b'PK\x03\x04', 6), (b'PK\x01\x02', 8)):                       # 把“已加密”标志位置上
            i = raw.find(sig); raw[i + at] |= 1
        with self.assertRaisesRegex(ValueError, '带密码'): m.parse(bytes(raw), 'x.zip')

    def test_snapshot_note_span_and_stale(self):
        rows = [{'t': '2026-09-03 08:00:00', 'paid': 1}, {'t': '2026-09-28 07:54:38', 'paid': 2}]
        self.assertEqual(m.snapshot_note('dy_settle', rows), {'span': ['2026-09-03', '2026-09-28'], 'stale': False})
        self.assertEqual(m.snapshot_note('dy_settle', rows + [{'t': '2026-09-01 00:00:00'}]), {'span': ['2026-09-01', '2026-09-28'], 'stale': True})
        self.assertEqual(m.snapshot_note('dy_ledger', [{'t': '2026-09-01 00:00:00'}]), {'span': ['2026-09-01', '2026-09-01'], 'stale': False})
        self.assertEqual(m.snapshot_note('dy_orders', []), {'span': [], 'stale': False})

    def test_bill_lines_group_entries_by_bill(self):
        rows = [['AR1', '标准应收单', 'CK20260922170', 'JY202609220185', '开发', '2026-09-22T09:11:20.353', '黄春艳', '2026-09-24T15:06:26.51', '300400008', '辣丝丝', '袋', 1.0, 49.9, 13.0, 44.1593, 5.7407, 49.9, 'XQLCK20260922170', 'SAL_OUTSTOCK'],
                ['AR1', '标准应收单', 'CK20260922170', 'JY202609220185', '开发', '2026-09-22T09:11:20.353', '黄春艳', '2026-09-24T15:06:26.51', '300400009', '赠品', '袋', 1.0, 0, 13.0, 0, 0, 0, 'XQLCK20260922170', 'SAL_OUTSTOCK'],
                ['AR1R', '标准应收单', 'JY202609220185', '6917926768823643751', '开发', '2026-09-27T15:42:15.89', ' ', None, '300400008', '辣丝丝', '袋', -1.0, 49.9, 13.0, -44.1593, -5.7407, -49.9, 'RK2609270047', 'SAL_RETURNSTOCK']]
        got = m.bill_lines(rows)
        self.assertEqual((list(got), len(got['AR1']['lines']), got['AR1']['jy'], got['AR1']['created'], got['AR1']['approver'], got['AR1']['approved']), (['AR1', 'AR1R'], 2, 'JY202609220185', '2026-09-22 09:11', '黄春艳', '2026-09-24 15:06'))
        self.assertEqual(got['AR1']['lines'][0], {'code': '300400008', 'name': '辣丝丝', 'unit': '袋', 'qty': 1.0, 'price': 49.9, 'rate': 13.0, 'net': 44.16, 'tax': 5.74, 'amount': 49.9, 'src': 'XQLCK20260922170', 'src_type': '销售出库单'})
        self.assertEqual((got['AR1R']['jy'], got['AR1R']['approver'], got['AR1R']['lines'][0]['qty'], got['AR1R']['lines'][0]['src_type']), ('JY202609220185', '', -1.0, '销售退货单'))
        self.assertEqual(m.bill_lines(None), {})

    def test_xlsx_reader_handles_shared_strings(self):
        import zipfile as zf
        sheet = ('<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
                 '<row r="1"><c r="A1" t="s"><v>0</v></c><c r="B1" s="1" t="s"><v>1</v></c><c r="C1"/><c r="AB1" t="s"><v>2</v></c></row>'
                 '<row r="2"><c r="A2" t="s"><v>3</v></c><c r="B2"><v>35.7</v></c><c r="AB2" t="inlineStr"><is><t>a&amp;b</t></is></c></row></sheetData></worksheet>')
        shared = ('<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>订单编号</t></si><si><t xml:space="preserve">应收金额</t></si>'
                  '<si><r><t>备</t></r><r><t>注</t></r></si><si><t>JY1</t></si></sst>')
        buf = io.BytesIO()
        with zf.ZipFile(buf, 'w') as z:
            z.writestr('[Content_Types].xml', '<Types/>'); z.writestr('xl/worksheets/sheet1.xml', sheet); z.writestr('xl/sharedStrings.xml', shared)
        self.assertEqual(list(m._xlsx_rows(buf.getvalue(), '筛不到的'.encode())), [{0: '订单编号', 1: '应收金额', 2: '', 27: '备注'}, {0: 'JY1', 1: '35.7', 27: 'a&b'}])

    def test_push_plan_takes_only_clean_blue_bills_and_balances_each_batch(self):
        s, l = self.rows()
        r = m.reconcile('2026-09', m.merge([], s), m.merge([], l), BILLS + [
            bill('P1', '2026-09-07', 10, O6),                                                         # O6 结算毛额 10：干净
            bill('P2', '2026-09-07', 30, '6930162745589511111', ds='A')])                             # 没结算 → 不在可核销里
        plan = m.push_plan('2026-09', m.merge([], s), None, r)
        self.assertEqual((plan['eligible'], plan['left'], plan['pushed']), ({'count': 3, 'amount': 90.0}, {'count': 3, 'amount': 90.0}, {'count': 0, 'amount': 0.0}))
        b = plan['batch']
        self.assertEqual((b['bills'], b['total'], b['line_total'], b['first'], b['last']), (['AR1', 'AR2', 'P1'], 90.0, 90.0, 'AR1', 'P1'))
        self.assertEqual([(x['kind'], x['name'], x['amount']) for x in b['lines']], [('cash', '到账', 83.5), ('fee', '平台服务费', 3.5), ('fee', '佣金', 2.0), ('fee', '招商服务费', 0.5), ('fee', '站外推广费', 0.5)])
        self.assertIn('2026年09月结算单扣款项 佣金2.00元', [x['memo'] for x in b['lines']])
        self.assertEqual([(g['bills'], g['open'], g['cash'], g['fees']) for g in b['groups']][:2], [(['AR1'], 5000, 4850, {'平台服务费': 150}), (['AR2'], 3000, 2600, {'平台服务费': 100, '佣金': 200, '招商服务费': 50, '站外推广费': 50})])
        halves = m.push_split('2026-09', b['groups'], [['AR1', 'P1'], ['AR2']])                          # 金蝶把一批拆成两张：各算各的，各自平
        self.assertEqual([(x['count'], x['total'], sum(l['amount'] for l in x['lines'])) for x in halves], [(2, 60.0, 60.0), (1, 30.0, 30.0)])
        self.assertEqual([(l['name'], l['amount']) for l in halves[1]['lines']], [('到账', 26.0), ('佣金', 2.0), ('平台服务费', 1.0), ('招商服务费', 0.5), ('站外推广费', 0.5)])
        self.assertIsNone(m.push_split('2026-09', b['groups'], [['AR1'], ['AR2']]))                         # 少了一张应收：不认
        one = m.push_plan('2026-09', m.merge([], s), None, r, size=1)['batch']
        self.assertEqual((one['bills'], one['total'], one['line_total']), (['AR1'], 50.0, 50.0))
        rest = m.push_plan('2026-09', m.merge([], s), None, r, taken={'AR1'}, size=1)
        self.assertEqual((rest['batch']['bills'], rest['pushed'], rest['left']['count']), (['AR2'], {'count': 1, 'amount': 50.0}, 2))
        self.assertEqual(m.push_plan('2026-09', m.merge([], s), None, r, taken={'AR1', 'AR2', 'P1'})['batch']['bills'], [])
        late = m.reconcile('2026-10', m.merge([], s), [], [bill('AR1', '2026-09-05', 50, O1)])                       # 9 月结算的钱，10 月才来核
        p10 = m.push_plan('2026-10', m.merge([], s), None, late)
        self.assertEqual((late['bills'][0]['cat'], p10['batch']['bills'], p10['skipped']['earlier']), ('ok', [], {'count': 1, 'amount': 50.0}))
        # 订单里有红字、核销过一部分、没审核的都不推
        X = '6930162745589522222'
        extra = m.parse(settle_csv(f"2026-09-10 08:00:00,'Q1,入账,25.48,聚合账户,货款结算入账,小店自卖,'{X},'{X},120.8,0,0,-94.8,-0.52,0,0,0,订单结算"), 'x.csv')[0]['rows']
        dirty = [bill('D1', '2026-09-05', 120.8, X), bill('D1R', '2026-09-09', -94.8, X), bill('AR1', '2026-09-05', 50, O1, written=0, ws='A', ds='A'),
                 dict(bill('AR2', '2026-09-05', 40, O2, written=10, ws='B'))]
        r2 = m.reconcile('2026-09', m.merge(s, extra), [], dirty)
        self.assertEqual({b['no']: b['cat'] for b in r2['bills']}, {'D1': 'ok', 'D1R': 'ok', 'AR1': 'ok', 'AR2': 'ok'})
        p2 = m.push_plan('2026-09', m.merge(s, extra), None, r2)
        self.assertEqual((p2['batch']['bills'], {k: v['count'] for k, v in p2['skipped'].items()}), ([], {'red': 2, 'partial': 1, 'unaudited': 1, 'earlier': 0, 'cents': 0}))
        self.assertEqual(p2['held'], {'D1': 'red', 'D1R': 'red', 'AR1': 'unaudited', 'AR2': 'partial'})
        for b in r2['bills']: b['hold'] = p2['held'].get(b['no'], '')
        self.assertEqual(sorted(b['no'] for b in m.select(r2, 'hold')), ['AR1', 'AR2', 'D1', 'D1R'])
        self.assertIn('可核销但系统不推', __import__('openpyxl').load_workbook(io.BytesIO(m.export(r2, '抖音店'))).sheetnames)


if __name__ == '__main__':
    unittest.main()
