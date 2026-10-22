# [Change Log] Date: 2026-10-22 | Author: Kimi / Reginald Zeng | Version: V2.999
# V2.999 报表导出·序时簿百万行改造的两个正确性闸：
#   ① _write_journal 分表页：表页数/各页行数/表头/冻结/筛选范围对（不连金蝶，假数据）；
#   ② SQLite ORDER BY 与旧 Python 元组序【凭证号+分录行号+内码】逐项等价——排序口径不许变。
import os
import sqlite3
import sys
import tempfile
import unittest
import unittest.mock
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1] / "01_Current_Deliverables/app/backend"
sys.path.insert(0, str(BACKEND))

import kingdee_client as kc  # noqa: E402
import report_export as rx  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

NCOL = len(kc.JOURNAL_COLS)


def _row(i):
    """造一行 14 列假分录：金额列给 float（ci 10/11），期间列给 int（ci 3），其余文本。"""
    r = ["账簿A", "2026/9/1", 9, "记", str(i), "摘要%d" % i, "6601", "管理费用",
         "0011301/永续质量中心", float(i), None, "制单人", "审核人", "总账"]
    assert len(r) == NCOL
    return r


def _build(journal_iter, total):
    wb = Workbook(write_only=True)
    rx._write_journal(wb, journal_iter, total)
    fd, path = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    wb.save(path)
    return path


class JournalSheetSplitTests(unittest.TestCase):
    def setUp(self):
        self._old = rx.JOURNAL_SHEET_MAX
        rx.JOURNAL_SHEET_MAX = 3                     # 小上限代替 100 万，逻辑同一套
        self._files = []

    def tearDown(self):
        rx.JOURNAL_SHEET_MAX = self._old
        for f in self._files:
            os.remove(f)

    def _save(self, rows, total):
        p = _build(iter(rows), total)
        self._files.append(p)
        return load_workbook(p)

    def test_split_into_three_sheets(self):
        wb = self._save([_row(i) for i in range(8)], 8)
        self.assertEqual(wb.sheetnames, ["序时账簿", "序时账簿_2", "序时账簿_3"])
        self.assertEqual([wb[n].max_row for n in wb.sheetnames], [4, 4, 3])   # 表头 + 3/3/2
        self.assertEqual(wb["序时账簿"].auto_filter.ref, "A1:%s4" % chr(64 + NCOL))
        self.assertEqual(wb["序时账簿_3"].auto_filter.ref, "A1:%s3" % chr(64 + NCOL))
        for n in wb.sheetnames:
            self.assertEqual(wb[n].freeze_panes, "A2")
            self.assertEqual([c.value for c in wb[n][1]], list(kc.JOURNAL_COLS))
        # 内容连贯：三个表页的数据行接起来就是原顺序
        got = [wb[n].cell(row=r, column=5).value for n in wb.sheetnames for r in range(2, wb[n].max_row + 1)]
        self.assertEqual(got, [str(i) for i in range(8)])

    def test_exact_multiple_no_empty_tail_sheet(self):
        wb = self._save([_row(i) for i in range(6)], 6)
        self.assertEqual(wb.sheetnames, ["序时账簿", "序时账簿_2"])

    def test_empty_journal_still_one_sheet(self):
        wb = self._save([], 0)
        self.assertEqual(wb.sheetnames, ["序时账簿"])
        self.assertEqual(wb["序时账簿"].max_row, 1)
        self.assertEqual(wb["序时账簿"].auto_filter.ref, "A1:%s1" % chr(64 + NCOL))


class JournalSortOrderTests(unittest.TestCase):
    """临时库的 ORDER BY 必须与旧实现 sorted(key=(_sort_num(vno), _sort_num(seq), eid)) 逐行一致。"""

    def test_sqlite_order_matches_python_tuple_order(self):
        keys = [("10", "2", 100), ("9", "1", 101), ("10", "1", 97), ("10", "1", 98),
                ("2", "5", 105), ("9", "10", 103), ("转", "1", 99), ("", "1", 96),
                ("9", "1", 102), ("记-7", "3", 104), ("10", "10", 95)]
        expect = sorted(keys, key=lambda k: (kc._sort_num(k[0]), kc._sort_num(k[1]), k[2]))
        fd, path = tempfile.mkstemp(suffix=".db")
        os.close(fd)
        try:
            db = sqlite3.connect(path)
            db.execute("CREATE TABLE j (vflag INT, vnum INT, vstr TEXT,"
                       " sflag INT, snum INT, sstr TEXT, eid INT, %s)"
                       % ", ".join("c%d" % i for i in range(1, NCOL + 1)))
            rows = []
            for vno, seq, eid in keys:
                vk, sk = kc._sort_num(vno), kc._sort_num(seq)
                payload = [str(eid)] + [None] * (NCOL - 1)          # c1 存内码当记号
                rows.append((vk[0], vk[1] if vk[0] == 0 else None, vk[1] if vk[0] else None,
                             sk[0], sk[1] if sk[0] == 0 else None, sk[1] if sk[0] else None,
                             eid, *payload))
            db.executemany("INSERT INTO j VALUES (%s)" % ",".join("?" * (7 + NCOL)), rows)
            db.commit()
            db.close()
            got = [int(r[0]) for r in kc.iter_journal_temp(path)]
            self.assertEqual(got, [k[2] for k in expect])
        finally:
            os.remove(path)


class DiskGuardTests(unittest.TestCase):
    """V2.1000 磁盘防护：预检报人话 + 泄漏 spool 兜底清理。"""

    def test_check_disk_raises_plain_message_when_short(self):
        usage = type("U", (), {"free": 100})()                       # 只剩 100 字节
        with unittest.mock.patch.object(rx.shutil, "disk_usage", return_value=usage):
            with self.assertRaises(RuntimeError) as cm:
                rx._check_disk(10000, tempfile.gettempdir())
        self.assertIn("磁盘空间不足", str(cm.exception))

    def test_check_disk_passes_when_enough(self):
        usage = type("U", (), {"free": 10**12})()
        with unittest.mock.patch.object(rx.shutil, "disk_usage", return_value=usage):
            rx._check_disk(10000, tempfile.gettempdir())             # 不抛即过

    def test_clean_stale_spools_only_deletes_old_openpyxl_files(self):
        d = tempfile.mkdtemp()
        try:
            old = os.path.join(d, "openpyxl.old")
            new = os.path.join(d, "openpyxl.new")
            other = os.path.join(d, "keep.txt")
            for f in (old, new, other):
                open(f, "w").write("x")
            os.utime(old, (0, 0))                                    # 1970 年，必算陈旧
            with unittest.mock.patch.object(rx.tempfile, "gettempdir", return_value=d):
                freed = rx._clean_stale_spools(max_age=1800)
            self.assertFalse(os.path.exists(old))
            self.assertTrue(os.path.exists(new))
            self.assertTrue(os.path.exists(other))
            self.assertEqual(freed, 1)
        finally:
            for f in os.listdir(d):
                os.remove(os.path.join(d, f))
            os.rmdir(d)


if __name__ == "__main__":
    unittest.main()
