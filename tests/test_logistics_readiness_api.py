# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-10-25 | Author: Kimi | Version: V2.1012
# Description: 物流 readiness/sign、外仓费用归口、主体回填、导出缓存与共享缓存测试；全部使用合成数据和桩。
import asyncio
import json
import sys
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.responses import JSONResponse
from sqlalchemy.sql import compiler as _sqla_compiler
from sqlalchemy.sql import operators as _sqla_operators

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "01_Current_Deliverables" / "app" / "backend"
sys.path.insert(0, str(BACKEND))


class _Conn:
    def __init__(self, rows=None):
        self.rows = list(rows or [])
        self.writes = []

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False

    def execute(self, stmt, *args, **kwargs):
        self.writes.append(stmt)
        row = self.rows.pop(0) if self.rows else []
        return _Result(row)


class _Result:
    def __init__(self, rows):
        self.rows = rows

    def mappings(self):
        return self

    def all(self):
        return self.rows


class _Engine:
    def __init__(self):
        self.rows = []
        self.connections = []

    def _conn(self):
        c = _Conn(self.rows)
        self.connections.append(c)
        return c

    def connect(self):
        return self._conn()

    def begin(self):
        return self._conn()


class _Request:
    def __init__(self, body=None):
        self.body = body or {}

    async def json(self):
        return self.body


def _load_module():
    engine = _Engine()
    db = types.SimpleNamespace(_engine=engine, settings={}, orgs=[{"full_name": "深圳市星期零食品科技有限公司", "short_name": "深圳星期零", "book_code": "100"},
                                                                       {"full_name": "深圳市星期九食品科技有限公司", "short_name": "深圳星期九", "book_code": "101"}],
                               suppliers=[{"short": "甲", "full": "甲方物流", "wt_lo": None, "wt_hi": None}])
    db.get_setting = lambda key, default=None: json.loads(json.dumps(db.settings.get(key, default)))
    db.set_setting = lambda key, value, *args: db.settings.__setitem__(key, json.loads(json.dumps(value)))
    db.list_orgs = lambda: db.orgs
    db.list_logi_suppliers = lambda: db.suppliers

    def save_supplier(value, _user):
        db.suppliers[0].update(value)
    db.save_logi_supplier = save_supplier
    db.audit = lambda *args: None
    core = types.ModuleType("core")
    core.JSONResponse = JSONResponse
    core._require_perm = lambda *_: {"name": "测试"}
    core.db = db
    kc = types.ModuleType("kingdee_client")
    kc.login = lambda: (object(), {})
    kc._query = lambda *args: []
    kc.GL_VOUCHER_FIELDS = []
    todo = types.ModuleType("todo_scenes")
    todo.fix_touch = lambda *args: None
    with patch.dict(sys.modules, {"core": core, "kingdee_client": kc, "todo_scenes": todo}):
        from routers import logistics_review as lr
    return lr, engine


class ReadinessApiTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def setUp(self):
        self.lr._ACCR_CACHE.clear()
        self.engine.rows = []

    async def test_sign_uses_readiness_and_snapshots_summary(self):
        steps = {k: {"status": "ok", "label": k, "count": 1, "blockers": []} for k in ("intake", "lines", "docs", "sign")}
        ready = {"ok": True, "carrier": "甲", "period": "2026-09", "mode": "normal", "steps": steps,
                 "can_sign": True, "blockers": [], "signed": None}
        lines = {"accr_total": 100, "bill_total": 100, "diff_total": 0, "n_unexplained": 0, "bill_src": "detail"}
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}), \
             patch.object(self.lr, "_signed", return_value=None), \
             patch.object(self.lr, "_collect_readiness", return_value=(ready, lines, {"req_total": 0})) as collect:
            out = await self.lr.review_sign_month(_Request({"carrier": "甲", "period": "2026-09", "note": "已核"}))
        self.assertTrue(out["ok"])
        collect.assert_called_once()
        self.assertEqual(collect.call_args.args[-1], "已核")
        snap = json.loads(out["signed"]["snap_json"])
        self.assertTrue(snap["readiness"]["can_sign"])
        self.assertEqual(snap["readiness"]["steps"], steps)

    async def test_sign_returns_409_with_steps_and_does_not_write(self):
        steps = {k: {"status": "ok", "label": k, "count": 0, "blockers": []} for k in ("intake", "docs", "sign")}
        steps["lines"] = {"status": "pending", "label": "登记制汇总已就绪", "count": 1, "blockers": ["差异必须写 note"]}
        ready = {"mode": "register", "steps": steps, "can_sign": False, "blockers": ["差异必须写 note"]}
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}), \
             patch.object(self.lr, "_signed", return_value=None), \
             patch.object(self.lr, "_collect_readiness", return_value=(ready, None, {"req_total": 100})):
            out = await self.lr.review_sign_month(_Request({"carrier": "甲", "period": "2026-09", "note": ""}))
        self.assertIsInstance(out, JSONResponse)
        self.assertEqual(out.status_code, 409)
        body = json.loads(out.body)
        self.assertFalse(body["ok"])
        self.assertEqual(body["steps"], steps)
        self.assertFalse(any(c.writes for c in self.engine.connections))

    async def test_readiness_endpoint_returns_shared_contract(self):
        ready = {"ok": True, "carrier": "甲", "period": "2026-09", "mode": "normal", "steps": {},
                 "can_sign": False, "blockers": [], "signed": None}
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}), \
             patch.object(self.lr, "_collect_readiness", return_value=(ready, None, {})):
            out = await self.lr.review_readiness(_Request(), carrier="甲", period="2026-09")
        self.assertEqual(out, ready)

    async def test_register_comparison_aggregates_requests_by_subject(self):
        overview = {"ok": True, "rows": [{"short": "甲", "cells": {"深圳星期零": {
            "accr": 100, "reqs": [{"amount": 60, "st": {"key": "approved"}},
                                   {"amount": 40, "st": {"key": "approved"}}]}}}]}
        with patch.object(self.lr, "review_overview", return_value=overview):
            data, total = self.lr._register_readiness(_Request(), "甲", "2026-09")
        self.assertEqual(total, 100)
        self.assertEqual(data["cells"][0]["diff"], 0)
        self.assertTrue(data["cells"][0]["has_accr"])


class RawDownloadApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def setUp(self):
        from sqlalchemy import create_engine
        self.old_engine = self.lr.db._engine
        self.db_engine = create_engine("sqlite:///:memory:")
        self.lr.BR.create(self.db_engine)
        self.lr.db._engine = self.db_engine
        with self.db_engine.begin() as c:
            for src, name, data in (
                ("线上", "旧账单.xlsx", b"old-xlsx"),
                ("线上", "2026年9月 线上&零售 最新账单.xlsx", b"new-xlsx"),
                ("旧版", "旧格式账单.xls", b"legacy-xls"),
            ):
                c.execute(self.lr.BR.insert().values(carrier="甲", period="2026-09", bill_src=src, name=name, data=data, size=len(data)))
            c.execute(self.lr.BR.insert().values(carrier="乙", period="2026-09", bill_src="线上", name="别家.xlsx", data=b"other-carrier"))
            c.execute(self.lr.BR.insert().values(carrier="甲", period="2026-08", bill_src="线上", name="别期.xlsx", data=b"other-period"))

    def tearDown(self):
        self.lr.db._engine = self.old_engine
        self.db_engine.dispose()

    def test_latest_excel_content_and_original_filename(self):
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}):
            out = self.lr.review_raw(_Request(), carrier="甲", period="2026-09", src="线上")
        self.assertEqual(out.status_code, 200)
        self.assertEqual(out.body, b"new-xlsx")
        self.assertEqual(out.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        from urllib.parse import quote
        self.assertIn("filename*=UTF-8''" + quote("2026年9月 线上&零售 最新账单.xlsx", safe=""), out.headers["content-disposition"])
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}):
            old = self.lr.review_raw(_Request(), carrier="甲", period="2026-09", src="旧版")
        self.assertEqual(old.body, b"legacy-xls")
        self.assertEqual(old.media_type, "application/vnd.ms-excel")
        self.assertIn(quote("旧格式账单.xls", safe=""), old.headers["content-disposition"])

    def test_permission_denied(self):
        with patch.object(self.lr, "_perm", return_value=None):
            out = self.lr.review_raw(_Request(), carrier="甲", period="2026-09", src="线上")
        self.assertIsInstance(out, JSONResponse)
        self.assertEqual(out.status_code, 403)

    def test_cross_scope_and_missing_src_are_404(self):
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}):
            for carrier, period, src in (("乙", "2026-09", "旧版"), ("甲", "2026-08", "旧版"), ("甲", "2026-09", "不存在")):
                with self.subTest(carrier=carrier, period=period, src=src):
                    out = self.lr.review_raw(_Request(), carrier=carrier, period=period, src=src)
                    self.assertIsInstance(out, JSONResponse)
                    self.assertEqual(out.status_code, 404)


class ReviewSettingsTests(unittest.IsolatedAsyncioTestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def setUp(self):
        self.lr.db.settings.clear()
        self.lr._ACCR_CACHE.clear()

    async def test_defaults_save_and_preserve_other_settings(self):
        default = self.lr.review_settings_read(_Request(), "甲")
        self.assertTrue(default["qty_enabled"])
        self.assertTrue(default["price_enabled"])
        self.assertEqual(default["scheme"], "")
        self.lr.db.settings[self.lr._REVIEW_SCOPE_KEY] = {"乙": {"qty_enabled": False, "price_enabled": True}}
        self.lr.db.settings[self.lr._VOL_KG_KEY] = {"乙": 250}
        out = await self.lr.review_settings_save(_Request({"carrier": "甲", "qty_enabled": False, "price_enabled": True,
                                                           "scheme": "volume", "vol_kg": 333.33, "wt_range": [0.9, 1.2],
                                                           "default_subject": "深圳星期零"}))
        self.assertTrue(out["ok"])
        self.assertEqual(out["scheme"], "volume")
        self.assertEqual(out["vol_kg"], 333.33)
        self.assertEqual(out["wt_range"], [0.9, 1.2])
        self.assertEqual(self.lr.db.settings[self.lr._REVIEW_SCOPE_KEY]["乙"]["qty_enabled"], False)
        self.assertEqual(self.lr.db.settings[self.lr._VOL_KG_KEY]["乙"], 250)
        out = await self.lr.review_settings_save(_Request({"carrier": "甲", "scheme": "box"}))
        self.assertEqual(out["vol_kg"], 333.33)

    async def test_validation(self):
        cases = [
            {"carrier": "甲", "qty_enabled": False, "price_enabled": False},
            {"carrier": "甲", "scheme": "bad"},
            {"carrier": "甲", "scheme": "volume", "vol_kg": None},
            {"carrier": "甲", "scheme": "pallet", "pallet_kg": 650, "pallet_box": None},
            {"carrier": "甲", "default_subject": "不存在"},
            {"carrier": "甲", "wt_range": [1.2, 0.9]},
        ]
        for body in cases:
            with self.subTest(body=body):
                out = await self.lr.review_settings_save(_Request(body))
                self.assertIsInstance(out, JSONResponse)
                self.assertEqual(out.status_code, 400)
        self.assertNotIn(self.lr._REVIEW_SCOPE_KEY, self.lr.db.settings)


class SubjectBackfillTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def test_kingdee_default_and_manual_precedence(self):
        rows = [{"doc_no": "XSCKD1", "subject": ""}, {"doc_no": "CGRK1", "subject": ""},
                {"doc_no": "FBDR1", "subject": ""}, {"doc_no": "QTCK1", "subject": ""},
                {"doc_no": "XSCKD2", "subject": "深圳星期九"}, {"doc_no": "XSCKD3", "subject": "", "subj_ovr": "深圳星期九"}]
        mats = {no: [{"_form": form, "往来": "甲仓→乙仓"}] for no, form in
                (("XSCKD1", "SAL_OUTSTOCK"), ("CGRK1", "STK_InStock"), ("FBDR1", "STK_TransferIn"), ("QTCK1", "STK_MisDelivery"))}
        heads = {"XSCKD1": {"销售组织": "深圳市星期零食品科技有限公司"}, "CGRK1": {"需求组织": "深圳星期九"},
                 "FBDR1": {"调入组织": "深圳星期零"}, "QTCK1": {"库存组织": "深圳星期零"}}
        with patch.object(self.lr, "_review_scope", return_value=({"default_subject": "深圳星期九"}, True)), \
             patch.object(self.lr, "_fetch_doc_heads", return_value=heads):
            self.lr._backfill_subjects(rows, "甲", mats=mats, session=(object(), {}))
        self.assertEqual(self.lr._eff_subject(rows[0]), "深圳星期零")
        self.assertEqual(rows[0]["_subject_source"], "kingdee")
        self.assertEqual(self.lr._eff_subject(rows[1]), "深圳星期九")
        self.assertEqual(self.lr._eff_subject(rows[2]), "深圳星期零")
        self.assertEqual(self.lr._eff_subject(rows[3]), "深圳星期九")
        self.assertEqual(rows[3]["_subject_source"], "supplier_default")
        self.assertEqual(rows[3]["subject"], "")
        self.assertEqual(self.lr._eff_subject(rows[4]), "深圳星期九")
        self.assertEqual(self.lr._eff_subject(rows[5]), "深圳星期九")
        with patch.object(self.lr, "_review_scope", return_value=({"default_subject": ""}, True)), \
             patch.object(self.lr, "_fetch_doc_heads", return_value=heads):
            other = {"doc_no": "QTCK1", "subject": ""}
            self.lr._backfill_subjects([other], "甲", mats=mats, session=(object(), {}))
        self.assertEqual(self.lr._eff_subject(other), "")
        self.assertTrue(other["_subject_pending"])


class SchemeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def test_volume_missing_dimensions_cannot_pass_by_net_weight(self):
        row = {"doc_no": "XSCKD1", "subject": "深圳星期零", "qty": 1, "charge_wt": 10,
               "amount": 10, "unit": "千克", "sub_fees": "{}", "fee": "销售出库费用"}
        mats = {"XSCKD1": [{"_form": "SAL_OUTSTOCK", "编码": "M1", "名称": "货品", "基本数量": 10,
                            "基本单位": "千克", "数量件": 10, "计价单位": "袋", "规格": "10袋/箱"}]}
        scope = {"qty_enabled": True, "price_enabled": True, "scheme": "volume", "default_subject": ""}
        with patch.object(self.lr, "_review_scope", return_value=(scope, True)), \
             patch.object(self.lr, "_fetch_doc_materials", return_value=mats), \
             patch.object(self.lr, "_load_spec", return_value={}), \
             patch.object(self.lr, "_vol_kg", return_value=333.33):
            docs = self.lr._box_docs([row], "甲")
        self.assertEqual(docs[0]["state"], "qtydiff")
        self.assertIn("箱子尺寸", docs[0]["mode_cn"])
        self.assertTrue(docs[0]["qty_enabled"])

    def test_volume_and_pallet_use_configured_algorithms(self):
        base = {"doc_no": "XSCKD1", "subject": "深圳星期零", "qty": 10, "amount": 10,
                "unit": "千克", "fee": "销售出库费用"}
        mats = {"XSCKD1": [{"_form": "SAL_OUTSTOCK", "编码": "M1", "名称": "货品", "基本数量": 1,
                            "基本单位": "千克", "数量件": 100, "计价单位": "袋", "规格": "10袋/箱"}]}
        vol_row = dict(base, charge_wt=10, sub_fees=json.dumps({"箱子尺寸": "10×10×10"}, ensure_ascii=False))
        vol_scope = {"qty_enabled": True, "price_enabled": True, "scheme": "volume", "default_subject": ""}
        with patch.object(self.lr, "_review_scope", return_value=(vol_scope, True)), \
             patch.object(self.lr, "_fetch_doc_materials", return_value=mats), \
             patch.object(self.lr, "_load_spec", return_value={}), \
             patch.object(self.lr, "_vol_kg", return_value=1000):
            vol_doc = self.lr._box_docs([vol_row], "甲")[0]
        self.assertEqual(vol_doc["state"], "ok")
        self.assertEqual(vol_doc["chk"]["kind"], "vol")
        pal_row = dict(base, charge_wt=650, sub_fees="{}")
        pal_scope = dict(vol_scope, scheme="pallet")
        with patch.object(self.lr, "_review_scope", return_value=(pal_scope, True)), \
             patch.object(self.lr, "_fetch_doc_materials", return_value=mats), \
             patch.object(self.lr, "_load_spec", return_value={}), \
             patch.object(self.lr, "_pallet_kg", return_value=650), \
             patch.object(self.lr, "_pallet_box", return_value=10):
            pal_doc = self.lr._box_docs([pal_row], "甲")[0]
        self.assertEqual(pal_doc["state"], "ok")
        self.assertEqual(pal_doc["chk"]["kind"], "pallet")

    def test_pallet_missing_rule_cannot_pass_by_net_weight(self):
        row = {"doc_no": "XSCKD1", "subject": "深圳星期零", "qty": 10, "charge_wt": 650,
               "amount": 10, "unit": "千克", "sub_fees": "{}", "fee": "销售出库费用"}
        mats = {"XSCKD1": [{"_form": "SAL_OUTSTOCK", "编码": "M1", "名称": "货品", "基本数量": 650,
                            "基本单位": "千克", "数量件": 100, "计价单位": "袋", "规格": "10袋/箱"}]}
        scope = {"qty_enabled": True, "price_enabled": True, "scheme": "pallet", "default_subject": ""}
        with patch.object(self.lr, "_review_scope", return_value=(scope, True)), \
             patch.object(self.lr, "_fetch_doc_materials", return_value=mats), \
             patch.object(self.lr, "_load_spec", return_value={}), \
             patch.object(self.lr, "_pallet_kg", return_value=650), \
             patch.object(self.lr, "_pallet_box", return_value=0):
            doc = self.lr._box_docs([row], "甲")[0]
        self.assertEqual(doc["state"], "qtydiff")
        self.assertIn("按托核不了", doc["mode_cn"])


class FailureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def setUp(self):
        self.lr._ACCR_CACHE.clear()
        self.engine.rows = []

    def test_accrual_login_failure_is_explicit(self):
        with patch.object(self.lr.kc, "login", side_effect=RuntimeError("连接超时")):
            with self.assertRaisesRegex(RuntimeError, "金蝶读取逐笔计提失败.*连接超时"):
                self.lr._accr_entries_raw("甲", "2026-09")

    def test_lines_failure_returns_502_and_does_not_cache(self):
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}), \
             patch.object(self.lr, "_build_lines", side_effect=RuntimeError("金蝶失败")):
            out = self.lr.review_lines(_Request(), "甲", "2026-09")
        self.assertEqual(out.status_code, 502)
        self.assertIn("金蝶失败", json.loads(out.body)["msg"])
        self.assertNotIn(("lines", "甲", "2026-09"), self.lr._ACCR_CACHE)

    def test_box_docs_material_failure_propagates(self):
        row = {"doc_no": "XSCKD1"}
        with patch.object(self.lr.kc, "login", side_effect=RuntimeError("物料服务不可用")):
            with self.assertRaisesRegex(RuntimeError, "金蝶读取单据物料失败"):
                self.lr._box_docs([row], "丰源")
        with patch.object(self.lr.kc, "login", return_value=(object(), {})), \
             patch.object(self.lr.kc, "_query", side_effect=RuntimeError("查询失败")):
            with self.assertRaisesRegex(RuntimeError, "金蝶查询SAL_OUTSTOCK物料失败"):
                self.lr._box_docs([row], "丰源")

    def test_result_material_failure_is_502_and_not_cached(self):
        self.engine.rows = [[{"id": 1, "doc_no": "XSCKD1", "grain": "detail", "amount": 10, "qty": 1}], [], []]
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}), \
             patch.object(self.lr, "_load_card", return_value=({}, 0)), \
             patch.object(self.lr, "_doc_ok", return_value={}), \
             patch.object(self.lr, "_code_maps", return_value={}), \
             patch.object(self.lr, "_box_docs", side_effect=RuntimeError("金蝶读取单据物料失败")):
            out = self.lr.review_result(_Request(), carrier="丰源", period="2026-09", group="all")
        self.assertEqual(out.status_code, 502)
        self.assertIn("物料", json.loads(out.body)["msg"])
        self.assertNotIn(("view", "丰源", "2026-09"), self.lr._ACCR_CACHE)


class MaterialCacheTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def setUp(self):
        self.lr._ACCR_CACHE.clear()
        self.lr._DOC_MAT_CACHE.clear()

    def test_shared_material_cache_fresh_and_bust(self):
        calls = []

        def query(session, conf, form, fields, condition):
            calls.append((form, condition))
            return [{"单号": "XSCKD1", "编码": "M1", "基本数量": 2}]

        docs = {"SAL_OUTSTOCK": {"XSCKD1"}}
        with patch.object(self.lr.kc, "_query", side_effect=query):
            first = self.lr._fetch_doc_materials(object(), {}, docs)
            second = self.lr._fetch_doc_materials(object(), {}, docs)
            self.assertEqual(len(calls), 1)
            self.assertEqual(first, second)
            self.lr._fetch_doc_materials(object(), {}, docs, fresh=True)
            self.assertEqual(len(calls), 2)
            self.lr._bust("丰源", "2026-09")
            self.assertEqual(len(self.lr._DOC_MAT_CACHE), 0)
            self.lr._fetch_doc_materials(object(), {}, docs)
            self.assertEqual(len(calls), 3)


class ExportCacheTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def test_export_reuses_lines_cache_instead_of_rebuilding(self):
        from io import BytesIO
        from openpyxl import load_workbook
        lines = {"ok": True, "rows": [], "fixes": [], "suppliers": []}
        payreq = types.ModuleType("routers.logistics_payreq")
        payreq.inv_blocks = lambda carrier, period: []
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}), \
             patch.object(self.lr, "review_result", return_value={"material": False, "detail": []}), \
             patch.object(self.lr, "_cached_lines", return_value=lines) as cached, \
             patch.object(self.lr, "_build_lines", side_effect=AssertionError("导出不得绕过缓存重算")), \
             patch.object(self.lr, "_raw_json_exists", return_value=True), \
             patch.dict(sys.modules, {"routers.logistics_payreq": payreq}):
            out = self.lr.review_export(_Request(), carrier="甲", period="2026-09")
        cached.assert_called_once()
        self.assertEqual(out.status_code, 200)
        wb = load_workbook(BytesIO(out.body), read_only=True)
        self.assertEqual(wb.sheetnames, ["复核表", "复核明细"])


class SheetRoundingApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def setUp(self):
        from sqlalchemy import create_engine
        from kernels import logistics_review_store as store
        self.old_engine = self.lr.db._engine
        self.engine = create_engine("sqlite:///:memory:")
        store._md.create_all(self.engine)
        self.lr.db._engine = self.engine
        amounts = [("蜜雪", 18222.50), ("植物肉", 547.779711), ("小料", 23384.9028), ("分步", 286.114155)]
        with self.engine.begin() as c:
            for sheet, amount in amounts:
                c.execute(self.lr.BL.insert().values(
                    period="2026-09", carrier="甲", grain="detail", subject="深圳星期零", doc_no="无单据",
                    fee_item="入库运费", annot="采购入库-植物肉", annot_source="default", bizline="植物肉",
                    amount=amount, bill_src="天鹰", src_sheet=sheet))

    def tearDown(self):
        self.lr.db._engine = self.old_engine
        self.engine.dispose()

    def _lines(self):
        entry = {"subject": "深圳星期零", "fee_norm": "入库运费", "fee": "入库运费", "amt": 42441.29,
                 "biz": "植物肉", "key": "记-1|深圳星期零", "vno": "记-1", "book_code": "100"}
        got = {"entries": [entry], "adj": [], "prior": []}
        with patch.object(self.lr, "_accr_entries", return_value=got), \
             patch.object(self.lr, "_dept_biz_rules", return_value=({}, {})), \
             patch.object(self.lr, "_backfill_subjects", lambda *args, **kwargs: None), \
             patch.object(self.lr, "_xsubj_hints", return_value={}), \
             patch.object(self.lr, "_subj_marks", return_value={}):
            return self.lr._build_lines(_Request(), "甲", "2026-09")

    def test_build_lines_totals_and_readiness_share_sheet_rounding(self):
        out = self._lines()
        self.assertEqual(round(18222.50 + 547.779711 + 23384.9028 + 286.114155, 2), 42441.30)
        self.assertEqual(out["bill_total"], 42441.29)
        totals = [r for r in out["rows"] if r.get("kind") == "gtotal"]
        self.assertEqual(len(totals), 1)
        self.assertEqual(totals[0]["bill"], 42441.29)
        self.assertEqual(totals[0]["diff"], 0)
        self.assertEqual(out["diff_total"], 0)
        totals[0]["mark"] = {"status": "ok", "dv": ""}
        step = self.lr.lready.evaluate("甲", "2026-09", lines=out)["steps"]["lines"]
        self.assertEqual(step["status"], "ok")

    def test_bills_api_uses_sheet_rounded_amount(self):
        with patch.object(self.lr, "_perm", return_value={"name": "测试"}), \
             patch.object(self.lr, "_signed", return_value=None), patch.object(self.lr, "_load_spec", return_value={}):
            out = self.lr.review_bills(_Request(), "甲", "2026-09")
        self.assertEqual(out["bills"][0]["amount"], 42441.29)
        self.assertEqual(out["bills"][0]["n_detail"], 4)

    def test_report_cache_key_contains_algorithm_version(self):
        core = types.ModuleType("core")
        core.JSONResponse = JSONResponse
        core._require_perm = lambda *_: {"name": "测试"}
        core.db = self.lr.db
        with patch.dict(sys.modules, {"core": core, "routers.logistics_review": self.lr}):
            from routers import logistics_intake_check as lic
            spec = {"sheets": []}
            old = lic._spec_sha(spec)
            with patch.object(lic, "_REPORT_ALGORITHM", "sheet-round-v2"):
                self.assertNotEqual(lic._spec_sha(spec), old)

    def test_classify_merge_keeps_annot_and_source_from_same_row(self):
        with self.engine.begin() as c:
            for source, annot in (("default", "销售单-植物肉"), ("raw", "销售单-小料")):
                c.execute(self.lr.BL.insert().values(period="2026-09", carrier="甲", grain="detail", subject="深圳星期零",
                                                     doc_no="XSCKD1", annot=annot, annot_source=source, amount=1,
                                                     bill_src="天鹰", src_sheet="明细"))
        core = types.ModuleType("core")
        core.JSONResponse = JSONResponse
        core.db = self.lr.db
        kc = types.ModuleType("kingdee_client")
        with patch.dict(sys.modules, {"core": core, "kingdee_client": kc, "routers.logistics_review": self.lr}):
            from routers import logistics_classify as lcr
            with patch.object(self.lr, "_backfill_subjects", lambda *args, **kwargs: None):
                docs = lcr._bill_docs("甲", "2026-09")
        self.assertEqual(len(docs), 1)
        self.assertIn((docs[0]["annot"], docs[0]["annot_source"]),
                      {("销售单-植物肉", "default"), ("销售单-小料", "raw")})

    def test_migrate_adds_annot_source_to_old_bill_lines(self):
        from sqlalchemy import create_engine, text
        from kernels import logistics_review_store as store
        engine = create_engine("sqlite:///:memory:")
        try:
            with engine.begin() as c:
                c.execute(text("CREATE TABLE logistics_bill_lines (id INTEGER PRIMARY KEY)"))
            store.migrate_cols(engine)
            with engine.connect() as c:
                cols = {r[1] for r in c.execute(text("PRAGMA table_info(logistics_bill_lines)")).all()}
            self.assertIn("annot_source", cols)
        finally:
            engine.dispose()


class FeeClassificationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.lr, cls.engine = _load_module()

    def test_hengmao_cold_storage_is_inbound_but_manual_override_wins(self):
        self.assertEqual(self.lr._eff_fee({"carrier": "恒茂", "fee_item": "冷藏费", "annot": "小料原料", "doc_no": "无单据"}), "入库运费")
        self.assertEqual(self.lr._eff_fee({"carrier": "恒茂", "fee_item": "入库处置费", "annot": "仓储费-小料", "doc_no": "FBDR1"}), "入库运费")
        self.assertEqual(self.lr._eff_fee({"carrier": "恒茂", "fee_item": "冷藏费", "annot": "仓储费-小料", "fee_ovr": "仓储费"}), "仓储费")
        self.assertEqual(self.lr._eff_fee({"carrier": "甲", "fee_item": "其他", "annot": "仓储费-小料", "doc_no": "无单据"}), "仓储费")


if __name__ == "__main__":
    unittest.main()
