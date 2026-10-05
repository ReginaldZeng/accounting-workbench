import datetime
import io
import unittest
from kernels import ec_open_items as m

TODAY = datetime.date(2026, 10, 5)


def ar(no, date, amount, written=0, shop='抖音店', ws='A', ds='C', t6='', t4=''):
    return dict(no=no, date=date + 'T00:00:00', shop=shop, ws=ws, ds=ds, amount=amount, written=written, t6=t6, t4=t4)


class OpenItemsTests(unittest.TestCase):
    def test_order_key_strips_split_suffix_and_ignores_non_orders(self):
        self.assertEqual(m.order_key('6951161373055259888-1'), '6951161373055259888')
        self.assertEqual(m.order_key('', 'JY202603080025', '6951161373055259888'), '6951161373055259888')
        self.assertEqual(m.order_key('JY202603080025', None), '')

    def test_piles_follow_today(self):
        self.assertEqual(m.pile_of('2026-10-05', TODAY), 'current')
        self.assertEqual(m.pile_of('2026-09-01', TODAY), 'current')
        self.assertEqual(m.pile_of('2026-08-31', TODAY), 'stray')
        self.assertEqual(m.pile_of('2025-10-01', TODAY), 'stray')
        self.assertEqual(m.pile_of('2025-09-30', TODAY), 'old')
        self.assertEqual(m.pile_of('2026-12-15', datetime.date(2027, 1, 3)), 'current')

    def test_open_amount_uses_written_off_and_merges_plan_rows(self):
        snap = m.build([ar('AR1', '2026-08-05', 60, 30, ws='B'), ar('AR1', '2026-08-05', 40, 10, ws='B')], [], TODAY)
        b = snap['ar'][0]
        self.assertEqual((b['amount'], b['written'], b['open'], b['hint']), (100.0, 40.0, 60.0, 'partial'))

    def test_red_blue_pairs_are_found_per_shop_and_order(self):
        o = '3315639003304031685'
        snap = m.build([ar('AR1', '2026-08-05', 89, 44.5, ws='B', t6=o), ar('AR2', '2026-08-20', -44.5, t4=o),
                        ar('AR3', '2026-08-06', 50, t6='5127624720015034733'), ar('AR4', '2026-08-21', -20, t4='5127624720015034733'),
                        ar('AR5', '2026-08-21', -48, t4='3316977912172104985'), ar('AR6', '2026-08-21', 48, shop='天猫店', t6='3316977912172104985'),
                        ar('AR7', '2026-10-02', 30, ds='A', t6='6930162745589595612')], [], TODAY)
        hint = {b['no']: (b['hint'], b['pair_net']) for b in snap['ar']}
        self.assertEqual(hint['AR1'], ('pair_even', 0.0)); self.assertEqual(hint['AR2'], ('pair_even', 0.0))
        self.assertEqual(hint['AR3'], ('pair_rest', 30.0)); self.assertIn('30.00', m.hint_text(snap['ar'][[b['no'] for b in snap['ar']].index('AR3')]))
        self.assertEqual(hint['AR5'], ('red_alone', None))          # 蓝字在另一家店，不算成对
        self.assertEqual(hint['AR6'], ('', None))
        self.assertEqual(hint['AR7'], ('unaudited', None))

    def test_summary_select_page_and_export(self):
        rows = [ar('AR%03d' % i, '2026-09-%02d' % (i % 28 + 1), 10, t6='69301627455895%05d' % i) for i in range(120)]
        rows += [ar('OLD1', '2023-03-01', 99.5, shop='天猫老店'), ar('RED1', '2026-07-01', -9.4, shop='天猫老店')]
        rec = [dict(no='SKD1', date='2026-09-21', shop='抖音店', ws='A', ds='C', amount=500, written=120, settle='电汇', account='宁波行')]
        snap = m.build(rows, rec, TODAY)
        s = m.summary(snap)
        self.assertEqual(s['total'], {'count': 122, 'amount': 1290.1})
        self.assertEqual(s['piles']['current']['count'], 120); self.assertEqual(s['piles']['old']['amount'], 99.5)
        self.assertEqual([r['shop'] for r in s['shops']], ['抖音店', '天猫老店'])
        self.assertEqual(s['shops'][1]['stray'], {'count': 1, 'amount': -9.4})
        self.assertEqual(s['receipts']['total'], {'count': 1, 'amount': 380.0})
        self.assertEqual(list(s['hints']), ['red_alone'])
        picked = m.select(snap, pile='current', shop='抖音店')
        p = m.page(picked, 3)
        self.assertEqual((p['total'], p['pages'], p['page'], len(p['rows']), p['amount']), (120, 3, 3, 20, 1200.0))
        self.assertEqual(p['months'], [{'month': '2026-09', 'count': 120, 'amount': 1200.0}])
        self.assertEqual(len(m.select(snap, q='ar0')), 100); self.assertEqual(len(m.select(snap, q='6930162745589500007')), 1)
        self.assertEqual(len(m.select(snap, hint='none')), 121); self.assertEqual(len(m.select(snap, side='rec', month='2026-09')), 1)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(m.export(snap, picked, 'ar', {'ts': '2026-10-05 15:00:00', 'operator': '测试'})))
        self.assertEqual(wb.sheetnames, ['汇总', '应收单明细']); self.assertEqual(wb['应收单明细'].max_row, 121)
        self.assertEqual(load_workbook(io.BytesIO(m.export(snap, snap['rec'], 'rec', None)))['收款单明细']['F2'].value, 380.0)


if __name__ == '__main__':
    unittest.main()
