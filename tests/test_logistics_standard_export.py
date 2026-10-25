# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-10-25 | Author: Kimi | Version: V2.1011
# Description: 物流标准复核表结构、同单汇总、合并单元格、金额勾稽和原账单链接测试；使用合成数据。
import json
import sys
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "01_Current_Deliverables" / "app" / "backend"))

from routers.logistics_standard_export import C, build_standard_workbook


class StandardExportTests(unittest.TestCase):
    def make_raw(self):
        wb = Workbook()
        inbound = wb.active
        inbound.title = "入库"
        cold = wb.create_sheet("冷藏")
        for ws in (inbound, cold):
            ws.append(["2026年9月对账单"])
            ws.append(["序号", "类型", "商品名称", "规格", "备注"])
            ws.append([1, "入库" if ws is inbound else "冷藏", "测试物料", "10袋/箱", "测试备注"])
        bio = BytesIO()
        wb.save(bio)
        return [("测试账单.xlsx", bio.getvalue())]

    def test_complete_workbook_and_reconciliation(self):
        bills = [
            {"id": 1, "doc_no": "FBDR1", "subject": "孝感星期九", "fee_item": "入库处置费", "amount": 20.17,
             "qty": 99, "charge_wt": 1039.5, "unit": "箱", "src_sheet": "入库", "src_row": 3, "annot": "仓储费-小料",
             "sub_fees": json.dumps({"处置费": 20.17, "标准": 20.16, "公式": "吨重×19.4元/吨"}, ensure_ascii=False)},
            {"id": 2, "doc_no": "无单据", "subject": "孝感星期九", "fee_item": "冷藏费", "amount": 58.2,
             "qty": 5, "charge_wt": None, "unit": "板", "src_sheet": "冷藏", "src_row": 3, "annot": "小料原料",
             "sub_fees": json.dumps({"冷藏费": 58.2, "标准": 58.2, "公式": "天数×计费数量×冷藏单价"}, ensure_ascii=False)},
        ]
        materials = [
            {"code": "M1", "name": "物料一", "spec": "10袋/箱", "base_kg": 10, "kg_unit": "千克", "kd": 10, "box_div": 10, "fee": 10.17},
            {"code": "M2", "name": "物料二", "spec": "20袋/箱", "base_kg": 20, "kg_unit": "千克", "kd": 20, "box_div": 20, "fee": 10},
        ]
        doc = dict(bills[0], lid=1, bizline="小料", doc_fee=20.17, bill_amt=30, bill_unit="千克", kd_sum=30,
                   kd_unit="千克", state="ok", chk={"kind": "weight"}, materials=materials)
        doc["sub_fees"] = json.loads(bills[0]["sub_fees"])
        res = {"total_bill": 78.37, "settings": {"qty_enabled": True, "price_enabled": True,
               "subjects": [{"code": "107", "name": "孝感星期九"}]}, "codes": {"biz": {"小料": "CPFL009"}}, "docs": [doc]}
        lines = {"accr_total": 78.37, "bill_total": 78.37, "diff_total": 0, "n_unexplained": 0,
                 "rows": [{"kind": "accr", "subject": "孝感星期九", "book_code": "107", "acct": "5101", "acct_name": "制造费用",
                           "fee_type": "入库运费", "amt": 78.37, "bill": 78.37, "diff": 0, "note": ""}]}
        wb = build_standard_workbook("恒茂", "2026-09", res, lines, bills, self.make_raw())
        self.assertEqual(wb.sheetnames, ["01复核说明书", "02单据运费", "03计提对账单", "原账单-入库", "原账单-冷藏"])
        ws = wb["02单据运费"]
        self.assertEqual(ws.max_row, 7)
        self.assertEqual(sum(ws.cell(r, C["账单金额"]).value or 0 for r in range(5, 8)), 78.37)
        self.assertEqual(sum(ws.cell(r, C["标准金额"]).value or 0 for r in range(5, 8)), 78.36)
        self.assertEqual(ws.cell(7, C["数量结论"]).value, "不适用·无运输单据")
        self.assertEqual(ws.cell(7, C["物料名称"]).value, "测试物料")
        self.assertTrue(ws.cell(6, C["父费用行ID"]).value)
        self.assertIn("0.01", wb["01复核说明书"]["D8"].value)
        for name in ("原账单-入库", "原账单-冷藏"):
            formulas = [c.value for row in wb[name].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
            self.assertTrue(any("02单据运费" in f for f in formulas))
        bad = [(ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and "#REF!" in c.value]
        self.assertEqual(bad, [])
        bio = BytesIO()
        wb.save(bio)
        reopened = load_workbook(BytesIO(bio.getvalue()), data_only=False)
        self.assertEqual(reopened.sheetnames, wb.sheetnames)

    def test_same_doc_is_compared_once_and_raw_quantity_cells_merge(self):
        bills = [
            {"id": 1, "doc_no": "FBDR1", "subject": "孝感星期九", "fee_item": "入库处置费", "amount": 20.17,
             "qty": 10, "charge_wt": 10, "src_sheet": "入库", "src_row": 3, "annot": "小料", "sub_fees": '{"标准": 20.17}'},
            {"id": 3, "doc_no": "FBDR1", "subject": "孝感星期九", "fee_item": "入库处置费", "amount": 10,
             "qty": 20, "charge_wt": 20, "src_sheet": "入库", "src_row": 4, "annot": "小料", "sub_fees": '{"标准": 10}'},
        ]
        materials = [{"code": "M1", "name": "物料一", "base_kg": 10, "kg_unit": "千克", "kd": 10, "fee": 20.17},
                     {"code": "M2", "name": "物料二", "base_kg": 20, "kg_unit": "千克", "kd": 20, "fee": 10}]
        docs = []
        for b in bills:
            docs.append(dict(b, lid=b["id"], doc_fee=b["amount"], bill_amt=30, doc_bill_all=30, bill_unit="千克",
                             kd_sum=30, kd_unit="千克", state="ok", chk={"kind": "weight"}, materials=materials,
                             sub_fees=json.loads(b["sub_fees"])))
        res = {"total_bill": 30.17, "settings": {"qty_enabled": True, "price_enabled": True, "subjects": []},
               "codes": {}, "docs": docs}
        lines = {"accr_total": 30.17, "bill_total": 30.17, "diff_total": 0, "rows": []}
        raw = load_workbook(BytesIO(self.make_raw()[0][1]))
        raw["入库"].append([2, "入库", "测试物料二", "20袋/箱", ""])
        bio = BytesIO()
        raw.save(bio)
        wb = build_standard_workbook("恒茂", "2026-09", res, lines, bills, [("测试账单.xlsx", bio.getvalue())])
        ws = wb["02单据运费"]
        self.assertEqual(ws.max_row, 6)
        self.assertEqual(ws.cell(5, C["账单数量"]).value, 30)
        self.assertEqual(ws.cell(5, C["账单金额"]).value, 30.17)
        self.assertEqual(ws.cell(5, C["标准金额"]).value, 30.17)
        self.assertIn("AC5-SUM(AA5:AA6)", ws.cell(5, C["数量差异"]).value)
        merged = {str(r) for r in wb["原账单-入库"].merged_cells.ranges}
        self.assertIn("G3:G4", merged)
        self.assertIn("H3:H4", merged)
        self.assertIn("I3:I4", merged)


if __name__ == "__main__":
    unittest.main()
