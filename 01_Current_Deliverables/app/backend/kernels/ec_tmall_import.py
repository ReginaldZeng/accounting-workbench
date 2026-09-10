# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-09-10 | Author: Codex | Version: V2.555
# Description: 宝贝标题仅在导入内存分类，保留U先/混合业务技术键，不持久化标题。
# Description: 解析天猫订单、宝贝、聚合资金和支付宝分片，生成汇总与逐单核算字段。
"""天猫平台导出导入：按业务白名单生成期间汇总与逐单核算字段，不接触个人信息。"""
import datetime
import decimal
import hashlib
import io
import re
import zipfile
from collections import Counter, defaultdict

from openpyxl import load_workbook


MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DATA_ROWS = 250_000
MAX_ALIPAY_FILES = 20
MAX_ALIPAY_FILE_BYTES = 30 * 1024 * 1024
MAX_ALIPAY_TOTAL_BYTES = 200 * 1024 * 1024
MAX_ALIPAY_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
MAX_ALIPAY_TOTAL_UNCOMPRESSED_BYTES = 800 * 1024 * 1024
MAX_ALIPAY_ROWS = 500_000
_PERIOD_RE = re.compile(r"^\d{4}-\d{2}$")
_MONEY = decimal.Decimal("0.01")

_ORDER_FIELDS = {
    "order_no": "订单编号",
    "payable": "买家应付货款",
    "shipping": "买家应付邮费",
    "list_total": "总金额",
    "paid": "买家实付金额",
    "payment_detail": "支付详情",
    "status": "订单状态",
    "created": "订单创建时间",
    "paid_at": "订单付款时间",
    "quantity": "宝贝总数量",
    "refund": "退款金额",
    "shipped": "发货时间",
    "confirmed": "确认收货时间",
    "confirmed_payout": "确认收货打款金额",
    "merchant_sku": "商家编码",
}
_ITEM_FIELDS = {
    "suborder_no": "子订单编号",
    "order_no": "主订单编号",
    "price": "商品价格",
    "quantity": "购买数量",
    "status": "订单状态",
    "merchant_sku": "商家编码",
    "payable": "买家应付货款",
    "paid": "买家实付金额",
    "refund_status": "退款状态",
    "refund": "退款金额",
    "created": "订单创建时间",
    "paid_at": "订单付款时间",
    "product_id": "商品ID",
}
_FUND_FIELDS = {
    "posted_at": "入账时间",
    "serial_no": "支付流水号",
    "order_no": "淘宝订单编号",
    "entry_type": "入账类型",
    "income": "收入金额（元）",
    "expense": "支出金额",
    "description": "业务描述",
}


class TmallImportError(ValueError):
    """可安全返回给前端的结构校验错误；消息中不得拼接业务单元格值。"""


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip().lstrip("'").strip()


def _number(value):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, decimal.Decimal)):
        return decimal.Decimal(str(value))
    text = str(value).strip().replace(",", "").replace("￥", "").replace("¥", "")
    if not text:
        return None
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return decimal.Decimal(text)
    except decimal.InvalidOperation:
        return None


def _money(value):
    number = _number(value)
    return number.quantize(_MONEY) if number is not None else None


def _date_text(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    text = str(value or "").strip()
    match = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})(?:\s+(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?)?", text)
    if not match:
        return ""
    year, month, day, hour, minute, second = match.groups()
    try:
        value = datetime.datetime(int(year), int(month), int(day), int(hour or 0), int(minute or 0), int(second or 0))
        return value.strftime("%Y-%m-%d %H:%M:%S")
    except ValueError:
        return ""


def _order_key(value):
    """稳定的不可逆技术键，用于跨资金来源勾稽。"""
    return hashlib.sha256(("tmall-order:" + _text(value)).encode("utf-8")).hexdigest()


def _payment_method(value):
    match = re.search(r"支付方式[：:]\s*([^|｜，,；;]+)", _text(value))
    return match.group(1).strip() if match else "未识别"


def _preflight_xlsx(data):
    if not data or len(data) > MAX_FILE_BYTES:
        raise TmallImportError("文件为空或超过 25 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 2_000 or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
                raise TmallImportError("工作簿解压后体积过大")
            if "xl/workbook.xml" not in {item.filename for item in members}:
                raise TmallImportError("文件不是有效的 XLSX 工作簿")
    except TmallImportError:
        raise
    except (zipfile.BadZipFile, OSError):
        raise TmallImportError("文件不是有效的 XLSX 工作簿")


def _open_sheet(data, fields):
    _preflight_xlsx(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise TmallImportError("工作簿无法读取，请确认文件未损坏")
    sheet = workbook[workbook.sheetnames[0]]
    if sheet.max_row == 1 and sheet.max_column == 1:
        sheet.reset_dimensions()
    first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None) or ()
    headers = [_text(value) for value in first]
    missing = [label for label in fields.values() if label not in headers]
    if missing:
        workbook.close()
        raise TmallImportError("缺少必需业务列：" + "、".join(missing))
    index = {key: headers.index(label) for key, label in fields.items()}
    if '商品标题' in headers:
        index['title'] = headers.index('商品标题')
    return workbook, sheet, index


def _selected(row, index):
    return {key: row[pos] if pos < len(row) else None for key, pos in index.items()}


def _amount(value):
    return value if value is not None else decimal.Decimal("0")


def parse_order_export(data, period):
    """订单主表，按订单创建时间归属期间。"""
    if not _PERIOD_RE.match(period or ""):
        raise TmallImportError("结算期间格式应为 YYYY-MM")
    workbook, sheet, index = _open_sheet(data, _ORDER_FIELDS)
    try:
        source_rows = in_period_rows = outside_period_rows = missing_created_rows = 0
        missing_order_number_rows = duplicate_order_rows = 0
        payment_time_missing = payment_in_period = payment_after_period = payment_before_period = 0
        order_ids = set()
        in_period_order_ids = set()
        statuses = Counter()
        payment_methods = Counter()
        match_index = {}
        detail_rows = []
        skus = set()
        missing_sku_rows = 0
        created_min = created_max = ""
        totals = {key: decimal.Decimal("0") for key in (
            "payable", "shipping", "list_total", "paid", "refund", "confirmed_payout", "quantity"
        )}

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_number > MAX_DATA_ROWS + 1:
                raise TmallImportError("数据行超过 250,000 行上限")
            item = _selected(row, index)
            if not any(item[key] not in (None, "") for key in ("order_no", "created", "paid")):
                continue
            source_rows += 1
            order_no = _text(item["order_no"])
            if not order_no:
                missing_order_number_rows += 1
            elif order_no in order_ids:
                duplicate_order_rows += 1
            else:
                order_ids.add(order_no)

            created = _date_text(item["created"])
            if not created:
                missing_created_rows += 1
                continue
            if created[:7] != period:
                outside_period_rows += 1
                continue
            in_period_rows += 1
            if order_no:
                in_period_order_ids.add(order_no)
            created_min = min(created_min, created) if created_min else created
            created_max = max(created_max, created)
            statuses[_text(item["status"]) or "空"] += 1
            payment_method = _payment_method(item["payment_detail"])
            payment_methods[payment_method] += 1
            if order_no:
                order_key = _order_key(order_no)
                match_index[order_key] = payment_method
                detail_rows.append({
                    "order_key": order_key,
                    "order_no": order_no,
                    "created_at": created,
                    "paid_at": _date_text(item["paid_at"]),
                    "shipped_at": _date_text(item["shipped"]),
                    "confirmed_at": _date_text(item["confirmed"]),
                    "status": _text(item["status"]) or "空",
                    "payment_method": payment_method,
                    "merchant_sku": _text(item["merchant_sku"]),
                    "quantity": float(_amount(_number(item["quantity"]))),
                    "current_paid": float(_amount(_money(item["paid"]))),
                    "refund": float(_amount(_money(item["refund"]))),
                    "confirmed_payout": float(_amount(_money(item["confirmed_payout"]))),
                })
            sku = _text(item["merchant_sku"])
            if sku:
                skus.add(sku)
            else:
                missing_sku_rows += 1
            for key in totals:
                totals[key] += _amount(_number(item[key]))

            paid_at = _date_text(item["paid_at"])
            if not paid_at:
                payment_time_missing += 1
            elif paid_at[:7] == period:
                payment_in_period += 1
            elif paid_at[:7] > period:
                payment_after_period += 1
            else:
                payment_before_period += 1

        if not source_rows:
            raise TmallImportError("工作簿没有可识别的订单数据")
        blocking_codes = []
        warning_codes = []
        if not in_period_rows:
            blocking_codes.append("NO_ROWS_IN_PERIOD")
        if missing_order_number_rows:
            blocking_codes.append("MISSING_ORDER_NUMBER")
        if duplicate_order_rows:
            blocking_codes.append("DUPLICATE_ORDER_NUMBER")
        if outside_period_rows:
            warning_codes.append("OUTSIDE_PERIOD_ROWS_EXCLUDED")
        return {
            "kind": "order",
            "period": period,
            "period_basis": "订单创建时间",
            "source_rows": source_rows,
            "in_period_rows": in_period_rows,
            "outside_period_rows": outside_period_rows,
            "missing_created_rows": missing_created_rows,
            "order_count": len(in_period_order_ids),
            "missing_order_number_rows": missing_order_number_rows,
            "duplicate_order_rows": duplicate_order_rows,
            "status_counts": dict(statuses),
            "payment_method_counts": dict(payment_methods),
            "payment_in_period_orders": payment_in_period,
            "payment_after_period_orders": payment_after_period,
            "payment_before_period_orders": payment_before_period,
            "payment_time_missing_orders": payment_time_missing,
            "merchant_sku_count": len(skus),
            "missing_merchant_sku_rows": missing_sku_rows,
            "goods_quantity": float(totals["quantity"]),
            "buyer_payable_total": float(totals["payable"].quantize(_MONEY)),
            "shipping_total": float(totals["shipping"].quantize(_MONEY)),
            "list_total": float(totals["list_total"].quantize(_MONEY)),
            "current_paid_total": float(totals["paid"].quantize(_MONEY)),
            "refund_total": float(totals["refund"].quantize(_MONEY)),
            "confirmed_payout_total": float(totals["confirmed_payout"].quantize(_MONEY)),
            "created_at_min": created_min,
            "created_at_max": created_max,
            "source_status": "blocked" if blocking_codes else ("warning" if warning_codes else "ready"),
            "blocking_codes": blocking_codes,
            "warning_codes": warning_codes,
            "_match_index": match_index,
            "_detail_rows": detail_rows,
        }
    finally:
        workbook.close()


def parse_item_export(data, period):
    """宝贝销售明细，按订单创建时间归属期间；金额是子订单成交态。"""
    if not _PERIOD_RE.match(period or ""):
        raise TmallImportError("结算期间格式应为 YYYY-MM")
    workbook, sheet, index = _open_sheet(data, _ITEM_FIELDS)
    try:
        source_rows = in_period_rows = outside_period_rows = missing_created_rows = 0
        missing_main_order_rows = missing_suborder_rows = duplicate_suborder_rows = 0
        payment_time_missing_rows = payment_in_period_rows = payment_after_period_rows = 0
        main_orders = Counter()
        suborders = set()
        in_period_suborders = set()
        match_index = set()
        business_kinds = defaultdict(set)
        statuses = Counter()
        refund_statuses = Counter()
        skus, product_ids = set(), set()
        missing_sku_rows = 0
        totals = {key: decimal.Decimal("0") for key in ("payable", "paid", "refund", "quantity")}
        created_min = created_max = ""

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_number > MAX_DATA_ROWS + 1:
                raise TmallImportError("数据行超过 250,000 行上限")
            item = _selected(row, index)
            if not any(item[key] not in (None, "") for key in ("suborder_no", "order_no", "created")):
                continue
            source_rows += 1
            main_order = _text(item["order_no"])
            suborder = _text(item["suborder_no"])
            if not main_order:
                missing_main_order_rows += 1
            if not suborder:
                missing_suborder_rows += 1
            elif suborder in suborders:
                duplicate_suborder_rows += 1
            else:
                suborders.add(suborder)

            created = _date_text(item["created"])
            if not created:
                missing_created_rows += 1
                continue
            if created[:7] != period:
                outside_period_rows += 1
                continue
            in_period_rows += 1
            if main_order:
                main_orders[main_order] += 1
                match_index.add(_order_key(main_order))
                from kernels.ec_month_fulfillment import title_kind
                business_kinds[_order_key(main_order)].add(title_kind(item.get('title')))
            if suborder:
                in_period_suborders.add(suborder)
            created_min = min(created_min, created) if created_min else created
            created_max = max(created_max, created)
            statuses[_text(item["status"]) or "空"] += 1
            refund_statuses[_text(item["refund_status"]) or "空"] += 1
            sku = _text(item["merchant_sku"])
            product_id = _text(item["product_id"])
            if sku:
                skus.add(sku)
            else:
                missing_sku_rows += 1
            if product_id:
                product_ids.add(product_id)
            for key in totals:
                totals[key] += _amount(_number(item[key]))
            paid_at = _date_text(item["paid_at"])
            if not paid_at:
                payment_time_missing_rows += 1
            elif paid_at[:7] == period:
                payment_in_period_rows += 1
            elif paid_at[:7] > period:
                payment_after_period_rows += 1

        if not source_rows:
            raise TmallImportError("工作簿没有可识别的宝贝销售明细")
        blocking_codes = []
        warning_codes = []
        if not in_period_rows:
            blocking_codes.append("NO_ROWS_IN_PERIOD")
        if missing_main_order_rows or missing_suborder_rows:
            blocking_codes.append("MISSING_ORDER_NUMBER")
        if duplicate_suborder_rows:
            blocking_codes.append("DUPLICATE_SUBORDER_NUMBER")
        if outside_period_rows:
            warning_codes.append("OUTSIDE_PERIOD_ROWS_EXCLUDED")
        if missing_sku_rows:
            warning_codes.append("MISSING_MERCHANT_SKU")
        return {
            "kind": "item",
            "period": period,
            "period_basis": "订单创建时间",
            "source_rows": source_rows,
            "in_period_rows": in_period_rows,
            "outside_period_rows": outside_period_rows,
            "missing_created_rows": missing_created_rows,
            "main_order_count": len(main_orders),
            "suborder_count": len(in_period_suborders),
            "multi_line_main_order_count": sum(1 for count in main_orders.values() if count > 1),
            "max_lines_per_main_order": max(main_orders.values(), default=0),
            "missing_main_order_rows": missing_main_order_rows,
            "missing_suborder_rows": missing_suborder_rows,
            "duplicate_suborder_rows": duplicate_suborder_rows,
            "status_counts": dict(statuses),
            "refund_status_counts": dict(refund_statuses),
            "payment_in_period_rows": payment_in_period_rows,
            "payment_after_period_rows": payment_after_period_rows,
            "payment_time_missing_rows": payment_time_missing_rows,
            "merchant_sku_count": len(skus),
            "missing_merchant_sku_rows": missing_sku_rows,
            "product_count": len(product_ids),
            "goods_quantity": float(totals["quantity"]),
            "buyer_payable_total": float(totals["payable"].quantize(_MONEY)),
            "gross_paid_total": float(totals["paid"].quantize(_MONEY)),
            "refund_total": float(totals["refund"].quantize(_MONEY)),
            "net_paid_after_refund": float((totals["paid"] - totals["refund"]).quantize(_MONEY)),
            "created_at_min": created_min,
            "created_at_max": created_max,
            "source_status": "blocked" if blocking_codes else ("warning" if warning_codes else "ready"),
            "blocking_codes": blocking_codes,
            "warning_codes": warning_codes,
            "_match_index": sorted(match_index),
            "_business_index": {key: ('mixed' if 'ufirst' in kinds and 'normal' in kinds
                                      else 'unknown' if 'unknown' in kinds
                                      else 'review' if 'review' in kinds
                                      else next(iter(kinds))) for key, kinds in business_kinds.items()},
        }
    finally:
        workbook.close()


def parse_fund_export(data, period):
    """聚合结算账户余额明细，按入账时间归属期间。"""
    if not _PERIOD_RE.match(period or ""):
        raise TmallImportError("结算期间格式应为 YYYY-MM")
    workbook, sheet, index = _open_sheet(data, _FUND_FIELDS)
    try:
        source_rows = in_period_rows = outside_period_rows = missing_posted_at_rows = 0
        missing_serial_rows = duplicate_serial_rows = 0
        serials, in_period_serials, order_ids = set(), set(), set()
        match_index = {}
        order_link_rows = 0
        types = {}
        descriptions = {}
        income_total = expense_total = decimal.Decimal("0")
        posted_min = posted_max = ""

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_number > MAX_DATA_ROWS + 1:
                raise TmallImportError("数据行超过 250,000 行上限")
            item = _selected(row, index)
            if not any(item[key] not in (None, "") for key in ("posted_at", "serial_no", "income", "expense")):
                continue
            source_rows += 1
            serial = _text(item["serial_no"])
            if not serial:
                missing_serial_rows += 1
            elif serial in serials:
                duplicate_serial_rows += 1
            else:
                serials.add(serial)
            posted_at = _date_text(item["posted_at"])
            if not posted_at:
                missing_posted_at_rows += 1
                continue
            if posted_at[:7] != period:
                outside_period_rows += 1
                continue
            in_period_rows += 1
            if serial:
                in_period_serials.add(serial)
            posted_min = min(posted_min, posted_at) if posted_min else posted_at
            posted_max = max(posted_max, posted_at)
            order_no = _text(item["order_no"])
            if order_no:
                order_link_rows += 1
                order_ids.add(order_no)
            income = _amount(_money(item["income"]))
            expense = _amount(_money(item["expense"]))
            if order_no:
                key = _order_key(order_no)
                matched = match_index.setdefault(key, {
                    "income": decimal.Decimal("0"), "expense": decimal.Decimal("0"),
                    "receipt_income": decimal.Decimal("0"), "trade_receipt_rows": 0,
                    "receipt_at_min": "", "receipt_at_max": "",
                })
                matched["income"] += income
                matched["expense"] += expense
                if _text(item["entry_type"]) == "交易收款":
                    matched["trade_receipt_rows"] += 1
                    matched["receipt_income"] += income
                    matched["receipt_at_min"] = min(matched["receipt_at_min"], posted_at) if matched["receipt_at_min"] else posted_at
                    matched["receipt_at_max"] = max(matched["receipt_at_max"], posted_at)
            income_total += income
            expense_total += expense
            entry_type = _text(item["entry_type"]) or "空"
            description = _text(item["description"])
            type_item = types.setdefault(entry_type, {"rows": 0, "income": decimal.Decimal("0"), "expense": decimal.Decimal("0")})
            type_item["rows"] += 1
            type_item["income"] += income
            type_item["expense"] += expense
            if description and "|" in description:
                code, _, label = description.partition("|")
                key = code.strip()[:30]
                desc_item = descriptions.setdefault(key, {"label": label.strip()[:80], "rows": 0, "expense": decimal.Decimal("0")})
                desc_item["rows"] += 1
                desc_item["expense"] += expense

        if not source_rows:
            raise TmallImportError("工作簿没有可识别的聚合结算账户明细")
        blocking_codes = []
        warning_codes = []
        if not in_period_rows:
            blocking_codes.append("NO_ROWS_IN_PERIOD")
        if missing_serial_rows:
            blocking_codes.append("MISSING_PAYMENT_SERIAL")
        if duplicate_serial_rows:
            blocking_codes.append("DUPLICATE_PAYMENT_SERIAL")
        if outside_period_rows:
            warning_codes.append("OUTSIDE_PERIOD_ROWS_EXCLUDED")
        return {
            "kind": "fund",
            "period": period,
            "period_basis": "入账时间",
            "account_scope": "聚合结算账户",
            "source_rows": source_rows,
            "in_period_rows": in_period_rows,
            "outside_period_rows": outside_period_rows,
            "missing_posted_at_rows": missing_posted_at_rows,
            "payment_serial_count": len(in_period_serials),
            "missing_payment_serial_rows": missing_serial_rows,
            "duplicate_payment_serial_rows": duplicate_serial_rows,
            "order_link_rows": order_link_rows,
            "linked_order_count": len(order_ids),
            "income_total": float(income_total.quantize(_MONEY)),
            "expense_total": float(expense_total.quantize(_MONEY)),
            "net_change": float((income_total - expense_total).quantize(_MONEY)),
            "entry_types": {
                key: {"rows": value["rows"], "income": float(value["income"].quantize(_MONEY)),
                      "expense": float(value["expense"].quantize(_MONEY))}
                for key, value in types.items()
            },
            "fee_codes": {
                key: {"label": value["label"], "rows": value["rows"],
                      "expense": float(value["expense"].quantize(_MONEY))}
                for key, value in descriptions.items()
            },
            "posted_at_min": posted_min,
            "posted_at_max": posted_max,
            "source_status": "blocked" if blocking_codes else ("warning" if warning_codes else "ready"),
            "blocking_codes": blocking_codes,
            "warning_codes": warning_codes,
            "_match_index": {
                key: {"income": float(value["income"].quantize(_MONEY)),
                      "expense": float(value["expense"].quantize(_MONEY)),
                      "receipt_income": float(value["receipt_income"].quantize(_MONEY)),
                      "trade_receipt_rows": value["trade_receipt_rows"],
                      "receipt_at_min": value["receipt_at_min"],
                      "receipt_at_max": value["receipt_at_max"]}
                for key, value in match_index.items()
            },
        }
    finally:
        workbook.close()


def _preflight_alipay_file(data):
    if not data or len(data) > MAX_ALIPAY_FILE_BYTES:
        raise TmallImportError("支付宝流水文件为空或超过 30 MB")
    if data[:2] != b"PK":
        return len(data)
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            unpacked = sum(item.file_size for item in members)
            if not members or len(members) > 2_000 or unpacked > MAX_ALIPAY_UNCOMPRESSED_BYTES:
                raise TmallImportError("支付宝流水解压后体积过大")
            if any(item.flag_bits & 0x1 for item in members):
                raise TmallImportError("不支持加密的支付宝流水包")
            return unpacked
    except TmallImportError:
        raise
    except (zipfile.BadZipFile, OSError):
        raise TmallImportError("支付宝流水压缩包无法读取")


def _alipay_row_key(row):
    serial = _text(row.get("serial"))
    if serial:
        return "serial:" + serial
    values = [_text(row.get(key)) for key in (
        "ts", "txn", "mch_no", "btype", "income", "outgo", "chan", "order_no", "desc"
    )]
    return "row:" + hashlib.sha256("\x1f".join(values).encode("utf-8")).hexdigest()


def _fee_code(value):
    raw = _text(value)
    code = raw.partition("|")[0].strip() if "|" in raw else ""
    return code if re.match(r"^[0-9A-Za-z]{1,20}$", code) else "NO_CODE"


def parse_alipay_exports(files, period):
    """多个支付宝 2088 分片原子导入：按流水号跨包去重，只保留汇总和不可逆订单键。"""
    if not _PERIOD_RE.match(period or ""):
        raise TmallImportError("结算期间格式应为 YYYY-MM")
    if not files or len(files) > MAX_ALIPAY_FILES:
        raise TmallImportError("支付宝流水文件数应在 1–20 个之间")
    if sum(len(data) for data in files) > MAX_ALIPAY_TOTAL_BYTES:
        raise TmallImportError("支付宝流水文件总体积超过 200 MB")

    unpacked_total = 0
    source_rows = 0
    rows_by_key = {}
    key_occurrences = Counter()
    segment_fingerprints = Counter()
    file_summaries = []
    from kernels import ec_settle

    for file_number, data in enumerate(files, start=1):
        unpacked_total += _preflight_alipay_file(data)
        if unpacked_total > MAX_ALIPAY_TOTAL_UNCOMPRESSED_BYTES:
            raise TmallImportError("支付宝流水解压后总体积超过上限")
        try:
            rows = ec_settle.parse_flow_any(data)
        except Exception:
            raise TmallImportError("第 %s 个文件不是可识别的支付宝账务组合流水" % file_number)
        if not rows:
            raise TmallImportError("第 %s 个文件没有可识别的支付宝流水数据" % file_number)
        source_rows += len(rows)
        if source_rows > MAX_ALIPAY_ROWS:
            raise TmallImportError("支付宝流水数据行超过 500,000 行上限")
        file_keys = []
        file_dates = []
        for row in rows:
            key = _alipay_row_key(row)
            file_keys.append(key)
            key_occurrences[key] += 1
            rows_by_key.setdefault(key, row)
            posted_at = _date_text(row.get("ts"))
            if posted_at:
                file_dates.append(posted_at)
        segment_digest = hashlib.sha256("\n".join(sorted(file_keys)).encode("utf-8")).hexdigest()
        segment_fingerprints[segment_digest] += 1
        file_summaries.append({
            "rows": len(rows),
            "posted_at_min": min(file_dates) if file_dates else "",
            "posted_at_max": max(file_dates) if file_dates else "",
        })

    unique_rows = list(rows_by_key.values())
    in_period_rows = outside_period_rows = missing_posted_at_rows = missing_serial_rows = 0
    income_total = expense_total = decimal.Decimal("0")
    receipt_income_total = refund_expense_total = decimal.Decimal("0")
    linked_order_rows = 0
    linked_orders = set()
    receipt_orders = set()
    match_index = {}
    fee_codes = defaultdict(lambda: {"rows": 0, "income": decimal.Decimal("0"), "expense": decimal.Decimal("0")})
    entry_types = defaultdict(lambda: {"rows": 0, "income": decimal.Decimal("0"), "expense": decimal.Decimal("0")})
    channel_zones = defaultdict(lambda: {"rows": 0, "income": decimal.Decimal("0"), "expense": decimal.Decimal("0")})
    posted_dates = []
    covered_days = set()
    allowed_entry_types = {"在线支付", "分账", "转账", "保证金", "退款（交易退款）", "其它"}

    for row in unique_rows:
        if not _text(row.get("serial")):
            missing_serial_rows += 1
        posted_at = _date_text(row.get("ts"))
        if not posted_at:
            missing_posted_at_rows += 1
            continue
        if posted_at[:7] != period:
            outside_period_rows += 1
            continue
        in_period_rows += 1
        posted_dates.append(posted_at)
        covered_days.add(posted_at[:10])
        income = _amount(_money(row.get("income")))
        expense = _amount(_money(row.get("outgo")))
        income_total += income
        expense_total += expense
        code = _fee_code(row.get("desc"))
        code_item = fee_codes[code]
        code_item["rows"] += 1
        code_item["income"] += income
        code_item["expense"] += expense
        entry_type = _text(row.get("btype"))
        entry_type = entry_type if entry_type in allowed_entry_types else "其他未识别"
        type_item = entry_types[entry_type]
        type_item["rows"] += 1
        type_item["income"] += income
        type_item["expense"] += expense
        zone = "aggregate_channel" if _text(row.get("chan")) == "聚合结算渠道" else "direct_alipay"
        zone_item = channel_zones[zone]
        zone_item["rows"] += 1
        zone_item["income"] += income
        zone_item["expense"] += expense
        if code == "0010001" and income > 0:
            receipt_income_total += income
        if code == "0020001" and expense > 0:
            refund_expense_total += expense
        order_no = _text(row.get("order_no"))
        if not order_no:
            continue
        linked_order_rows += 1
        order_key = _order_key(order_no)
        linked_orders.add(order_key)
        matched = match_index.setdefault(order_key, {
            "income": decimal.Decimal("0"), "expense": decimal.Decimal("0"),
            "receipt_income": decimal.Decimal("0"), "refund_expense": decimal.Decimal("0"),
            "receipt_rows": 0, "linked_rows": 0, "receipt_at_min": "", "receipt_at_max": "",
        })
        matched["income"] += income
        matched["expense"] += expense
        matched["linked_rows"] += 1
        if code == "0010001" and income > 0:
            matched["receipt_income"] += income
            matched["receipt_rows"] += 1
            matched["receipt_at_min"] = min(matched["receipt_at_min"], posted_at) if matched["receipt_at_min"] else posted_at
            matched["receipt_at_max"] = max(matched["receipt_at_max"], posted_at)
            receipt_orders.add(order_key)
        if code == "0020001" and expense > 0:
            matched["refund_expense"] += expense

    if not in_period_rows:
        raise TmallImportError("支付宝流水在所选期间内没有可识别数据")

    def summarise(data):
        return {
            key: {
                "rows": value["rows"],
                "income": float(value["income"].quantize(_MONEY)),
                "expense": float(value["expense"].quantize(_MONEY)),
            }
            for key, value in data.items()
        }

    warning_codes = []
    duplicate_occurrences = sum(count - 1 for count in key_occurrences.values())
    repeated_file_count = sum(count - 1 for count in segment_fingerprints.values())
    if duplicate_occurrences:
        warning_codes.append("DUPLICATE_ROWS_DEDUPED")
    if repeated_file_count:
        warning_codes.append("REPEATED_SEGMENT_DEDUPED")
    if missing_serial_rows:
        warning_codes.append("MISSING_SERIAL_FALLBACK_HASH")
    if outside_period_rows:
        warning_codes.append("OUTSIDE_PERIOD_ROWS_EXCLUDED")
    return {
        "kind": "alipay",
        "period": period,
        "period_basis": "入账时间",
        "account_scope": "支付宝 2088 账户",
        "file_count": len(files),
        "effective_segment_count": len(segment_fingerprints),
        "fully_repeated_file_count": repeated_file_count,
        "file_summaries": file_summaries,
        "source_rows": source_rows,
        "unique_rows": len(unique_rows),
        "in_period_rows": in_period_rows,
        "outside_period_rows": outside_period_rows,
        "duplicate_row_occurrences": duplicate_occurrences,
        "missing_posted_at_rows": missing_posted_at_rows,
        "missing_serial_rows": missing_serial_rows,
        "covered_day_count": len(covered_days),
        "posted_at_min": min(posted_dates) if posted_dates else "",
        "posted_at_max": max(posted_dates) if posted_dates else "",
        "linked_order_rows": linked_order_rows,
        "linked_order_count": len(linked_orders),
        "receipt_order_count": len(receipt_orders),
        "income_total": float(income_total.quantize(_MONEY)),
        "expense_total": float(expense_total.quantize(_MONEY)),
        "net_change": float((income_total - expense_total).quantize(_MONEY)),
        "platform_receipt_income_total": float(receipt_income_total.quantize(_MONEY)),
        "trade_refund_expense_total": float(refund_expense_total.quantize(_MONEY)),
        "entry_types": summarise(entry_types),
        "fee_codes": summarise(fee_codes),
        "channel_zones": summarise(channel_zones),
        "source_status": "warning" if warning_codes else "ready",
        "blocking_codes": [],
        "warning_codes": warning_codes,
        "_match_index": {
            key: {
                "income": float(value["income"].quantize(_MONEY)),
                "expense": float(value["expense"].quantize(_MONEY)),
                "receipt_income": float(value["receipt_income"].quantize(_MONEY)),
                "refund_expense": float(value["refund_expense"].quantize(_MONEY)),
                "receipt_rows": value["receipt_rows"],
                "linked_rows": value["linked_rows"],
                "receipt_at_min": value["receipt_at_min"],
                "receipt_at_max": value["receipt_at_max"],
            }
            for key, value in match_index.items()
        },
    }


PARSERS = {"order": parse_order_export, "item": parse_item_export, "fund": parse_fund_export}
