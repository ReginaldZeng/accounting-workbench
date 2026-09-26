# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-09-26 | Author: Claude Opus 4.8 | Version: V2.632
# Description: 【物流账单复核】通用解析器——按「取数说明」(intake_spec) 认列，不写死序号（同一家导出月间列会漂移，写死必错）。
#   一张 sheet 按 spec 的角色解析：detail=对账逐单(带单号)、accrual=计提口径(月结按费用项)、ignore=价目表跳过。
#   出中间表行(dict)。表名按前缀/正则匹配月度变动（如「*发货明细」）。落库/核价核量在 router 串起来。
import re
import json


def _s(v):
    return "" if v is None else str(v).strip()


def _f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def match_sheet(spec_name, real_name):
    """spec 里的 name 支持精确、前缀 * 通配、和 | 分隔的正则关键词。"""
    if spec_name == real_name:
        return True
    if "*" in spec_name or "|" in spec_name:
        pat = spec_name.replace("*", ".*")
        for part in pat.split("|"):
            if re.search(part, real_name):
                return True
    return False


def find_col(hdr, name_or_list):
    """按列名（或别名列表）在表头找列号，取首个命中。认名不认序号。"""
    names = name_or_list if isinstance(name_or_list, list) else [name_or_list]
    for n in names:
        for i, h in enumerate(hdr):
            if h == n:
                return i
    return None


def split_nos(no):
    """一格多单号 A+B / A/B / A，B → [A,B]。"""
    parts = re.split(r"[+/,，、;；\s]+", no or "")
    return [p for p in parts if p]


def _subject(spec_sub, row, hdr):
    if not spec_sub:
        return ""
    if "fixed" in spec_sub:
        return spec_sub["fixed"]
    if "column" in spec_sub:
        ci = find_col(hdr, spec_sub["column"])
        return _s(row[ci]) if ci is not None and ci < len(row) else ""
    return ""


def parse_detail_sheet(sp, ws, period, carrier):
    """detail 角色 sheet → 逐单据行。按 spec 的 doc_col/amount_cols/qty_col/wt_col/prov_col/carrier_sub_col 认列。"""
    rows = list(ws.iter_rows(values_only=True))
    hr = int(sp.get("header_row", 1)) - 1
    if hr >= len(rows):
        return []
    hdr = [_s(x) for x in rows[hr]]
    c_doc = find_col(hdr, sp["doc_col"]) if sp.get("doc_col") else None
    c_amts = [find_col(hdr, n) for n in sp["amount_cols"]] if sp.get("amount_cols") else \
             ([find_col(hdr, sp["amount_col"])] if sp.get("amount_col") else [])
    c_amts = [c for c in c_amts if c is not None]
    c_qty = find_col(hdr, sp["qty_col"]) if sp.get("qty_col") else None
    c_wt = find_col(hdr, sp["wt_col"]) if sp.get("wt_col") else None
    c_prov = find_col(hdr, sp["prov_col"]) if sp.get("prov_col") else None
    c_cs = find_col(hdr, sp["carrier_sub_col"]) if sp.get("carrier_sub_col") else None
    marker = sp.get("summary_marker")
    out = []
    for ri, r in enumerate(rows[hr + 1:], start=hr + 2):
        if not any(x is not None for x in r):
            continue
        if marker and _s(r[0]) and marker in _s(r[0]):
            continue
        doc = _s(r[c_doc]) if c_doc is not None and c_doc < len(r) else _s(sp.get("doc", ""))
        if c_doc is not None and not doc:
            continue
        base = sum((_f(r[c]) or 0) for c in c_amts) if c_amts else None
        row = {
            "period": period, "carrier": carrier, "grain": "detail",
            "subject": _subject(sp.get("subject"), r, hdr),
            "doc_no": "+".join(split_nos(doc)) if doc and doc != "无单据" else doc,
            "annot": sp.get("annot", ""), "fee_item": sp.get("fee_item", ""),
            "qty": _f(r[c_qty]) if c_qty is not None and c_qty < len(r) else None,
            "unit": sp.get("qty_unit", ""),
            "amount": round(base, 2) if base is not None else None,
            "base_amount": round(_f(r[c_amts[0]]) or 0, 2) if c_amts else None,
            "carrier_sub": _s(r[c_cs]) if c_cs is not None and c_cs < len(r) else "",
            "prov": _s(r[c_prov]) if c_prov is not None and c_prov < len(r) else "",
            "charge_wt": _f(r[c_wt]) if c_wt is not None and c_wt < len(r) else None,
            "src_sheet": ws.title, "src_row": ri,
        }
        out.append(row)
    return out


def parse_accrual_sheet(sp, ws, period, carrier):
    """accrual 角色 sheet（月结清单，半结构）→ 按费用项一行。定位费用项标签行，取数值单元格：金额=末值、数量=最大值。"""
    rows = list(ws.iter_rows(values_only=True))
    n = len(rows)
    fee_map = sp.get("fee_map", {})
    subj = sp.get("subject", {}).get("fixed", "")
    out = []
    for ri in range(n):
        cells = [_s(c) for c in rows[ri]]
        label = next((c for c in cells if c in fee_map), None)
        if not label:
            continue
        nums = [_f(c) for c in rows[ri] if _f(c) is not None]
        if len(nums) >= 2:
            # 行内费用（名字与金额同一行）：金额=末值，数量=其余最大值
            amt, qty = nums[-1], max(nums[:-1], default=None)
        else:
            # 段头费用（名字在段头，金额在本段「合计」行）：往下找到合计取数，遇到下一段费用名即止
            amt = qty = None
            for j in range(ri + 1, min(ri + 40, n)):
                jc = [_s(c) for c in rows[j]]
                if any(c in fee_map for c in jc):
                    break
                if jc and any("合计" in c for c in jc[:2]):
                    jn = [_f(c) for c in rows[j] if _f(c) is not None]
                    if jn:
                        amt, qty = jn[-1], (max(jn[:-1], default=None) if len(jn) > 1 else None)
                    break
        out.append({"period": period, "carrier": carrier, "grain": "accrual", "subject": subj,
                    "doc_no": "", "annot": fee_map[label], "fee_item": label,
                    "qty": qty, "unit": "", "amount": amt, "src_sheet": ws.title, "src_row": ri + 1})
    return out


def parse_bill(spec, data):
    """按取数说明解析整本账单 xlsx → {'detail':[...], 'accrual':[...], 'skipped':[表名], 'period':...}。
    data=bytes 或路径。period 从 spec 传入或调用方补。"""
    import openpyxl
    from io import BytesIO
    src = BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    carrier = spec.get("carrier", "")
    period = spec.get("period", "")
    detail, accrual, skipped = [], [], []
    for ws in wb.worksheets:
        sp = None
        for cand in spec.get("sheets", []):
            if match_sheet(cand["name"], ws.title):
                sp = cand
                break
        if sp is None or sp.get("role") == "ignore":
            skipped.append(ws.title)
            continue
        if sp["role"] == "detail":
            detail += parse_detail_sheet(sp, ws, period, carrier)
        elif sp["role"] == "accrual":
            accrual += parse_accrual_sheet(sp, ws, period, carrier)
    return {"detail": detail, "accrual": accrual, "skipped": skipped, "carrier": carrier, "period": period}
