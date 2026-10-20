"""核对表导出：只消费已核结果，不修改计薪口径或写金蝶。"""
import calendar
import json
from collections import defaultdict
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

HEADERS = ['姓名', '手机号', '归属方', '日期', '周', '班次', '上班卡', '无效卡', '下班卡',
           '计算工时（未扣减）', '扣减', '重算计薪工时', '上报工时', '差异', '计薪单价',
           '管理费', '应发薪资', '应发管理费', '合计', '核对结果', '人力备注']

def person_id(p):
    return json.dumps([str(p.get(k) or '') for k in ('姓名', '归属', '部门', '岗位')], ensure_ascii=False, separators=(',', ':'))

def money(v):
    return float(Decimal(str(v)).quantize(Decimal('.01'), rounding=ROUND_HALF_UP))

def select_people(res, mode='all', agencies=None, people=None):
    if mode not in ('all', 'agency', 'person'):
        raise ValueError('导出范围无效')
    keys = set(agencies or []) if mode == 'agency' else set(people or [])
    if mode != 'all' and not keys:
        raise ValueError('请至少选择一个服务方或人员')
    result = [p for p in res.get('people', []) if mode == 'all' or
              (p.get('归属') in keys if mode == 'agency' else person_id(p) in keys)]
    if not result:
        raise ValueError('所选范围没有人员，请重新选择')
    return sorted(result, key=lambda p: (p.get('归属', ''), p.get('姓名', ''), person_id(p)))

def build_data(res, month, mode='all', agencies=None, people=None, punch=None):
    from kernels import temp_attendance as ta
    year, mon = map(int, month.split('-'))
    first = date(year, mon, 1)
    last = calendar.monthrange(year, mon)[1]
    selected = select_people(res, mode, agencies, people)
    if len({person_id(p) for p in selected}) != len(selected):
        raise ValueError('存在姓名、服务方、部门、岗位均相同的重复结算行，请先核实，避免重复导出金额')
    index = defaultdict(list)
    for r in res.get('rows', []):
        index[person_id(r)].append(r)
    detail, summaries = [], []
    for p in selected:
        key = person_id(p)
        rr = index[key]
        daily = defaultdict(list)
        for r in rr:
            daily[int(r['日'])].append(r)
        tails = {str(r.get('手机尾号')) for r in rr if r.get('手机尾号')}
        rec = None
        if punch:
            candidates = punch.get('by_key', {}).get(ta.norm_name(p['姓名']), [])
            matches = [r for r in candidates if str(r.get('标识') or '') in tails]
            if len(matches) == 1:
                rec = matches[0]
            elif not tails:
                rec, _ = ta.match_punch(punch, p['姓名'])
        phone = p.get('手机号') or (rec or {}).get('手机号') or ''
        if not phone and len(tails) == 1:
            ident = next(iter(tails))
            phone = ident if len(ident) == 11 and ident.isdigit() else '尾号 ' + ident
        note = '' if phone else '手机号未提供；'
        start = len(detail)
        for day in range(0, last + 2):
            dt = first + timedelta(days=day-1)
            base = [p['姓名'], str(phone), p.get('归属', ''), dt, '一二三四五六日'[dt.weekday()]]
            if day in (0, last + 1):
                times = ((rec or {}).get('bnd') or {}).get(ta._BND_PREV if day == 0 else ta._BND_NEXT, [])
                raw = ' '.join(ta.fmt_hm(t) for t in times)
                detail.append((base + ['边界参考', raw, '', ''] + [None]*10 +
                               [('边界原始卡集中列于上班卡列（未切班，仅供参考）' if times else '边界原表未提供或无打卡') + '；不计本月', ''], 'boundary', key))
                continue
            for r in daily.get(day) or [None]:
                if r is None:
                    # 内核会省略休息日及已归前夜的卡；不擅自把它们重判为缺卡。
                    raw = ((rec or {}).get('days') or {}).get(day, [])
                    detail.append((base + ['-', '', ' '.join(ta.fmt_hm(t) for t in raw), ''] +
                                   [None, None, 0, 0, 0, None, None, 0, 0, 0,
                                    note + ('无本日上报；原始卡供参考，跨日卡以归属班日为准' if raw else '无本日上报；无可用打卡' if rec else '无本日上报；原始打卡不可用'), ''], 'neutral', key))
                    continue
                hours = float(r.get('上报工时') or 0)
                shift = r.get('班型')
                if shift == '白夜混合':
                    # 同日班型以已切出的首卡窗口判断；缺卡不猜计价班型。
                    clock = r.get('上班打卡') or ''
                    if r.get('跨度') is not None and ':' in clock:
                        hh, mm = map(int, clock.replace('次日', '').split(':'))
                        shift = '夜班' if hh*60+mm >= (res.get('params') or ta.DEFAULT_PARAMS).get('night_start_from', 1080) else '白班'
                prefix = '夜班' if shift == '夜班' else '白班'
                missing = bool(p.get('合同缺档')) or shift not in ('白班', '夜班')
                wage = None if missing else p.get(prefix+'工资单价')
                mgmt = None if missing else p.get(prefix+'管理费单价')
                a = money(hours*wage) if wage is not None else (0 if hours == 0 else None)
                b = money(hours*mgmt) if mgmt is not None else (0 if hours == 0 else None)
                result = note + (r.get('判定') or '')
                if r.get('跨度') is not None and r.get('规则扣减') is None:
                    result += '；旧留档未保存扣减，请按当前参数重跑补齐'
                if missing and hours:
                    result += '；合同价缺档或班型待核，金额未计算'
                ack = r.get('已认定')
                if ack:
                    result += f"；已认定：{ack.get('认定人', '')} {ack.get('时间', '')} {ack.get('理由', '')}"
                if p.get('部门'):
                    result += '；部门：' + p['部门']
                detail.append((base + [shift, r.get('上班打卡', ''), ' '.join(r.get('无效卡') or []),
                               r.get('下班打卡', ''), r.get('跨度'), r.get('规则扣减'), r.get('重算工时', 0),
                               hours, r.get('差异', 0), wage, mgmt, a, b,
                               money(a+b) if a is not None and b is not None else None, result, ''],
                               'acked' if ack else r.get('档', ''), key))
        own = [v for v, band, _ in detail[start:] if band != 'boundary']
        wage = None if any(v[16] is None for v in own) else money(sum(v[16] for v in own))
        mgmt = None if any(v[17] is None for v in own) else money(sum(v[17] for v in own))
        # 金额按日展示，与既有按月计算可能有分币差；显式列出校验而不改原结算。
        mismatch = abs(sum(v[12] or 0 for v in own)-float(p.get('上报总工时') or 0)) > .001
        check = '工时与原结算不符' if mismatch else '一致'
        for val, old in ((wage, p.get('应付工资')), (mgmt, p.get('应付管理费'))):
            if val is None or old is None:
                check = '金额未完整计算'
            elif abs(val-old) > .011:
                check = '逐日与月结金额不符，请核实班型/分币差'
        summaries.append([p['姓名'], str(phone), p.get('归属', ''), p.get('部门', ''), p.get('岗位', ''),
                          sum(v[12] or 0 for v in own), sum(v[11] or 0 for v in own),
                          sum(v[13] or 0 for v in own), wage, mgmt,
                          money(wage+mgmt) if wage is not None and mgmt is not None else None,
                          p.get('补贴奖罚'), p.get('表上合计'),
                          sum(1 for _, band, k in detail[start:] if band in ('over_out', 'hard', 'ambig', 'thin', 'mixed')),
                          check, ''])
    return detail, summaries

def export_xlsx(res, month, mode='all', agencies=None, people=None, punch=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    detail, people_rows = build_data(res, month, mode, agencies, people, punch)
    wb = Workbook(); wb.remove(wb.active)
    agency_rows = []
    for agency in sorted({r[2] for r in people_rows}):
        group = [r for r in people_rows if r[2] == agency]
        complete = all(r[10] is not None for r in group)
        agency_rows.append([agency, len({(r[0],r[1]) for r in group}), len(group),
                            sum(r[5] for r in group), sum(r[6] for r in group), sum(r[7] for r in group),
                            money(sum(r[8] for r in group)) if complete else None,
                            money(sum(r[9] for r in group)) if complete else None,
                            money(sum(r[10] for r in group)) if complete else None,
                            money(sum(r[11] or 0 for r in group)),
                            sum(r[13] for r in group), sum(r[14] != '一致' for r in group), ''])
    def sheet(name, headers, rows, widths, freeze):
        ws = wb.create_sheet(name)
        scope_label = {'all': '全量', 'agency': '按服务方', 'person': '按人员'}[mode]
        ws.append([f'临时工考勤核对表（核实用） | {month} | {scope_label} | 导出时间：{__import__("datetime").datetime.now():%Y-%m-%d %H:%M}'])
        ws.append(['金额=上报工时×合同价；缺价留空。管理费为元/小时。差异=重算−上报；计算工时为扣减前原始跨度，扣减后按本期参数取整。'])
        ws.append(['手机号仅有尾号时明确标注；边界日只作参考不计薪。汇总金额不含合同外奖惩；人力备注供线下回填，不自动写回系统。'])
        ws.append(headers)
        for row in rows:
            ws.append(row)
        for row in ws:
            for cell in row:
                cell.font = Font(name='微软雅黑', size=10, color='172D49')
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                # 所有外部文本作为文本保存，防止 Excel 将姓名/备注等解释成公式。
                if isinstance(cell.value, str): cell.data_type = 's'
        for cell in ws[4]:
            cell.fill = PatternFill('solid', fgColor='172D49'); cell.font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
        for row in (1,2,3):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(headers)); ws.row_dimensions[row].height=30
        ws.row_dimensions[4].height=34
        ws.freeze_panes=freeze; ws.auto_filter.ref=f'A4:{get_column_letter(len(headers))}{ws.max_row}'
        ws.print_title_rows='1:4'; ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape'; ws.page_setup.paperSize=ws.PAPERSIZE_A3; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.print_options.horizontalCentered=True
        ws.print_area=f'A1:{get_column_letter(len(headers))}{ws.max_row}'
        ws.sheet_view.showGridLines=False
        return ws
    a=sheet('服务方汇总', ['归属方','人数（姓名+标识）','结算行数','上报工时','重算计薪工时','差异','应发薪资','应发管理费','合计（不含奖惩）','合同外奖惩（原表）','待核日次','月结校验待核行数','人力备注'], agency_rows, [22,20,12,14,16,12,18,18,22,24,14,22,36], 'B5')
    total = ['本次导出合计', len({(r[0], r[1]) for r in people_rows}), len(people_rows)]
    for j in range(3,12):
        values=[r[j] for r in agency_rows]
        total.append(None if any(v is None for v in values) else money(sum(values)))
    total.append('')
    a.append(total)
    for cell in a[a.max_row]:
        cell.font=Font(name='微软雅黑',size=10,bold=True,color='172D49')
        cell.fill=PatternFill('solid',fgColor='EDEBFF')
    a.print_area=f'A1:M{a.max_row}'
    b=sheet('逐人汇总', ['姓名','手机号','归属方','部门','岗位','上报工时','重算计薪工时','差异','应发薪资','应发管理费','合计（不含奖惩）','合同外奖惩（原表）','原表请款合计','待核日次','月结校验','人力备注'],people_rows,[15,19,18,18,12,14,16,12,18,18,22,22,20,14,40,36],'D5')
    d=sheet('逐日核对',HEADERS,[r for r,_,_ in detail],[15,19,18,14,6,12,14,30,17,20,12,20,14,12,14,14,18,18,18,65,40],'D5')
    fills={'boundary':'EEF0F3','over_out':'FDE8E8','hard':'FDE8E8','ambig':'FFF4D6','thin':'FFF4D6','mixed':'FFF4D6','acked':'EAF2FF'}
    for i,(row,band,_) in enumerate(detail,5):
        for cell in d[i]:
            if band in fills: cell.fill=PatternFill('solid',fgColor=fills[band])
        d.cell(i,2).number_format='@'; d.cell(i,4).number_format='yyyy-mm-dd'
        for j in range(10,20):d.cell(i,j).number_format='#,##0.00;[Red]-#,##0.00;0.00'
        d.row_dimensions[i].height=42
        d.cell(i,21).fill=PatternFill('solid',fgColor='FFFBE6')
    for ws, cols in ((a,range(4,11)),(b,range(6,14))):
        for row in ws.iter_rows(min_row=5):
            for j in cols: row[j-1].number_format='#,##0.00;[Red]-#,##0.00;0.00'
            ws.row_dimensions[row[0].row].height=30
    for row in b.iter_rows(min_row=5):row[1].number_format='@'
    out=BytesIO();wb.save(out);return out.getvalue()
