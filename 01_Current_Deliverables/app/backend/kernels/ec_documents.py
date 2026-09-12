"""Cross-period financial document registry; no core/Kingdee imports or original-file writes."""
import hashlib
import io
import json
import threading
from datetime import datetime
from collections import defaultdict
from decimal import Decimal
import openpyxl
from sqlalchemy import MetaData, Table, Column, String, Text, Integer, Index, UniqueConstraint, select, insert, update, delete
from sqlalchemy.dialects.mysql import LONGTEXT
from kernels import ec_flow_ledger as ledger
from kernels import ec_tmall_import as tm

md=MetaData()
ec_document_files=Table('ec_document_files',md,
    Column('id',Integer,primary_key=True),Column('shop',String(120),nullable=False),
    Column('digest',String(64),nullable=False),Column('filename',String(180)),Column('kinds',String(120)),
    Column('row_count',Integer),Column('operator',String(50)),Column('ts',String(20)),
    UniqueConstraint('shop','digest',name='uq_ec_doc_file'))
ec_document_versions=Table('ec_document_versions',md,
    Column('id',Integer,primary_key=True),Column('file_id',Integer,nullable=False),
    Column('document_key',String(64),nullable=False),Column('fingerprint',String(64)),
    Column('sheet',String(128)),Column('row_number',Integer),Column('payload',Text().with_variant(LONGTEXT(),'mysql')),
    Index('ix_ec_doc_versions','document_key'))
ec_document_heads=Table('ec_document_heads',md,
    Column('document_key',String(64),primary_key=True),Column('shop',String(120),nullable=False),
    Column('kind',String(20)),Column('document_no',String(128)),Column('order_no',String(128)),
    Column('period',String(7)),Column('fingerprint',String(64)),Column('conflict',Integer,default=0),
    Column('payload',Text().with_variant(LONGTEXT(),'mysql')),Column('file_id',Integer),
    Column('sheet',String(128)),Column('row_number',Integer),
    Index('ix_ec_doc_order','order_no','shop'),Index('ix_ec_doc_no','document_no'),Index('ix_ec_doc_scope','shop','period','kind'))
ec_flow_matches=Table('ec_flow_matches',md,
    Column('flow_id',Integer,primary_key=True),Column('account_id',String(32)),
    Column('raw_order_no',String(128)),Column('order_no',String(128)),Column('shop',String(120)),
    Column('status',String(32)),Column('message',String(200)),Column('payload',Text().with_variant(LONGTEXT(),'mysql')),
    Column('ts',String(20)),Index('ix_ec_match_status','status','account_id'),
    Index('ix_ec_match_order','order_no'),Index('ix_ec_match_raw','raw_order_no'))
TABLES=tuple(md.tables.values())
LOCK=threading.RLock()
STATUSES={'amount_equal':'金额一致','linked':'已关联，金额待核对','amount_pending':'金额待核对',
    'missing_document':'有订单号但缺单据','unresolved':'无法确定订单号','conflict':'关联或单据版本冲突',
    'not_order_based':'无需逐单关联','pending_index':'待建立关联'}
RULES=[('item',{'主订单编号','子订单编号'}),('refund',{'退款编号','退款总额'}),
       ('wdt',{'出库单编号','原始单号'}),('order',{'订单编号','买家实付金额','订单创建时间'})]
FIELDS={
 'order':{'订单编号','订单创建时间','订单付款时间','发货时间','确认收货时间','订单状态','买家实付金额','退款金额','确认收货打款金额','总金额','支付详情','店铺名称'},
 'item':{'主订单编号','子订单编号','订单创建时间','订单付款时间','确认收货时间','订单状态','买家实付金额','退款金额','商品价格','购买数量','商家编码','商品ID'},
 'refund':{'订单编号','退款编号','支付宝交易号','订单付款时间','退款申请时间','退款完结时间','退款状态','退款总额','售后类型','货物状态','退给买家金额','退给平台金额'},
 'wdt':{'原始单号','子单原始单号','原始子订单号','出库单编号','订单编号','店铺','发货时间','付款时间','下单时间','应收金额','货品数量','货品编号','货品成交总价','出库单状态','出库状态','物流单号','物流公司'}}
def text(v):return ledger.string(v).strip()
def pack(v):return json.dumps(v,ensure_ascii=False,sort_keys=True,default=str)
def digest(v):return hashlib.sha256(pack(v).encode()).hexdigest()
def chunks(rows,size=400):
    for i in range(0,len(rows),size):yield rows[i:i+size]

def parse_documents(blob,filename,shop,shop_aliases=()):
    tm._preflight_xlsx(blob)
    wb=openpyxl.load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    docs=[];skipped=0;seen=set()
    try:
        for ws in wb:
            ws.reset_dimensions();it=ws.iter_rows(values_only=True);kind=None
            for number,row in enumerate(it,1):
                heads=[text(v) for v in row]
                kind=next((k for k,required in RULES if required<=set(heads)),None)
                if kind:break
                if number>=20:break
            if not kind:raise ValueError('无法按列识别资料类型：'+ws.title+'；支持平台订单、子订单、退款及旺店通出库')
            if len([h for h in heads if h])!=len(set(h for h in heads if h)):raise ValueError('表头重名，无法可靠识别')
            ix={h:i for i,h in enumerate(heads) if h in FIELDS[kind]}
            for line,row in enumerate(it,number+1):
                if line>250001:raise ValueError('单表超过250,000行')
                if not any(v is not None for v in row):continue
                raw={h:text(row[i]) if i<len(row) else '' for h,i in ix.items()}
                if kind=='wdt' and raw.get('店铺')!=shop:skipped+=1;continue
                if raw.get('店铺名称') and raw['店铺名称'] not in {shop,*shop_aliases}:
                    raise ValueError('文件店铺名称与所选店铺不一致，请先核实基础资料中的店铺别名')
                for h,i in ix.items():
                    if h in {'订单编号','主订单编号','子订单编号','退款编号','原始单号','子单原始单号','原始子订单号','物流单号'} and i<len(row) and isinstance(row[i],(int,float)):
                        raise ValueError('单据编号为数值格式，请按文本重新导出，避免精度损失')
                no=raw.get('子订单编号') if kind=='item' else raw.get('退款编号') if kind=='refund' else raw.get('出库单编号') if kind=='wdt' else raw.get('订单编号')
                order=raw.get('主订单编号') if kind=='item' else raw.get('原始单号') if kind=='wdt' else raw.get('订单编号')
                if not no or not order:raise ValueError('缺少单据编号或订单编号，文件未入库')
                if kind=='wdt':
                    # Line-grain warehouse exports have multiple materials per shipment.
                    no=no+':'+digest([raw.get(h,'') for h in ('原始单号','子单原始单号','原始子订单号','货品编号')])[:20]
                key=digest([shop,kind,no]);fp=digest(raw)
                if (key,fp) in seen:continue
                seen.add((key,fp))
                date=raw.get('退款申请时间') if kind=='refund' else raw.get('下单时间') if kind=='wdt' else raw.get('订单创建时间')
                stamp=tm._date_text(date)
                if kind in ('order','item','refund') and not stamp:raise ValueError('业务日期缺失或无法识别')
                for h in ('买家实付金额','退款金额','确认收货打款金额','退款总额','应收金额'):
                    if h=='退款金额' and raw.get(h)=='无退款申请':continue
                    if raw.get(h):ledger.decimal(raw[h])
                # Classify U先 without retaining buyer-facing product text or contact columns.
                if '商品标题' in heads:
                    title=text(row[heads.index('商品标题')])
                    raw['业务类型']='ufirst' if 'U先' in title or 'u先' in title else 'normal'
                docs.append(dict(document_key=key,shop=shop,kind=kind,document_no=no,order_no=order,
                    period=stamp[:7],fingerprint=digest(raw),payload=pack(raw),sheet=ws.title,row_number=line,conflict=0))
        if not docs:raise ValueError('没有属于所选店铺的有效业务记录')
        return {'filename':filename,'digest':hashlib.sha256(blob).hexdigest(),'docs':docs,'skipped':skipped}
    finally:wb.close()

def store_documents(engine,shop,parsed,operator):
    added=duplicates=conflicts=0;affected=set();now=datetime.now().isoformat(' ',timespec='seconds')
    with LOCK,engine.begin() as cx:
        for batch in parsed:
            if cx.execute(select(ec_document_files.c.id).where(ec_document_files.c.shop==shop,ec_document_files.c.digest==batch['digest'])).first():
                duplicates+=len(batch['docs']);continue
            fid=cx.execute(insert(ec_document_files).values(shop=shop,digest=batch['digest'],filename=batch['filename'],
                kinds=','.join(sorted({r['kind'] for r in batch['docs']})),row_count=len(batch['docs']),operator=operator,ts=now)).inserted_primary_key[0]
            previous={}
            for keys in chunks([r['document_key'] for r in batch['docs']]):
                previous.update({r.document_key:r for r in cx.execute(select(ec_document_heads).where(ec_document_heads.c.document_key.in_(keys)))})
            versions=[];new=[];changed=[]
            for r in batch['docs']:
                affected.update([r['order_no'],r['document_no']])
                versions.append(dict(file_id=fid,document_key=r['document_key'],fingerprint=r['fingerprint'],sheet=r['sheet'],row_number=r['row_number'],payload=r['payload']))
                old=previous.get(r['document_key'])
                if old:
                    if old.fingerprint==r['fingerprint']:duplicates+=1
                    else:changed.append(r['document_key']);conflicts+=1
                else:
                    new.append(dict(r,file_id=fid));added+=1
                    # Same document with different rows in one export is also a version conflict.
                    from types import SimpleNamespace
                    previous[r['document_key']]=SimpleNamespace(fingerprint=r['fingerprint'])
            for part in chunks(versions):cx.execute(insert(ec_document_versions),part)
            for part in chunks(new):cx.execute(insert(ec_document_heads),part)
            for part in chunks(changed):cx.execute(update(ec_document_heads).where(ec_document_heads.c.document_key.in_(part)).values(conflict=1))
    return {'added':added,'duplicates':duplicates,'conflicts':conflicts,'affected':sorted(affected),'skipped':sum(b['skipped'] for b in parsed)}

def reconcile(engine,flow_table,account_table,affected=None):
    """Persist cross-period matches. Queries never parse exports or rebuild all orders."""
    now=datetime.now().isoformat(' ',timespec='seconds')
    with LOCK,engine.begin() as cx:
        accounts={r.id:dict(kind=r.kind,shops=json.loads(r.shops or '[]')) for r in cx.execute(select(account_table))}
        heads=list(cx.execute(select(ec_document_heads)))
        by_order=defaultdict(list);by_child=defaultdict(set)
        for r in heads:
            by_order[r.order_no].append(r)
            if r.kind=='item':by_child[r.document_no].add((r.shop,r.order_no))
        known=set(by_order)|set(by_child)
        previous={r.flow_id:(r.raw_order_no,r.order_no) for r in cx.execute(select(ec_flow_matches.c.flow_id,ec_flow_matches.c.raw_order_no,ec_flow_matches.c.order_no))}
        # All imported cash dates contribute to an order's lifetime receipts, while UI period totals stay on the original flow table.
        rows=cx.execute(select(flow_table.c.id,flow_table.c.account_id,flow_table.c.payload).execution_options(stream_results=True))
        linked=[];receipts=defaultdict(Decimal)
        for r in rows:
            acc=accounts.get(r.account_id,{'shops':[],'kind':''});v=ledger.resolve_order(json.loads(r.payload),known)
            link=v['order_link'];no=link['order_no'];raw_no=no or v['supplement']['order_candidate']
            aliases={p for sh,p in by_child.get(no,set()) if sh in acc['shops']}
            candidate_orders={no}|aliases if no else set()
            docs=[d for o in candidate_orders for d in by_order.get(o,[]) if d.shop in acc['shops']]
            shops={d.shop for d in docs}
            parent=next(iter(aliases)) if len(aliases)==1 else no
            status='conflict' if link['status']=='conflict' or len(shops)>1 or len(aliases)>1 or any(d.conflict for d in docs) or any('同一流水号内容冲突' in flag for flag in v.get('flags',[])) else 'linked' if docs else 'missing_document' if no else 'unresolved'
            if status=='unresolved' and '推广佣金返还' in str(v.get('remark','')):status='not_order_based'
            shop=next(iter(shops)) if len(shops)==1 else ''
            code=v.get('code') or ledger.es._code_of(v.get('desc'))[0]
            is_receipt=code=='0010001' or v.get('bucket')=='receipt'
            copy=acc['kind']=='alipay' and v.get('chan')=='聚合结算渠道'
            if is_receipt and status=='linked' and not copy:receipts[(shop,parent)]+=ledger.decimal(v.get('income'))-ledger.decimal(v.get('outgo'))
            # Do not retain the full raw statement payload a second time in memory.
            linked.append((r.id,r.account_id,v['order_link'],raw_no,parent,shop,status,docs,is_receipt,copy))
        changes=[];counts=defaultdict(int);affected=set(affected) if affected is not None else None;updated=0
        def save(part):
            if not part:return
            cx.execute(delete(ec_flow_matches).where(ec_flow_matches.c.flow_id.in_([r['flow_id'] for r in part])))
            cx.execute(insert(ec_flow_matches),part)
        for rid,aid,link,raw_no,no,shop,status,docs,is_receipt,copy in linked:
            if affected is not None and rid in previous and not ({raw_no,no,*previous[rid]}&affected):continue
            message=STATUSES[status];extra={}
            main=[d for d in docs if d.kind=='order']
            if status=='linked' and is_receipt and main and not copy:
                raw=json.loads(main[0].payload);expected=raw.get('确认收货打款金额')
                if expected not in ('',None):
                    wanted=ledger.decimal(expected);actual=receipts[(shop,no)]
                    status='amount_equal' if actual==wanted and actual!=0 else 'amount_pending'
                    message='订单累计交易收款与平台确认打款金额一致' if status=='amount_equal' else '订单累计交易收款与平台确认打款金额待核对'
                    extra={'expected':str(wanted),'received':str(actual),'difference':str(actual-wanted),'basis':'跨期累计交易收款；费用及退款单独核对，非银行到账'}
            elif status=='linked' and not main:message='已找到历史子订单或其他业务单据，金额仍待核对'
            if copy:message='聚合渠道映射流水，不重复计入订单收款'
            evidence=[{'kind':d.kind,'document_no':d.document_no,'order_no':d.order_no,'period':d.period,
                'file_id':d.file_id,'sheet':d.sheet,'row':d.row_number,'conflict':bool(d.conflict)} for d in docs]
            changes.append(dict(flow_id=rid,account_id=aid,raw_order_no=raw_no,order_no=no,shop=shop,status=status,message=message,
                payload=pack({'documents':evidence,'amount_check':extra,'link':link}),ts=now));counts[status]+=1;updated+=1
            if len(changes)>=400:save(changes);changes=[]
        save(changes)
        return {'updated':updated,'statuses':dict(counts)}
