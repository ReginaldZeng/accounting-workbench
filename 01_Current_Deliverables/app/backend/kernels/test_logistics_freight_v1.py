# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-10-26 | Author: Kimi | Version: V-draft（逐单运费 V1）
# Description: 合成内存 Excel 与内存 SQLite 测试；不读取真实附件、不连接金蝶。
import hashlib
import unittest
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from openpyxl import Workbook
from sqlalchemy import create_engine, select

from kernels import logistics_freight_v1 as freight
from kernels import logistics_freight_v1_store as store

CONTEXT = {"carrier": "测试承运", "subject": "测试主体", "biz_line": "测试业务"}


def excel(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows, merges in sheets:
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
        for merged in merges:
            ws.merge_cells(merged)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheet_rows(kind="sale"):
    if kind == "transfer":
        return [["分步调拨"], ["物料编码", "金蝶单号", "含税金额（18.5元/吨）", "吨数(T)"]]
    if kind == "inbound":
        return [["原辅料"], ["序号", "物料编码", "金蝶入库单号", "含税金额（18.5元/吨）"]]
    if kind == "return":
        return [["退货"], ["序号", "物料编码", "金蝶入库单号", "含税金额（18.5元/吨）"]]
    return [["销售装货"], ["序号", "物料编码", "金蝶单据编号", "含税金额（18.5元/吨）"]]


class ParserTests(unittest.TestCase):
    def parse(self, data, context=None):
        return freight.parse_workbook(data, "合成.xlsx", context or CONTEXT)

    def test_merged_doc_does_not_copy_amount_and_same_doc_accumulates(self):
        rows = sheet_rows() + [[1, "M1", "XSCKD1", 10], [2, "M2", None, 20], [3, "M3", "XSCKD2", 0],
                                [None, None, "合计", 30]]
        result = self.parse(excel([("销售成品装货明细", rows, ["C3:C4"])]))
        self.assertEqual(result["summary"]["line_count"], 3)
        self.assertEqual(result["summary"]["doc_count"], 2)
        self.assertEqual(result["summary"]["total_amount"], "30")
        self.assertEqual(result["summary"]["difference"], "0")
        self.assertEqual([x["amount"] for x in result["lines"]], ["10", "20", "0"])
        self.assertEqual(result["lines"][1]["doc_no"], "XSCKD1")
        self.assertEqual(result["lines"][1]["merged_range"], "C3:C4")

    def test_header_names_not_fixed_columns_and_special_total(self):
        rows = [["装货"], ["含税金额（18.5元/吨）", "金蝶单据编号", "物料编码"],
                [1.25, "XSCKD9", "M1"], [2.75, "XSCKD8", "M2"], [4, None, "6%含税金额"]]
        result = self.parse(excel([("小料装货明细", rows, [])]))
        self.assertEqual(result["summary"]["line_count"], 2)
        self.assertEqual(result["summary"]["total_amount"], "4")
        self.assertEqual(result["summary"]["source_total"], "4")
        self.assertEqual(result["summary"]["status"], "ready")

    def test_transfer_skips_sheet1_and_reconciles(self):
        main = sheet_rows("transfer") + [["M1", "FBDC1", 1.234567], ["M2", None, 2], [None, "合计", 3.234567]]
        result = self.parse(excel([("分步调拨卸货明细 ", main, ["B3:B4"]),
                                   ("Sheet1", [["辅助"], ["FBDC1", 100]], [])]))
        self.assertEqual(result["summary"]["line_count"], 2)
        self.assertEqual(result["summary"]["total_amount"], "3.234567")
        self.assertEqual(result["lines"][1]["doc_no"], "FBDC1")
        self.assertEqual(len(result["summary"]["sheets"]), 1)

    def test_inbound_duplicate_zero_return_and_reconciliation_difference(self):
        main = sheet_rows("inbound") + [[1, "M1", "CGRK183089", 100], [2, "M1", "CGRK183089", 100],
                                       [None, None, "合计", 150]]
        zero_return = sheet_rows("return") + [[None, None, None, 0], [None, None, "合计", 0]]
        result = self.parse(excel([("小料卸货明细 (2)", main, []),
                                   ("植物肉小料退货明细 (3)", zero_return, [])]))
        self.assertEqual(result["summary"]["line_count"], 2)
        self.assertEqual(result["summary"]["total_amount"], "200")
        self.assertEqual(result["summary"]["source_total"], "150")
        self.assertEqual(result["summary"]["difference"], "50")
        self.assertEqual([x["doc_no"] for x in result["lines"]], ["CGRK183089", "CGRK183089"])
        self.assertEqual({x["code"] for x in result["issues"]},
                         {"duplicate_rows", "zero_return_skipped", "reconciliation_diff"})

    def test_nonzero_return_and_required_fields_block(self):
        rows = sheet_rows("return") + [[1, "M1", "CGRK1", 3]]
        result = self.parse(excel([("小料退货明细", rows, [])]))
        self.assertEqual(result["summary"]["status"], "blocked")
        self.assertIn("nonzero_return_blocked", {x["code"] for x in result["issues"]})
        sale = sheet_rows() + [[1, "M1", None, 0]]
        result = self.parse(excel([("销售成品装货明细", sale, [])]), {"carrier": "", "subject": ""})
        self.assertEqual({x["code"] for x in result["issues"]},
                         {"missing_doc_no", "missing_carrier", "missing_subject"})
        with patch.object(freight, "_fee_type", return_value=""):
            result = self.parse(excel([("销售成品装货明细", sheet_rows() + [[1, "M1", "D1", 0]], [])]))
        self.assertIn("missing_fee_type", {x["code"] for x in result["issues"]})

    def test_kingdee_other_and_sam_layouts_take_freight_once(self):
        class FakeSheet:
            def __init__(self, name, rows):
                self.name, self.rows = name, rows
                self.nrows = len(rows)

            def row_values(self, row):
                return self.rows[row]

        class FakeBook:
            def __init__(self, sheets):
                self._sheets = sheets

            def sheet_names(self):
                return [x.name for x in self._sheets]

            def sheets(self):
                return self._sheets

        other = FakeSheet("其他", [["日期", "单据编号", "物料编码", "物流商", "运费", "销售组织"],
                                  ["2026-09-01", "XSCKD1", "M1", "甲", 12.5, "主体A"],
                                  [None, None, "M2", None, None, None]])
        sam = FakeSheet("山姆", [["单据编号", "运费", "日期", "销售组织", "物流商", "物料编码"],
                                ["XSCKD2", 7.5, "2026-09-01", "主体A", "乙", "M3"],
                                [None, None, None, None, None, "M4"]])
        with patch.object(freight.xlrd, "open_workbook", return_value=FakeBook([other, sam])):
            result = freight.parse_workbook(bytes.fromhex("d0cf11e0a1b11ae1"), "合成.xls", CONTEXT)
        self.assertEqual(result["summary"]["doc_count"], 2)
        self.assertEqual(result["summary"]["total_amount"], "20")
        self.assertEqual([x["amount"] for x in result["lines"]], ["12.5", "0", "7.5", "0"])
        self.assertEqual(result["lines"][1]["doc_no"], "XSCKD1")
        self.assertEqual(result["lines"][1]["carrier"], "甲")
        self.assertEqual(result["lines"][3]["subject"], "主体A")

    def test_carrier_exact_keyword_unknown_and_conflict(self):
        catalog = [{"name": "天鹰", "full_name": "天鹰物流有限公司"},
                   {"name": "天鹰华南", "full_name": "天鹰华南供应链"},
                   {"name": "跨越", "full_name": "跨越速运"}]
        self.assertEqual(freight.normalize_carrier_name("天鹰物流有限公司", catalog)["status"], "exact")
        keyword = freight.normalize_carrier_name("天鹰华南货运", catalog)
        self.assertEqual(keyword["status"], "conflict")
        only = freight.normalize_carrier_name("跨越物流公司", catalog)
        self.assertEqual((only["status"], only["carrier"]), ("keyword", "跨越"))
        unknown = freight.normalize_carrier_name("陌生物流", catalog)
        self.assertEqual(unknown["status"], "unknown")
        generic = freight.normalize_carrier_name("物流", catalog)
        self.assertEqual(generic["status"], "unknown")

    def test_normalize_parsed_preserves_raw_and_blocks_unknown(self):
        parsed = {"carrier": "", "summary": {}, "issues": [], "lines": [
            {"doc_no": "D1", "carrier": "跨越物流公司", "sheet": "S", "row": 3},
            {"doc_no": "D2", "carrier": "陌生物流", "sheet": "S", "row": 4},
        ]}
        freight.normalize_parsed_carriers(parsed, [{"name": "跨越", "full_name": "跨越速运"}])
        self.assertEqual(parsed["lines"][0]["carrier"], "跨越")
        self.assertEqual(parsed["lines"][1]["carrier"], "")
        self.assertEqual(parsed["summary"]["status"], "blocked")
        self.assertEqual(parsed["issues"][0]["raw_carrier"], "跨越物流公司")
        self.assertEqual(parsed["issues"][1]["code"], "unknown_carrier")

    def test_filename_and_document_period_mismatch(self):
        parsed = {"summary": {}, "issues": [], "lines": [{"doc_no": "D1", "doc_date": "46267"}]}
        freight.validate_parsed_period(parsed, "9月份天鹰明细.xlsx", "2026-08")
        issue = parsed["issues"][0]
        self.assertEqual(issue["code"], "period_mismatch")
        self.assertIn("2026-08", issue["message"])
        self.assertIn("2026-09", issue["message"])
        self.assertEqual(parsed["summary"]["status"], "blocked")
        parsed = {"summary": {}, "issues": [], "lines": [{"doc_no": "D1", "doc_date": "2026-09-30"}]}
        freight.validate_parsed_period(parsed, "9月份天鹰明细.xlsx", "2026-09")
        self.assertEqual(parsed["summary"]["status"], "ready")


class StoreTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite://")
        store.md.create_all(self.engine)

    def test_schema_blob_sha_and_numeric(self):
        data = excel([("销售成品装货明细", sheet_rows() + [[1, "M1", "D1", "1.2345678"]], [])])
        parsed = freight.parse_workbook(data, "a.xlsx", CONTEXT)
        batch = store.create_draft(self.engine, "a.xlsx", data, parsed, "测试")
        saved = store.get_file(self.engine, batch["id"])
        self.assertEqual(saved["sha256"], hashlib.sha256(data).hexdigest())
        self.assertEqual(bytes(saved["content"]), data)
        with self.engine.connect() as cx:
            amount = cx.execute(select(store.lines.c.amount)).scalar()
        self.assertEqual(amount, Decimal("1.234568"))
        self.assertEqual(store.batches.c.status.type.length, 12)
        self.assertEqual(str(store.lines.c.amount.type), "NUMERIC(18, 6)")


if __name__ == "__main__":
    unittest.main()
