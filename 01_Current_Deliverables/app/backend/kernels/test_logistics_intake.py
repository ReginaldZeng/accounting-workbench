# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-10-08 | Author: Claude Opus 5.5 | Version: V2.864
# Description: 物流账单取数(logistics_intake)单测——单号列合并单元格：合并块下半截沿用首行单号(不看这行写没写日期序号)。合成数据。
import io
import unittest

import openpyxl

from kernels import logistics_intake as li

SPEC = {"carrier": "天鹰物流", "period": "2026-08", "sheets": [{
    "name": "蜜雪装货明细*", "role": "detail", "header_row": 2, "row_re": {"col": "产品名称", "re": "^(?!合计|总计|小计).+"},
    "doc_col": ["金蝶单据编号", "金蝶单号"], "doc_ffill": True, "doc_default": "无单据", "merge_doc": True,
    "amount_col": "含税金额*", "wt_col": "吨数(T)", "wt_scale": 1000, "amount_dp": 6, "doc_ffill_if_blank": ["装货日期"]}]}


def _book(merge=True):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "蜜雪装货明细 (3)"
    ws.append(["蜜雪装货记录表"])
    ws.append(["序号", "产品名称", "吨数(T)", "装货日期", "金蝶单据编号", "含税金额（18.5元/吨）"])
    ws.append([1, "冷冻米麻薯", 8, 46265, "XSCKD219507", 148])
    ws.append([2, "冷冻米麻薯", 8, 46265, None, 148])            # 第 4 行：单号格和上一行合并，自己写了序号、日期
    ws.append([3, "冷冻米麻薯", 8, 46265, "XSCKD219505", 148])
    ws.append([4, "冷冻米麻薯", 8, 46265, None, 148])            # 第 6 行：没合并、日期齐全却没单号 → 漏填
    ws.append([5, "冷冻米麻薯", 15, 46265, "XSCKD219513", 277.5])
    ws.append([None, "冷冻米麻薯", 2, None, None, 37])           # 第 8 行：没合并、没日期 → 续行(老规则)
    ws.merge_cells("A1:F1")
    if merge:
        ws.merge_cells("E3:E4")
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


class MergedDocCell(unittest.TestCase):
    def test_merged_rows(self):
        self.assertEqual(li.merged_rows(_book()), {"蜜雪装货明细 (3)": {4: {4: 3}}})    # 横向合并(标题行)不算
        self.assertEqual(li.merged_rows(_book(False)), {})
        self.assertEqual(li.merged_rows(b"not a zip"), {})

    def test_merged_cell_inherits_doc(self):
        by = {}
        for r in li.parse_bill(dict(SPEC), _book())["detail"]:
            by.setdefault(r["doc_no"], []).append(r)
        self.assertAlmostEqual(by["XSCKD219507"][0]["charge_wt"], 16000)     # 两行并成一单：8 吨 + 8 吨
        self.assertAlmostEqual(by["XSCKD219507"][0]["amount"], 296)
        self.assertAlmostEqual(by["XSCKD219505"][0]["charge_wt"], 8000)      # 没合并的那行不并进来
        self.assertEqual(len(by["无单据"]), 1)                               # 它还是按漏填单列
        self.assertAlmostEqual(by["XSCKD219513"][0]["charge_wt"], 17000)     # 没日期的续行照旧沿用上一单

    def test_without_merge_same_as_before(self):
        by = {}
        for r in li.parse_bill(dict(SPEC), _book(False))["detail"]:
            by.setdefault(r["doc_no"], []).append(r)
        self.assertAlmostEqual(by["XSCKD219507"][0]["charge_wt"], 8000)
        self.assertEqual(len(by["无单据"]), 2)


class DimCol(unittest.TestCase):
    """V2.888 按体积计费：备注列里的箱子尺寸读进分项(文字，不进金额)；不是尺寸的备注不收。合成数据。"""

    def test_dim_col(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "零担"
        ws.append(["合成账单"])
        ws.append(["金蝶单号", "总件数", "计重", "运费", "送货费", "备注"])
        ws.append(["XSCKD000001", 41, 158.91, 286.05, 100, "38*25.5*12"])
        ws.append(["XSCKD000002", 6, 30, 165, None, "起步价"])
        ws.append(["XSCKD000003", 2, 7.75, 20, None, " 38 × 25.5 × 12 cm "])
        bio = io.BytesIO()
        wb.save(bio)
        spec = {"carrier": "合成", "period": "2026-08", "sheets": [{
            "name": "零担", "role": "detail", "header_row": 2, "doc_col": ["金蝶单号"], "qty_col": "总件数", "wt_col": "计重",
            "fee_parts": {"运费": ["运费"], "送货费": ["送货费"]}, "dim_col": "备注"}]}
        got = {r["doc_no"]: r for r in li.parse_bill(spec, bio.getvalue())["detail"]}
        import json as _j
        sf = {k: _j.loads(v.get("sub_fees") or "{}") for k, v in got.items()}
        self.assertEqual(sf["XSCKD000001"].get("箱子尺寸"), "38*25.5*12")
        self.assertEqual(got["XSCKD000001"]["amount"], 386.05)          # 尺寸是文字，不进金额
        self.assertNotIn("箱子尺寸", sf["XSCKD000002"])
        self.assertEqual(sf["XSCKD000003"].get("箱子尺寸"), "38×25.5×12cm")
        self.assertTrue(li.DIM_RE.match(sf["XSCKD000003"]["箱子尺寸"]))


class TotalCol(unittest.TestCase):
    """V2.987 跨越 9 月：账单新加了 夜间取货费/退仓费/资源调配费 三列，取数说明只认老列名，拆项求和悄悄少读 81 元。
    配了 total_col(账单自己的应付合计)就以它为准，对不上的差额记「其他费用（未拆）」。合成数据。"""

    def _bill(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "帐单-深圳星期零"
        ws.append(["合成"])
        ws.append(["金蝶单号", "折后运费", "保费", "退仓费", "应付金额\n（元）"])
        ws.append(["CGRK000001", 54.78, 2, 60, 116.78])
        ws.append(["QTCK000002", 89.6, 2, 0, 91.6])
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue()

    def _spec(self, **kw):
        sp = {"name": "深圳星期零|孝感星期九", "role": "detail", "header_row": 2, "doc_col": "金蝶单号",
              "fee_parts": {"运费": ["折后运费"], "保费": ["保费"]}}
        sp.update(kw)
        return {"carrier": "合成", "period": "2026-09", "sheets": [sp]}

    def test_without_total_drops_new_column(self):
        got = {r["doc_no"]: r for r in li.parse_bill(self._spec(), self._bill())["detail"]}
        self.assertEqual(got["CGRK000001"]["amount"], 56.78)            # 老行为：新列没进拆项就丢了

    def test_total_col_keeps_money_and_shows_gap(self):
        import json as _j
        got = {r["doc_no"]: r for r in li.parse_bill(self._spec(total_col="应付金额*"), self._bill())["detail"]}
        self.assertEqual(got["CGRK000001"]["amount"], 116.78)
        self.assertEqual(_j.loads(got["CGRK000001"]["sub_fees"])["其他费用（未拆）"], 60.0)
        self.assertNotIn("其他费用（未拆）", _j.loads(got["QTCK000002"]["sub_fees"]))
        self.assertEqual(got["QTCK000002"]["amount"], 91.6)


class AnnotSource(unittest.TestCase):
    def _detail_book(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "明细"
        ws.append(["金蝶单号", "金额", "标注"])
        ws.append(["XSCKD1", 1.234567, "原词"])
        ws.append(["XSCKD2", 2.345678, " "])
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue()

    def _spec(self, **kw):
        sp = {"name": "明细", "role": "detail", "doc_col": "金蝶单号", "amount_col": "金额", "amount_dp": 6,
              "annot_col": "标注", "annot_map": {"原词": "映射"}, "annot_fmt": "仓储费-{}"}
        sp.update(kw)
        return {"carrier": "合成", "period": "2026-09", "sheets": [sp]}

    def test_raw_default_none_and_amount_dp(self):
        got = li.parse_bill(self._spec(annot="采购入库-植物肉"), self._detail_book())["detail"]
        self.assertEqual((got[0]["annot"], got[0]["annot_source"]), ("仓储费-映射", "raw"))
        self.assertEqual((got[1]["annot"], got[1]["annot_source"]), ("采购入库-植物肉", "default"))
        self.assertEqual(got[0]["amount"], 1.234567)
        self.assertEqual(got[1]["amount"], 2.345678)
        no_default = li.parse_bill(self._spec(annot_fmt="{}"), self._detail_book())["detail"]
        self.assertEqual((no_default[1]["annot"], no_default[1]["annot_source"]), ("", "none"))
        no_column = li.parse_bill(self._spec(annot_col=None), self._detail_book())["detail"]
        self.assertEqual(no_column[0]["annot_source"], "none")

    def test_accrual_fee_map_and_default_annot_are_default(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "汇总"
        ws.append(["快递费", 2, 10])
        ws.append([1, "卸货", 3, 20])
        sp = {"fee_map": {"快递费": "销售出库单-电商"}, "default_annot": "销售出库单-电商"}
        got = li.parse_accrual_sheet(sp, ws, "2026-09", "合成")
        self.assertEqual([r["annot_source"] for r in got], ["default", "default"])


class SheetRoundedTotal(unittest.TestCase):
    def test_real_four_sheet_tail_difference(self):
        amounts = [("蜜雪", 18222.50), ("植物肉", 547.779711), ("小料", 23384.9028), ("分步", 286.114155)]
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        sheets = []
        for name, amount in amounts:
            ws = wb.create_sheet(name)
            ws.append(["金蝶单号", "金额"])
            ws.append(["无单据", amount])
            sheets.append({"name": name, "role": "detail", "doc_col": "金蝶单号", "amount_col": "金额", "amount_dp": 6})
        bio = io.BytesIO()
        wb.save(bio)
        spec = {"carrier": "天鹰物流", "period": "2026-09", "sheets": sheets}
        full = li.parse_bill(spec, bio.getvalue())
        self.assertEqual([r["amount"] for r in full["detail"]], [x[1] for x in amounts])
        self.assertEqual(round(sum(x[1] for x in amounts), 2), 42441.30)
        self.assertEqual(li.sheet_round_sum(full["detail"]), 42441.29)
        self.assertEqual(li.sheet_report(spec, bio.getvalue(), "2026-09")["amount"], 42441.29)
        grouped = li.sheet_round_groups(full["detail"], lambda r: "同一业务")
        self.assertEqual(grouped, {"同一业务": 42441.29})
        self.assertEqual(round(sum(grouped.values()), 2), li.sheet_round_sum(full["detail"], lambda r: "同一业务"))

    def test_blank_src_sheet_is_a_physical_bucket(self):
        rows = [{"amount": 0.004, "src_sheet": ""}, {"amount": 0.004, "src_sheet": None}]
        self.assertEqual(li.sheet_round_sum(rows), 0.01)
        rows[1]["src_sheet"] = "另一页"
        self.assertEqual(li.sheet_round_sum(rows), 0.0)


class SheetReport(unittest.TestCase):
    """V2.989 复核台第①步「读全没有」：算钱的列加起来对账单自己的合计列。合成数据(仿跨越 9 月新加退仓费那一页)。"""

    def _bill(self, with_total=True):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "帐单-深圳星期零"
        ws.append(["合成"])
        ws.append(["序号", "金蝶单号", "计费重量（公斤）", "运单运费", "折后运费", "保费", "退仓费"] + (["应付金额\n（元）"] if with_total else []))
        ws.append([1, "CGRK000001", 30, 78.26, 54.78, 2, 60] + ([116.78] if with_total else []))
        ws.append([2, "QTCK000002", 50, 128, 89.6, 2, 0] + ([91.6] if with_total else []))
        ws.append(["合计", None, 80, 206.26, 144.38, 4, 60] + ([208.38] if with_total else []))
        ws2 = wb.create_sheet("Sheet1")
        ws2.append(["透视"])
        bio = io.BytesIO()
        wb.save(bio)
        return bio.getvalue()

    def _spec(self, **kw):
        sp = {"name": "深圳星期零|孝感星期九", "role": "detail", "header_row": 2, "doc_col": "金蝶单号", "wt_col": "计费重量（公斤）", "summary_marker": "合计",
              "fee_parts": {"运费": ["折后运费"], "保费": ["保费"]}}
        sp.update(kw)
        return {"carrier": "合成", "sheets": [sp]}

    def test_missing_column_shows_gap(self):
        r = li.sheet_report(self._spec(), self._bill(), "2026-09")
        s = r["sheets"][0]
        self.assertEqual(s["verdict"], "gap")
        self.assertEqual(s["total"]["how"], "guess")                 # 没配合计列：按表头认出「应付金额」只拿来比
        self.assertEqual(s["total"]["sum"], 208.38)                   # 合计行不算进去
        self.assertEqual(s["gap"], -60.0)
        self.assertEqual([u["name"] for u in s["unread"]], ["运单运费", "退仓费"])
        self.assertEqual(r["sheets"][1]["verdict"], "skip")
        self.assertEqual(r["verdict"], "gap")

    def test_all_columns_ok(self):
        r = li.sheet_report(self._spec(fee_parts={"运费": ["折后运费"], "保费": ["保费"], "退仓费": ["退仓费"]}), self._bill(), "2026-09")
        s = r["sheets"][0]
        self.assertEqual((s["verdict"], s["amount"], s["n"]), ("ok", 208.38, 2))
        self.assertNotIn("计费重量（公斤）", [u["name"] for u in s["unread"] + s["others"]])   # 用到的列不算没读
        self.assertEqual(r["verdict"], "ok")

    def test_total_col_keeps_money_but_flags_unsplit(self):
        s = li.sheet_report(self._spec(total_col="应付金额*"), self._bill(), "2026-09")["sheets"][0]
        self.assertEqual(s["verdict"], "unsplit")
        self.assertEqual((s["amount"], s["unsplit"]), (208.38, 60.0))

    def test_single_amount_col_is_its_own_total(self):
        s = li.sheet_report(self._spec(fee_parts=None, amount_col="应付金额*"), self._bill(), "2026-09")["sheets"][0]
        self.assertEqual((s["verdict"], s["total"]["how"], s["amount"]), ("ok", "self", 208.38))

    def test_no_total_but_unread_fee(self):
        s = li.sheet_report(self._spec(), self._bill(with_total=False), "2026-09")["sheets"][0]
        self.assertEqual(s["verdict"], "maybe")
        self.assertIn("退仓费", [u["name"] for u in s["unread"]])

    def test_rows_merged_into_one_doc_still_counted(self):
        """天鹰：一张单拆几行，解析时并成一行(只剩首行行号)。续行的钱也要加进来比，不能报对不上。"""
        s = li.sheet_report(dict(SPEC), _book(), "2026-08")["sheets"][0]
        self.assertEqual(s["verdict"], "ok")
        self.assertEqual(s["total"]["sum"], s["amount"])

    def test_configured_column_missing_this_month(self):
        s = li.sheet_report(self._spec(fee_parts={"运费": ["折后运费"], "保费": ["保费"], "退仓费": ["退仓费"], "回单费": ["回单费"]}), self._bill(), "2026-09")["sheets"][0]
        self.assertEqual([m["name"] for m in s["money"] if not m["found"]], ["回单费"])
        self.assertEqual(s["verdict"], "ok")


if __name__ == "__main__":
    unittest.main()
