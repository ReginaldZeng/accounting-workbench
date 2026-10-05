# V2.810: 金蝶未核销清单——电商应收单 / 收款单里还没核销完的，按店铺、账龄分堆。纯计算：不连金蝶、不碰库。
import datetime
import io
import math
import re
from collections import defaultdict

# 金蝶取数字段（只读查询）。未核销额 = 应收金额 − 已核销金额（收款计划行，含红蓝对冲；
# 不用「关联收款金额」——它漏掉红蓝对冲，部分核销的单会被算多）。平台订单号在 Text6，红字应收在 Text4。
AR_FIELDS = [('FBillNo', 'no'), ('FDATE', 'date'), ('FCUSTOMERID.FName', 'shop'), ('FWRITTENOFFSTATUS', 'ws'),
             ('FDOCUMENTSTATUS', 'ds'), ('FPAYAMOUNTFOR', 'amount'), ('FWRITTENOFFAMOUNTFOR_P', 'written'),
             ('F_ora_Text6', 't6'), ('F_ora_Text4', 't4')]
REC_FIELDS = [('FBillNo', 'no'), ('FDATE', 'date'), ('FCONTACTUNIT.FName', 'shop'), ('FWRITTENOFFSTATUS', 'ws'),
              ('FDOCUMENTSTATUS', 'ds'), ('FRECTOTALAMOUNTFOR', 'amount'), ('FWRITTENOFFAMOUNTFOR_D', 'written'),
              ('FSETTLETYPEID.FName', 'settle'), ('FACCOUNTID.FNumber', 'account')]

PILES = {'current': '当期', 'stray': '零星', 'old': '老账'}
PILE_HELP = {'current': '本月和上月的单，月结正在核', 'stray': '更早、一年以内，平时核销漏下的', 'old': '挂了一年以上'}
HINTS = {'pair_even': '同一订单红字蓝字都挂着·互冲即平', 'pair_rest': '同一订单红字蓝字都挂着·互冲后还有差额',
         'red_alone': '红字应收单独挂着', 'unaudited': '应收单还没审核', 'partial': '只核了一部分'}
WS = {'A': '未核销', 'B': '部分核销', 'C': '已核销'}
DS = {'A': '未审核', 'B': '审核中', 'C': '已审核', 'D': '重新审核', 'Z': '暂存'}
SIDES = {'ar': '应收单', 'rec': '收款单'}
_ORDER = re.compile(r'^(\d{15,})(?:-\d+)?$')


def _text(value):
    return str(value or '').strip()


def _num(value):
    try: return float(value or 0)
    except (TypeError, ValueError): return 0.0


def order_key(*candidates):
    """平台订单号：15 位以上纯数字，去掉拆单后缀（-1、-2）。取不到返回空串。"""
    for value in candidates:
        m = _ORDER.match(_text(value))
        if m: return m.group(1)
    return ''


def month_shift(day, months):
    """day 所在月往前/后挪 months 个月的月初，返回 YYYY-MM-01。"""
    index = day.year * 12 + day.month - 1 + months
    return '%04d-%02d-01' % (index // 12, index % 12 + 1)


def pile_of(date, today):
    if date >= month_shift(today, -1): return 'current'
    return 'old' if date < month_shift(today, -12) else 'stray'


def _merge(rows, extra):
    """金蝶按明细行返回，同一张单可能多行：金额累加，表头字段取第一行。"""
    bills = {}
    for r in rows:
        no = _text(r.get('no'))
        if not no: continue
        b = bills.get(no)
        if b is None:
            b = bills[no] = dict(no=no, date=_text(r.get('date'))[:10], shop=_text(r.get('shop')) or '（未填客户）',
                                 ws=_text(r.get('ws')), ds=_text(r.get('ds')), _a=[], _w=[], **extra(r))
        b['_a'].append(_num(r.get('amount'))); b['_w'].append(_num(r.get('written')))
    out = []
    for b in bills.values():
        b['amount'] = round(math.fsum(b.pop('_a')), 2); b['written'] = round(math.fsum(b.pop('_w')), 2)
        b['open'] = round(b['amount'] - b['written'], 2)
        out.append(b)
    return out


def build(ar_rows, rec_rows, today=None):
    """金蝶原始行 → 快照 {ar:[...], rec:[...]}。每张单带 pile（哪一堆）与 hint（只靠金蝶就能判的提示）。"""
    today = today or datetime.date.today()
    ar = _merge(ar_rows, lambda r: dict(order=order_key(r.get('t6'), r.get('t4'))))
    rec = _merge(rec_rows, lambda r: dict(settle=_text(r.get('settle')), account=_text(r.get('account'))))
    groups = defaultdict(list)
    for b in ar:
        b['pile'] = pile_of(b['date'], today)
        if b['order']: groups[(b['shop'], b['order'])].append(b)
    paired = {}
    for key, group in groups.items():
        if any(b['open'] < 0 for b in group) and any(b['open'] > 0 for b in group):
            paired[key] = round(math.fsum(b['open'] for b in group), 2)
    for b in ar:
        net = paired.get((b['shop'], b['order'])) if b['order'] else None
        b['pair_net'] = net
        b['hint'] = (('pair_even' if abs(net) < 0.005 else 'pair_rest') if net is not None else
                     'red_alone' if b['open'] < 0 else 'unaudited' if b['ds'] != 'C' else 'partial' if b['ws'] == 'B' else '')
    for b in rec:
        b['pile'] = pile_of(b['date'], today); b['order'] = ''; b['pair_net'] = None
        b['hint'] = 'unaudited' if b['ds'] != 'C' else 'partial' if b['ws'] == 'B' else ''
    key = lambda b: (b['date'], b['no'])
    return {'ar': sorted(ar, key=key, reverse=True), 'rec': sorted(rec, key=key, reverse=True)}


def _cell():
    return {'count': 0, 'amount': 0.0}


def _add(cell, bill):
    cell['count'] += 1; cell['amount'] += bill['open']


def _round(cell):
    cell['amount'] = round(cell['amount'], 2); return cell


def summary(snapshot):
    """首屏汇总：三堆合计、店铺×三堆、提示分布、收款单一侧合计。"""
    total, piles, hints, shops = _cell(), {k: _cell() for k in PILES}, {k: _cell() for k in HINTS}, {}
    for b in snapshot['ar']:
        _add(total, b); _add(piles[b['pile']], b)
        if b['hint']: _add(hints[b['hint']], b)
        row = shops.setdefault(b['shop'], dict(shop=b['shop'], total=_cell(), **{k: _cell() for k in PILES}))
        _add(row['total'], b); _add(row[b['pile']], b)
    for row in shops.values():
        for k in ('total', *PILES): _round(row[k])
    rec_total, rec_shops = _cell(), {}
    for b in snapshot['rec']:
        _add(rec_total, b); _add(rec_shops.setdefault(b['shop'], dict(shop=b['shop'], total=_cell()))['total'], b)
    return {'total': _round(total), 'piles': {k: _round(v) for k, v in piles.items()},
            'hints': {k: _round(v) for k, v in hints.items() if v['count']},
            'shops': sorted(shops.values(), key=lambda r: -abs(r['total']['amount'])),
            'receipts': {'total': _round(rec_total),
                         'shops': sorted((dict(r, total=_round(r['total'])) for r in rec_shops.values()), key=lambda r: -abs(r['total']['amount']))}}


def select(snapshot, side='ar', pile='', shop='', month='', hint='', q=''):
    """按条件筛出单据（已按日期倒序）。q 模糊匹配单号 / 平台订单号。"""
    q = _text(q).lower()
    def keep(b):
        if pile and b['pile'] != pile: return False
        if shop and b['shop'] != shop: return False
        if month and b['date'][:7] != month: return False
        if hint and (b['hint'] != hint if hint != 'none' else b['hint']): return False
        return not q or q in b['no'].lower() or q in b['order']
    return [b for b in snapshot['rec' if side == 'rec' else 'ar'] if keep(b)]


def page(rows, number=1, size=50):
    pages = max(1, -(-len(rows) // size)); number = min(max(1, number), pages)
    months = defaultdict(_cell)
    for b in rows: _add(months[b['date'][:7]], b)
    return {'rows': rows[(number - 1) * size:number * size], 'total': len(rows), 'page': number, 'pages': pages,
            'amount': round(math.fsum(b['open'] for b in rows), 2),
            'months': [dict(month=k, **_round(v)) for k, v in sorted(months.items(), reverse=True)]}


def hint_text(bill):
    if bill['hint'] == 'pair_rest': return '%s %.2f' % (HINTS['pair_rest'], bill['pair_net'])
    return HINTS.get(bill['hint'], '')


def export(snapshot, rows, side, meta):
    """导出 Excel：汇总 + 当前筛选下的明细。只写模式，几万行不吃内存。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    def head(ws, values):
        from openpyxl.cell import WriteOnlyCell
        cells = []
        for v in values:
            c = WriteOnlyCell(ws, value=v); c.font = bold; cells.append(c)
        ws.append(cells)
    s = summary(snapshot)
    ws = wb.create_sheet('汇总')
    ws.append(['金蝶未核销清单（电商）', '', '同步时间', (meta or {}).get('ts', ''), '同步人', (meta or {}).get('operator', '')])
    ws.append(['未核销金额 = 应收金额 − 已核销金额；红字应收为负数。只读金蝶，不改任何单据。'])
    ws.append([])
    head(ws, ['店铺（金蝶客户）', '合计张数', '合计未核销', *[x for k in PILES for x in (PILES[k] + '张数', PILES[k] + '未核销')]])
    for r in s['shops']:
        ws.append([r['shop'], r['total']['count'], r['total']['amount'], *[x for k in PILES for x in (r[k]['count'], r[k]['amount'])]])
    ws.append(['合计', s['total']['count'], s['total']['amount'], *[x for k in PILES for x in (s['piles'][k]['count'], s['piles'][k]['amount'])]])
    ws.append([])
    head(ws, ['只靠金蝶就能判的提示', '张数', '未核销'])
    for k, v in s['hints'].items(): ws.append([HINTS[k], v['count'], v['amount']])
    ws.append([])
    head(ws, ['收款单未核销·往来单位', '张数', '未核销'])
    for r in s['receipts']['shops']: ws.append([r['shop'], r['total']['count'], r['total']['amount']])
    ws = wb.create_sheet(SIDES[side] + '明细')
    if side == 'rec':
        head(ws, ['收款单号', '日期', '往来单位', '收款金额', '已核销', '未核销', '核销状态', '单据状态', '结算方式', '收款账户', '哪一堆', '提示'])
        for b in rows:
            ws.append([b['no'], b['date'], b['shop'], b['amount'], b['written'], b['open'], WS.get(b['ws'], b['ws']), DS.get(b['ds'], b['ds']),
                       b.get('settle', ''), b.get('account', ''), PILES[b['pile']], hint_text(b)])
    else:
        head(ws, ['应收单号', '业务日期', '店铺（金蝶客户）', '应收金额', '已核销', '未核销', '核销状态', '单据状态', '平台订单号', '哪一堆', '提示'])
        for b in rows:
            ws.append([b['no'], b['date'], b['shop'], b['amount'], b['written'], b['open'], WS.get(b['ws'], b['ws']), DS.get(b['ds'], b['ds']),
                       b['order'], PILES[b['pile']], hint_text(b)])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()
