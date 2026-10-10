# V2.811: 抖音月结——两份动账明细的解析、与金蝶应收逐单对、收款单草稿。纯计算：不连金蝶、不碰库。
import calendar
import codecs
import csv
import datetime
import html
import io
import math
import re
import unicodedata
import zipfile
from collections import defaultdict

KINDS = {'dy_settle': '抖音动账明细（订单维度）', 'dy_ledger': '抖音账户流水（带余额）', 'dy_orders': '旺店通订单明细（认合单）',
         'dy_platform': '抖音平台订单明细（抖店订单导出）', 'dy_insure': '抖音运费险保费明细（逐单）',
         'dy_returns': '旺店通退换单（退款不退货等）'}
INSURE_SCENE = '退换货运费险'               # 账户流水里这一种，一笔流水是一批保单；逐单的在同一个压缩包的「保费支出」里
SETTLE = '货款结算入账'
# 结算时从货款里直接扣掉的费用列（文件里是负数）
FEE_COLUMNS = ['平台服务费', '佣金', '服务商佣金', '渠道分成', '招商服务费', '站外推广费', '其他分成']
SUBSIDY_COLUMNS = ['实际平台补贴_运费', '实际平台补贴', '其他平台补贴', '以旧换新抵扣', '政府补贴平台垫资', '实际达人补贴', '实际抖音支付补贴', '实际抖音月付营销补贴', '银行补贴']
FORMAT = 4                                  # 解析出的字段变了就加一：让同一份文件重新上传时能覆盖旧快照
MAX_ROWS = 1000000                          # 一份资料最多这么多行：平台导出一个月几万行，超过这个数不像导出的原始文件
MAX_UNZIPPED = 150 << 20                    # 压缩包里单个 csv 解开后的上限（真实最大的一份解开 26 MB）
MAX_ZIP_FILES = 20                          # 一个压缩包里最多处理几个 csv
_TIME = re.compile(r'20\d{2}-\d{2}-\d{2}')    # 平台原始导出的时间都是这个开头；被 Excel 另存过会变成 2026/9/22
_SETTLE_HEAD = {'动帐流水号', '动账方向', '动账金额', '动账场景', '订单号', '订单实付应结', '订单退款', '平台服务费'}
_LEDGER_HEAD = {'动账流水号', '账户方向', '动账金额(元)', '动账场景', '账户余额(元)'}
_INSURANCE_HEAD = {'保险单号', '动账流水号', '金额(元)'}
_INSURE_FULL = _INSURANCE_HEAD | {'关联子订单号', '动账时间'}        # 列齐了才逐单接入；不齐的照旧跳过
# 旺店通「退换管理」导出：退款不退货没有退货入库，金蝶不会自动出红字——这张表说的是"为什么没红字"
_RETURNS_HEAD = {'退换单号', '类型', '退款阶段', '退换原因', '店铺', '原始单号', '原始子订单号', '货品编号', '入库数量', '分摊退款金额', '登记时间'}
_ORDERS_HEAD = {'订单编号', '店铺', '子单原始单号', '订单状态', '应收金额', '分摊后总价'}
_PLATFORM_HEAD = {'主订单编号', '子订单编号', '订单应付金额', '订单状态', '售后状态', '订单提交时间', '平台实际承担优惠金额'}
# 结算后把钱退给买家 / 退回补贴：冲的是应收；其余结算后场景（分账、退分账）是费用的返还
REFUND_SCENES = {'退款-结算后退款-退用户', '退款-订单退款触发-退补贴'}
CATEGORIES = {'ok': '钱已到账·金额一致，可核销', 'mismatch': '钱已到账·金额对不上', 'pair': '结算前已退款·红字蓝字互冲即平',
              'overdue': '发货超过 {days} 天还没结算', 'transit': '在途·月底前还没结算', 'no_order': '应收单上没有平台订单号'}
BUCKETS = {'ok': '金额一致', 'refunded': '先结算后又退款（两边实际一致）', 'mismatch': '金额对不上', 'written': '应收已核销过', 'missing': '金蝶里没有这单的应收'}


def _text(value):
    return str(value or '').strip().lstrip("'`").strip()


def _num(value):
    try: return float(str(value or 0).replace(',', '') or 0)
    except (TypeError, ValueError): return 0.0


def _decode(blob):
    for encoding in ('utf-8-sig', 'gb18030'):
        try: return blob.decode(encoding)
        except UnicodeDecodeError: continue
    raise ValueError('文件编码无法识别')


def _encoding(blob):
    """分块把整份验一遍，返回能读通的编码；不把整份转成字符串留在内存里。"""
    for encoding in ('utf-8-sig', 'gb18030'):
        check = codecs.getincrementaldecoder(encoding)()
        try:
            for i in range(0, len(blob), 1 << 20): check.decode(blob[i:i + (1 << 20)])
            check.decode(b'', final=True)
            return encoding
        except UnicodeDecodeError: continue
    raise ValueError('文件编码无法识别')


def _tables(blob, filename):
    """上传的可能是 csv，也可能是装着多个 csv 的 zip。逐个产出 (文件名, 行字典列表, 表头)。"""
    if blob[:4] == b'\xd0\xcf\x11\xe0': raise ValueError('「%s」是老版 Excel（.xls），系统读不了；请传平台导出的原始文件' % filename)
    if zipfile.is_zipfile(io.BytesIO(blob)):
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            found = 0
            for info in z.infolist():
                if info.is_dir() or not info.filename.lower().endswith('.csv'): continue
                try: name = info.filename.encode('cp437').decode('gbk')
                except (UnicodeEncodeError, UnicodeDecodeError): name = info.filename
                if info.flag_bits & 1: raise ValueError('压缩包带密码，系统打不开；请先解压，把里面的 csv 传上来')
                if info.file_size > MAX_UNZIPPED: raise ValueError('压缩包里的「%s」解开超过 %d MB，不像平台导出的文件' % (name.rsplit('/', 1)[-1], MAX_UNZIPPED >> 20))
                found += 1
                if found > MAX_ZIP_FILES: raise ValueError('压缩包里的 csv 超过 %d 个，不像平台导出的文件' % MAX_ZIP_FILES)
                yield from _tables(z.read(info), name.rsplit('/', 1)[-1])
        return
    # 不整表读进内存（26MB 的订单导出整表读要吃 400 多 MB）：把读取器交出去，边读边转成要用的那几列
    reader = csv.DictReader(io.TextIOWrapper(io.BytesIO(blob), encoding=_encoding(blob), newline=''))
    head = set(reader.fieldnames or [])
    yield filename, reader, head


def _settle_row(r):
    sign = 1 if _text(r.get('动账方向')) == '入账' else -1
    fees = {c: round(-_num(r.get(c)), 2) for c in FEE_COLUMNS if _num(r.get(c))}
    return {'id': _text(r.get('动帐流水号')), 't': _text(r.get('动账时间')), 'amt': round(sign * _num(r.get('动账金额')), 2),
            'scene': _text(r.get('动账场景')), 'btype': _text(r.get('计费类型')), 'order': _text(r.get('订单号')),
            'refund': round(_num(r.get('订单退款')), 2), 'fees': fees, 'memo': _text(r.get('备注'))[:60],
            # V2.816 订单全链路要用：下单时间、商品、买家实付、各项补贴合计（实付 + 补贴 = 结算前的订单金额；运费不随货款结算，不计入）
            'ot': _text(r.get('下单时间')), 'goods': _text(r.get('商品名称'))[:60],
            'paid': round(_num(r.get('订单实付应结')), 2),
            'subsidy': round(math.fsum(_num(r.get(c)) for c in SUBSIDY_COLUMNS), 2)}


def _ledger_row(r):
    return {'id': _text(r.get('动账流水号')), 't': _text(r.get('动账时间')), 'amt': round(_num(r.get('动账金额(元)')), 2),
            'scene': _text(r.get('动账场景')), 'order': _text(r.get('关联订单号')), 'bal': round(_num(r.get('账户余额(元)')), 2),
            'memo': _text(r.get('备注'))[:60]}


def _insure_row(r):
    """账户流水压缩包里的「保费支出」：一张保单一行，带着它所属的那笔账户流水和关联的子订单。文件里金额是正数（支出），这里记成负数，和账户流水同向。"""
    return {'id': _text(r.get('保险单号')), 'flow': _text(r.get('动账流水号')), 'order': _text(r.get('关联子订单号')), 't': _text(r.get('动账时间')),
            'amt': round(-_num(r.get('金额(元)')), 2), 'memo': _text(r.get('摘要描述'))[:30]}


def _cell(value):
    """旺店通 csv 把长数字和时间包成 ="…"（防 Excel 改格式），取里面的。"""
    text = _text(value)
    return text[2:-1].strip() if text.startswith('="') and text.endswith('"') else text


def _return_row(r, shop=''):
    """旺店通退换单：一张退换单一个子订单一种货品一行。只留对账用的列；客户网名、地址、各种备注不要。
    金额用「分摊退款金额」：一张退换单分几行时，「退款总额」每行都重复写整张单的数。别的店的行、末尾的合计行不要。"""
    tk = _cell(r.get('退换单号'))
    if tk in ('', 'NA') or (shop and _cell(r.get('店铺')) != shop): return {'id': '', 't': ''}
    sub = _cell(r.get('原始子订单号'))
    return {'id': '%s|%s|%s' % (tk, sub, _cell(r.get('货品编号'))), 'tk': tk, 't': _cell(r.get('登记时间')), 'order': _platform_order(_cell(r.get('原始单号'))), 'sub': sub,
            'type': _cell(r.get('类型')), 'stage': _cell(r.get('退款阶段')), 'why': _cell(r.get('退换原因'))[:30], 'state': _cell(r.get('处理状态')),
            'pstate': _cell(r.get('平台退款状态')), 'goods': _cell(r.get('货品名称'))[:40], 'qty': _num(_cell(r.get('登记数量'))), 'back': _num(_cell(r.get('入库数量'))),
            'amt': round(_num(_cell(r.get('分摊退款金额'))), 2), 'done': _cell(r.get('退款成功时间'))}


def _platform_order(value):
    """旺店通的「子单原始单号」→ 平台订单号：去掉补发单的 A 后缀和拆单的 -1。不是平台订单号（手工单等）返回空串。"""
    text = _text(value)
    for _ in range(2):
        if text[-1:] in ('A', 'a'): text = text[:-1]
        head, _, tail = text.rpartition('-')
        if head and tail.isdigit() and len(tail) <= 2: text = head
    return text if text.isdigit() and len(text) >= 15 else ''


_CELL = re.compile(rb'<c r="([A-Z]+)\d+"([^>]*?)(?:/>|>(?:<is><t[^>]*>([^<]*)</t></is>|<v>([^<]*)</v>)?</c>)')
_SHARED = re.compile(rb'<si>(.*?)</si>', re.S)
_TEXT = re.compile(rb'<t[^>]*>([^<]*)</t>')


def _xlsx_rows(blob, must=b''):
    """流式读 xlsx 第一张表：按 </row> 切块，每行用正则取单元格，返回 {列序号: 文本}。
    旺店通一个月的订单明细解开有 200 多 MB，用 openpyxl 逐格读要 7 分钟，这样读约 15 秒。
    旺店通原样导出的是内联字符串；用 Excel 另存过的是共享字符串，两种都认。
    must 只对内联字符串的文件有用：给了就只要含这串字节的行（表头行总是给）。"""
    z = zipfile.ZipFile(io.BytesIO(blob))
    shared = []
    if 'xl/sharedStrings.xml' in z.namelist():
        shared = [html.unescape(b''.join(_TEXT.findall(m.group(1))).decode('utf-8')) for m in _SHARED.finditer(z.read('xl/sharedStrings.xml'))]
    if shared: must = b''
    sheets = sorted(n for n in z.namelist() if n.startswith('xl/worksheets/') and n.endswith('.xml'))
    if not sheets: raise ValueError('文件不是 Excel 表格（里面没有工作表）')
    sheet = sheets[0]
    buf, first = b'', True
    with z.open(sheet) as f:
        while True:
            chunk = f.read(4 << 20)
            if not chunk: break
            buf += chunk; parts = buf.split(b'</row>'); buf = parts.pop()
            for part in parts:
                if first and b'&#' in part: must = b''                # 有的导出把汉字写成 &#数字; 实体，按字节筛行会全筛掉
                if not first and must and must not in part: continue
                first = False
                row = {}
                for m in _CELL.finditer(part):
                    n = 0
                    for ch in m.group(1): n = n * 26 + (ch - 64)
                    if b't="s"' in m.group(2): row[n - 1] = shared[int(m.group(4))] if m.group(4) else ''
                    else: row[n - 1] = html.unescape((m.group(3) or m.group(4) or b'').decode('utf-8'))
                yield row


def _orders(blob, filename, shop):
    """旺店通订单明细（xlsx）。只取核对要用的列，收件人、电话、地址等一概不留。
    一个「旺店通订单 × 平台订单」汇成一条：金蝶应收是照旺店通订单开的，一张里可能合了几个平台订单。"""
    it = _xlsx_rows(blob, html.escape(shop).encode('utf-8') if shop else b'')
    first = next(it, None) or {}
    at = {_text(name): i for i, name in first.items()}
    if not _ORDERS_HEAD <= set(at): raise ValueError('「%s」表头不是旺店通订单明细' % filename)
    cell = lambda row, k: row.get(at[k]) if k in at else None
    when = lambda row, k: _text(cell(row, k))[:19] if _text(cell(row, k))[:2] == '20' else ''
    agg = {}
    for i, row in enumerate(it):
        if shop and _text(cell(row, '店铺')) != shop: continue
        jy, order = _text(cell(row, '订单编号')), _platform_order(cell(row, '子单原始单号'))
        t = when(row, '交易时间') or when(row, '付款时间')
        if not jy or not order or not t: continue
        r = agg.get(jy + '|' + order)
        if r is None:
            r = agg[jy + '|' + order] = {'id': jy + '|' + order, 'jy': jy, 'order': order, 't': t, 'n': i, 'status': _text(cell(row, '订单状态')),
                                         'refund': _text(cell(row, '订单退款状态')), 'recv': round(_num(cell(row, '应收金额')), 2),
                                         'paid': 0.0, 'ship': when(row, '发货时间'), 'amount': 0.0, 'qty': 0.0, 'goods': []}
        r['amount'] += _num(cell(row, '分摊后总价')); r['qty'] += _num(cell(row, '实发数量'))
        r['paid'] += _num(cell(row, '买家实付'))                                          # 买家实付在明细里是按货品行分摊的（带三位小数），逐行加完再取整
        goods = _text(cell(row, '货品名称'))[:30]
        if goods and goods not in r['goods'] and len(r['goods']) < 4: r['goods'].append(goods)
    if not agg:
        raise ValueError(('「%s」里没有「%s」的订单，请确认导出时包含了这家店' % (filename, shop)) if shop else '「%s」里没有带平台订单号的订单' % filename)
    for r in agg.values(): r['amount'], r['paid'] = round(r['amount'], 2), round(r['paid'], 2)
    return list(agg.values())


_PHONE = re.compile(r'(?<!\d)(?:\+?86[- ]?)?1[3-9]\d(?:[-. ]?\d){8}(?!\d)|(?<!\d)0\d{2,3}[- ]?\d{7,8}(?!\d)|(?<!\d)\d{17}[\dXx](?!\d)')
_PLACE = re.compile(r'[\u4e00-\u9fa5]{2}(?:省|自治区)|[\u4e00-\u9fa5]{2}市[\u4e00-\u9fa5]{1,12}(?:区|县|镇)|(?:路|街|巷|弄|道)\d+号|\d+(?:栋|幢|单元|室|号楼)')
_PRIVATE_WORDS = ('地址', '收件', '收货', '电话', '手机', '改址', '门牌', '姓名', '改寄', '转寄', '寄到', '发到', '微信', '身份证')
MEMO_HIDDEN = '（备注里有收货信息，未保留）'


def _memo(value):
    """商家备注是客服手写的：退款原因、补发单号有用，但改地址时会把买家的电话、地址、姓名写进去。
    先拿整段判断、再截断（先截的话号码可能正好被切断认不出）；有一点迹象就整条不留。"""
    text = _text(value)
    if not text or text == '-': return ''
    seen = unicodedata.normalize('NFKC', text[:300])                  # 全角数字转成半角再判断；只存前 80 字，看前 300 字足够
    if _PHONE.search(seen) or _PLACE.search(seen) or any(w in seen for w in _PRIVATE_WORDS): return MEMO_HIDDEN
    return text[:80]


def _platform_row(r):
    """抖店后台「订单导出」的一行＝一个子订单。只留对账用的列；收件人、电话、地址、买家留言一概不留，
    商家备注里带电话、地址迹象的整条不留（见 _memo）。"""
    text = lambda k, n=40: '' if _text(r.get(k)) == '-' else _text(r.get(k))[:n]
    money = lambda k: round(_num(_text(r.get(k))), 2)
    return {'id': text('子订单编号'), 'order': text('主订单编号'), 't': text('订单提交时间'), 'goods': text('选购商品', 60), 'spec': text('商品规格'),
            'code': text('商家编码'), 'qty': _num(_text(r.get('商品数量'))), 'pay': money('订单应付金额'), 'plat': money('平台实际承担优惠金额'),
            'kol': money('达人实际承担优惠金额'), 'shop': money('商家实际承担优惠金额'), 'freight': money('运费'),
            'status': text('订单状态'), 'after': text('售后状态'), 'ship': text('发货时间'), 'done': text('订单完成时间'),
            'cancel': text('取消原因', 30), 'memo': _memo(r.get('商家备注'))}


def export_time(filename):
    """抖店导出的 csv 文件名以导出时刻（10 位秒数）开头，取出来写成北京时间；认不出返回空串。"""
    m = re.match(r'(1[6-9]\d{8})_', _text(filename).rsplit('/', 1)[-1])
    if not m: return ''
    return datetime.datetime.fromtimestamp(int(m.group(1)), datetime.timezone(datetime.timedelta(hours=8))).strftime('%Y-%m-%d %H:%M')


_AFTER_DONE = {'售后关闭', '售后已拒绝', '换货成功', '补寄成功'}      # 售后已经办完、不退钱的：不影响结算
_AFTER_OPEN = {'待收退货', '待退货', '售后待处理'}                    # 售后还在办的


def _unshipped(r):
    """没发货就退款关闭的子订单：金蝶不给它开应收，它的售后也不该算到这张应收头上。"""
    return r['status'] == '已关闭' and not r['ship']


def platform_index(rows):
    """平台订单明细按主订单汇总。gross＝下单时这单值多少钱（买家应付 ＋ 平台、达人承担的优惠）；unshipped＝其中没发货就关闭的子订单。
    金蝶蓝字照发货开，所以该等于 gross，或等于 gross − unshipped（2026 年 8、9 月实测：没退款的 16,599 笔结算，
    动账明细的 实付＋补贴 与导出逐个子订单相等；按订单看，不等的都是有子订单没发货就关闭的）。
    状态、售后只看发过货的子订单；整单都没发货才看全部。x＝这单的状态是哪次导出时的（各子订单里最早的那次）。"""
    out = {}
    for r in rows or []:
        o = out.setdefault(r['order'], {'t': r['t'], 'gross': 0.0, 'pay': 0.0, 'subsidy': 0.0, 'unshipped': 0.0, 'freight': 0.0,
                                        'status': [], 'after': [], 'done': '', 'ship': '', 'x': r.get('x', ''), 'subs': []})
        value = r['pay'] + r['plat'] + r['kol']
        o['gross'] += value; o['pay'] += r['pay']; o['subsidy'] += r['plat'] + r['kol']; o['freight'] += r.get('freight', 0)
        if _unshipped(r): o['unshipped'] += value
        o['t'] = min(o['t'], r['t']); o['x'] = min(o['x'], r.get('x', '')); o['subs'].append(r)
    for o in out.values():
        live = [r for r in o['subs'] if not _unshipped(r)] or o['subs']
        live = [r for r in live if r['pay'] + r['plat'] + r['kol'] > 0] or live       # 有带金额的子订单时，0 元子订单（赠品）不参与判断状态
        for r in live:
            if r['status'] not in o['status']: o['status'].append(r['status'])
            if r['after'] and r['after'] not in o['after']: o['after'].append(r['after'])
            o['done'] = max(o['done'], r['done']); o['ship'] = max(o['ship'], r['ship'])
        for k in ('gross', 'pay', 'subsidy', 'unshipped', 'freight'): o[k] = round(o[k], 2)
    return out


def platform_closed(o):
    return bool(o) and set(o['status']) == {'已关闭'}


def platform_state(o, end):
    """一句话说这单在平台上什么情况——给还没结算的蓝字应收看：这笔钱还会不会来。
    状态是订单导出那一刻的，不是月底的；知道是哪天导出的就写在句尾。"""
    if not o: return ''
    status, after = set(o['status']), o['after']
    tail = '（按 %s 的订单导出）' % o['x'][5:10] if o.get('x') else ''
    if status == {'已关闭'}:
        why = '（%s）' % '、'.join(after) if after else ''
        if o.get('freight') and o['ship']: return '平台订单已关闭%s，货款不会结算，运费 %.2f 元可能照结%s' % (why, o['freight'], tail)
        return '平台订单已关闭%s，这笔钱不会结算了%s' % (why, tail)
    if '退款成功' in after: return '平台上有退款成功的售后，会少结或不结' + tail
    doing = [a for a in after if a in _AFTER_OPEN]
    other = [a for a in after if a not in _AFTER_OPEN and a not in _AFTER_DONE]
    if doing: return '平台上售后还在处理（%s）%s' % ('、'.join(doing), tail)
    if other: return '平台上有售后记录（%s）%s' % ('、'.join(other), tail)      # 不认识的售后状态：不猜它办没办完
    if status <= {'已完成', '已关闭'} and o['done']:
        return '买家 %s 确认收货%s，等平台结算%s' % (o['done'][5:10], '（月底时还没确认）' if o['done'][:10] > end else '', tail)
    if '已发货' in status: return '买家还没确认收货' + tail
    return '平台订单状态：' + '、'.join(o['status']) + tail


def platform_states(index, orders, end, group=0):
    """几个订单各自在平台上的情况拼成一句。group＝这组（合单）一共几个订单：
    说法不止一种、或说到的只是合单里的一部分时，句子前带订单尾号，免得读成整张应收都是这个情况。
    合单里下单金额是 0 的订单（赠品单）不会有结算，不说。"""
    if max(group, len(orders)) > 1: orders = [o for o in orders if not (index.get(o) and index[o]['gross'] == 0)]
    said = [(o, platform_state(index.get(o), end)) for o in orders]
    said = [(o, s) for o, s in said if s]
    if not said: return ''
    sentences = list(dict.fromkeys(s for _, s in said))
    if len(sentences) == 1 and len(said) >= max(group, 1): return sentences[0]
    return '；'.join('尾号 %s：%s' % ('、'.join(o[-4:] for o, x in said if x == s), s) for s in sentences)


def parse(blob, filename, shop=''):
    """只看表头认资料类型，不看文件名。返回 [{name, kind, rows|skip}]；认不出抛 ValueError。
    shop 只对旺店通订单明细有用：那份文件是全店铺的，只取这家店的行。"""
    if zipfile.is_zipfile(io.BytesIO(blob)) and '[Content_Types].xml' in zipfile.ZipFile(io.BytesIO(blob)).namelist():
        return [{'name': filename, 'kind': 'dy_orders', 'rows': _orders(blob, filename, shop)}]
    out = []
    for name, rows, head in _tables(blob, filename):
        if _SETTLE_HEAD <= head: kind, make = 'dy_settle', _settle_row
        elif _LEDGER_HEAD <= head: kind, make = 'dy_ledger', _ledger_row
        elif _PLATFORM_HEAD <= head: kind, make = 'dy_platform', _platform_row
        elif _INSURE_FULL <= head: kind, make = 'dy_insure', _insure_row
        elif _RETURNS_HEAD <= head: kind, make = 'dy_returns', lambda r: _return_row(r, shop)
        elif _INSURANCE_HEAD <= head:
            out.append({'name': name, 'kind': '', 'skip': '保费明细已含在账户流水里，不单独接入'}); continue
        else: raise ValueError('「%s」表头不是抖音动账明细、带余额的账户流水、抖店订单导出，也不是旺店通订单明细、退换单' % name)
        parsed, bad = [], 0
        for i, x in enumerate(map(make, rows)):                                      # n=文件内行序
            if i >= MAX_ROWS: raise ValueError('「%s」超过 %d 万行，不像平台导出的原始文件' % (name, MAX_ROWS // 10000))
            if not x['id'] or not x['t']: continue
            if not _TIME.match(x['t']) or 'E+' in x['id'].upper(): bad += 1; continue
            parsed.append(dict(x, n=i))
        # 有一行不对就整份不收：收一半的话，另一半会被存进一个不存在的月份，页面上还显示成功
        if bad: raise ValueError('「%s」有 %d 行的日期或单号被改过格式（像是用 Excel 打开后另存的），请传平台导出的原始文件' % (name, bad))
        if not parsed and kind == 'dy_insure':                                         # 这个月没买运费险：不算错，别连累同包的账户流水
            out.append({'name': name, 'kind': '', 'skip': '保费明细是空的'}); continue
        if not parsed and kind == 'dy_returns': raise ValueError('「%s」里没有%s的退换单' % (name, '「%s」' % shop if shop else '可用'))
        if not parsed: raise ValueError('「%s」里没有可用的流水行' % name)
        out.append({'name': name, 'kind': kind, 'rows': parsed})
    if not out: raise ValueError('压缩包里没有 csv 文件')
    return out


def snapshot_note(kind, rows):
    """数据准备清单上要看的两件事：这份资料盖到哪几天；动账明细是不是旧版解析的（缺买家实付 / 补贴，订单抽屉的算式要用）。"""
    days = sorted(r['t'][:10] for r in rows if r.get('t'))
    return {'span': [days[0], days[-1]] if days else [], 'stale': kind == 'dy_settle' and any('paid' not in r for r in rows)}


def by_month(rows):
    months = defaultdict(list)
    for r in rows: months[r['t'][:7]].append(r)
    return dict(months)


_PLATFORM_RANK = {'已支付': 1, '待发货': 1, '已发货': 2, '已完成': 3, '已关闭': 3}


def older(kind, old, new):
    """后传的这一行是不是比库里那行还旧（误传了一份更早的导出）。只凭看得出先后的迹象判断，看不出就当它是新的。"""
    if kind == 'dy_platform':
        if old.get('x') and new.get('x') and new['x'] < old['x']: return True
        return bool(old.get('done') and not new.get('done')) or _PLATFORM_RANK.get(new.get('status'), 9) < _PLATFORM_RANK.get(old.get('status'), 0)
    if kind == 'dy_orders':
        return (old.get('status') == '已取消' and new.get('status') != '已取消') or (old.get('status') == '已完成' and new.get('status') == '已发货')
    return False


def merge(old, new, kind='', skipped=None):
    """同一流水号只留一份（新文件覆盖旧的），按时间排好。
    订单类资料（kind 给 dy_platform / dy_orders）：新文件里明显比库里更旧的行不覆盖；skipped 给个列表就把没覆盖的号记进去。"""
    rows = {r['id']: r for r in old}
    for r in new:
        if kind and r['id'] in rows and older(kind, rows[r['id']], r):
            if skipped is not None: skipped.append(r['id'])
            continue
        rows[r['id']] = r
    return sorted(rows.values(), key=lambda r: (r['t'], r.get('n', 0), r['id']))


def chain(rows):
    """账户流水按余额链排好：同一秒内常有多笔，时间排不出先后，就按"上一笔余额 + 本笔金额 = 本笔余额"接龙。
    返回 (排好的行, 接不上的断点数)。"""
    out, breaks, prev, i = [], 0, None, 0
    while i < len(rows):
        j = i
        while j < len(rows) and rows[j]['t'] == rows[i]['t']: j += 1
        group = rows[i:j]; i = j
        if prev is None and len(group) > 1:                              # 开头：找没有前一笔的那行当起点
            ends = [round(r['bal'], 2) for r in group]
            start = next((r for r in group if round(r['bal'] - r['amt'], 2) not in ends), group[0])
            prev = round(start['bal'] - start['amt'], 2)
        while group:
            pick = next((r for r in group if prev is None or abs(r['bal'] - r['amt'] - prev) < 0.006), None)
            if pick is None: pick = group[0]; breaks += 1
            group.remove(pick); out.append(pick); prev = pick['bal']
    return out, breaks


def period_end(period):
    y, m = int(period[:4]), int(period[5:7])
    return '%s-%02d' % (period, calendar.monthrange(y, m)[1])


def _days_before(day, days):
    import datetime
    d = datetime.date(int(day[:4]), int(day[5:7]), int(day[8:10])) - datetime.timedelta(days=days)
    return d.isoformat()


def _gross(row):
    """这笔结算对应该冲掉多少应收 = 到账 + 结算时扣的费用（即 实付 + 各项补贴 − 结算时退款）。"""
    return row['amt'] + math.fsum(row['fees'].values())


def order_groups(orders):
    """旺店通会把同一买家的几个平台订单合成一张单发货，金蝶应收照合并后的单开、只记其中一个平台订单号。
    这里把"同一张旺店通订单里的平台订单"连成一组，对账按组对。返回 (取组函数, 订单→旺店通记录)。"""
    parent, info = {}, defaultdict(list)
    def find(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]; x = parent[x]
        return x
    by_jy = defaultdict(list)
    for r in orders or []:
        info[r['order']].append(r)
        if r['status'] != '已取消': by_jy[r['jy']].append(r['order'])
    for members in by_jy.values():
        for o in members[1:]: parent[find(o)] = find(members[0])
    members = defaultdict(set)
    for o in list(parent): members[find(o)].add(o)
    return (lambda o: members[find(o)] if o in parent else {o}), info


def _why(open_net, expected, opened, refund_at, after, tolerance):
    """金额对不上的订单，属于哪一种。只描述两边的事实，不替人下结论。"""
    red = any(b['open'] < 0 for b in opened); diff = round(open_net - expected, 2)
    if abs(diff) <= 1: return '尾差 %.2f' % diff
    if red and abs(open_net) <= tolerance and abs(refund_at) < 0.005: return '金蝶红字蓝字已全冲掉，抖音没退款、全额结了 %.2f' % expected
    if red and abs(open_net) <= tolerance: return '金蝶红字蓝字已全冲掉，抖音只退了 %.2f、仍结了 %.2f' % (abs(refund_at), expected)
    if not red and abs(refund_at) >= 0.005: return '抖音结算时退了 %.2f，金蝶没开红字' % abs(refund_at)
    if abs(after) >= 0.005: return '抖音结算后又有退款／调整 %.2f，和金蝶红字对不上' % after
    if red: return '金蝶红字金额和抖音退的不一样，差 %.2f' % diff
    return '金蝶应收和抖音结算金额不一样，差 %.2f' % diff


def reconcile(period, settle, ledger, bills, orders=None, overdue_days=10, tolerance=0.01):
    """period 月结：settle / ledger 为截至该期末的流水（可含更早月份），bills 为该店金蝶应收（全部状态），
    orders 为旺店通订单明细（可没有：没有就认不出合单，合单的订单会落到"对不上"）。"""
    end = period_end(period); late = _days_before(end, overdue_days)
    settle = [r for r in settle if r['t'][:7] <= period]
    ledger, breaks = chain([r for r in ledger if r['t'][:7] <= period])
    in_period = [r for r in settle if r['t'][:7] == period]
    group, _ = order_groups(orders)
    gid = lambda o: min(group(o))
    # 每组订单：本期结算应冲额、截至期末累计应冲额，以及依据（到账 / 扣费 / 结算时间 / 结算时退款）
    now, total, settled_orders = defaultdict(float), defaultdict(float), set()
    proof = {True: defaultdict(lambda: [0.0, 0.0, '', 0.0]), False: defaultdict(lambda: [0.0, 0.0, '', 0.0])}
    after = defaultdict(float); refunded = defaultdict(float)
    for r in settle:
        if not r['order']: continue
        g = gid(r['order'])
        if r['scene'] != SETTLE:
            after[g] += r['amt']
            if r['scene'] in REFUND_SCENES: refunded[g] += r['amt']
            continue
        total[g] += _gross(r); settled_orders.add(r['order'])
        this = r['t'][:7] == period
        if this: now[g] += _gross(r)
        for scope in ((True, False) if this else (False,)):
            p = proof[scope][g]; p[0] += r['amt']; p[1] += math.fsum(r['fees'].values()); p[2] = max(p[2], r['t']); p[3] += r['refund']
    # 金蝶应收：只看业务日期在期末及以前的
    bills = [b for b in bills if b['date'] <= end]
    by_group = defaultdict(list)
    for b in bills:
        if b['order']: by_group[gid(b['order'])].append(b)
    rows, cats = [], {k: {'count': 0, 'orders': 0, 'amount': 0.0} for k in CATEGORIES}
    seen, verdict = set(), {}
    for b in bills:
        if b['ws'] == 'C': continue
        g = gid(b['order']) if b['order'] else ''
        opened = [x for x in by_group[g] if x['ws'] != 'C'] if g else [b]
        open_net = round(math.fsum(x['open'] for x in opened), 2)
        members = group(b['order']) if b['order'] else set()
        expected, cash, fee, settled_at, reason = None, None, None, '', ''
        if not b['order']: cat = 'no_order'
        elif g in total:
            scope = g in now
            expected = round(now[g] if scope else total[g], 2)
            p = proof[scope][g]; cash, fee, settled_at = round(p[0], 2), round(p[1], 2), p[2]
            if abs(open_net - expected) <= tolerance: cat = 'ok'
            elif refunded[g] and abs(open_net - (expected + refunded[g])) <= tolerance:
                cat, reason = 'pair', '抖音先结算、后来又把钱退回去了，两边实际都是 %.2f' % open_net
            elif len(members) > 1 and any(o not in settled_orders for o in members):
                cat, reason = 'transit', '合单发货的 %d 个订单里还有没结算的' % len(members)
            else: cat, reason = 'mismatch', _why(open_net, expected, opened, p[3], after[g], tolerance)
        elif len(opened) > 1 and abs(open_net) <= tolerance and any(x['open'] < 0 for x in opened): cat = 'pair'
        else: cat = 'overdue' if min(x['date'] for x in opened) <= late else 'transit'
        verdict[g or b['no']] = cat
        rows.append(dict(no=b['no'], date=b['date'], order=b['order'], amount=b['amount'], written=b['written'], open=b['open'],
                         ws=b['ws'], ds=b['ds'], cat=cat, order_open=open_net, flow=expected, cash=cash, fee=fee, settled_at=settled_at,
                         diff=None if expected is None else round(open_net - expected, 2), reason=reason, merged=len(members) if len(members) > 1 else 0,
                         pending=sorted(o for o in members if o not in settled_orders) if len(members) > 1 else []))
        c = cats[cat]; c['count'] += 1; c['amount'] += b['open']
        if (cat, g or b['no']) not in seen: seen.add((cat, g or b['no'])); c['orders'] += 1
    for c in cats.values(): c['amount'] = round(c['amount'], 2)
    rows.sort(key=lambda r: (r['date'], r['no']), reverse=True)
    # 本期结算的订单，流水这边应冲多少、金蝶那边挂着多少——把两边的差逐桶摆出来（各桶流水之和 = 收款明细合计）
    buckets = {k: {'orders': 0, 'flow': 0.0, 'ar': 0.0} for k in BUCKETS}
    missing = []
    per_order = defaultdict(float)
    for r in in_period:
        if r['scene'] == SETTLE and r['order']: per_order[r['order']] += _gross(r)
    for g, flow in now.items():
        own = by_group.get(g, []); opened = [x for x in own if x['ws'] != 'C']
        ar = round(math.fsum(x['open'] for x in opened), 2)
        key = ('missing' if not own else 'written' if not opened else 'ok' if abs(ar - flow) <= tolerance
               else 'refunded' if verdict.get(g) == 'pair' else 'mismatch')
        members = [o for o in group(g) if o in per_order]
        b = buckets[key]; b['orders'] += len(members); b['flow'] += flow; b['ar'] += ar
        if key in ('missing', 'written'): missing += [{'order': o, 'flow': round(per_order[o], 2), 'bucket': key} for o in members]
    for b in buckets.values(): b['flow'] = round(b['flow'], 2); b['ar'] = round(b['ar'], 2)
    missing.sort(key=lambda r: -abs(r['flow']))
    # 收款单草稿：到账行 + 一种扣款一行。到账行优先用带余额的账户流水，没有则退回订单维度明细
    ledger_now = [r for r in ledger if r['t'][:7] == period]
    source = ledger_now or in_period
    net = round(math.fsum(r['amt'] for r in source), 2)
    lines = defaultdict(float)
    for r in in_period:
        if r['scene'] == SETTLE:
            for name, value in r['fees'].items(): lines[name] += value
    for r in source:
        if r['scene'] != SETTLE: lines[r['scene'] or '未注明场景'] -= r['amt']
    deductions = [{'name': k, 'amount': round(v, 2), 'kind': 'fee' if k in FEE_COLUMNS else 'other'}
                  for k, v in sorted(lines.items(), key=lambda x: -abs(x[1])) if round(v, 2)]
    draft = {'net': net, 'deductions': deductions, 'deduction_total': round(math.fsum(d['amount'] for d in deductions), 2),
             'from_ledger': bool(ledger_now)}
    draft['total'] = round(net + draft['deduction_total'], 2)
    draft['fee_total'] = round(math.fsum(d['amount'] for d in deductions if d['kind'] == 'fee'), 2)
    draft['other_total'] = round(math.fsum(d['amount'] for d in deductions if d['kind'] == 'other'), 2)
    draft['settle_cash'] = round(net + draft['other_total'], 2)          # 货款结算实际进账户的钱（平台已扣完费用）
    draft['settled_gross'] = round(math.fsum(now.values()), 2)          # 本期结算订单应冲额合计，应等于 total
    # 余额：期初 / 期末 / 链是否连续
    balance = None
    if ledger_now:
        balance = {'open': round(ledger_now[0]['bal'] - ledger_now[0]['amt'], 2), 'close': ledger_now[-1]['bal'],
                   'net': net, 'breaks': breaks, 'rows': len(ledger_now), 'first': ledger_now[0]['t'], 'last': ledger_now[-1]['t']}
    merged = {g for g in set(now) | set(by_group) if g and len(group(g)) > 1}
    return {'period': period, 'end': end, 'overdue_days': overdue_days, 'bills': rows, 'categories': cats, 'buckets': buckets,
            'missing': missing, 'draft': draft, 'balance': balance,
            'coverage': {'settle_rows': len(in_period), 'settled_orders': len(per_order), 'ledger_rows': len(ledger_now),
                         'settle_first': in_period[0]['t'] if in_period else '', 'settle_last': in_period[-1]['t'] if in_period else '',
                         'open_bills': len(rows), 'order_rows': len(orders or []), 'merged_groups': len(merged)}}


def order_detail(period, settle, ledger, bills, order, orders=None, platform=None):
    """一个订单的全部依据：抖音这边每一笔动账（到账、各项扣费、当时余额），金蝶那边每一张应收单。
    合单发货的，把同组的平台订单一起列出来——金蝶应收是按合并后的那张单开的。"""
    end = period_end(period)
    group, info = order_groups(orders)
    members = sorted(group(order))
    balance = {r['id']: r['bal'] for r in ledger}
    flows = []
    for r in settle:
        if r['order'] not in members or r['t'][:7] > period: continue
        settled = r['scene'] == SETTLE
        flows.append({'id': r['id'], 't': r['t'], 'order': r['order'], 'scene': r['scene'], 'btype': r.get('btype', ''), 'amt': r['amt'], 'fees': r['fees'],
                      'fee': round(math.fsum(r['fees'].values()), 2), 'refund': r['refund'], 'gross': round(_gross(r), 2) if settled else None,
                      'bal': balance.get(r['id']), 'memo': r['memo'], 'in_period': r['t'][:7] == period})
    own = sorted((b for b in bills if b['order'] in members and b['date'] <= end), key=lambda b: (b['date'], b['no']))
    opened = [b for b in own if b['ws'] != 'C']
    settled = [f for f in flows if f['gross'] is not None]
    raw = [r for r in settle if r['order'] in members and r['t'][:7] <= period and r['scene'] == SETTLE]
    known = bool(raw) and all('paid' in r for r in raw)                 # 旧版解析存的快照没有实付 / 补贴，要重新上传才有
    total = lambda rows, key: round(math.fsum(r[key] for r in rows), 2)
    wdt = {}
    for o in members:
        for r in info.get(o, []):
            w = wdt.setdefault(r['jy'], {'jy': r['jy'], 'status': r['status'], 'refund': r['refund'], 'recv': r['recv'], 'paid': 0.0, 'ship': r['ship'], 'orders': [], 'goods': []})
            w['orders'].append(o); w['goods'] += [x for x in r['goods'] if x not in w['goods']]
            w['paid'] = round(w['paid'] + r['paid'], 2)                 # 合单：这张旺店通单里几个平台订单的买家实付加起来（应收金额本来就是整张单的数）
    # 结算前退了一部分：平台补贴按同样比例收回，动账明细里的补贴是收回以后剩下的，不能直接和蓝字比。
    # 蓝字 − 实付 就是下单时的补贴；按退款比例折下来等于明细里的补贴，说明蓝字没错，该看的是红字冲了多少。
    # 平台订单明细（抖店订单导出）：下单时这单值多少、现在什么状态。合单的把同组订单都列上；有一个订单没在导出里就不给合计，免得拿半个数去比蓝字
    index = platform_index([r for r in platform or [] if r['order'] in members])
    whole = bool(index) and all(o in index for o in members)
    subs = [dict(r, order=o) for o in members for r in index.get(o, {}).get('subs', [])]
    p_gross = round(math.fsum(index[o]['gross'] for o in members), 2) if whole else None
    p_unshipped = round(math.fsum(index[o]['unshipped'] for o in members), 2) if whole else None
    back = None
    blue, red = total([b for b in own if b['amount'] > 0], 'amount'), total([b for b in own if b['amount'] < 0], 'amount')
    paid, left, refund = (total(raw, 'paid'), total(raw, 'subsidy'), abs(total(settled, 'refund'))) if known else (0, 0, 0)
    # 有平台订单明细时，「蓝字没错」先要和平台下单金额对得上（全部子订单，或去掉没发货就关闭的）；
    # 不然退款接近全额时，下面按比例的那条判断几乎什么蓝字都放行
    blue_fits = p_gross is None or abs(blue - p_gross) <= 0.02 or abs(blue - (p_gross - p_unshipped)) <= 0.02
    if known and refund > 0 and paid > 0 and blue > paid + left + 0.01 and blue_fits:
        before = round(blue - paid, 2)
        if abs(before * (1 - refund / paid) - left) <= 0.02 * len(raw) + 0.01:
            back = {'before': before, 'back': round(before - left, 2), 'blue': blue, 'red': red, 'keep': round(blue - refund - (before - left), 2),
                    'red_should': round(-(refund + before - left), 2)}
    # 平台状态只对「还没结算、金蝶还挂着没核销蓝字」的订单说：已结算的不说，只剩红字的也不说
    waiting = [o for o in members if o not in {f['order'] for f in settled}]
    priced = [o for o in waiting if not (index.get(o) and index[o]['gross'] == 0)] or waiting   # 合单里 0 元的赠品单不会有结算，判关没关不看它
    # 平台状态是订单导出那天的，应收只看到期末：月底发货、下月初退货的，红字记在下个月，这里单独带出来给卡片说明
    later = [b for b in bills if b['order'] in members and b['date'] > end and b['amount'] < 0]
    flow_total = total([f for f in settled if f['in_period']] or settled, 'gross')
    even = bool(settled) and abs(total(opened, 'open') - flow_total) <= 0.01      # 钱已到账、和未核销应收一分不差：没什么还欠着
    owed = any(b['ws'] != 'C' and b['open'] > 0 for b in own) and not even
    return {'order': order, 'members': members, 'wdt': sorted(wdt.values(), key=lambda w: w['jy']), 'flows': flows, 'bills': own, 'settled': bool(settled), 'subsidy_back': back,
            'platform': subs, 'platform_gross': p_gross, 'platform_unshipped': p_unshipped,
            'platform_pay': round(math.fsum(index[o]['pay'] for o in members), 2) if whole else None,
            'platform_state': platform_states(index, waiting, end, len(members)) if owed else '',
            'platform_closed': owed and bool(waiting) and all(platform_closed(index.get(o)) for o in priced),
            'later_red': total(later, 'amount'), 'later_red_at': max((b['date'] for b in later), default=''),
            'platform_at': min((index[o]['t'] for o in index), default=''), 'blue_total': blue,
            'unsettled': [o for o in members if o not in {f['order'] for f in settled}] if len(members) > 1 else [],
            'flow_total': flow_total,
            'cash_total': total(settled, 'amt'), 'fee_total': total(settled, 'fee'), 'refund_total': abs(total(settled, 'refund')),
            'after_total': total([f for f in flows if f['gross'] is None], 'amt'),
            'paid_total': total(raw, 'paid') if known else None, 'subsidy_total': total(raw, 'subsidy') if known else None,
            'ordered_at': min((r['ot'] for r in raw if r.get('ot')), default=''), 'goods': list(dict.fromkeys(r['goods'] for r in raw if r.get('goods'))),
            'settled_at': max((f['t'] for f in settled), default=''), 'btype': next((f['btype'] for f in settled if f['btype']), ''),
            'ar_total': total(own, 'amount'), 'written_total': total(own, 'written'), 'open_total': total(opened, 'open')}


# 订单抽屉里应收单的分录明细：打开抽屉时现查金蝶（只读），不进快照。顺序就是 bill_lines 拆行的顺序。
BILL_LINE_KEYS = ['FBillNo', 'FBillTypeID.FName', 'F_ora_Text3', 'F_ora_Text4', 'FCreatorId.FName', 'FCreateDate', 'FAPPROVERID.FName', 'FAPPROVEDATE',
                  'FMATERIALID.FNumber', 'FMATERIALID.FName', 'FPRICEUNITID.FName', 'FPriceQty', 'FTaxPrice', 'FEntryTaxRate',
                  'FNoTaxAmountFor_D', 'FTAXAMOUNTFOR_D', 'FALLAMOUNTFOR_D', 'FSourceBillNo', 'FSOURCETYPE']
SOURCE_TYPES = {'SAL_OUTSTOCK': '销售出库单', 'SAL_RETURNSTOCK': '销售退货单'}


def bill_lines(rows):
    """金蝶应收单分录（列序同 BILL_LINE_KEYS）→ {应收单号: 单据类型、旺店通单号、生成人、审核人 + 各行物料／数量／单价／税／来源单据}。"""
    when = lambda v: _text(v).replace('T', ' ')[:16]
    out = {}
    for r in rows or []:
        no, kind, t3, t4, creator, created, approver, approved, code, name, unit, qty, price, rate, net, tax, amount, src, src_type = (list(r) + [None] * 19)[:19]
        bill = out.setdefault(_text(no), {'type': _text(kind), 'jy': next((x for x in (_text(t3), _text(t4)) if x.startswith('JY')), ''),
                                          'creator': _text(creator), 'created': when(created), 'approver': _text(approver), 'approved': when(approved), 'lines': []})
        bill['lines'].append({'code': _text(code), 'name': _text(name), 'unit': _text(unit), 'qty': round(_num(qty), 4), 'price': round(_num(price), 4), 'rate': round(_num(rate), 2),
                              'net': round(_num(net), 2), 'tax': round(_num(tax), 2), 'amount': round(_num(amount), 2),
                              'src': _text(src), 'src_type': SOURCE_TYPES.get(_text(src_type), _text(src_type))})
    return out


def flow_rows(period, ledger, settle, insure, scene, known=(), sub2main=None):
    """「账户进出汇总」里某一项（货款结算以外的一种进出）的逐笔明细。返回 (哪种明细, 行)。
    保费类的：这一项的每笔账户流水都配得上保费明细、金额也相等，就按保单逐单列（一张保单对一个子订单）；其余按账户流水逐笔列，没有账户流水时退回订单维度动账明细。
    保单跟着它所在的那笔账户流水走，不看保单自己的时间和摘要：平台改过场景名（权益保险 → 退换货运费险），保单摘要却一律写运费险，按流水归才和汇总那一行相等。
    known＝认得的平台订单号（有结算或有应收的）：关联单号在里面才值得点开看订单。"""
    sub2main = sub2main or {}
    now = [r for r in ledger if r['t'][:7] == period]
    mine = [r for r in now or [r for r in settle if r['t'][:7] == period] if r['scene'] != SETTLE and (r['scene'] or '未注明场景') == scene]
    if now:
        ids = {r['id'] for r in mine}
        rows = [r for r in insure or [] if r['flow'] in ids]
        # 有流水没配上保单，或者保单加起来不等于流水：逐单列出来会和汇总对不上，整项退回按流水列
        if {r['flow'] for r in rows} != ids or abs(math.fsum(r['amt'] for r in rows) - math.fsum(r['amt'] for r in mine)) >= 0.005: rows = []
    else: rows = [r for r in insure or [] if scene == INSURE_SCENE and r['t'][:7] == period]
    if rows:
        main = lambda r: sub2main.get(r['order'], r['order'])
        return 'insure', [{'t': r['t'], 'id': r['id'], 'flow': r['flow'], 'order': main(r), 'amt': r['amt'], 'memo': r['memo'], 'known': main(r) in known}
                          for r in sorted(rows, key=lambda r: (r['t'], r['id']))]
    return 'ledger', [{'t': r['t'], 'id': r['id'], 'order': r['order'], 'amt': r['amt'], 'memo': r['memo'], 'bal': r.get('bal'), 'known': r['order'] in known} for r in mine]


def returns_index(rows, period):
    """平台订单 → 旺店通到本期末为止登记的退换单（合计、类型、原因、有没有退货入库）。"""
    out = {}
    for r in rows or []:
        if not r['order'] or r['t'][:7] > period: continue
        o = out.setdefault(r['order'], {'amt': 0.0, 'back': 0.0, 't': '', 'tks': [], 'types': [], 'why': []})
        o['amt'] = round(o['amt'] + r['amt'], 2); o['back'] += r.get('back') or 0; o['t'] = max(o['t'], r['t'])
        for k, v in (('tks', r['tk']), ('types', r['type']), ('why', r['why'])):
            if v and v not in o[k]: o[k].append(v)
    return out


def returns_notes(bills, index, orders=None):
    """给应收清单的每张蓝字写一句：旺店通登记了什么退换单、金蝶有没有对应的红字、要不要人补。只加说明（rnote），不改分类，不影响下推。
    合单发货的按同组订单一起看。没有退货入库又没有红字的，才说"要手工补"——有退货入库的金蝶会跟着入库单自动出红字。"""
    group, _ = order_groups(orders)
    members = {}
    for b in bills:
        if b['order']: members.setdefault(b['order'], sorted(group(b['order'])))
    blue, red = defaultdict(float), defaultdict(float)
    for b in bills:
        if not b['order']: continue
        key = members[b['order']][0]
        if b['amount'] > 0: blue[key] += b['amount']
        else: red[key] += b['amount']
    for b in bills:
        b['rnote'] = ''
        if not b['order'] or b['amount'] <= 0: continue
        hits = [index[o] for o in members[b['order']] if o in index]
        if not hits: continue
        key = members[b['order']][0]
        amt = round(math.fsum(h['amt'] for h in hits), 2); back = sum(h['back'] for h in hits)
        types = '、'.join(dict.fromkeys(t for h in hits for t in h['types'])); why = next((w for h in hits for w in h['why']), '')
        when = max(h['t'] for h in hits)[5:10]
        said = '旺店通 %s 登记了%s %.2f%s' % (when, types or '退换单', amt, '（%s）' % why if why else '')
        settled = b.get('flow') is not None
        if red[key] or back: tail = ''                                   # 已有红字，或者货退回来了（金蝶会跟着入库单出红字）：只说登记了什么
        elif not settled and amt >= blue[key] - 0.005: tail = '：全额退了、没有退货入库，这单不会再结算，金蝶也不会自动出红字，要手工补红字冲掉'
        elif not settled: tail = '：没有退货入库，金蝶不会自动出红字；结算时会少结这一块，到时要手工补红字'
        elif b['cat'] == 'mismatch' and (b.get('diff') or 0) > 0:
            tail = '：没有退货入库，金蝶不会自动出红字；要手工补红字 %.2f（就是两边的差额%s）' % (b['diff'], '，含平台按比例收回的补贴' if b['diff'] > amt + 0.005 else '')
        else: tail = '：结算时没扣这笔退款，金蝶也没有对应的红字；要是结算之后才退的，钱在「货款结算以外的账户进出」里，红字要另外补'
        b['rnote'] = said + tail


def push_export(shop, period, batches, bills):
    """推了什么：每张应收一行（源单编号、本次收款金额、平台订单号），另附每批收款单的收款明细。
    batches＝下推记录里还算数的批（没撤回的）；bills＝金蝶应收快照。只推整张蓝字，所以本次收款金额＝应收金额。"""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    wb = Workbook(write_only=True); bold = Font(bold=True)
    def head(ws, values):
        cells = []
        for v in values:
            c = WriteOnlyCell(ws, value=v); c.font = bold; cells.append(c)
        ws.append(cells)
    by_no = {b['no']: b for b in bills}
    ws = wb.create_sheet('下推清单')
    ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 26
    for col in 'DEF': ws.column_dimensions[col].width = 20
    head(ws, ['单据编号/源单编号', '本次收款金额', '旺店通原始单号/订单号', '应收业务日期', '下推时间', '操作人', '金蝶收款单内码'])
    count, total = 0, []
    for x in batches:
        for no in x['bills']:
            b = by_no.get(no, {})
            ws.append([no, b.get('amount'), b.get('order', ''), b.get('date', ''), x.get('at', ''), x.get('by', ''), x.get('fid')])
            count += 1; total.append(b.get('amount') or 0)
    ws.append(['合计 %d 张' % count, round(math.fsum(total), 2)])
    ws = wb.create_sheet('收款明细')
    ws.column_dimensions['A'].width = 22; ws.column_dimensions['B'].width = 14; ws.column_dimensions['C'].width = 14; ws.column_dimensions['D'].width = 60
    ws.append(['%s · %s 已下推的收款单：每张的收款明细（到账一行 + 一种扣款一行）' % (shop, period)])
    for x in batches:
        ws.append([])
        head(ws, ['%s 下推 · %s · %d 张应收 · 内码 %s' % (x.get('at', ''), x.get('by', ''), x.get('count') or len(x['bills']), x.get('fid'))])
        head(ws, ['行', '结算方式', '金额', '摘要'])
        for i, line in enumerate(x.get('lines') or [], 1):
            ws.append([i, '支付宝' if line.get('kind') == 'cash' else '内部转销', line.get('amount'), line.get('memo', '')])
        ws.append(['收款明细合计', '', round(math.fsum(l.get('amount') or 0 for l in x.get('lines') or []), 2), '应等于这一批应收合计 %.2f' % (x.get('total') or 0)])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def category_label(cat, days):
    return CATEGORIES[cat].format(days=days)


def select(result, cat='', q=''):
    q = _text(q).lower()
    hold = cat == 'hold'                                                # 可核销、但系统不下推的（订单带红字等）
    return [b for b in result['bills'] if (not cat or (b.get('hold') if hold else b['cat'] == cat)) and (not q or q in b['no'].lower() or q in b['order'])]


def page(rows, number=1, size=50):
    pages = max(1, -(-len(rows) // size)); number = min(max(1, number), pages)
    return {'rows': rows[(number - 1) * size:number * size], 'total': len(rows), 'page': number, 'pages': pages,
            'amount': round(math.fsum(b['open'] for b in rows), 2)}


def memo(period, name, amount):
    """扣款行摘要，照现行收款单的写法。"""
    return '%s年%s月结算单扣款项 %s%.2f元' % (period[:4], period[5:7], name, amount)


def export(result, shop, book=None):
    """导出 Excel：收款单草稿 + 各类应收清单（可核销的那页就是下推用的源单）。"""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    wb = Workbook(write_only=True); bold = Font(bold=True)
    def head(ws, values):
        cells = []
        for v in values:
            c = WriteOnlyCell(ws, value=v); c.font = bold; cells.append(c)
        ws.append(cells)
    days, d, period = result['overdue_days'], result['draft'], result['period']
    ws = wb.create_sheet('本月账户进出汇总')
    ws.append(['%s · %s 账户进出汇总（用来和账户余额核对，不是收款单）' % (shop, period)])
    ws.append([])
    head(ws, ['项目', '金额', '说明'])
    ws.append(['货款结算到账', d['settle_cash'], '平台已扣完费用后进账户的钱'])
    for x in d['deductions']:
        if x['kind'] == 'other': ws.append(['　' + x['name'], -x['amount'], '货款结算以外的账户进出，不在任何一张下推的收款单里'])
    ws.append(['本月账户净变动', d['net'], '应等于账户期末余额 − 期初余额'])
    ws.append([])
    head(ws, ['结算时平台直接扣掉的费用（不经过账户）', '金额'])
    for x in d['deductions']:
        if x['kind'] == 'fee': ws.append([x['name'], x['amount']])
    ws.append(['合计', d['fee_total']])
    ws.append([])
    head(ws, ['本月结算的订单', '订单数', '流水应冲应收', '金蝶挂着的应收', '两边差'])
    for k, label in BUCKETS.items():
        b = result['buckets'][k]; ws.append([label, b['orders'], b['flow'], b['ar'], round(b['flow'] - b['ar'], 2)])
    if result['balance']:
        bal = result['balance']; ws.append([])
        head(ws, ['账户余额', '流水', '金蝶账面'])
        ws.append(['期初', bal['open'], (book or {}).get('open')]); ws.append(['期末', bal['close'], (book or {}).get('close')])
        ws.append(['本月净变动', bal['net'], '']); ws.append(['余额链断点', bal['breaks'], ''])
    COLUMNS = ['应收单号', '业务日期', '平台订单号', '应收金额', '已核销', '未核销', '该订单(组)未核销合计', '抖音结算时间', '到账金额', '结算时扣费', '流水应冲额(到账+扣费)', '差额', '合单订单数', '情况说明']
    line = lambda b: [b['no'], b['date'], b['order'], b['amount'], b['written'], b['open'], b['order_open'], b['settled_at'], b['cash'], b['fee'], b['flow'], b['diff'], b['merged'] or '', b['reason']]
    todo = [b for b in result['bills'] if b['cat'] in ('mismatch', 'overdue', 'no_order')]
    if todo:
        ws = wb.create_sheet('交人工处理')
        ws.append(['系统判断不了、要人看的应收单：金额对不上的写明了两边各是什么情况；该结没结的要查买家是否一直没确认或已退款。'])
        head(ws, ['分类'] + COLUMNS)
        for b in sorted(todo, key=lambda b: (b['cat'], b['reason'][:8], b['date'])): ws.append([category_label(b['cat'], days)] + line(b))
    hold = [b for b in result['bills'] if b.get('hold')]
    if hold:
        ws = wb.create_sheet('可核销但系统不推')
        ws.append(['金额对得上，但系统不替你下推的应收单（订单带红字的要先红蓝对冲）。'])
        head(ws, ['原因'] + COLUMNS)
        for b in sorted(hold, key=lambda b: (b['hold'], b['order'], b['no'])): ws.append([PUSH_SKIP[b['hold']]] + line(b))
    titles = {'ok': '可核销(下推源单)', 'mismatch': '金额对不上', 'pair': '红蓝互冲', 'overdue': '该结没结', 'transit': '在途', 'no_order': '无订单号'}
    for cat, title in titles.items():
        rows = select(result, cat)
        if not rows: continue
        ws = wb.create_sheet(title)
        ws.append([category_label(cat, days)])
        head(ws, COLUMNS)
        for b in rows: ws.append(line(b))
    if result['missing']:
        ws = wb.create_sheet('流水有·应收对不上号')
        head(ws, ['平台订单号', '流水应冲额', '情况'])
        for m in result['missing']: ws.append([m['order'], m['flow'], BUCKETS[m['bucket']]])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ---------------- 下推收款单（V2.821）：只挑最干净的应收，分批备成金蝶暂存收款单 ----------------
PUSH_SKIP = {'red': '订单里有红字应收（要先红蓝对冲，留给人）', 'partial': '已经核销过一部分', 'unaudited': '应收单还没审核',
             'earlier': '钱是更早月份结算的（那个月的到账已经入过账，不能再算一次）', 'cents': '两边差几分钱'}


def push_lines(period, groups):
    """一张收款单的收款明细：到账一行 + 一种扣款一行（金额按分累加，不走浮点）。"""
    fees = defaultdict(int)
    for g in groups:
        for k, v in g['fees'].items(): fees[k] += v
    lines = [{'kind': 'cash', 'name': '到账', 'amount': sum(g['cash'] for g in groups) / 100, 'memo': ''}]
    return lines + [{'kind': 'fee', 'name': k, 'amount': v / 100, 'memo': memo(period, k, v / 100)} for k, v in sorted(fees.items(), key=lambda x: -x[1]) if v]


def push_split(period, groups, receipts):
    """金蝶一次下推可能自己拆成几张收款单（实测 16,070 张应收被拆成 2 张）。receipts 是每张收款单里的应收单号；
    给每张各算一份收款明细。同一个订单的几张应收被拆到不同收款单里就算不了，返回 None。"""
    where = {}
    for i, nos in enumerate(receipts):
        for no in nos: where[no] = i
    parts = [[] for _ in receipts]
    for g in groups:
        homes = {where.get(no) for no in g['bills']}
        if len(homes) != 1 or None in homes: return None
        parts[homes.pop()].append(g)
    if sum(len(g['bills']) for p in parts for g in p) != sum(len(nos) for nos in receipts): return None
    return [{'lines': push_lines(period, p), 'total': sum(g['open'] for g in p) / 100, 'count': sum(len(g['bills']) for g in p)} for p in parts]


def push_plan(period, settle, orders, result, taken=(), size=500):
    """从"可核销"里挑能直接下推的，排出下一批。只推整张蓝字应收、本次收款＝应收全额：
    订单（组）里不能有红字、不能核销过、必须已审核、两边分毫不差。
    返回 {eligible, skipped, pushed, batch}；batch 里是这一批的应收单号和收款明细（到账一行 + 一种扣款一行），批内两边相等。"""
    taken = set(taken)
    group, _ = order_groups(orders)
    gid = lambda o: min(group(o))
    by_group = defaultdict(list)
    for b in result['bills']:
        if b['cat'] == 'ok' and b['order']: by_group[gid(b['order'])].append(b)
    flows = defaultdict(list)
    for r in settle:
        if r['order'] and r['scene'] == SETTLE and r['t'][:7] == period: flows[gid(r['order'])].append(r)      # 只认本期结算的：每批的到账行是本期进账户的钱
    cents = lambda v: int(round(v * 100))
    eligible, skipped, held = [], {k: {'count': 0, 'amount': 0.0} for k in PUSH_SKIP}, {}
    for g, rows in by_group.items():
        why = ('red' if any(b['open'] < 0 for b in rows) else 'partial' if any(b['ws'] != 'A' or b['written'] for b in rows)
               else 'unaudited' if any(b['ds'] != 'C' for b in rows) else '')
        own = flows.get(g, [])
        if not why and not own: why = 'earlier'
        cash =sum(cents(r['amt']) for r in own); fee = sum(cents(v) for r in own for v in r['fees'].values())
        if not why and cash + fee != sum(cents(b['open']) for b in rows): why = 'cents'
        if why:
            skipped[why]['count'] += len(rows); skipped[why]['amount'] += math.fsum(b['open'] for b in rows)
            held.update({b['no']: why for b in rows}); continue
        eligible.append({'gid': g, 'bills': sorted(rows, key=lambda b: b['no']), 'flows': own, 'date': min(b['date'] for b in rows)})
    for v in skipped.values(): v['amount'] = round(v['amount'], 2)
    eligible.sort(key=lambda e: (e['date'], e['bills'][0]['no']))
    total = lambda es: round(math.fsum(b['open'] for e in es for b in e['bills']), 2)
    count = lambda es: sum(len(e['bills']) for e in es)
    left = [e for e in eligible if not any(b['no'] in taken for b in e['bills'])]
    batch, n = [], 0
    for e in left:
        if batch and n + len(e['bills']) > size: break
        batch.append(e); n += len(e['bills'])
    groups = []
    for e in batch:
        fees = defaultdict(int)
        for r in e['flows']:
            for k, v in r['fees'].items(): fees[k] += cents(v)
        groups.append({'bills': [b['no'] for b in e['bills']], 'open': sum(cents(b['open']) for b in e['bills']),
                       'cash': sum(cents(r['amt']) for r in e['flows']), 'fees': dict(fees)})
    lines = push_lines(period, groups)
    nos = [b['no'] for e in batch for b in e['bills']]
    return {'eligible': {'count': count(eligible), 'amount': total(eligible)}, 'skipped': skipped, 'held': held,
            'pushed': {'count': count(eligible) - count(left), 'amount': round(total(eligible) - total(left), 2)},
            'left': {'count': count(left), 'amount': total(left)},
            'batch': {'bills': nos, 'groups': groups, 'count': len(nos), 'total': total(batch), 'lines': lines, 'line_total': round(math.fsum(l['amount'] for l in lines), 2),
                      'first': nos[0] if nos else '', 'last': nos[-1] if nos else '',
                      'from': batch[0]['date'] if batch else '', 'to': max((e['date'] for e in batch), default='')}}
