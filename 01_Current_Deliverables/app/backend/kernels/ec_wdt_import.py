# -*- coding: utf-8 -*-
# [Change Log] Date: 2026-09-10 | Author: Codex | Version: V2.553
# Description: 旺店通销售出库只读导入，为天猫逐单核算补充企业内部履约来源。
"""旺店通销售出库导入：只生成月度业务汇总，不返回或持久化逐单/个人信息。"""
import datetime
import decimal
import io
import re
import zipfile

from openpyxl import load_workbook


MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_DATA_ROWS = 250_000

_PERIOD_RE = re.compile(r"^\d{4}-\d{2}$")
_FIELDS = {
    "order_no": "订单编号",
    "shipment_no": "出库单编号",
    "shop": "店铺",
    "warehouse": "仓库",
    "sku": "货品编号",
    "quantity": "货品数量",
    "order_pay": "订单支付金额",
    "receivable": "应收金额",
    "total_cost": "货品总成本",
    "ship_time": "发货时间",
}
_MONEY = decimal.Decimal("0.01")


class WdtImportError(ValueError):
    """可安全返回给前端的结构校验错误；消息中不得拼接单元格值。"""


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, decimal.Decimal)):
        return decimal.Decimal(str(value))
    text = str(value).strip().replace(",", "").replace("￥", "").replace("¥", "")
    if not text:
        return None
    if text.startswith("(") and text.endswith(")"):
        text = "-" + text[1:-1]
    try:
        return decimal.Decimal(text)
    except decimal.InvalidOperation:
        return None


def _money(value):
    number = _number(value)
    return number.quantize(_MONEY) if number is not None else None


def _date_text(value):
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    match = re.match(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
    if not match:
        return ""
    try:
        return datetime.date(*map(int, match.groups())).isoformat()
    except ValueError:
        return ""


def _preflight_xlsx(data):
    if not data or len(data) > MAX_FILE_BYTES:
        raise WdtImportError("文件为空或超过 25 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 2_000 or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
                raise WdtImportError("工作簿解压后体积过大")
            names = {item.filename for item in members}
            if "xl/workbook.xml" not in names:
                raise WdtImportError("文件不是有效的 XLSX 工作簿")
    except WdtImportError:
        raise
    except (zipfile.BadZipFile, OSError):
        raise WdtImportError("文件不是有效的 XLSX 工作簿")


def parse_wdt_sales_export(data, period):
    """按发货时间归属期间，返回不含订单号、店铺名和个人信息的汇总。"""
    if not _PERIOD_RE.match(period or ""):
        raise WdtImportError("结算期间格式应为 YYYY-MM")
    _preflight_xlsx(data)

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise WdtImportError("工作簿无法读取，请确认文件未损坏")

    try:
        sheet = workbook[workbook.sheetnames[0]]
        if sheet.max_row == 1 and sheet.max_column == 1:
            sheet.reset_dimensions()
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None) or ()
        headers = [_text(value) for value in first]
        missing = [label for label in _FIELDS.values() if label not in headers]
        if missing:
            raise WdtImportError("缺少必需业务列：" + "、".join(missing))
        index = {key: headers.index(label) for key, label in _FIELDS.items()}

        source_rows = in_period_rows = outside_period_rows = missing_ship_time_rows = 0
        missing_order_number_rows = cost_blank_rows = cost_nonzero_rows = 0
        goods_quantity = decimal.Decimal("0")
        cost_total = decimal.Decimal("0")
        orders = {}
        order_months = {}
        shipments, shops, warehouses, skus = set(), set(), set(), set()
        ship_date_min = ship_date_max = ""

        for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_number > MAX_DATA_ROWS + 1:
                raise WdtImportError("数据行超过 250,000 行上限")
            selected = {key: row[pos] if pos < len(row) else None for key, pos in index.items()}
            if not any(selected[key] not in (None, "") for key in ("order_no", "shipment_no", "sku")):
                continue
            source_rows += 1
            order_no = _text(selected["order_no"])
            ship_date = _date_text(selected["ship_time"])
            if not ship_date:
                missing_ship_time_rows += 1
                continue
            if order_no:
                order_months.setdefault(order_no, set()).add(ship_date[:7])
            if ship_date[:7] != period:
                outside_period_rows += 1
                continue
            in_period_rows += 1
            ship_date_min = min(filter(None, (ship_date_min, ship_date))) if ship_date_min else ship_date
            ship_date_max = max(ship_date_max, ship_date)

            if not order_no:
                missing_order_number_rows += 1
            else:
                order = orders.setdefault(order_no, {"lines": 0, "pay": set(), "receivable": set()})
                order["lines"] += 1
                pay = _money(selected["order_pay"])
                receivable = _money(selected["receivable"])
                if pay is not None:
                    order["pay"].add(pay)
                if receivable is not None:
                    order["receivable"].add(receivable)

            shipment_no = _text(selected["shipment_no"])
            shop = _text(selected["shop"])
            warehouse = _text(selected["warehouse"])
            sku = _text(selected["sku"])
            if shipment_no:
                shipments.add(shipment_no)
            if shop:
                shops.add(shop)
            if warehouse:
                warehouses.add(warehouse)
            if sku:
                skus.add(sku)

            quantity = _number(selected["quantity"])
            if quantity is not None:
                goods_quantity += quantity
            cost = _money(selected["total_cost"])
            if cost is None:
                cost_blank_rows += 1
            else:
                cost_total += cost
                if cost != 0:
                    cost_nonzero_rows += 1

        if not source_rows:
            raise WdtImportError("工作簿没有可识别的销售出库明细")

        missing_amount_orders = sum(1 for item in orders.values()
                                    if not item["pay"] or not item["receivable"])
        conflicting_amount_orders = sum(1 for item in orders.values()
                                        if len(item["pay"]) > 1 or len(item["receivable"]) > 1)
        order_pay_total = sum((next(iter(item["pay"])) for item in orders.values()
                               if len(item["pay"]) == 1), decimal.Decimal("0"))
        receivable_total = sum((next(iter(item["receivable"])) for item in orders.values()
                                if len(item["receivable"]) == 1), decimal.Decimal("0"))
        multi_line_order_count = sum(1 for item in orders.values() if item["lines"] > 1)
        max_lines_per_order = max((item["lines"] for item in orders.values()), default=0)
        cross_period_order_count = sum(1 for months in order_months.values()
                                       if period in months and len(months) > 1)

        blocking_codes = []
        warning_codes = []
        if not in_period_rows:
            blocking_codes.append("NO_ROWS_IN_PERIOD")
        if missing_order_number_rows:
            blocking_codes.append("MISSING_ORDER_NUMBER")
        if missing_amount_orders:
            blocking_codes.append("MISSING_ORDER_AMOUNT")
        if conflicting_amount_orders:
            blocking_codes.append("CONFLICTING_ORDER_AMOUNT")
        if cross_period_order_count:
            blocking_codes.append("CROSS_PERIOD_ORDER_AMOUNT")
        if cost_nonzero_rows == 0:
            blocking_codes.append("COST_ALL_ZERO")
            cost_status = "blocked_all_zero"
        elif cost_blank_rows:
            blocking_codes.append("COST_PARTIAL_MISSING")
            cost_status = "blocked_partial_missing"
        else:
            cost_status = "ready"
        if outside_period_rows:
            warning_codes.append("CROSS_PERIOD_ROWS_EXCLUDED")
        if missing_ship_time_rows:
            warning_codes.append("MISSING_SHIP_TIME_EXCLUDED")

        return {
            "period": period,
            "period_basis": "发货时间",
            "source_rows": source_rows,
            "in_period_rows": in_period_rows,
            "outside_period_rows": outside_period_rows,
            "missing_ship_time_rows": missing_ship_time_rows,
            "order_count": len(orders),
            "shipment_count": len(shipments),
            "multi_line_order_count": multi_line_order_count,
            "max_lines_per_order": max_lines_per_order,
            "cross_period_order_count": cross_period_order_count,
            "store_count": len(shops),
            "warehouse_count": len(warehouses),
            "sku_count": len(skus),
            "goods_quantity": float(goods_quantity),
            "order_pay_total": float(order_pay_total.quantize(_MONEY)),
            "receivable_total": float(receivable_total.quantize(_MONEY)),
            "missing_order_number_rows": missing_order_number_rows,
            "missing_amount_orders": missing_amount_orders,
            "conflicting_amount_orders": conflicting_amount_orders,
            "cost_total": float(cost_total.quantize(_MONEY)),
            "cost_nonzero_rows": cost_nonzero_rows,
            "cost_blank_rows": cost_blank_rows,
            "cost_status": cost_status,
            "ship_date_min": ship_date_min,
            "ship_date_max": ship_date_max,
            "source_status": "blocked" if blocking_codes else ("warning" if warning_codes else "ready"),
            "blocking_codes": blocking_codes,
            "warning_codes": warning_codes,
        }
    finally:
        workbook.close()
