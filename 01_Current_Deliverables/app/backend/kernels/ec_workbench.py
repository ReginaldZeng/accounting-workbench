# V2.556 · Unified ecommerce evidence and metrics. Pure functions; no Kingdee writes.
import io
import re
import hashlib
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from kernels import ec_tmall_import as tm, ec_settle as es, ec_month_fulfillment as fulfil
from kernels import ec_flow_ledger as ledger

KINDS = {'order': '平台订单', 'item': '商品与子订单', 'wdt': '旺店通销售出库',
         'alipay': '支付宝流水', 'fund': '聚合账户流水', 'refund': '退款售后明细'}
REQUIRED = tuple(KINDS)
LABELS = {'normal': '正常销售', 'ufirst': 'U先试用装', 'mixed': '混合订单', 'review': '待确认分类', 'unknown': '待分类'}
FEE_LABELS = {c: label for c, label, _ in es.FEE_MAP_SEED}


def amount(value):
    try:
        cleaned = str(value or 0).strip().replace(',', '').replace('￥', '').replace('¥', '')
        d = Decimal(cleaned or '0')
        if not d.is_finite():
            raise ValueError('金额必须为有限数值')
        return float(d.quantize(Decimal('.01')))
    except (InvalidOperation, ValueError):
        raise ValueError('金额格式错误，不能将缺失或错误金额视为零')


def total(values):
    return amount(sum((Decimal(str(v or 0)) for v in values), Decimal(0)))


def text(value, limit=240):
    return str(value or '').strip()[:limit]


def identity(value):
    if isinstance(value, (float, int)):
        raise ValueError('订单编号必须按文本导出，避免长编号精度丢失')
    return text(value, 64).strip("'\"")


def sheet_rows(data):
    tm._preflight_alipay_file(data)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        if ws.max_row == 1 and ws.max_column == 1:
            ws.reset_dimensions()
        iterator = ws.values
        header = [text(v) for v in next(iterator, ())]
        if not any(header):
            raise ValueError('没有可识别表头')
        for i, row in enumerate(iterator, 2):
            if i > tm.MAX_DATA_ROWS + 1:
                raise ValueError('数据行超过上限')
            yield {h: row[j] if j < len(row) else None for j, h in enumerate(header) if h}
    finally:
        wb.close()


def choose(row, *names):
    for name in names:
        if row.get(name) not in (None, ''):
            return row[name]
    return None


def fee_kind(code, label, income, expense):
    if code == '0010001' or label == '交易收款':
        return 'receipt'
    if code == '0020001' or '交易退款' in label or '售后退款' in label:
        return 'refund'
    if any(s in label for s in ('提现', '划转', '缴存', '保证金', '补齐', '充值')):
        return 'transfer'
    if any(s in label for s in ('补贴', '理赔')):
        return 'adjustment'
    if code in FEE_LABELS or any(s in label for s in ('服务费', '佣金', '积分', '运费险')):
        return 'fee'
    return 'unclassified'


def event(row, channel):
    row = ledger.resolve_order(row)
    code, label = es._code_of(row.get('desc'))
    label = FEE_LABELS.get(code, label or row['supplement']['business_label'] or text(row.get('btype')))
    # Do not persist arbitrary descriptions, payment serials or customer identifiers.
    label = re.sub(r'\d{8,}', '[编号省略]', text(label, 120))
    income, expense = amount(row.get('income')), amount(row.get('outgo'))
    return {'key': tm._alipay_row_key(row), 'date': tm._date_text(row.get('ts')),
            'channel': channel, 'code': code, 'label': label,
            'income': income, 'expense': expense, 'kind': fee_kind(code, label, income, expense)}


def parse_source(kind, payloads, period, shop):
    if kind not in KINDS:
        raise ValueError('不支持的数据类别')
    if kind != 'alipay' and len(payloads) != 1:
        raise ValueError('此类别每次导入一个完整文件；支付宝支持多个分片')
    if not payloads or sum(map(len, payloads)) > tm.MAX_ALIPAY_TOTAL_BYTES:
        raise ValueError('文件为空或总体积超过 200 MB')
    out = {'version': 1, 'kind': kind, 'period': period, 'shop': shop, 'rows': [], 'index': {},
           'business': {}, 'events': {}, 'status': 'ready', 'warnings': [], 'balance': None}
    data = payloads[0]
    if kind in ('order', 'item', 'fund'):
        parsed = tm.PARSERS[kind](data, period)
        if parsed.get('blocking_codes'):
            raise ValueError('文件存在阻断项：' + '、'.join(parsed['blocking_codes']))
        out['status'] = parsed['source_status']
        out['warnings'] = parsed['warning_codes']
        out['index'] = parsed.pop('_match_index', {})
        out['business'] = parsed.pop('_business_index', {})
        out['rows'] = parsed.pop('_detail_rows', [])
        out['summary'] = parsed
    if kind == 'order':
        mapping = {r['order_no']: r for r in out['rows']}
        for r in sheet_rows(data):
            if not r.get('订单编号'):
                continue
            no = identity(r['订单编号'])
            if no in mapping:
                mapping[no]['order_amount'] = amount(choose(r, '总金额', '买家应付货款'))
    elif kind == 'item':
        for r in sheet_rows(data):
            if tm._date_text(r.get('订单创建时间'))[:7] != period:
                continue
            no, sub = identity(r.get('主订单编号')), identity(r.get('子订单编号'))
            if not no or not sub:
                continue
            out['rows'].append({'order_no': no, 'suborder_no': sub,
                'name': text(choose(r, '商品标题', '宝贝标题', '商品名称', '标题')),
                'sku': text(r.get('商家编码'), 120), 'quantity': amount(r.get('购买数量')),
                'paid': amount(r.get('买家实付金额')), 'refund': 0 if text(r.get('退款金额'))=='无退款申请' else amount(r.get('退款金额')),
                'refund_status': text(r.get('退款状态')), 'type': fulfil.title_kind(choose(r, '商品标题', '宝贝标题', '商品名称', '标题'))})
        kinds = defaultdict(set)
        for r in out['rows']:
            kinds[tm._order_key(r['order_no'])].add(r['type'])
        for key, ks in kinds.items():
            if key not in out['business'] or out['business'][key] == 'unknown':
                out['business'][key] = 'mixed' if {'normal', 'ufirst'} <= ks else 'review' if len(ks) > 1 else next(iter(ks))
    elif kind == 'wdt':
        for r in sheet_rows(data):
            if text(r.get('店铺')) != shop or not r.get('出库单编号'):
                continue
            no = identity(r.get('子单原始单号'))
            if not no:
                continue
            out['rows'].append({'order_no': no, 'original_order': text(r.get('原始单号')),
                'internal_order': text(r.get('订单编号')), 'shipment_no': text(r.get('出库单编号')),
                'date': tm._date_text(r.get('发货时间')), 'status': text(r.get('出库单状态')),
                'sku': text(r.get('货品编号')), 'name': text(choose(r, '货品名称', '商家货品名称')),
                'spec': text(choose(r, '规格名称', '规格')), 'quantity': amount(r.get('货品数量')),
                'receivable': amount(r['应收金额']) if r.get('应收金额') not in (None, '') else None,
                'logistics_no': text(choose(r, '物流单号', '快递单号')),
                'logistics_company': text(choose(r, '物流公司', '物流公司名称'))})
        if not out['rows']:
            raise ValueError('没有该店铺的可关联出库明细，请核对店铺及表头')
    elif kind == 'refund':
        seen = set()
        for r in sheet_rows(data):
            no = identity(choose(r, '主订单编号', '订单编号', '订单号', '淘宝订单编号'))
            if not no:
                continue
            rid = text(choose(r, '退款编号', '退款单号', '售后单号'))
            if not rid:
                raise ValueError('售后明细缺少退款编号，无法可靠去重')
            if rid in seen:
                raise ValueError('售后明细存在重复退款编号，请提供一份完整最终状态报表')
            seen.add(rid)
            state = text(choose(r, '退款状态', '售后状态'))
            typ = text(choose(r, '售后类型', '退款类型', '服务类型'))
            refund_amount = choose(r, '退款成功金额', '退款总额', '退款金额', '退款金额（元）')
            if state in ('退款成功', '退款完成', '已退款') and refund_amount is None:
                raise ValueError('退款成功记录缺少金额，不能按零处理')
            out['rows'].append({'order_no': no, 'refund_no': rid,
                'type': 'return_refund' if typ in ('退货退款', '退款退货') else 'refund_only' if typ in ('仅退款', '退款不退货') else 'unknown',
                'status': state, 'success': state in ('退款成功', '退款完成', '已退款'),
                'amount': amount(refund_amount),
                'date': tm._date_text(choose(r, '退款成功时间', '退款完结时间', '退款完成时间', '完成时间'))})
        if not out['rows']:
            raise ValueError('没有可识别售后明细，需订单号、退款编号、类型、状态及金额')
    elif kind == 'alipay':
        parsed = tm.parse_alipay_exports(payloads, period)
        out['index'] = parsed.pop('_match_index')
        out.update(summary=parsed, status=parsed['source_status'], warnings=parsed['warning_codes'])
        seen = set()
        for blob in payloads:
            for r in es.parse_flow_any(blob):
                ev = event(r, 'alipay')
                if ev['key'] in seen or ev['date'][:7] != period:
                    continue
                seen.add(ev['key'])
                linked = ledger.resolve_order(r)
                no = linked['order_link']['order_no']
                ev['order_link'] = linked['order_link']
                if linked['order_link']['status']=='conflict':
                    out.setdefault('conflicted_keys',[]).extend(tm._order_key(n) for n in
                        [r.get('order_no','')]+[e['order_no'] for e in linked['supplement']['evidence']] if n)
                    continue
                ev['order_key'] = tm._order_key(no) if no else ''
                # Aggregation-channel rows are evidence copies, not an additional wallet receipt.
                ev['aggregate_copy'] = text(r.get('chan')) == '聚合结算渠道'
                out['rows'].append(ev)
    elif kind == 'fund':
        for r in sheet_rows(data):
            date = tm._date_text(r.get('入账时间'))
            if date[:7] != period:
                continue
            no = identity(r.get('淘宝订单编号'))
            raw = {'serial': text(r.get('支付流水号')), 'ts': date, 'order_no': no,
                   'desc': r.get('业务描述'), 'btype': r.get('入账类型'),
                   'income': r.get('收入金额（元）'), 'outgo': r.get('支出金额')}
            ev = event(raw, 'fund')
            if text(r.get('入账类型')) == '交易收款':
                ev['kind'] = 'receipt'
            elif '退款' in text(r.get('入账类型')):
                ev['kind'] = 'refund'
            ev['order_key'] = tm._order_key(no) if no else ''
            out['rows'].append(ev)
            balance = choose(r, '账户余额', '账户余额（元）', '余额（元）', '余额')
            if balance is not None and (not out['balance'] or date > out['balance']['as_of']):
                out['balance'] = {'amount': amount(balance), 'as_of': date, 'source': '聚合账户余额明细'}
    out['summary'] = out.get('summary') or {'source_rows': len(out['rows']), 'in_period_rows': len(out['rows'])}
    out['summary']['file_count'] = len(payloads)
    return out


def build_orders(sources, ar_index=None, recognition='shipment'):
    orders = sources.get('order', {}).get('rows', [])
    items, shipments, refunds, events = (defaultdict(list) for _ in range(4))
    for r in sources.get('item', {}).get('rows', []): items[r['order_no']].append(r)
    internal_owners = defaultdict(set)
    for r in sources.get('wdt', {}).get('rows', []):
        shipments[r['order_no']].append(r)
        internal_owners[r.get('internal_order') or r['shipment_no']].add(r['order_no'])
    for r in sources.get('refund', {}).get('rows', []): refunds[r['order_no']].append(r)
    for channel in ('alipay', 'fund'):
        for r in sources.get(channel, {}).get('rows', []):
            if r.get('order_key'): events[r['order_key']].append(r)
    business = sources.get('item', {}).get('business', {})
    out = []
    for source_order in orders:
        r = dict(source_order)
        no, key = r['order_no'], r['order_key']
        period = sources.get('order',{}).get('period') or r.get('period','')
        eligible = lambda date: bool(date) and (not period or str(date)[:7] <= period)
        line_items = items[no]
        typ = business.get(key, 'unknown')
        paid = amount(r.get('current_paid'))
        # Item report preserves gross transaction paid amount; master export can be current-state.
        if line_items and all('paid' in x for x in line_items): paid = total(x['paid'] for x in line_items)
        refund = amount(r.get('refund'))
        if line_items: refund = total(x['refund'] for x in line_items)
        paid_success = bool(r.get('paid_at')) or r.get('status') in ('交易成功', '买家已付款', '卖家已发货')
        actual = [e for e in events[key] if not e.get('aggregate_copy')]
        destinations = sorted({e['channel'] for e in actual if e['kind'] == 'receipt' and e['income'] > 0})
        fee_events = [e for e in actual if e['kind'] == 'fee']
        cash_present = any(k in sources for k in ('alipay', 'fund'))
        fee = total(e['expense'] - e['income'] for e in fee_events) if cash_present else None
        net = total(e['income']-e['expense'] for e in actual if e['kind'] != 'transfer') if actual else None
        ship = shipments[no]
        shipped = [x for x in ship if eligible(x.get('date')) and x.get('status') in ('已完成', '已发货')]
        ar = (ar_index or {}).get(no)
        issues = []
        if any(x.get('refund',0)>0 and x.get('refund_status') not in ('退款成功','退款完成','已退款') for x in line_items):
            issues.append('退款状态与累计退款金额不一致：核对售后及跨期流水')
        if typ in ('unknown', 'review', 'mixed'): issues.append('业务分类待确认' if typ != 'mixed' else '混合订单待拆分')
        pending_ship = typ != 'ufirst' and paid_success and refund < paid and not r.get('shipped_at') and not shipped and r.get('status') not in ('交易关闭', '交易成功')
        if pending_ship and r.get('status') not in ('买家已付款', '等待卖家发货'):
            pending_ship = False
        should_ar = bool(shipped or eligible(r.get('shipped_at'))) if recognition == 'shipment' else eligible(r.get('confirmed_at'))
        if period and any(x.get('date','')[:7]>period for x in ship): issues.append('跨期发货：按实际发货期间确认')
        ship_amounts = {}
        for x in shipped:
            ship_amounts.setdefault(x.get('internal_order') or x['shipment_no'], set()).add(x.get('receivable'))
        shared_internal = any(len(internal_owners[k])>1 for k in ship_amounts)
        expected = total(next(iter(v)) for v in ship_amounts.values()) if ship_amounts and not shared_internal and all(len(v)==1 and None not in v for v in ship_amounts.values()) else None
        if shared_internal and typ=='normal': issues.append('合并出库订单：应收金额需分摊核对')
        ar_diff = amount(ar['amount']-expected) if ar and expected is not None else None
        if typ == 'normal' and should_ar:
            if ar_index is None: issues.append('金蝶应收待同步')
            elif not ar: issues.append('已发货未匹配应收')
            elif ar_diff not in (None, 0): issues.append('应收金额差异')
            if not shipped: issues.append('旺店通出库未匹配')
        if len(destinations)>1:
            issues.append('多个账户收款待核对')
            net = None  # Possible aggregation copies or split receipts: do not double-count a conclusion.
        if any(e['kind']=='unclassified' for e in actual): issues.append('流水收支待分类')
        successful = [x for x in refunds[no] if x.get('success')]
        classified = total(x['amount'] for x in successful if x['type'] != 'unknown')
        refund_only = total(x['amount'] for x in successful if x['type']=='refund_only')
        return_refund = total(x['amount'] for x in successful if x['type']=='return_refund')
        if classified > refund: issues.append('售后与订单退款金额不一致')
        unknown_refund = amount(max(0, refund-classified))
        if classified > refund:
            refund_only = return_refund = 0
            unknown_refund = refund
        if unknown_refund > 0: issues.append('退款类型待分类')
        r.update(paid=paid, paid_success=paid_success, refund=refund, gsv=amount(paid-refund),
            business_type=typ, business_label=LABELS.get(typ, typ), items=line_items, shipments=ship,
            refunds=refunds[no], events=events[key], fees=fee, fee_events=fee_events, net_receipt=net,
            destination='、'.join({'alipay':'支付宝','fund':'聚合账户'}[d] for d in destinations) or '待结算 / 待补流水',
            refund_only=refund_only, return_refund=return_refund, unknown_refund=unknown_refund,
            unshipped=pending_ship, unshipped_amount=amount(max(0,paid-refund)) if pending_ship else 0,
            ar_documents=ar['documents'] if ar else [], ar_amount=ar['amount'] if ar else None,
            ar_expected=expected, ar_diff=ar_diff, should_ar=should_ar and typ=='normal',
            issues=list(dict.fromkeys(issues)), manual=bool(issues),
            receipt_amount=total(e['income'] for e in actual if e['kind']=='receipt'))
        out.append(r)
    return out


def metrics(rows, cash_available=True):
    paid = [r for r in rows if r['paid_success']]
    result = {'orders':len(rows), 'paid_orders':len(paid), 'gmv':total(r['paid'] for r in paid),
        'refund':total(r['refund'] for r in paid), 'gsv':total(r['gsv'] for r in paid),
        'fees':total(r['fees'] for r in rows) if cash_available and all(r['fees'] is not None for r in rows) else None,
        'refund_only':total(r['refund_only'] for r in paid), 'return_refund':total(r['return_refund'] for r in paid),
        'unknown_refund':total(r['unknown_refund'] for r in paid),
        'unshipped_count':sum(r['unshipped'] for r in rows), 'unshipped_amount':total(r['unshipped_amount'] for r in rows),
        'manual':sum(r['manual'] for r in rows), 'ar_matched':sum(bool(r['ar_documents']) for r in rows),
        'receipts':total(r['receipt_amount'] for r in rows), 'net_receipts':total(r['net_receipt'] for r in rows),
        'ufirst':sum(r['business_type']=='ufirst' for r in rows)}
    for kind in ('refund_only', 'return_refund', 'unknown_refund'):
        result[kind+'_rate'] = result[kind]/result['gmv'] if result['gmv'] else None
        result[kind+'_count'] = sum(r[kind]>0 for r in paid)
    return result
