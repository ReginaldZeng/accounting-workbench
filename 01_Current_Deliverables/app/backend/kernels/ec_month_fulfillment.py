# [Change Log] Date: 2026-09-10 | Author: Codex | Version: V2.555
# Description: Scoped WDT shipment evidence and U-first routing, without asserting posting correctness.
import io
import re
from collections import defaultdict
from decimal import Decimal
from openpyxl import load_workbook
from kernels import ec_wdt_import as wdt
from kernels.ec_tmall_import import _order_key

TARGET_SHOP = '星期零STARFIELD 天猫官旗店'


def title_kind(title):
    text = str(title or '').strip()
    if not text:
        return 'unknown'
    if re.search(r'[【\[]\s*(?:天猫\s*)?U\s*先\s*[】\]]', text, re.I):
        return 'ufirst'
    if re.search(r'U\s*先|试用', text, re.I):
        return 'review'
    return 'normal'


def identity(value):
    if not isinstance(value, str):
        return ''
    text = value.strip().strip("'\"").strip()
    return text if re.fullmatch(r'\d{16,30}', text) else ''


def parse_shipments(data, period):
    """Only business whitelist fields leave memory. Keep out-of-month rows as explicit evidence."""
    wdt._preflight_xlsx(data)
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        if sheet.max_row == 1 and sheet.max_column == 1:
            sheet.reset_dimensions()
        iterator = sheet.values
        headers = next(iterator, ())
        labels = ['子单原始单号', '出库单编号', '订单编号', '店铺', '发货时间',
                  '出库单状态', '货品编号', '货品数量', '单品支付金额', '应收金额', '订单类型']
        if any(label not in headers for label in labels):
            return {'version': 1, 'available': False, 'reason': '缺少逐单关联列，请重新导出完整销售出库明细', 'orders': {}}
        positions = {label: headers.index(label) for label in labels}
        groups, internal = defaultdict(list), defaultdict(list)
        skipped = target_rows = 0
        for number, row in enumerate(iterator, 2):
            if number > wdt.MAX_DATA_ROWS + 1:
                raise wdt.WdtImportError('数据行超过 250,000 行上限')
            r = {label: row[pos] if pos < len(row) else None for label, pos in positions.items()}
            if r['店铺'] != TARGET_SHOP:
                continue
            target_rows += 1
            key = identity(r['子单原始单号'])
            date = wdt._date_text(r['发货时间'])
            if not key or not date or not r['出库单编号'] or r['出库单状态'] not in ('已完成', '已发货'):
                skipped += 1
                continue
            def number_value(label):
                value = wdt._number(r[label])
                return value if value is not None and value.is_finite() else None
            record = {'shipment_no': str(r['出库单编号']), 'date': date,
                      'sku': str(r['货品编号'] or ''), 'quantity': number_value('货品数量'),
                      'line_paid': number_value('单品支付金额'), 'receivable': number_value('应收金额'),
                      'order_type': str(r['订单类型']), 'key': _order_key(key)}
            groups[record['key']].append(record)
            internal[str(r['订单编号'])].append(record)
        # Order-level amounts repeat on every SKU. Attribute only unambiguous, single-platform orders.
        order_amounts = defaultdict(list)
        for records in internal.values():
            keys = {r['key'] for r in records}
            amounts = {r['receivable'] for r in records}
            if len(keys) == 1 and len(amounts) == 1 and None not in amounts:
                order_amounts[next(iter(keys))].append(next(iter(amounts)))
            else:
                for key in keys:
                    order_amounts[key].append(None)
        output = {}
        for key, records in groups.items():
            documents = defaultdict(list)
            for record in records:
                documents[record['shipment_no']].append(record)
            docs = []
            for bill, lines in sorted(documents.items()):
                docs.append({'shipment_no': bill, 'dates': sorted({r['date'] for r in lines}),
                             'quantity': float(sum((r['quantity'] for r in lines), Decimal(0))) if all(r['quantity'] is not None for r in lines) else None,
                             'items': [{'sku': r['sku'], 'quantity': float(r['quantity']) if r['quantity'] is not None else None} for r in lines]})
            amounts = order_amounts[key]
            dates = sorted({r['date'] for r in records})
            output[key] = {'documents': docs, 'quantity': float(sum((r['quantity'] for r in records), Decimal(0))) if all(r['quantity'] is not None for r in records) else None,
                           'receivable': float(sum(amounts, Decimal(0))) if amounts and None not in amounts else None,
                           'line_paid': float(sum((r['line_paid'] for r in records), Decimal(0))) if all(r['line_paid'] is not None for r in records) else None,
                           'cross_period': any(date[:7] != period for date in dates),
                           'in_period': any(date[:7] == period for date in dates),
                           'order_types': sorted({r['order_type'] for r in records})}
        return {'version': 1, 'available': True, 'orders': output, 'target_rows': target_rows,
                'skipped_rows': skipped, 'linked_orders': len(output)}
    finally:
        book.close()


def enrich(order, shipment, business, ar_available, wdt_available):
    kind = business or 'unknown'
    labels = {'normal': '正常销售', 'ufirst': 'U先试用装', 'mixed': '混合订单', 'review': '试用待确认', 'unknown': '待分类'}
    order.update(business_type=kind, business_label=labels[kind], wdt_documents=(shipment or {}).get('documents', []),
                 wdt_quantity=(shipment or {}).get('quantity'), wdt_receivable=(shipment or {}).get('receivable'),
                 wdt_line_paid=(shipment or {}).get('line_paid'), wdt_cross_period=bool((shipment or {}).get('cross_period')),
                 push_status='未提供推送日志', gl_status='待核对总账凭证' if kind in ('ufirst', 'mixed') else '')
    if kind == 'ufirst':
        status = 'U先：总账凭证核对'
        shipment_status = '不要求旺店通出库'
        reason = '不要求金蝶应收；仍保留平台收退款及资金记录'
        if shipment or order.get('kingdee_ar_no'):
            reason += '；发现出库或应收记录，需复核业务分类'
    elif kind in ('unknown', 'review', 'mixed'):
        status = '混合订单待拆分核对' if kind == 'mixed' else '待确认业务类型'
        shipment_status = '已找到出库记录' if shipment else '暂不判断缺出库'
        reason = '混合订单不整单豁免；正常商品核对出库与应收，U先部分核对总账' if kind == 'mixed' else '请重导宝贝销售明细完成分类；疑似试用装需人工确认'
    elif not wdt_available:
        status, shipment_status, reason = '待导入旺店通', '未提供出库明细', '来源与勾稽中导入销售出库明细'
    elif not shipment:
        status, shipment_status, reason = '未找到出库记录', '未找到出库记录', '仅当前导出范围未命中，不直接判定未发货或漏单'
    elif shipment['cross_period']:
        status, shipment_status, reason = '跨月发货待核对', '跨月发货', '出库明细含所选月之外的发货；不计作本月应收缺失'
    elif not ar_available:
        status, shipment_status, reason = '待同步金蝶', '已找到出库记录', '尚无可用金蝶应收快照'
    elif not order.get('kingdee_ar_no'):
        status, shipment_status, reason = '已发货未找到应收', '已找到出库记录', '仅当前金蝶取数范围未命中；不等同推送失败'
    else:
        status, shipment_status, reason = '应收已关联·金额待核对', '已找到出库记录', '旺店通应收仅作参考；优惠、运费、退货及数量单位口径未确认，不自动判定正确'
    order.update(ar_check_status=status, wdt_status=shipment_status, ar_check_reason=reason)
    return order
