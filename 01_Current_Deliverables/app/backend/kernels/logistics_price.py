# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-09-26 | Author: Claude Opus 4.8 | Version: V2.632
# Description: 【物流账单复核】价格卡内核（核价的地基）。价一律来自合同，不拿账单金额倒算（否则核价是循环、永远零差）。
#   ① 从承运商合同价目表 xlsx 导入价格卡行（迅鸽云仓格式：快递费按省×重量档、基础/增值服务费、包装材料）。
#   ② 查价：std_freight(快递公司,省,计费重量)→标准快递费（首重+续重+<3kg分档）；unit_price(费用项)→简单单价。
#   实证：迅鸽 2026-08 逐单重算 5711 单，账单基价 15972.80 vs 标准 15982.95，5706 一致、0 多收（2026-09-25 验）。
import json
import math

_KG_UNITS = {"千克", "公斤", "kg", "KG", "Kg"}


def _f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _norm_prov(s):
    if not s:
        return ""
    s = str(s).strip()
    for suf in ("维吾尔自治区", "壮族自治区", "回族自治区", "特别行政区", "自治区", "省", "市"):
        if s.endswith(suf):
            s = s[:-len(suf)]
            break
    return s


# ---------- 从合同价目表导入（迅鸽云仓格式；每家合同布局或不同，按承运商扩） ----------
_EXPRESS_SHEETS = {"快递费（圆通）": ("圆通", 3), "快递费（中通）": ("中通", 1),
                   "快递费（韵达）": ("韵达", 1), "快递费（邮政）": ("邮政", 1)}
_FIRST_KG = {"圆通": 3, "中通": 1, "韵达": 1, "邮政": 1}


def parse_contract_xlsx(data, carrier="迅鸽"):
    """解析迅鸽云仓合同价目表 → [价格卡行 dict]。data=bytes 或路径。行含 fee_item/sub/unit/price/tier_json/first_kg/source。"""
    import openpyxl
    from io import BytesIO
    src = BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = []

    def add(fee_item, sub, unit, price=None, tier=None, first_kg=None, src=""):
        rows.append({"carrier": carrier, "fee_item": fee_item, "sub": sub or "", "unit": unit or "",
                     "price": price, "tier_json": json.dumps(tier, ensure_ascii=False) if tier else None,
                     "first_kg": first_kg, "source": src, "status": "已签署"})

    # 快递费：省 → 档位数组
    for sheet, (name, fk) in _EXPRESS_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        s = wb[sheet]
        tier = {}
        for r in list(s.iter_rows(values_only=True))[4:]:
            prov = _norm_prov(r[1])
            vals = [_f(x) for x in r[2:8]]
            if prov and any(v is not None for v in vals):
                tier[prov] = vals
        if tier:
            add("快递费", name, "元/单", tier=tier, first_kg=fk, src=f"附件二·{sheet}")
    if "快递费（顺丰）" in wb.sheetnames:
        s = wb["快递费（顺丰）"]
        tier = {}
        for r in list(s.iter_rows(values_only=True))[4:]:
            prov = _norm_prov(r[0])
            vals = [_f(x) for x in r[1:5]]
            if prov and any(v is not None for v in vals):
                tier[prov] = vals
        if tier:
            add("快递费", "顺丰", "元/单", tier=tier, first_kg=None, src="附件二·快递费（顺丰）")
    # 基础/增值服务费
    for sheet in ("基础服务费", "增值服务费"):
        if sheet not in wb.sheetnames:
            continue
        for r in wb[sheet].iter_rows(min_row=4, values_only=True):
            item = r[0]
            if item and _f(r[3]) is not None:
                add(str(item).strip(), "", r[4] or "", price=_f(r[3]), src=f"附件二·{sheet}")
    # 包装材料
    if "包装材料" in wb.sheetnames:
        for r in wb["包装材料"].iter_rows(min_row=4, values_only=True):
            if r[0] and _f(r[3]) is not None:
                add("包材·" + str(r[0]).strip(), str(r[1] or ""), r[2] or "", price=_f(r[3]), src="附件二·包装材料")
    return rows


# ---------- 查价 ----------
def load_card(rows):
    """价格卡行 → 结构：{'express':{快递公司:{'tier':{省:[..]},'first_kg':n}}, 'unit':{费用项:price}}。"""
    card = {"express": {}, "unit": {}}
    for r in rows:
        if r["fee_item"] == "快递费" and r.get("tier_json"):
            card["express"][r["sub"]] = {"tier": json.loads(r["tier_json"]), "first_kg": r.get("first_kg")}
        elif r.get("price") is not None:
            card["unit"][r["fee_item"]] = r["price"]
    return card


def express_table(name):
    """账单「快递公司」名 → 价卡表名。空运/快运/自提无档，返回 None。"""
    n = str(name or "")
    if "圆通" in n:
        return "圆通"
    if "中通快运" in n or "云驰" in n:
        return None
    if "中通" in n:
        return "中通"
    if "韵达" in n:
        return "韵达"
    if "顺丰" in n and ("空运" in n or "到付" in n):
        return None
    if "顺丰" in n:
        return "顺丰"
    if "邮政" in n:
        return "邮政"
    return None


def std_freight(card, express_name, prov, wt):
    """按合同重算标准快递费。返回 (标准价, 计价档) 或 (None, 原因)——不倒算。"""
    wt = _f(wt)
    if wt is None or wt <= 0:
        return None, "无计费重量"
    tab = express_table(express_name)
    if not tab:
        return None, "价卡无此快递（空运/快运/自提）"
    ex = card["express"].get(tab)
    if not ex:
        return None, "价卡无 %s" % tab
    row = ex["tier"].get(_norm_prov(prov))
    if not row:
        return None, "价卡无省份 %s" % _norm_prov(prov)
    if tab == "顺丰":
        v1, v2, v3, cont = row
        if wt <= 1:
            return v1, "1kg档"
        if wt <= 2:
            return v2, "2kg档"
        if wt <= 3:
            return v3, "3kg档"
        return round(v3 + math.ceil(wt - 3) * (cont or 0), 2), "3kg+续重"
    lt05, lt1, lt2, lt3, first, cont = row
    fk = ex.get("first_kg") or 3
    if wt < 0.5 and lt05 is not None:
        return lt05, "<0.5kg"
    if wt < 1 and lt1 is not None:
        return lt1, "0.5-1kg"
    if wt < 2 and lt2 is not None:
        return lt2, "1-2kg"
    if wt < 3 and lt3 is not None:
        return lt3, "2-3kg"
    if first is not None:
        return round(first + math.ceil(max(0, wt - fk)) * (cont or 0), 2), "首重%dkg+续重" % fk
    return None, "价卡档缺"


def unit_price(card, fee_item):
    """简单单价查询（操作费/退货/仓储/装卸/包材）。命中返回单价，否则 None。"""
    return card["unit"].get(fee_item)
