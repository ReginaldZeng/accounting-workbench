# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-10-25 | Author: Kimi | Version: V2.1010
# Description: 物流账单标准复核表：说明书、单据运费、计提对账单、原始账单；只呈现已有复核结果。
from copy import copy
from io import BytesIO
import json
import re

import openpyxl
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

NAVY = "17324D"
TEAL = "1F6E8C"
WHITE = "FFFFFF"
TEXT = "1B2733"
MUTED = "5E6B78"
BLUE = "DCEAF3"
PALE = "F5F8FA"
YELLOW = "FFF2CC"
GREEN = "E7F3EB"
RED = "FCE8E6"
GRAY = "E7ECEF"
THIN = Side(style="thin", color="CBD5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DOC_HEADERS = [
    "主体", "科目", "部门", "产品线", "产品项目", "来源sheet页", "源行号", "父费用行ID", "费用行ID",
    "核对方案", "核对字段", "计费标准", "单据类型", "单据编码", "费用类型", "物料分组", "物料编码", "物料名称", "规格",
    "基本数量", "基本单位", "销售额", "长", "宽", "高", "单位换算系数", "折合计价数量", "折合计价单位",
    "账单数量", "计费标准数量", "数量差异", "毛净重比", "数量结论", "账单金额", "标准金额", "金额差异",
    "单位费用", "费比", "备注",
]
DOC_WIDTHS = [22, 18, 18, 18, 18, 12, 9, 18, 18, 15, 13, 20, 16, 17, 16, 13, 15, 28, 20,
              12, 11, 12, 8, 8, 8, 14, 14, 13, 13, 14, 13, 12, 28, 13, 13, 13, 13, 11, 36]
C = {name: i + 1 for i, name in enumerate(DOC_HEADERS)}
CL = {name: get_column_letter(i) for name, i in C.items()}


def _money(v):
    return round(float(v or 0), 2)


def _json(v):
    try:
        return json.loads(v) if isinstance(v, str) else (v or {})
    except Exception:
        return {}


def _dim(code, name):
    return "%s\n%s" % (code, name or "") if code else str(name or "")


def _doc_type(no):
    prefix = "".join(x for x in str(no or "") if x.isalpha())
    return {"FBDR": "调拨入库单", "FBDC": "调拨出库单", "CGRK": "采购入库单", "CGTL": "采购退料单",
            "XSCKD": "销售出库单", "QTCK": "其他出库单", "RK": "销售退货单"}.get(prefix, "金蝶单据" if prefix else "无单据")


def _scheme(x, no_doc=False):
    if no_doc:
        return "按批次", "计费数量", "不适用"
    kind = (x.get("chk") or {}).get("kind")
    if kind == "vol":
        return "按体积", "长×宽×高", "体积重"
    if kind == "pallet":
        return "按托数", "托数", "托规"
    if kind == "box":
        return "按箱数", "箱数", "箱规"
    if kind == "boxwt":
        return "按重量", "计费吨重", "每箱计费千克"
    return "按重量", "重量", "毛净重比"


def _new_sheet(wb, name, title, subtitle, headers, widths):
    ws = wb.create_sheet(name)
    n = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n)
    cell = ws.cell(1, 1, title)
    cell.font = Font(size=17, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    for col in range(1, n + 1):
        ws.cell(1, col).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    cell = ws.cell(2, 1, subtitle)
    cell.font = Font(size=10, color=MUTED)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32
    for i, header in enumerate(headers, 1):
        cell = ws.cell(4, i, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[4].height = 36
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    return ws


def _append_rows(ws, items, start=5, money_cols=(), pct_cols=()):
    for ri, values in enumerate(items, start):
        for ci, value in enumerate(values, 1):
            cell = ws.cell(ri, ci, value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=ci <= 5, shrink_to_fit=ci > 5)
            cell.font = Font(size=10, color=TEXT)
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
            if ci in money_cols:
                cell.number_format = "#,##0.00;[Red]-#,##0.00"
            if ci in pct_cols:
                cell.number_format = "0.00%"
        ws.row_dimensions[ri].height = 44
    if items:
        ws.auto_filter.ref = "A4:%s%d" % (get_column_letter(ws.max_column), ws.max_row)


def _print_setup(ws, width=1):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = width
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.print_title_rows = "1:4"
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "物流账单标准复核表 | 第 &P 页 / 共 &N 页"
    ws.oddFooter.center.size = 9


def _dimension_map(lines, res):
    entries = [x for x in lines.get("rows", []) if x.get("kind") == "accr"]
    subjects = {x.get("name"): x.get("code") for x in (res.get("settings", {}).get("subjects") or [])}
    biz_codes = res.get("codes", {}).get("biz", {}) or {}

    def pick(x):
        fee = x.get("fee_item") or ""
        exact = [e for e in entries if e.get("fee_type") == fee]
        related = [e for e in entries if (fee.startswith("入库") and e.get("fee_type") == "入库运费") or
                   (fee.startswith("出库") and e.get("fee_type") == "出库运费")]
        e = (exact or related or entries or [{}])[0]
        subject = x.get("subject") or e.get("subject") or ""
        biz = x.get("bizline") or e.get("biz") or ""
        return {
            "subject": _dim(subjects.get(subject) or e.get("book_code"), subject),
            "acct": _dim(e.get("acct"), e.get("acct_name")),
            "dept": _dim(e.get("dept_code"), e.get("dept")),
            "biz": _dim(biz_codes.get(biz) or e.get("biz_code"), biz),
            "proj": _dim(e.get("proj_code"), e.get("proj")),
        }
    return pick


def _build_doc_sheet(wb, carrier, period, res, lines, bills, source_rows):
    ws = _new_sheet(wb, "02单据运费", "单据运费复核",
                    "一行为一个物料或一个无单据费用批次；单据维度与金额结论可合并，不同产品分行。未核、缺参和不适用均明确写明。",
                    DOC_HEADERS, DOC_WIDTHS)
    docs = {x.get("lid"): x for x in res.get("docs", [])}
    dimension = _dimension_map(lines, res)
    rowno = 5
    links = {}
    qty_enabled = res.get("settings", {}).get("qty_enabled", True)
    price_enabled = res.get("settings", {}).get("price_enabled", True)
    for b in bills:
        lid = b.get("id")
        src_sheet, src_row = b.get("src_sheet") or "", b.get("src_row")
        original = source_rows.get((src_sheet, src_row), {})
        no = (b.get("doc_no") or "").split("+")[0]
        no_doc = not no or no == "无单据"
        x = docs.get(lid) or {
            "lid": lid, "doc_no": "" if no_doc else no, "subject": b.get("subject"), "fee_item": b.get("fee_item"),
            "bizline": b.get("bizline") or ("原料" if "原料" in str(b.get("annot") or "") else "小料"),
            "doc_fee": b.get("amount"), "bill_amt": b.get("charge_wt") if b.get("charge_wt") is not None else b.get("qty"),
            "bill_unit": "千克" if b.get("charge_wt") is not None else b.get("unit"), "materials": [],
            "sub_fees": _json(b.get("sub_fees")), "state": "info", "mode_cn": "无运输单据，按原账单批次核价",
        }
        sf = x.get("sub_fees") or _json(b.get("sub_fees"))
        fee_id = "HM-%s-%05d" % (period.replace("-", ""), int(lid))
        materials = x.get("materials") or []
        if no_doc or not materials:
            materials = [{
                "code": "", "name": original.get("商品名称") or b.get("fee_item"), "spec": original.get("规格") or "",
                "base_kg": None, "kg_unit": "", "kd": None, "box_div": None, "fee": None, "sales": None,
            }]
        start = rowno
        scheme, check_field, price_std = _scheme(x, no_doc)
        for mi, m in enumerate(materials, 1):
            rr = rowno
            rowno += 1
            dim = dimension(x)
            values = {k: "" for k in DOC_HEADERS}
            values.update({
                "主体": dim["subject"], "科目": dim["acct"], "部门": dim["dept"], "产品线": dim["biz"], "产品项目": dim["proj"],
                "来源sheet页": src_sheet, "源行号": src_row, "父费用行ID": "" if mi == 1 else fee_id,
                "费用行ID": fee_id if mi == 1 else "%s-M%d" % (fee_id, mi), "核对方案": scheme,
                "核对字段": check_field, "计费标准": sf.get("公式") or price_std, "单据类型": _doc_type(no),
                "单据编码": no if not no_doc else "无单据", "费用类型": x.get("fee_item") or b.get("fee_item"),
                "物料分组": "原料" if "原料" in str(b.get("annot") or "") else (x.get("bizline") or ""),
                "物料编码": m.get("code") or "", "物料名称": m.get("name") or "", "规格": m.get("spec") or "",
                "基本数量": m.get("base_kg"), "基本单位": m.get("kg_unit") or "", "销售额": m.get("sales"),
                "长": "", "宽": "", "高": "", "单位换算系数": m.get("box_div") or m.get("kg_per"),
                "折合计价数量": m.get("kd") if qty_enabled and not no_doc else None,
                "折合计价单位": x.get("kd_unit") or "", "账单数量": x.get("bill_amt") if mi == 1 and not no_doc else (b.get("qty") if no_doc else None),
                "计费标准数量": (b.get("charge_wt") if b.get("charge_wt") is not None else b.get("qty")) if mi == 1 else None,
                "数量结论": ("未核量" if not qty_enabled else "不适用·无运输单据" if no_doc else "待补每箱计费重量" if x.get("bw_lack") else
                           "一致" if x.get("state") == "ok" else "金蝶查无" if x.get("state") == "miss" else "数量不符" if x.get("state") == "qtydiff" else "待确认"),
                "账单金额": b.get("amount") if mi == 1 else None,
                "标准金额": sf.get("标准") if mi == 1 and price_enabled else None,
                "单位费用": (round(float(m.get("fee")) / float(m.get("base_kg")), 6)
                           if m.get("fee") is not None and m.get("base_kg") and str(m.get("kg_unit") or "").lower() in ("千克", "kg") else "不适用"),
                "费比": m.get("ratio"), "备注": "；".join(str(v) for v in [x.get("mode_short") or x.get("mode_cn"), original.get("备注"), b.get("note")] if v),
            })
            if not price_enabled:
                values["备注"] = "未核价；" + values["备注"]
            chk = x.get("chk") or {}
            if chk.get("kind") == "vol" and chk.get("lwh"):
                values["长"], values["宽"], values["高"] = chk["lwh"]
            for header, value in values.items():
                ws.cell(rr, C[header], value)
            if mi == 1:
                if not no_doc and qty_enabled:
                    end = start + len(materials) - 1
                    kd_range = "%s%d:%s%d" % (CL["折合计价数量"], start, CL["折合计价数量"], end)
                    ws.cell(rr, C["数量差异"], "=IF(OR(%s%d=\"\",COUNT(%s)=0),\"\",%s%d-SUM(%s))" % (
                        CL["账单数量"], rr, kd_range, CL["账单数量"], rr, kd_range))
                    ws.cell(rr, C["毛净重比"], "=IF(OR(%s%d=\"\",SUM(%s)=0),\"\",%s%d/SUM(%s))" % (
                        CL["账单数量"], rr, kd_range, CL["账单数量"], rr, kd_range))
                if price_enabled:
                    ws.cell(rr, C["金额差异"], "=IF(OR(%s%d=\"\",%s%d=\"\"),\"\",%s%d-%s%d)" % (
                        CL["账单金额"], rr, CL["标准金额"], rr, CL["账单金额"], rr, CL["标准金额"], rr))
        end = rowno - 1
        links[(src_sheet, src_row)] = {"first": start, "last": end, "fee_id": fee_id}
        if end > start:
            merge_headers = DOC_HEADERS[:7] + DOC_HEADERS[9:15] + DOC_HEADERS[28:]
            for header in merge_headers:
                ci = C[header]
                ws.merge_cells(start_row=start, start_column=ci, end_row=end, end_column=ci)
        for rr in range(start, end + 1):
            for ci in range(1, len(DOC_HEADERS) + 1):
                cell = ws.cell(rr, ci)
                cell.border = BORDER
                cell.font = Font(size=10, color=TEXT)
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=ci <= 5, shrink_to_fit=ci > 5)
                if rr % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=PALE)
                if ci in (C["账单金额"], C["标准金额"], C["金额差异"], C["基本数量"], C["折合计价数量"], C["账单数量"], C["计费标准数量"], C["数量差异"]):
                    cell.number_format = "#,##0.00;[Red]-#,##0.00"
                if ci == C["毛净重比"]:
                    cell.number_format = "0.0000"
                if ci == C["费比"]:
                    cell.number_format = "0.00%"
            ws.row_dimensions[rr].height = 52
        for ci in range(C["数量差异"], C["备注"] + 1):
            ws.cell(start, ci).fill = PatternFill("solid", fgColor=YELLOW)
    ws.auto_filter.ref = "A4:%s%d" % (get_column_letter(len(DOC_HEADERS)), ws.max_row)
    ws.freeze_panes = "J5"
    _print_setup(ws, 2)
    ws.print_title_cols = "A:C"
    return ws, links


def _build_lines_sheet(wb, lines):
    headers = ["费用主体", "科目", "部门", "产品线", "产品项目", "费用类型", "凭证号", "计提金额(含税)", "税率", "账单金额", "计提差异\n(计提-账单)", "差异解释", "处理结论"]
    widths = [25, 18, 19, 18, 18, 18, 12, 17, 10, 15, 17, 34, 20]
    ws = _new_sheet(wb, "03计提对账单", "计提对账单",
                    "金蝶计提与原始账单归口勾稽；总额对平不等于主体、费用类型和明细都一致。", headers, widths)
    items = []
    for r in lines.get("rows", []):
        if r.get("kind") == "gtotal":
            continue
        note = r.get("note") or ""
        if r.get("kind") == "bill_only":
            note = note or "账单有这一归口，金蝶计提未匹配到同一主体/费用类型"
        elif not r.get("bill"):
            note = note or "金蝶计提未匹配到同归口账单"
        items.append([_dim(r.get("book_code"), r.get("subject")), _dim(r.get("acct"), r.get("acct_name")),
                      _dim(r.get("dept_code"), r.get("dept")), _dim(r.get("biz_code"), r.get("biz")),
                      _dim(r.get("proj_code"), r.get("proj")), r.get("fee") or r.get("fee_type"), r.get("vno") or "",
                      r.get("amt"), r.get("tax_rate"), r.get("bill"), r.get("diff"), note,
                      "待解释" if abs(float(r.get("diff") or 0)) >= 0.01 and not r.get("note") else "已说明"])
    items.append(["合计", "", "", "", "", "", "", lines.get("accr_total"), None, lines.get("bill_total"), lines.get("diff_total"), "", ""])
    _append_rows(ws, items, money_cols=(8, 10, 11), pct_cols=(9,))
    for ci in range(1, len(headers) + 1):
        ws.cell(ws.max_row, ci).font = Font(bold=True)
        ws.cell(ws.max_row, ci).fill = PatternFill("solid", fgColor=BLUE)
    _print_setup(ws)
    return ws


def _build_cover(wb, carrier, period, res, lines, bills, raw_names):
    ws = _new_sheet(wb, "01复核说明书", "%s账单复核说明书" % carrier,
                    "用途：业务提交付款 → 会计复核 → 经理审批。异常未闭环时为复核草稿，不作为最终付款依据。",
                    ["项目", "说明", "项目", "说明"], [24, 60, 24, 60])
    settings = res.get("settings", {})
    docs = res.get("docs", [])
    abnormal = len({x.get("doc_no") for x in docs if x.get("state") in ("miss", "qtydiff", "price") and not x.get("confirmed")})
    unexplained = lines.get("n_unexplained") or 0
    price_rows = [(b, _json(b.get("sub_fees"))) for b in bills]
    price_missing = sum(1 for b, sf in price_rows if sf.get("标准") is None) if settings.get("price_enabled", True) else 0
    price_diff = round(sum(float(b.get("amount") or 0) - float(sf["标准"]) for b, sf in price_rows if sf.get("标准") is not None), 2)
    price_bad = sum(1 for b, sf in price_rows if sf.get("标准") is not None and abs(float(b.get("amount") or 0) - float(sf["标准"])) >= 0.005)
    doc_n = len({(b.get("doc_no") or "").split("+")[0] for b in bills if b.get("doc_no") and b.get("doc_no") != "无单据"})
    nodoc_n = sum(1 for b in bills if not b.get("doc_no") or b.get("doc_no") == "无单据")
    pending = abnormal or unexplained or price_missing or price_bad
    status = "复核草稿·异常待处理" if pending else "复核完成·可提交审批"
    price_conclusion = "未核价" if not settings.get("price_enabled", True) else "标准金额差 %s 元；%d 行公式差、%d 行缺核价依据" % (f"{price_diff:,.2f}", price_bad, price_missing)
    items = [
        ["一、复核对象", "%s · %s" % (carrier, period), "原始金额", "%s 元；%d 行费用明细" % (f"{_money(res.get('total_bill')):,.2f}", len(bills))],
        ["二、复核设置", "核量%s · 核价%s" % ("☑" if settings.get("qty_enabled", True) else "□", "☑" if settings.get("price_enabled", True) else "□"),
         "复核范围", "%d 个金蝶单号；另 %d 行无运输单费用" % (doc_n, nodoc_n)],
        ["三、数量口径", "账单数量与金蝶折合计价数量比较；不同产品分行展示", "数量结论", "%d 个单号存在数量差异、查无或缺参，详见单据运费" % abnormal],
        ["四、价格口径", "标准金额来自已配置的账单公式或合同价；无依据时留空并写未核价", "价格结论", price_conclusion],
        ["五、计提勾稽", "计提 %s 元；账单 %s 元；总额差 %s 元" % (f"{_money(lines.get('accr_total')):,.2f}", f"{_money(lines.get('bill_total')):,.2f}", f"{_money(lines.get('diff_total')):,.2f}"),
         "归口结论", "%d 笔差异未解释；总额对平不等于归口一致" % unexplained],
        ["六、异常说明", "数量缺参、归口差异等须逐条写明原因和处理结论", "异常处理", "待处理" if pending else "已闭环"],
        ["七、附件情况", "原始账单：%s；金蝶物料已接入单据运费" % "、".join(raw_names), "欠票提示", "本表未单独取得上月末往来及欠票证据，不作欠票判断"],
        ["八、提交结论", status, "建议付款", "待异常闭环后确定" if pending else "按复核结果提交"],
        ["业务提交", "以钉钉请款单为准", "会计复核", (lines.get("signed") or {}).get("reviewer") or "待完成"],
        ["经理审批", "以钉钉审批为准", "批准金额", "待会计复核完成后填写"],
    ]
    _append_rows(ws, items, start=5)
    for rr in range(5, 5 + len(items)):
        for ci in (1, 3):
            ws.cell(rr, ci).font = Font(bold=True, color=NAVY, size=10)
            ws.cell(rr, ci).fill = PatternFill("solid", fgColor=BLUE)
    ws.auto_filter.ref = None
    _print_setup(ws)
    return ws


def _copy_raw_sheet(wb, source, title):
    ws = wb.create_sheet(title)
    cells = list(getattr(source, "_cells", {}).values())
    styled = len(cells) <= 40000
    for cell in cells:
        if isinstance(cell, MergedCell) or (cell.value is None and not (styled and cell.has_style)):
            continue
        target = ws.cell(cell.row, cell.column, cell.value)
        target.number_format = cell.number_format
        if styled and cell.has_style:
            target.font = copy(cell.font)
            target.fill = copy(cell.fill)
            target.border = copy(cell.border)
            target.alignment = copy(cell.alignment)
    for rg in source.merged_cells.ranges:
        try:
            ws.merge_cells(str(rg))
        except Exception:
            pass
    for dim in source.column_dimensions.values():
        if dim.width and dim.min and dim.max:
            for ci in range(dim.min, min(dim.max, 300) + 1):
                ws.column_dimensions[get_column_letter(ci)].width = dim.width
    for ri, dim in source.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[ri].height = dim.height
    ws.freeze_panes = source.freeze_panes
    return ws


def _append_raw_review(ws, source, links):
    sheet_rows = {key[1]: value for key, value in links.items() if key[0] == source.title}
    if not sheet_rows:
        return
    used_cells = [c for c in getattr(source, "_cells", {}).values() if c.value is not None]
    first = max((c.column for c in used_cells), default=source.max_column) + 2
    headers = ["费用行ID", "金蝶核对数量", "数量差异", "数量结论", "标准金额", "金额差异", "复核说明"]
    ws.cell(1, first, "系统复核加工（左侧原账单不变）")
    ws.cell(1, first).font = Font(bold=True, color=WHITE)
    ws.cell(1, first).fill = PatternFill("solid", fgColor=NAVY)
    for offset, header in enumerate(headers):
        ci = first + offset
        cell = ws.cell(2, ci, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = [19, 16, 14, 28, 14, 14, 45][offset]
    for src_row, link in sheet_rows.items():
        first_row, last_row = link["first"], link["last"]
        refs = [link["fee_id"],
                "=SUM('02单据运费'!%s%d:%s%d)" % (CL["折合计价数量"], first_row, CL["折合计价数量"], last_row),
                "='02单据运费'!%s%d" % (CL["数量差异"], first_row),
                "='02单据运费'!%s%d" % (CL["数量结论"], first_row),
                "='02单据运费'!%s%d" % (CL["标准金额"], first_row),
                "='02单据运费'!%s%d" % (CL["金额差异"], first_row),
                "='02单据运费'!%s%d" % (CL["备注"], first_row)]
        for offset, value in enumerate(refs):
            cell = ws.cell(src_row, first + offset, value)
            cell.fill = PatternFill("solid", fgColor=YELLOW)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if offset in (1, 2, 4, 5):
                cell.number_format = "#,##0.00;[Red]-#,##0.00"


def build_standard_workbook(carrier, period, res, lines, bills, raw_files):
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    raw_books = []
    source_rows = {}
    for fname, data in raw_files:
        try:
            src = openpyxl.load_workbook(BytesIO(data), data_only=False)
        except Exception:
            continue
        raw_books.append((fname, src))
        for ws in src.worksheets:
            headers = {cell.value: cell.column for cell in ws[2] if cell.value}
            wanted = {name: headers.get(name) for name in ("商品名称", "规格", "备注")}
            for row in range(3, ws.max_row + 1):
                source_rows[(ws.title, row)] = {name: ws.cell(row, ci).value if ci else None for name, ci in wanted.items()}
    doc_ws, links = _build_doc_sheet(wb, carrier, period, res, lines, bills, source_rows)
    _build_lines_sheet(wb, lines)
    raw_names = []
    used = set(wb.sheetnames)
    for fname, src in raw_books:
        raw_names.append(fname)
        for source in src.worksheets:
            if source.sheet_state != "visible":
                continue
            base = "原账单-" + re.sub(r"[\\/*?:\[\]]", "_", source.title)
            title = base[:31]
            n = 2
            while title in used:
                suffix = "(%d)" % n
                title = base[:31 - len(suffix)] + suffix
                n += 1
            used.add(title)
            target = _copy_raw_sheet(wb, source, title)
            _append_raw_review(target, source, links)
    _build_cover(wb, carrier, period, res, lines, bills, raw_names)
    ordered = [wb["01复核说明书"], wb["02单据运费"], wb["03计提对账单"]]
    ordered += [ws for ws in wb.worksheets if ws not in ordered]
    wb._sheets = ordered
    return wb
