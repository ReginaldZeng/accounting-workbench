# -*- coding: utf-8 -*-
# [Change Log]
# Date: 2026-09-28 | Author: Claude Opus 4.8 | Version: V2.676
# Description: 【物流账单复核】路由（新工具线在后端的落点）。复核=核价(合同价格卡)×核量(金蝶数量)→归一态，接计提。
#   端点：取数说明读 / 导入合同价格卡 / 价格卡读 / 上传账单解析落中间表 / 接金蝶回填数量 / 出复核结果(费用项汇总+逐单)。
#   算法在 kernels/logistics_price + logistics_review + logistics_intake；表在 kernels/logistics_review_store；金蝶只读走 kingdee_client。
#   pilot=迅鸽（取数说明已种子，价格卡导《附件二》，核量取 XQLCK 出库数量）。计提行=复用 logistics_bills 聚合，另接。
import json
import re
import calendar
import time
from datetime import datetime

from fastapi import APIRouter, Request, Response
from sqlalchemy import select, insert, delete, update, func

from core import JSONResponse, _require_perm, db
import kingdee_client as kc
import todo_scenes
from kernels import logistics_review_store as store
from kernels import logistics_price as lp
from kernels import logistics_review as lr
from kernels import logistics_intake as intake
from kernels import logistics_recon as lrc

router = APIRouter()
BL, PC, SP = store.bill_lines, store.price_card, store.intake_spec
BR = store.bill_raw            # 账单原件(V2.843)
SG, LN, CP = store.review_sign, store.review_line_note, store.review_carrier_pts   # 复核登记 / 逐笔差异解释 / 供应商复核要点
FX = store.review_line_fix   # 计提更正(只登记应改为什么，打印交专人去金蝶改)
DK = store.review_doc_ok     # 逐单已确认(复核人核过没问题的单据)
SJ = store.review_subj       # 按主体复核结论(通过/有疑问)


def _subj_marks(period, carrier=None):
    """{(承运商, 主体): {status, note, by, at}}"""
    q = select(SJ).where(SJ.c.period == period)
    if carrier:
        q = q.where(SJ.c.carrier == carrier)
    with db._engine.connect() as c:
        return {(r["carrier"], r["subject"]): {"status": r["status"], "note": r["note"] or "", "by": r["updated_by"] or "",
                                              "at": r["updated_at"] or ""} for r in c.execute(q).mappings().all()}


def _doc_ok(carrier, period):
    """本月已确认的单据 {单号: {by, at}}。"""
    with db._engine.connect() as c:
        return {r[0]: {"by": r[1] or "", "at": r[2] or ""} for r in c.execute(select(
            DK.c.doc_no, DK.c.confirmed_by, DK.c.confirmed_at).where((DK.c.carrier == carrier) & (DK.c.period == period))).all()}

_DETAIL_KEYS = ("period", "carrier", "grain", "subject", "doc_no", "annot", "fee", "bizline",
                "fee_item", "qty", "unit", "amount", "carrier_sub", "prov", "charge_wt",
                "sub_fees", "src_sheet", "src_row", "note")


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _perm(request):
    return _require_perm(request, "logistics_upload")


async def _read_upload(request):
    ctype = request.headers.get("content-type", "")
    if "multipart/form-data" in ctype:
        form = await request.form()
        f = form.get("file") or form.get("files")
        if f is not None and hasattr(f, "read"):
            try:
                request.state.upload_name = getattr(f, "filename", "") or ""     # 存账单原件时记原文件名(V2.843)
            except Exception:
                pass
            return await f.read()
    return await request.body()


def _load_spec(carrier):
    with db._engine.connect() as c:
        r = c.execute(select(SP.c.spec_json).where(SP.c.carrier == carrier)).first()
    return json.loads(r[0]) if r else None


def _load_card(carrier):
    with db._engine.connect() as c:
        rows = c.execute(select(PC).where(PC.c.carrier == carrier)).mappings().all()
    return lp.load_card([dict(r) for r in rows]), len(rows)


# ---------- 取数说明 ----------
@router.get("/api/logistics-review/spec")
def review_spec(request: Request, carrier: str = "迅鸽"):
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    return {"ok": True, "carrier": carrier, "spec": _load_spec(carrier)}


# ---------- 本月有计提的承运商（金蝶 2241 供应商往来·计提凭证汇总）----------
@router.get("/api/logistics-review/carriers")
def review_carriers(request: Request, period: str = ""):
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    accr = {}
    if period and "-" in period:
        y, m = period.split("-")[:2]
        try:
            accr = lrc.accrued_by_carrier(kc.fetch_gl_voucher(int(y), int(m), prefix="2241"))
        except Exception:
            accr = {}
    sup = db.list_logi_suppliers() or []
    full2short = {s.get("full"): s.get("short") for s in sup if s.get("full")}

    def match_short(name):
        # 金蝶摘要提取的承运商名常截在「供应链/物流」，与档案全名「…有限公司」不全等 → 前缀/包含兜底
        if name in full2short:
            return full2short[name]
        for sfull, sshort in full2short.items():
            if sfull and (sfull.startswith(name) or name.startswith(sfull) or name in sfull or sfull in name):
                return sshort
        return None

    with db._engine.connect() as c:
        specs = {r[0] for r in c.execute(select(SP.c.carrier)).all()}
    out = []
    for full, amt in accr.items():
        short = match_short(full) or full
        out.append({"short": short, "full": full, "accrued": round(amt or 0, 2), "has_spec": short in specs})
    # 有计提但没配取数说明的也列出来（灰示"未配"）；再补上已配却本月无计提的（如 pilot 迅鸽），排在后
    listed = {x["short"] for x in out}
    for sc in specs:
        if sc not in listed:
            out.append({"short": sc, "full": "", "accrued": None, "has_spec": True})
    out.sort(key=lambda x: (-(x["accrued"] or -1), x["short"]))
    return {"ok": True, "period": period, "carriers": out, "kd_ok": bool(accr)}


# ---------- 第一页总览：承运商 × 主体 的 计提/付款/差异（金蝶 2241）----------
_SUBJECTS = ["深圳星期零", "深圳星期九", "孝感星期九"]   # 固定三列（其余账簿归「其它」不单列）


def _ov_kd_rows(period, fresh=False):
    """金蝶本期 2241 凭证行(总表与请款单认月份共用，按账期缓存 30 分钟) → (rows, fetched_at, cached)。"""
    import time as _t
    y, m = period.split("-")[:2]
    # 供应商维度(FFLEX4)：名称用金蝶全称、带编码(原来从摘要截名，截在「供应链/物流」不全)，V2.725
    fields = list(kc.GL_VOUCHER_FIELDS) + [("FACCOUNTBOOKID.FName", "账簿"),
                                           ("FDetailID.FFLEX4.FNumber", "供应商码"), ("FDetailID.FFLEX4.FName", "供应商")]
    ck = ("ov2", period)
    cc = _ACCR_CACHE.get(ck)
    if cc and not fresh and _t.time() - cc[2] < 1800:
        return cc[0], cc[1], True
    rows = []
    try:
        s, conf = kc.login()
        rows = kc._query(s, conf, "GL_VOUCHER", fields,
                         "FAccountID.FNumber like '2241%%' and FYear=%d and FPeriod=%d" % (int(y), int(m)))
    except Exception:
        rows = []
    fetched_at = _now()
    if rows:
        _ACCR_CACHE[ck] = (rows, fetched_at, _t.time())
    return rows, fetched_at, False


def _xmonth_picks():
    """付款做账里人工选定、且计提不在账单那个月的(跨月核销) → [{bill 账单月, accr 计提月, vno, subject, code, payee, carrier}]。
    V2.820(用户 2026-10-05「这里应该接的是计提和请款一起」)：复核按「一张请款单 + 它核销的计提」看——诚煜 6 月账单核销 6 月记-518 + 7 月记-522，
    记-522 就算在 6 期诚煜名下、从 7 期拿掉。只读本库设置，不碰金蝶。"""
    picks = db.get_setting("logi_voucher_pick", None) or {}
    if not picks:
        return []
    PRT = store.payreq
    with db._engine.connect() as c:
        rs = [dict(r) for r in c.execute(select(PRT.c.inst_id, PRT.c.period, PRT.c.subject, PRT.c.sup_code, PRT.c.payee, PRT.c.carrier).where(
            PRT.c.inst_id.in_(list(picks.keys())))).mappings().all()]
    out = []
    for r in rs:
        if not r.get("period"):
            continue
        for p in (picks.get(r["inst_id"]) or {}).get("picks") or []:
            try:
                ym = "%04d-%02d" % (int(p["year"]), int(p["month"]))
            except (KeyError, TypeError, ValueError):
                continue
            if ym != r["period"]:
                out.append({"bill": r["period"], "accr": ym, "vno": str(p["vno"]), "subject": r.get("subject"), "code": r.get("sup_code"),
                            "payee": r.get("payee"), "carrier": r.get("carrier")})
    return out


def _ov_xmonth(accr, carriers, period, book2short):
    """总表的计提按跨月核销调整：别的月份、被本期账单选定核销的计提加进来；本期计提被别的月份账单选走的减掉。→ 说明列表。"""
    notes = []
    for x in _xmonth_picks():
        if x["bill"] != period and x["accr"] != period:
            continue
        rows, _, _ = _ov_kd_rows(x["accr"])
        mine = [r for r in rows if str(r.get("FVOUCHERGROUPNO") or "").strip() == x["vno"] and float(r.get("FCREDIT") or 0) > 0
                and str(r.get("供应商码") or "").strip() == (x["code"] or "") and book2short.get(str(r.get("账簿") or "")) == x["subject"]]
        amt_accr = sum(float(r["FCREDIT"]) for r in mine if "计提" in str(r.get("FEXPLANATION") or ""))
        # 直接做账的费用凭证(摘要「××提起支付…」，不带「计提」，如禾享)：本来不算在哪个月的计提里，被账单选定后算到账单那个月；它所在的月份没东西可减
        amt_direct = sum(float(r["FCREDIT"]) for r in mine if "计提" not in str(r.get("FEXPLANATION") or "") and "提起支付" in str(r.get("FEXPLANATION") or ""))
        key = x["code"] or x["payee"]
        if x["bill"] == period and (amt_accr or amt_direct):
            accr[(x["subject"], key)] = accr.get((x["subject"], key), 0.0) + amt_accr + amt_direct
            carriers.setdefault(key, (x["payee"], x["code"] or ""))
            notes.append({"key": key, "subject": x["subject"], "dir": "in", "vno": x["vno"], "month": x["accr"], "amt": round(amt_accr + amt_direct, 2),
                          "direct": bool(amt_direct)})
        elif x["accr"] == period and amt_accr:
            accr[(x["subject"], key)] = accr.get((x["subject"], key), 0.0) - amt_accr
            notes.append({"key": key, "subject": x["subject"], "dir": "out", "vno": x["vno"], "month": x["bill"], "amt": round(amt_accr, 2)})
    return notes


def _ov_accr(rows, period, full2short=None):
    """2241 行 → 计提 {(主体简称, 供应商键): 含税} 与 承运商 {键: (金蝶全称, 编码)}。键=供应商编码，没有则摘要名。"""
    if full2short is None:
        full2short = {s.get("full"): s.get("short") for s in (db.list_logi_suppliers() or []) if s.get("full")}
    book2short = {o.get("full_name"): o.get("short_name") for o in (db.list_orgs() or []) if o.get("full_name")}
    # 计提=贷方(摘要「计提…运费/仓储费/装卸/搬运/物流」)
    import logi_scope
    logi_codes = logi_scope.data()["codes"]      # 物流供应商名单：读金蝶供应商分组(V2.846，原来认编码前缀「物流运输服务」)
    accr, carriers = {}, {}
    for r in rows:
        z = str(r.get("FEXPLANATION") or "")
        book = book2short.get(str(r.get("账簿") or ""), None)
        cr = r.get("FCREDIT") or 0
        sname, scode = str(r.get("供应商") or "").strip(), str(r.get("供应商码") or "").strip()
        # 物流供应商(编码 物流运输服务…/在物流供应商档案里)的计提：摘要有「计提」即算，不再要求带费用关键词
        # (记-562「计提武汉顺鸽…8月线上kikiherb-花草茶」没写仓储费，原来被漏)
        is_logi = scode in logi_codes or (sname and sname in full2short)
        if cr and "计提" in z and _accr_is_current(z, period) and (is_logi or any(k in z for k in lrc._ACCR_KW)):
            mo = lrc._ACCR_RE.search(z)
            if not sname and not mo:
                continue
            cf = sname or mo.group(1)
            key = scode or cf
            carriers[key] = (cf, scode)
            accr[(book, key)] = accr.get((book, key), 0.0) + float(cr)
    try:
        xm = _ov_xmonth(accr, carriers, period, book2short)
    except Exception:
        xm = []
    _OV_XM[period] = xm
    return accr, carriers


_OV_XM = {}     # period → 这期总表的跨月核销调整(给总表格子里写一句)


def _register_carriers():
    """走登记制的承运商(取数说明里 review_mode=register)：没有可逐单核的账单——卡板租赁(路凯)、国际快递寄样(禾享)、议价报销(货拉拉)。
    V2.826(用户 2026-10-05 同意路凯、禾享走登记制)：总表上只核对 计提 / 请款 两个数，点「登记」即可，不进逐单核价核量。"""
    out = set()
    with db._engine.connect() as c:
        for cr, sj in c.execute(select(SP.c.carrier, SP.c.spec_json)).all():
            try:
                if (json.loads(sj or "{}") or {}).get("review_mode") == "register":
                    out.add(cr)
            except Exception:
                pass
    return out


@router.get("/api/logistics-review/overview")
def review_overview(request: Request, period: str = "", fresh: int = 0):
    """总表。金蝶 2241 凭证按账期缓存 30 分钟(fresh=1 强制重取)；账单应付与复核状态每次从本库现算(便宜、改了即时)。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    if not period or "-" not in period:
        return JSONResponse({"ok": False, "msg": "缺账期"}, status_code=400)
    sup = db.list_logi_suppliers() or []
    full2short = {s.get("full"): s.get("short") for s in sup if s.get("full")}

    def sup_short(name):
        if name in full2short:
            return full2short[name]
        for sfull, sshort in full2short.items():
            if sfull and (sfull.startswith(name) or name.startswith(sfull) or name in sfull or sfull in name):
                return sshort
        return name

    rows, fetched_at, cached = _ov_kd_rows(period, fresh)
    accr, carriers = _ov_accr(rows, period, full2short)
    # 本月付款（按复核结果）：本月已复核账单应付合计，按承运商×主体（权责发生制·同期间比，非金蝶跨月现金借方）
    with db._engine.connect() as c:
        specs = {r[0] for r in c.execute(select(SP.c.carrier)).all()}
        brs = [dict(r) for r in c.execute(select(BL.c.carrier, BL.c.grain, BL.c.subject, BL.c.subj_ovr, BL.c.amount).where(
            (BL.c.period == period) & (BL.c.grain.in_(("accrual", "detail"))))).mappings().all()]
        signed = {r["carrier"]: {"reviewer": r["reviewer"], "signed_at": r["signed_at"]} for r in c.execute(
            select(SG).where((SG.c.period == period) & (SG.c.status == "signed"))).mappings().all()}
        # 复核进度(总表看得出做到哪)：差异解释/确认单据/计提更正 各几笔
        prog = {}
        for tb, k in ((LN, "notes"), (DK, "docs_ok"), (FX, "fixes")):
            q = select(tb.c.carrier, func.count()).where(tb.c.period == period)
            if tb is LN:
                q = q.where(func.length(func.coalesce(LN.c.note, "")) > 0)
            for cr, n in c.execute(q.group_by(tb.c.carrier)).all():
                prog.setdefault(cr, {})[k] = int(n)
    # 账单应付与逐笔复核同口径：有费用项汇总行(迅鸽)用汇总行，否则逐单；主体走覆盖列+全称→简称(极鲜达账单写全称)
    has_acc = {r["carrier"] for r in brs if r["grain"] == "accrual" and (r["amount"] or 0)}
    billmap = {}
    for r in brs:
        if (r["grain"] == "accrual") != (r["carrier"] in has_acc):
            continue
        k = (_eff_subject(r), str(r["carrier"]))
        billmap[k] = round(billmap.get(k, 0.0) + float(r["amount"] or 0), 2)
    smarks = _subj_marks(period)
    out = {}
    for key, (cf, scode) in carriers.items():
        short = sup_short(cf)
        cells = {}
        tot_accr = 0.0
        for subj in _SUBJECTS:
            a = round(accr.get((subj, key), 0.0), 2)
            p = billmap.get((subj, short), 0.0)      # 本月复核应付（该承运商本月账单复核后金额）
            cells[subj] = {"accr": a, "paid": p, "diff": round(a - p, 2)}
            tot_accr += a
        pg = prog.get(short, {})
        has_bill = any(k[1] == short for k in billmap)
        live = [sj for sj in _SUBJECTS if cells[sj]["accr"] or cells[sj]["paid"]]
        for sj in _SUBJECTS:
            mk = smarks.get((short, sj))
            if mk:
                cells[sj]["mark"] = mk
        n_ok = sum(1 for sj in live if (smarks.get((short, sj)) or {}).get("status") == "ok")
        n_q = sum(1 for sj in live if (smarks.get((short, sj)) or {}).get("status") == "question")
        # 通过了但计提≠复核应付：账单核过了，说明计提记错，要红冲更正(用户 2026-10-01)
        n_fix = sum(1 for sj in live if (smarks.get((short, sj)) or {}).get("status") == "ok" and abs(cells[sj]["diff"]) >= 0.01)
        if n_ok or n_q:
            pg = {**pg, "subj_ok": n_ok, "subj_q": n_q, "subj_n": len(live), "subj_fix": n_fix}
        # 状态：已登记 > 复核中(做过解释/确认/更正) > 账单已传 > 未传账单 > 未配取数说明
        st = ("signed" if signed.get(short) else ("allfix" if pg.get("subj_fix") else "allok") if (pg.get("subj_n") and pg.get("subj_ok") == pg.get("subj_n"))
              else "doing" if any(pg.values()) else "billed" if has_bill
              else "nobill" if short in specs else "nospec")
        for n in _OV_XM.get(period) or []:          # 跨月核销：格子里写明这格的计提加了/减了哪张别的月份的凭证
            if n["key"] == key and n["subject"] in cells:
                cells[n["subject"]].setdefault("xm", []).append({k: n.get(k) for k in ("dir", "vno", "month", "amt", "direct")})
        out[key] = {"carrier": cf, "code": scode, "short": short, "full": cf, "has_spec": short in specs, "cells": cells,
                    "total_accr": round(tot_accr, 2), "signed": signed.get(short), "status": st, "progress": pg}

    def _natkey(x):   # 按供应商编码自然排序(物流运输服务009 < 物流运输服务063)，没编码的排最后
        c = x.get("code") or ""
        return (0 if c else 1, [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", c)], x["carrier"])
    rowlist = list(out.values())
    try:     # 钉钉请款单(V2.730)：每个主体格挂进度、没计提但有请款的补一行
        from routers import logistics_payreq as LPQ
        payreq = LPQ.overview_merge(period, rowlist, u)
    except Exception as e:
        payreq = {"err": str(e)[:200]}
    try:     # 登记制：没有账单可复核，「付款(复核)」那一格用请款金额顶上，和计提比
        reg = _register_carriers()
        for x in rowlist:
            if x.get("short") not in reg:
                continue
            x["register"] = True
            for cl in (x.get("cells") or {}).values():
                live = [v for v in cl.get("reqs") or [] if v["st"]["key"] != "void"]
                if live and not cl.get("paid"):
                    cl["paid"] = round(sum(float(v.get("amount") or 0) for v in live), 2)
                    cl["diff"] = round(float(cl.get("accr") or 0) - cl["paid"], 2)
                    cl["reg"] = True
            if not x.get("signed"):
                # 登记制的多半不计提、直接做费用凭证：金蝶里没「计提」不是问题，但要在付款做账里把它的费用凭证选定了，这里才有数可比
                has_accr = any(float(cl.get("accr") or 0) for cl in (x.get("cells") or {}).values())
                x["status"] = "register" if has_accr else "regnoacc"
    except Exception:
        pass
    rowlist = sorted(rowlist, key=_natkey)
    return {"ok": True, "period": period, "subjects": _SUBJECTS, "rows": rowlist, "kd_ok": bool(rows),
            "fetched_at": fetched_at, "cached": cached, "payreq": payreq}


# ---------- 价格卡：导入合同价目表 / 读 ----------
@router.post("/api/logistics-review/price-card/import")
async def price_card_import(request: Request, carrier: str = "迅鸽"):
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    data = await _read_upload(request)
    if not data:
        return JSONResponse({"ok": False, "msg": "未收到文件"}, status_code=400)
    try:
        rows = lp.parse_contract_xlsx(data, carrier)
    except Exception:
        return JSONResponse({"ok": False, "msg": "价目表解析失败，请确认是合同价目表原格式"}, status_code=400)
    with db._engine.begin() as c:
        c.execute(delete(PC).where(PC.c.carrier == carrier))
        for r in rows:
            c.execute(insert(PC).values(updated_by=u["name"], updated_at=_now(), **r))
    db.audit(u["name"], "物流复核-导入价格卡", carrier, "%d 行；源合同价目表（已签署）" % len(rows))
    return {"ok": True, "rows": len(rows)}


@router.get("/api/logistics-review/price-card")
def price_card_read(request: Request, carrier: str = "迅鸽"):
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    with db._engine.connect() as c:
        rows = [dict(r) for r in c.execute(select(PC).where(PC.c.carrier == carrier)).mappings().all()]
    return {"ok": True, "carrier": carrier, "rows": rows}


# ---------- 上传账单解析 → 中间表 ----------
@router.post("/api/logistics-review/parse")
async def review_parse(request: Request, carrier: str = "迅鸽", period: str = ""):
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    spec = _load_spec(carrier)
    if not spec:
        return JSONResponse({"ok": False, "msg": "该承运商还没配取数说明"}, status_code=400)
    data = await _read_upload(request)
    if not data:
        return JSONResponse({"ok": False, "msg": "未收到文件"}, status_code=400)
    r = import_bill(carrier, period, data, u["name"], fname=getattr(request.state, "upload_name", "") or "")
    if not r.get("ok"):
        return JSONResponse(r, status_code=400)
    return r


def _save_raw_bill(carrier, period, src, data, operator, origin, fname):
    """导入账单时把原文件存一份(同一份账单重导就换掉)，导出复核结果时附在后面。存不成不影响导入。V2.843"""
    import hashlib
    try:
        data = bytes(data)
        with db._engine.begin() as c:
            scope = (BR.c.carrier == carrier) & (BR.c.period == period)
            if src:      # 一家一月几份账单：只换同一份(没标份的旧件一并换)，和 bill_lines 的替换范围一致
                scope = scope & ((BR.c.bill_src == src) | (BR.c.bill_src.is_(None)) | (BR.c.bill_src == ""))
            c.execute(delete(BR).where(scope))
            c.execute(insert(BR).values(carrier=carrier, period=period, bill_src=src or "", name=(fname or "")[:200], origin=(origin or "")[:200],
                                        sha256=hashlib.sha256(data).hexdigest(), size=len(data), data=data, created_by=operator, created_at=_now()))
    except Exception:
        pass


def _raw_bill_files(carrier, period):
    """这家这月的账单原件 → [(文件名, bytes)]。先取导入时存下的(V2.843 起)；之前导的没存，退回钉钉请款单里已导入的那份账单附件(按内容去重)。只读。"""
    out, seen = [], set()
    try:
        with db._engine.connect() as c:
            for n, sha, data in c.execute(select(BR.c.name, BR.c.sha256, BR.c.data).where(
                    (BR.c.carrier == carrier) & (BR.c.period == period)).order_by(BR.c.id)).all():
                if data and sha not in seen:
                    seen.add(sha); out.append((n or "", bytes(data)))
    except Exception:
        out = []
    if out:
        return out
    PRT, PFT = store.payreq, store.payreq_file
    try:
        with db._engine.connect() as c:
            for n, sha, data in c.execute(select(PFT.c.name, PFT.c.sha256, PFT.c.data).select_from(PFT.join(PRT, PRT.c.inst_id == PFT.c.inst_id)).where(
                    (PRT.c.carrier == carrier) & (PRT.c.period == period) & (PRT.c.bill_state == "imported") & (PFT.c.role == "bill") &
                    ((PRT.c.excluded.is_(None)) | (PRT.c.excluded == ""))).order_by(PFT.c.id)).all():
                if data and sha not in seen:
                    seen.add(sha); out.append((n or "", bytes(data)))
    except Exception:
        pass
    return out


def _attach_raw_sheets(wb, files):
    """把账单原件逐表原样附到导出里(页名「原账单-表名」)：值、合并格、列宽行高、字体底色边框数字格式照搬(图片/印章带不过来)；
    老格式 .xls 只搬值；隐藏的表不搬。返回附上的页名。"""
    from copy import copy
    from io import BytesIO
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.utils import get_column_letter
    used, added, total = set(wb.sheetnames), [], 0

    def title(base):
        t = ("原账单-" + re.sub(r"[\\/*?:\[\]]", "_", str(base or "表")))[:31]
        k = 2
        while t in used:
            suf = "(%d)" % k
            t = t[:31 - len(suf)] + suf
            k += 1
        used.add(t)
        return t

    for _fname, data in files:
        try:
            if bytes(data[:4]) == bytes.fromhex("d0cf11e0"):       # 老 .xls
                for sh in intake.open_book(data).worksheets:
                    ws = wb.create_sheet(title(sh.title))
                    for row in sh.iter_rows(values_only=True):
                        try:
                            ws.append(list(row))
                        except Exception:
                            ws.append([None if v is None else str(v) for v in row])
                    added.append(ws.title)
                continue
            src = openpyxl.load_workbook(BytesIO(data), data_only=True)
        except Exception:
            continue
        for sh in src.worksheets:
            if sh.sheet_state != "visible":
                continue
            ws = wb.create_sheet(title(sh.title))
            # 只走表里真有的格子(_cells)：有的账单把格式刷到第 100 万行，按行列范围 iter_rows 会把整片空格子都建出来，导出卡死。
            cells = list(getattr(sh, "_cells", {}).values())
            styled = len(cells) <= 40000 and total <= 200000       # 特别大的表只搬值和数字格式，免得导出太慢
            total += len(cells)
            for cell in cells:
                if isinstance(cell, MergedCell) or (cell.value is None and not (styled and cell.has_style)):
                    continue
                try:
                    nc = ws.cell(row=cell.row, column=cell.column, value=cell.value)
                except Exception:
                    nc = ws.cell(row=cell.row, column=cell.column, value=re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(cell.value)))
                nc.number_format = cell.number_format
                if styled and cell.has_style:
                    nc.font = copy(cell.font); nc.fill = copy(cell.fill); nc.border = copy(cell.border); nc.alignment = copy(cell.alignment)
            for rg in sh.merged_cells.ranges:
                try:
                    ws.merge_cells(str(rg))
                except Exception:
                    pass
            for dim in sh.column_dimensions.values():
                if dim.width and dim.min and dim.max:
                    for i in range(dim.min, min(dim.max, 300) + 1):
                        ws.column_dimensions[get_column_letter(i)].width = dim.width
            for k, dim in list(sh.row_dimensions.items())[:20000]:
                if dim.height:
                    ws.row_dimensions[k].height = dim.height
            try:
                ws.freeze_panes = sh.freeze_panes
            except Exception:
                pass
            added.append(ws.title)
    return added


def import_bill(carrier, period, data, operator, origin="手工上传", fname=""):
    """解析一份账单落 bill_lines（上传页与钉钉请款单自动导入共用）→ {ok, bill_src, detail, accrual, skipped, sources} / {ok:False,msg}。"""
    spec = _load_spec(carrier)
    if not spec:
        return {"ok": False, "msg": "该承运商还没配取数说明"}
    spec["period"] = period
    try:
        res = intake.parse_bill(spec, data)
    except Exception:
        return {"ok": False, "msg": "账单解析失败，请核对取数说明与账单格式"}
    batch = "%s|%s|%s" % (carrier, period, datetime.now().strftime("%Y%m%d%H%M%S"))
    src = (res.get("bill_src") or "").strip()
    with db._engine.begin() as c:
        # 一家一月可有几份账单(货主不同，如迅鸽 starfield/kikiherb)：只替换同一份；没标份的旧数据视同本份一并替换(V2.715)
        scope = (BL.c.carrier == carrier) & (BL.c.period == period) & ((BL.c.review_mode.is_(None)) | (BL.c.review_mode != "register"))
        if src:
            scope = scope & ((BL.c.bill_src == src) | (BL.c.bill_src.is_(None)) | (BL.c.bill_src == ""))
        # V2.788：人工改过的归类(特批主体/费用类型+原因)按单号带到新导入的行上——原来重导一次就全没了(顺丰冷运 8 月因此一直不敢重导)
        keep = {}
        for o in c.execute(select(BL.c.doc_no, BL.c.subj_ovr, BL.c.fee_ovr, BL.c.ovr_reason).where(scope)).all():
            if o[0] and (o[1] or o[2] or o[3]):
                keep.setdefault(o[0], (o[1], o[2], o[3]))
        c.execute(delete(BL).where(scope))
        for grain in ("detail", "accrual"):
            for r in res[grain]:
                vals = {k: r.get(k) for k in _DETAIL_KEYS}
                if r.get("doc_no") in keep:
                    vals["subj_ovr"], vals["fee_ovr"], vals["ovr_reason"] = keep[r.get("doc_no")]
                vals["bill_src"] = src or None
                vals["batch_id"] = batch
                vals["review_mode"] = "audit"
                vals["created_at"] = _now()
                c.execute(insert(BL).values(**vals))
    db.audit(operator, "物流复核-解析账单", "%s %s" % (carrier, period),
             "%s：明细 %d 行 / 计提 %d 行 / 跳过 %d 表" % (origin, len(res["detail"]), len(res["accrual"]), len(res["skipped"])))
    _bust(carrier, period)   # 账单重解析 → 逐笔/逐单视图缓存作废
    _save_raw_bill(carrier, period, src, data, operator, origin, fname)
    with db._engine.connect() as c:   # 本月现有几份账单(按货主)：汇总行合计，没有汇总行用明细合计
        srcs = {}
        for r in c.execute(select(BL.c.bill_src, BL.c.grain, func.count(), func.sum(BL.c.amount)).where(
                (BL.c.carrier == carrier) & (BL.c.period == period)).group_by(BL.c.bill_src, BL.c.grain)).all():
            d = srcs.setdefault(r[0] or "（未标货主）", {})
            d[r[1]] = (int(r[2]), round(float(r[3] or 0), 2))
    sources = [{"src": k, "amount": (v.get("accrual") or v.get("detail") or (0, 0))[1],
                "n_detail": (v.get("detail") or (0, 0))[0], "n_accrual": (v.get("accrual") or (0, 0))[0]} for k, v in sorted(srcs.items())]
    return {"ok": True, "bill_src": src, "detail": len(res["detail"]), "accrual": len(res["accrual"]),
            "skipped": res["skipped"], "sources": sources}


# ---------- 接金蝶回填出库数量（核量）----------
# 按重量核量的承运商（干线/冷运：核量比金蝶出库重量kg，不是件数）
_WEIGHT_CARRIERS = {"顺丰冷运", "天鹰物流", "顺丰速运"}   # 顺丰速运(V2.732)：快递寄样按计费重量比金蝶出库重量


def _kd_weight_by_doc(s, conf, docs):
    """按单号→金蝶单据取货物出库重量kg（千克计量物料的基本数量之和）。返回 {单号: kg}。只读。"""
    by_form = {}
    for no in docs:
        d0 = (no or "").split("+")[0]
        if not d0 or d0 == "无单据":
            continue
        pre = "".join(ch for ch in d0 if ch.isalpha())
        for form in _FORM_BY_PREFIX.get(pre, ["SAL_OUTSTOCK"])[:1]:
            by_form.setdefault(form, set()).add(d0)
    mats = _fetch_doc_materials(s, conf, by_form) if by_form else {}
    out = {}
    for no, ms in mats.items():
        kg = 0.0
        for m in ms:
            u = str(m.get("基本单位") or "")
            if "千克" in u or "kg" in u.lower():
                try:
                    kg += float(m.get("基本数量") or 0)
                except (TypeError, ValueError):
                    pass
        out[no] = round(kg, 2)
    return out


@router.post("/api/logistics-review/kingdee-qty")
def review_kingdee_qty(request: Request, carrier: str = "迅鸽", period: str = ""):
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    if not period or "-" not in period:
        return JSONResponse({"ok": False, "msg": "缺账期"}, status_code=400)
    _bust(carrier, period)   # 「接金蝶核量」＝强制重取：逐单视图缓存作废，下次按金蝶最新物料重建
    # 取本批 detail 单号前缀（迅鸽=XQLCK），按月拉出库单聚合数量（货品，剔包装）
    with db._engine.connect() as c:
        docs = [r[0] for r in c.execute(select(BL.c.doc_no).where(
            (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain == "detail"))).all()]
    # 按重量核量的承运商（顺丰冷运/天鹰）：逐单取金蝶出库重量kg 回填 kd_qty
    if carrier in _WEIGHT_CARRIERS:
        try:
            s, conf = kc.login()
            wmap = _kd_weight_by_doc(s, conf, docs)
        except Exception:
            return JSONResponse({"ok": False, "msg": "金蝶取数失败，稍后重试；不影响已存复核结果"}, status_code=502)
        hit = _fill_kd(carrier, period, wmap)
        db.audit(u["name"], "物流复核-接金蝶核量(重量)", "%s %s" % (carrier, period),
                 "只读取数；单据 %d，回填 %d 行(kg)" % (len(wmap), hit))
        return {"ok": True, "kd_docs": len(wmap), "filled": hit, "by": "weight"}
    # 按件数核量(迅鸽等)：单号前缀→对应金蝶单据(XQLCK 销售出库 / RK 销售退货 / QTCK 其他出库 / FBDR 调拨…)，
    # 按单号直接查、不限日期(跨月发货也取得到)，件数=货品数量件之和(剔包装)。V2.717：原来只查当月销售出库、按单号前缀回填(短号会套到长号上)。
    PACK = ("纸箱", "电商专供袋", "拉链", "气泡", "胶带", "气枕", "葫芦膜", "编织袋", "文件封")
    by_form = {}
    for no in docs:
        d0 = (no or "").split("+")[0].strip()
        if not d0 or d0 == "无单据":
            continue
        pre = "".join(ch for ch in d0 if ch.isalpha())
        by_form.setdefault(_FORM_BY_PREFIX.get(pre, ["SAL_OUTSTOCK"])[0], set()).add(d0)
    try:
        s, conf = kc.login()
        mats = _fetch_doc_materials(s, conf, by_form) if by_form else {}
    except Exception:
        return JSONResponse({"ok": False, "msg": "金蝶取数失败，稍后重试；不影响已存复核结果"}, status_code=502)
    qty = {}
    for no, ms in mats.items():
        q = 0.0
        for m in ms:
            if any(p in str(m.get("名称") or "") for p in PACK):
                continue
            try:
                q += float(m.get("数量件") if m.get("数量件") not in (None, "") else (m.get("基本数量") or 0))
            except (TypeError, ValueError):
                pass
        qty[no] = round(q, 2)
    hit = _fill_kd(carrier, period, qty)
    db.audit(u["name"], "物流复核-接金蝶核量", "%s %s" % (carrier, period),
             "只读取数；金蝶单据 %d 张(%s)，回填 %d 行" % (len(qty), "/".join(sorted(by_form)), hit))
    return {"ok": True, "kd_docs": len(qty), "filled": hit, "want_docs": sum(len(v) for v in by_form.values())}


def _fill_kd(carrier, period, kmap):
    """金蝶量回填中间表：先清空本月旧值，再按「首个单号完全相等」回填(不用前缀匹配)。返回回填行数。"""
    with db._engine.begin() as c:
        rows = c.execute(select(BL.c.id, BL.c.doc_no).where(
            (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain == "detail"))).all()
        ids_by = {}
        for rid, no in rows:
            ids_by.setdefault((no or "").split("+")[0].strip(), []).append(rid)
        c.execute(update(BL).where((BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain == "detail")).values(kd_qty=None))
        hit = 0
        for no, v in kmap.items():
            ids = ids_by.get(no)
            if ids:
                c.execute(update(BL).where(BL.c.id.in_(ids)).values(kd_qty=v))
                hit += len(ids)
    return hit


# ---------- 登记制（议价/报销：货拉拉等）：单据运费·其他单据 ----------
_FORM_BY_PREFIX = {
    "FBDR": ["STK_TransferIn", "STK_TransferOut"], "FBDC": ["STK_TransferOut", "STK_TransferIn"],
    "CGRK": ["STK_InStock"], "XSCKD": ["SAL_OUTSTOCK"], "XQLCK": ["SAL_OUTSTOCK", "STK_MisDelivery"],   # 迅鸽 XQLCK 也有其他出库单(2026-09 实证 268 张)
    "QTCK": ["STK_MisDelivery"], "RK": ["SAL_RETURNSTOCK"], "CGTL": ["PUR_MRB"],
}
# 各单据的「基本单位数量」字段 Key 大小写不同（逐单据写死，缺列降级）
_QTYFIELD_BY_FORM = {
    "STK_TransferIn": "FBaseQty", "STK_TransferOut": "FBaseQty", "STK_InStock": "FBaseUnitQty",
    "SAL_OUTSTOCK": "FBaseUnitQty", "STK_MisDelivery": "FBaseQty", "SAL_RETURNSTOCK": "FBaseunitQty",
    "PUR_MRB": "FBASEUNITQTY",
}


@router.post("/api/logistics-review/register")
async def review_register_add(request: Request):
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    try:
        b = await request.json()
    except Exception:
        return JSONResponse({"ok": False, "msg": "请求格式错误"}, status_code=400)
    carrier, period, doc_no = (b.get("carrier") or "").strip(), (b.get("period") or "").strip(), (b.get("doc_no") or "").strip()
    amount = b.get("amount")
    if not carrier or not period or amount in (None, ""):
        return JSONResponse({"ok": False, "msg": "承运商、账期、金额必填"}, status_code=400)
    try:
        amount = round(float(amount), 2)
    except Exception:
        return JSONResponse({"ok": False, "msg": "金额须为数字"}, status_code=400)
    note = _reg_note(b.get("dept"), b.get("source"), b.get("date"), b.get("note"))
    with db._engine.begin() as c:
        c.execute(insert(BL).values(
            batch_id="register|%s|%s" % (carrier, period), period=period, carrier=carrier,
            grain="detail", review_mode="register", subject=(b.get("subject") or "").strip(),
            doc_no="+".join([p for p in re.split(r"[+/,，、;；\s]+", doc_no) if p]) if doc_no else "",
            annot=(b.get("annot") or "").strip(), fee_item=(b.get("fee_item") or "运费").strip(),
            amount=amount, unit=(b.get("unit") or "").strip(), note=note, created_at=_now()))
    db.audit(u["name"], "物流复核-登记单据运费", "%s %s" % (carrier, period), "%s %.2f元" % (doc_no or "无单号", amount))
    return {"ok": True}


@router.post("/api/logistics-review/register/delete")
async def review_register_delete(request: Request):
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    rid = b.get("id")
    with db._engine.begin() as c:
        c.execute(delete(BL).where((BL.c.id == rid) & (BL.c.review_mode == "register")))
    return {"ok": True}


_REG_TPL_COLS = ["日期", "费用主体", "费用类型", "业务线", "ERP单据号", "需求部门", "承运商", "运费", "备注"]


def _reg_note(dept, source, date, extra=""):
    return "；".join(x for x in [
        ("需求部门 " + dept) if dept else "",
        ("来源 " + source) if source else "",
        ("日期 " + date) if date else "", (extra or "")] if x)


def _cell_str(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


@router.get("/api/logistics-review/register/template")
def review_register_template(request: Request):
    """下载「其他单据（登记制）」批量导入模板 xlsx。列＝用户现有表：日期/费用主体/费用类型/业务线/ERP单据号/需求部门/承运商/运费/备注。"""
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from io import BytesIO
    wb = Workbook(); ws = wb.active; ws.title = "单据运费登记"
    ws.append(_REG_TPL_COLS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F6E8C")
    ws.append(["2026-08-29", "深圳星期零", "销售出库运费", "星期零电商", "FBDR075727", "电商部门", "货拉拉", 347.4, "示例行，可删"])
    ws.append(["2026-08-06", "深圳星期零", "销售出库运费", "星期零电商", "FBDR074944", "电商部门", "货拉拉", 197.99, ""])
    for i, wd in enumerate([12, 12, 14, 14, 16, 12, 10, 10, 18], start=1):
        ws.column_dimensions[chr(64 + i)].width = wd
    bio = BytesIO(); wb.save(bio)
    return Response(content=bio.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=logistics_register_template.xlsx"})


@router.post("/api/logistics-review/register/import")
async def review_register_import(request: Request, period: str = ""):
    """批量导入其他单据运费（登记制）。解析用户模板：承运商+运费必填，账期按「日期」列取 YYYY-MM（缺则用当前账期）。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    data = await _read_upload(request)
    if not data:
        return JSONResponse({"ok": False, "msg": "未收到文件"}, status_code=400)
    try:
        import openpyxl
        from io import BytesIO
        wb = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        grid = [list(row) for row in ws.iter_rows(values_only=True)]
    except Exception:
        return JSONResponse({"ok": False, "msg": "Excel 读取失败，请用下载的模板另存 .xlsx 再传"}, status_code=400)
    # 找表头行（含「承运商」且含「运费」或「金额」）
    hdr_i, hdr = -1, []
    for i, row in enumerate(grid[:8]):
        txt = [_cell_str(c) for c in row]
        if any("承运商" in t for t in txt) and any(("运费" in t or "金额" in t) for t in txt):
            hdr_i, hdr = i, txt; break
    if hdr_i < 0:
        return JSONResponse({"ok": False, "msg": "没找到表头行（需含「承运商」和「运费」列），请用下载模板"}, status_code=400)

    def col(*keys):
        for j, h in enumerate(hdr):
            if any(k in h for k in keys):
                return j
        return -1
    ci = {"date": col("日期"), "subject": col("费用主体", "主体"), "fee_item": col("费用类型", "类型"),
          "bizline": col("业务线"), "doc_no": col("单据号", "单号", "ERP"), "dept": col("需求部门", "部门"),
          "carrier": col("承运商"), "amount": col("运费", "金额"), "note": col("备注")}
    if ci["carrier"] < 0 or ci["amount"] < 0:
        return JSONResponse({"ok": False, "msg": "表头缺「承运商」或「运费」列"}, status_code=400)

    def get(row, key):
        j = ci[key]
        return _cell_str(row[j]) if 0 <= j < len(row) else ""
    added, skipped, errs = 0, 0, []
    with db._engine.begin() as c:
        for r in grid[hdr_i + 1:]:
            if not any(_cell_str(x) for x in r):
                continue
            carrier = get(r, "carrier"); amt_s = get(r, "amount")
            if not carrier or amt_s == "":
                skipped += 1; continue
            try:
                amount = round(float(str(amt_s).replace(",", "").replace("￥", "").replace("¥", "")), 2)
            except Exception:
                skipped += 1; errs.append("金额非数字：%s" % amt_s); continue
            date = get(r, "date")
            per = period
            m = re.match(r"(\d{4})[-/年.](\d{1,2})", date or "")
            if m:
                per = "%s-%s" % (m.group(1), m.group(2).zfill(2))
            if not per:
                skipped += 1; errs.append("缺账期且日期不可解析"); continue
            doc_no = get(r, "doc_no")
            c.execute(insert(BL).values(
                batch_id="register|%s|%s" % (carrier, per), period=per, carrier=carrier,
                grain="detail", review_mode="register", subject=get(r, "subject"),
                doc_no="+".join([p for p in re.split(r"[+/,，、;；\s]+", doc_no) if p]) if doc_no else "",
                annot=get(r, "bizline"), fee_item=get(r, "fee_item") or "运费",
                amount=amount, unit="", note=_reg_note(get(r, "dept"), "", date, get(r, "note")), created_at=_now()))
            added += 1
    db.audit(u["name"], "物流复核-导入单据运费", period or "多期", "新增 %d 行 / 跳过 %d 行" % (added, skipped))
    return {"ok": True, "added": added, "skipped": skipped, "errs": errs[:5]}


@router.get("/api/logistics-review/register-list")
def review_register_list(request: Request, period: str = ""):
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    with db._engine.connect() as c:
        rows = [dict(r) for r in c.execute(select(BL).where(
            (BL.c.review_mode == "register") & (BL.c.period == period)).order_by(BL.c.id.desc())).mappings().all()]
    lr.review_details(rows, {"express": {}, "unit": {}}, {})
    view = [{k: r.get(k) for k in ("id", "carrier", "subject", "doc_no", "annot", "fee_item",
             "amount", "unit", "kd_kg", "qty_state", "verdict", "note")} for r in rows]
    total = round(sum((r.get("amount") or 0) for r in rows), 2)
    real = sum(1 for r in rows if r.get("qty_state") == "ok")
    miss = sum(1 for r in rows if r.get("qty_state") == "miss")
    return {"ok": True, "period": period, "rows": view, "count": len(rows),
            "total": total, "doc_real": real, "doc_miss": miss}


# 各单据的物料明细字段（编码/名称/基本单位数量/基本单位/往来/销售额），大小写逐单据写死，缺列降级
_DOC_MAT_FIELDS = {
    "SAL_OUTSTOCK": [("FBillNo", "单号"), ("FMaterialID.FNumber", "编码"), ("FMaterialID.FName", "名称"),
                     ("FBaseUnitQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"),
                     ("FCustomerID.FName", "往来"), ("FAllAmount", "销售额"),
                     ("FMaterialID.FSpecification", "规格"), ("FRealQty", "数量件"), ("FUnitID.FName", "计价单位")],
    "STK_TransferIn": [("FBillNo", "单号"), ("FMaterialId.FNumber", "编码"), ("FMaterialId.FName", "名称"),
                       ("FBaseQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"),
                       ("FSrcStockId.FName", "源仓"), ("FDestStockId.FName", "目的仓"),
                       ("FMaterialId.FSpecification", "规格"), ("FQty", "数量件"), ("FUnitId.FName", "计价单位")],
    "STK_TransferOut": [("FBillNo", "单号"), ("FMaterialId.FNumber", "编码"), ("FMaterialId.FName", "名称"),
                        ("FBaseQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"),
                        ("FSrcStockId.FName", "源仓"), ("FDestStockId.FName", "目的仓"),
                        ("FMaterialId.FSpecification", "规格"), ("FQty", "数量件"), ("FUnitId.FName", "计价单位")],
    "STK_InStock": [("FBillNo", "单号"), ("FMaterialId.FNumber", "编码"), ("FMaterialId.FName", "名称"),
                    ("FBaseUnitQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"), ("FSupplierId.FName", "往来"),
                    ("FMaterialId.FSpecification", "规格"), ("FRealQty", "数量件"), ("FUnitId.FName", "计价单位")],
    "STK_MisDelivery": [("FBillNo", "单号"), ("FMaterialID.FNumber", "编码"), ("FMaterialID.FName", "名称"),
                        ("FBaseQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"), ("FDeptId.FName", "往来"),
                        ("FQty", "数量件")],
    # 销售退货单(RK 开头，迅鸽退件表)：字段 2026-09-30 真机实测
    "SAL_RETURNSTOCK": [("FBillNo", "单号"), ("FMaterialId.FNumber", "编码"), ("FMaterialId.FName", "名称"),
                        ("FBaseunitQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"), ("FRetcustId.FName", "往来"),
                        ("FRealQty", "数量件"), ("FUnitID.FName", "计价单位")],
}


def _fetch_doc_materials(s, conf, docs_by_form):
    """按 单据前缀→form 分组，查金蝶物料明细。返回 {单号: [{编码,名称,基本数量,基本单位,往来,销售额}]}。只读，缺列降级。
    同一前缀可能对应几种单据(XQLCK 既有销售出库也有其他出库、FBDR 调入/调出)：首选单据查不到的，按 _FORM_BY_PREFIX 依次再试(V2.718)。"""
    out = _fetch_doc_materials_once(s, conf, docs_by_form)
    tried = {f: set(str(d) for d in ds) for f, ds in docs_by_form.items()}
    left = {str(d) for ds in docs_by_form.values() for d in ds if str(d) not in out}
    for _ in range(3):
        nxt = {}
        for d in left:
            pre = "".join(ch for ch in d if ch.isalpha())
            for form in _FORM_BY_PREFIX.get(pre, ["SAL_OUTSTOCK"]):
                if d not in tried.get(form, set()):
                    nxt.setdefault(form, set()).add(d)
                    break
        if not nxt:
            break
        for f, ds in nxt.items():
            tried.setdefault(f, set()).update(ds)
        out.update(_fetch_doc_materials_once(s, conf, nxt))
        left = {d for d in left if d not in out}
    return out


def _fetch_doc_materials_once(s, conf, docs_by_form):
    out = {}
    for form, docs in docs_by_form.items():
        fields = _DOC_MAT_FIELDS.get(form)
        if not fields:
            continue
        docs = list(docs)
        for i in range(0, len(docs), 200):
            inlist = ",".join("'%s'" % d.replace("'", "") for d in docs[i:i + 200])
            cols = list(fields)
            rr = []
            while cols:
                try:
                    rr = kc._query(s, conf, form, cols, "FBillNo in (%s)" % inlist)
                    break
                except Exception:
                    cols = cols[:-1]
            for r in rr:
                no = r.get("单号")
                if not no:
                    continue
                lai = r.get("往来")
                if r.get("源仓") or r.get("目的仓"):   # 调拨单：往来＝源仓(调出)→目的仓(调入)
                    lai = "%s→%s" % (r.get("源仓") or "?", r.get("目的仓") or "?")
                out.setdefault(str(no), []).append({
                    "编码": r.get("编码"), "名称": r.get("名称"),
                    "基本数量": r.get("基本数量"), "基本单位": r.get("基本单位"),
                    "往来": lai, "销售额": r.get("销售额"),
                    "规格": r.get("规格"), "数量件": r.get("数量件"), "计价单位": r.get("计价单位")})
    return out


_BOX_CARRIERS = {"丰源", "极鲜达", "恒茂", "链盟", "跨越物流", "中通快运", "易风达", "诚煜物流"}  # 按件数/箱核对(有账单重量则按重量)：金蝶数量(袋)÷规格箱规=箱数，整车比箱、打托倒算托规
_QTY_CARRIERS = {"迅鸽"}  # 快递按件数核：账单件数 vs 金蝶出库件数(剔包装)，不做箱规换算
_PACK_KW = ("纸箱", "包装袋", "包材", "运输袋", "编织袋", "拉链", "气泡", "胶带", "气枕", "葫芦膜", "文件封", "缠绕膜", "打托", "托盘", "护角")  # 包材(不摊运费、不进kg基数)
# 费用类型按单据前缀通用推导（所有承运商共用，不再每家写死）
_FEE_BY_PREFIX = {"FBDR": "调拨运费", "FBDC": "调拨运费", "XSCKD": "销售出库运费", "XQLCK": "销售出库运费",
                  "CGRK": "采购入库运费", "QTCK": "其他出库运费", "RK": "销售退货运费", "CGTL": "采购退料运费"}


def _fee_of(doc_no, fallback):
    pre = "".join(ch for ch in (doc_no or "").split("+")[0] if ch.isalpha())
    return _FEE_BY_PREFIX.get(pre, fallback or "运输费")


def _fee_norm(s):
    """把账单侧(采购入库运费/销售出库运费/调拨运费…)与计提侧(入库运费/出库运费…)费用名归一到同一类，
    才能主体×费用类型对上账。识别不了的原样返回。"""
    s = str(s or "").strip()
    if "研发" in s:
        return "研发外购"        # 计提费用项目"研发外购" ↔ 账单标注"研发费用TOB/TOC"
    if "搬运" in s or "设备调拨" in s or "设备转移" in s:
        return "搬运费"          # 设备/托盘调拨不是货物调拨：计提记 5101 搬运费(易风达 8 月记-565 托盘费用，V2.764)，要排在「调拨→入库运费」前面
    if "仓储" in s:
        return "仓储费"
    if "入库" in s:
        return "入库运费"
    if "出库" in s:
        return "出库运费"
    if "调拨" in s or "调入" in s or "调出" in s:
        return "入库运费"        # 业务口径(2026-09-30)：成品调拨、原料调拨都记入库运费
    if "退货" in s or "退料" in s:
        return "退货运费"
    return s or "运费"


_PRIOR_KW = ("红冲", "更正", "冲回", "冲销", "核销")


def _accr_is_current(expl, period):
    """是不是本月的计提分录。摘要带红冲/更正/冲回/冲销/核销(上期计提的更正与核销)，或写明的是别的月份，
    都算上期更正、不计入本月——否则 8 月里做的"红冲 7 月 + 按发票更正 7 月"会把 7 月金额混进 8 月。摘要没写月份的算本月。"""
    z = str(expl or "")
    if any(k in z for k in _PRIOR_KW):
        return False
    mo = re.search(r"(\d{1,2})月", z)
    if mo and period and "-" in period:
        return int(mo.group(1)) == int(period.split("-")[1])
    return True


def _is_outbound_fee(fee):
    return "出库" in str(fee or "")


_ORG_MAP = {"ts": 0.0, "m": {}}


def _short_subject(s):
    """主体归一：账单里写全称(深圳市星期零食品科技有限公司)的，按组织档案 full_name→short_name 转成简称(深圳星期零)，
    与计提侧(账簿全称→简称)同口径，否则同一家会被拆成"全称行只有账单、简称行只有计提"两行对不上。档案5分钟缓存。"""
    import time as _t
    s = str(s or "").strip()
    if not s:
        return s
    if _t.time() - _ORG_MAP["ts"] > 300:
        m = {}
        for o in (db.list_orgs() or []):
            f, sh = str(o.get("full_name") or "").strip(), str(o.get("short_name") or "").strip()
            if f and sh:
                m[f] = sh
        _ORG_MAP["m"], _ORG_MAP["ts"] = m, _t.time()
    return _ORG_MAP["m"].get(s, s)


def _eff_subject(r):
    """账单主体：人工覆盖(subj_ovr) 优先，否则原 subject；再全称→简称归一。"""
    return _short_subject(str(r.get("subj_ovr") or "").strip() or str(r.get("subject") or "").strip())


_CANON_FEES = {"入库运费", "出库运费", "退货运费", "仓储费", "研发外购", "搬运费"}


def _eff_fee(r):
    """账单费用类型(已归一到 入库/出库/调拨/退货运费/仓储费)，优先级：
    ① 人工覆盖 fee_ovr；② 账单原费用名 fee_item 本就是规范类(采购入库/销售出库/调拨/仓储…)就直接用——
       也让老机制里直接改到 fee_item 的归类生效；③ 费用标注 annot 带单据类型(销售出库单-电商/成品仓储-电商)，
       汇总行(无单号)靠它归口——迅鸽 第三方快递费/B2C操作费/物料费 都标"销售出库单"→出库运费；
    ④ 否则(顺丰=冷运运费 这类非规范名)按单号前缀 _fee_of 归口。"""
    if r.get("fee_ovr"):
        fo = _fee_norm(r["fee_ovr"])
        if fo in _CANON_FEES:      # 手改值不在规范类里(如把业务线"TOC"填进了费用类型)当作没填
            return fo
    fn = _fee_norm(r.get("fee_item"))
    if fn in _CANON_FEES:
        return fn
    fa = _fee_norm(r.get("annot"))
    if fa in _CANON_FEES:
        return fa
    return _fee_norm(_fee_of(r.get("doc_no"), r.get("fee_item")))


# —— 费用类型「物流部填」vs「系统判」(V2.735，用户 2026-10-01：他们填归他们填、系统做归系统做，两边对比才能复核对不对) ——
# 比对在归一类上做(出库运费/入库运费/研发外购/仓储费/搬运费)：规范附表 F 的 销售出库费用、出库装卸费用 都归出库运费，
# 成品/原料入库费用、成品/原料调拨费用、入库装卸费用 都归入库运费——科目层面(6601/6401/5101)由计提对账管，这里只核类别。
def _fee_cmp(x):
    x = str(x or "")
    return "出库运费" if x == "退货运费" else x      # 附表 H：销售退货单按销售出库费用


def _fill_fee(r):
    """物流部填的费用类型 → 归一类；没填返回 ''。新账单读「费用类型」列(存 bill_lines.fee)；
    老账单(8 月)从「类别」标注推：研发→研发外购、仓储→仓储费、销售单/样品单/福利领用/退货→出库运费、调拨/入库→入库运费。"""
    v = str(r.get("fee") or "").strip()
    if v:
        return "搬运费" if "设备" in v else _fee_cmp(_fee_norm(v))
    a = str(r.get("annot") or "")
    if not a:
        return ""
    fa = _fee_norm(a)
    if fa in _CANON_FEES:
        return _fee_cmp(fa)
    if any(k in a for k in ("销售单", "样品单", "福利", "退货", "销售出库")):
        return "出库运费"
    if any(k in a for k in ("调拨", "入库")):
        return "入库运费"
    return ""


def _sys_fee(doc_no, parties):
    """系统按金蝶单据判费用类型(规范附表 H)：销售出库/电商出库/销售退货→出库运费；采购入库/退料、分步式调拨→入库运费；
    其他出库单看领料部门——研发中心领的→研发外购，其余(销售中心样品、福利领用…)→出库运费。没单号判不了返回 ''。"""
    pre = "".join(ch for ch in (doc_no or "").split("+")[0] if ch.isalpha())
    if pre in ("XSCKD", "XQLCK", "RK"):
        return "出库运费"
    if pre in ("CGRK", "CGTL", "FBDC", "FBDR"):
        return "入库运费"
    if pre == "QTCK":
        return "研发外购" if any("研发" in str(p or "") for p in (parties or [])) else "出库运费"
    return ""


def _fee_ovr(r):
    """人工定的费用类型(归一到规范类)；没定或填的不是规范类返回 ''。"""
    fo = _fee_norm(r.get("fee_ovr")) if r.get("fee_ovr") else ""
    return fo if fo in _CANON_FEES else ""


def _fee_check(r, doc_no, parties):
    """→ (系统判, 物流部填, 结论 ok/diff/nofill/nosys/na/manual)。仓储、搬运(设备调拨)这类不看单据的，系统不判(na)。
    manual(V2.788)＝人工定过费用类型(特批：如采购入库单的运费经特批由深圳星期零按出库运费承担)，以人工定的为准，不再报「物流部填≠系统判」。"""
    fs, ff = _sys_fee(doc_no, parties), _fill_fee(r)
    if _fee_ovr(r):
        return fs, ff, "manual"
    if ff in ("仓储费", "搬运费"):
        return fs, ff, "na"
    if not fs:
        return fs, ff, "nosys"
    if not ff:
        return fs, ff, "nofill"
    return fs, ff, ("ok" if _fee_cmp(ff) == _fee_cmp(fs) else "diff")


def _box_reg(spec):
    """从规格型号解析箱规（N袋/箱）。返回 int 或 None。"""
    m = re.search(r"(\d+)\s*(?:袋|盒|包|瓶|罐|桶|支|个|件|盒装)\s*/\s*箱", str(spec or ""))   # 10袋/箱、12盒/箱…
    return int(m.group(1)) if m else None


def _box_div(m):
    """金蝶数量折成箱的除数：计价单位本来是箱→1；规格 N袋/箱→N；规格 20kg/箱、10L/件 且金蝶数量单位是千克/升→按每箱千克/升折(恒茂采购原料)。"""
    u = str(m.get("计价单位") or m.get("基本单位") or "").strip()
    if u == "箱":
        return 1
    br = _box_reg(m.get("规格"))
    if br:
        return br
    sp = str(m.get("规格") or "")
    # 每箱千克/升：20kg/箱、10L/件，或 1L*12/箱、1KG*10/箱(=单件×件数)
    mw = re.search(r"(\d+(?:\.\d+)?)\s*(kg|千克|公斤|l|升)\s*(?:[*×xX]\s*(\d+)\s*(?:袋|盒|包|瓶|罐|桶|支|个|件)?)?\s*/\s*(?:箱|件)", sp, re.I)
    if mw:
        su = mw.group(2).lower()
        if (su in ("kg", "千克", "公斤") and u in ("千克", "kg", "公斤")) or (su in ("l", "升") and u in ("升", "l", "L")):
            return float(mw.group(1)) * (float(mw.group(3)) if mw.group(3) else 1.0)
    return None


def _bizline_of(annot):
    """从费用标注取业务线段（植物肉/鲜食/零售/小料/豆蛋制品/电商/山姆零售/kikiherb/海外）。"""
    for b in ("植物肉", "鲜食", "山姆", "零售", "小料", "豆蛋制品", "电商", "kikiherb", "海外"):
        if b in (annot or ""):
            return "山姆零售" if b == "山姆" else b
    return ""


@router.get("/api/logistics-review/doc-freight")
def review_doc_freight(request: Request, period: str = "", mode: str = "other", q: str = "", page: int = 1, size: int = 80,
                       carrier: str = "", ffee: str = ""):
    """物料级单据运费（两 tab 同一套列）：一行=单据的一个物料行，运费按基本数量摊，单位运费=摊得运费/基本数量，费比=运费/销售额。
    mode='other' 登记制(其他单据)；
    mode='sales' 账单复核(已登记)——V2.773(用户 2026-10-03「核对确认后应该补充到逐单运费上」)：
      某家某月在复核台「确认通过并登记已复核」后，它的逐单账单明细(带金蝶单号的)自动进来；撤销登记就退出。不复制数据，直接读中间表、按登记状态过滤
      (登记后当月归类与备注锁定，读到的就是定稿)。原来只认账单行 fee='销售出库费用' 的，各家都进不来。
      所有单据类型都进(销售出库/采购入库/调拨/其他出库)，费用类型＝复核台归口后的(出库运费/入库运费/研发外购…)。carrier/ffee 可筛。"""
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    facets, pending, nodoc = None, [], {"n": 0, "amt": 0.0}
    with db._engine.connect() as c:
        if mode == "sales":
            signed = {r[0]: (r[1], r[2]) for r in c.execute(select(SG.c.carrier, SG.c.reviewer, SG.c.signed_at).where(
                (SG.c.period == period) & (SG.c.status == "signed"))).all()}
            allsrc = [dict(r) for r in c.execute(select(BL).where(
                (BL.c.period == period) & (BL.c.grain == "detail") & (BL.c.review_mode == "audit"))).mappings().all()]
            pend = {}
            for r in allsrc:
                if r.get("carrier") not in signed:
                    pend[r.get("carrier")] = pend.get(r.get("carrier"), 0.0) + float(r.get("amount") or 0)
            pending = [{"carrier": k, "amount": round(v, 2)} for k, v in sorted(pend.items(), key=lambda kv: -kv[1])]
            src = [r for r in allsrc if r.get("carrier") in signed]
            for r in src:
                r["_fee_type"] = _eff_fee(r)
                r["subject"] = _eff_subject(r)
                r["_biz"] = _bill_biz(r)
            fc, ff = {}, {}
            for r in src:
                a = float(r.get("amount") or 0)
                x = fc.setdefault(r.get("carrier"), {"carrier": r.get("carrier"), "amount": 0.0, "by": signed[r.get("carrier")][0], "at": signed[r.get("carrier")][1]})
                x["amount"] = round(x["amount"] + a, 2)
                ff[r["_fee_type"]] = round(ff.get(r["_fee_type"], 0.0) + a, 2)
            facets = {"carriers": sorted(fc.values(), key=lambda x: -x["amount"]), "fees": sorted(ff.items(), key=lambda kv: -kv[1])}
            if carrier:
                src = [r for r in src if r.get("carrier") == carrier]
            if ffee:
                src = [r for r in src if r["_fee_type"] == ffee]
            keep = []
            for r in src:           # 没有单号的(仓储费、账单调整)到不了单据，另计
                if not (r.get("doc_no") or "").strip() or r.get("doc_no") == "无单据":
                    nodoc["n"] += 1
                    nodoc["amt"] += float(r.get("amount") or 0)
                else:
                    keep.append(r)
            src = keep
        else:
            src = [dict(r) for r in c.execute(select(BL).where(
                (BL.c.period == period) & (BL.c.review_mode == "register")).order_by(BL.c.id.desc())).mappings().all()]
            for r in src:
                r["_src"], r["_signed"] = "登记", True
            # V2.784(用户「其它出库单不是应该出现在其它单据类那里吗」)：账单里的非销售出库单据也进这个页签——
            #   其他出库单(QTCK、以及 XQLCK 开头其实是其他出库的)、采购入库(CGRK)、调拨(FBDR/FBDC)、退料等。销售出库单、退货单在第一个页签，不重复。
            signed_o = {r[0] for r in c.execute(select(SG.c.carrier).where((SG.c.period == period) & (SG.c.status == "signed"))).all()}
            aud = [dict(r) for r in c.execute(select(BL).where(
                (BL.c.period == period) & (BL.c.grain == "detail") & (BL.c.review_mode == "audit") &
                (BL.c.doc_no != "") & (BL.c.doc_no != "无单据") & BL.c.doc_no.isnot(None))).mappings().all()]
            for r in aud:
                r["_src"], r["_signed"] = "账单", r.get("carrier") in signed_o
                r["_fee_type"] = _eff_fee(r)
                r["subject"] = _eff_subject(r)
                r["_biz"] = _bill_biz(r)
            src = src + aud
    # 其他单据页签：哪些单号算「销售出库/退货」(不进本页签)——本月出库单里有的；或单号像销售出库(XSCKD/XQLCK/RK)、金蝶里查到是别的月份的销售出库/退货
    sales_nos, where_o = set(), {}
    if mode != "sales":
        try:
            docs_m, _ts = _month_outstock(period)
        except Exception:
            docs_m = {}
        cand = set()
        for r in src:
            if r.get("_src") != "账单":
                continue
            for no in [p for p in (r.get("doc_no") or "").split("+") if p]:
                if no in docs_m:
                    sales_nos.add(no)
                elif "".join(ch for ch in no if ch.isalpha()) in ("XSCKD", "XQLCK", "RK"):
                    cand.add(no)
        where_o = _orphan_where(period, sorted(cand)) if cand else {}
        for no in cand:
            w = where_o.get(no)
            if w and w[0] in ("out", "ret"):
                sales_nos.add(no)
            elif not w and not docs_m:
                sales_nos.add(no)        # 金蝶取不到出库单时退回按单号前缀判，宁可少列也不把销售出库单混进来
    # 按单据汇总运费（一单号多行=累加）。账单复核页同一张单可能有两家的费用(链盟运费 + 天鹰装货费)，按「承运商+单号」分开
    docfee = {}
    for r in src:
        for no in [p for p in (r.get("doc_no") or "").split("+") if p] or ["（无单号）"]:
            if mode != "sales" and r.get("_src") == "账单" and no in sales_nos:
                continue                  # 销售出库/退货单在第一个页签
            key = ((r.get("carrier") or "") + "|" + no) if (mode == "sales" or r.get("_src") == "账单") else no
            g = docfee.setdefault(key, {"src": r.get("_src") or "", "signed": bool(r.get("_signed")),
                                        "运费": 0.0, "no": no, "subject": r.get("subject"), "annot": r.get("annot"),
                                        "fee_item": r.get("fee_item"), "carrier": r.get("carrier"),
                                        "fee_type": r.get("_fee_type") or "", "biz": r.get("_biz"),
                                        "dept": (r.get("note") or ""), "qty_state": r.get("qty_state"),
                                        "reg_id": r.get("id") if (mode != "sales" and r.get("_src") != "账单") else None})   # 只有手工登记的能删
            n = len([p for p in (r.get("doc_no") or "").split("+") if p]) or 1
            g["运费"] += (r.get("amount") or 0) / n
    # 全量口径：总运费/单据数在取金蝶物料前算好（销售出库单可达数千张，不能每次翻页全量拉金蝶）
    all_total = round(sum(g["运费"] for g in docfee.values()), 2)
    # 单据号搜索先按单号过滤（物料名/编码搜索仅在本页已取物料内二次过滤）
    doc_items = sorted(docfee.items(), key=lambda kv: kv[0])
    if q:
        doc_items = [kv for kv in doc_items if q in (kv[1]["no"] or "")]
    doc_total = len(doc_items)
    page = max(1, int(page))
    page_items = doc_items[(page - 1) * size: page * size]  # 只取本页这一批单据
    # 只对本页单据取金蝶物料明细
    by_form = {}
    for _k, _g in page_items:
        no = _g["no"]
        if no == "（无单号）":
            continue
        pre = "".join(ch for ch in no if ch.isalpha())
        for form in _FORM_BY_PREFIX.get(pre, ["SAL_OUTSTOCK"])[:1]:
            by_form.setdefault(form, set()).add(no)
    mats = {}
    if by_form:
        try:
            s, conf = kc.login()
            mats = _fetch_doc_materials(s, conf, by_form)
        except Exception:
            mats = {}
    rows = []
    for dkey, g in page_items:
        no = g["no"]
        fee = round(g["运费"], 2)
        biz = g.get("biz") or _bizline_of(g["annot"])
        lines = mats.get(no) or []
        kgsum = sum(float(m["基本数量"] or 0) for m in lines if m.get("基本数量") not in (None, ""))
        if not lines:
            rows.append({"key": dkey, "fee_type": g.get("fee_type") or "", "src": g.get("src") or "", "signed": g.get("signed"),
                         "subject": g["subject"], "carrier": g["carrier"], "fee_item": g["fee_item"], "bizline": biz, "doc_no": no,
                         "party": g["dept"] if mode == "other" else "", "code": "", "name": "（金蝶无此单据物料）",
                         "baseqty": None, "baseunit": "", "fee": fee, "unitfee": None, "sales": None, "ratio": None,
                         "reg_id": g.get("reg_id")})
            continue
        for m in lines:
            bq = None
            try:
                bq = float(m["基本数量"]) if m.get("基本数量") not in (None, "") else None
            except (TypeError, ValueError):
                bq = None
            share = (bq / kgsum) if (kgsum and bq) else (1.0 / len(lines))
            fline = round(fee * share, 2)
            sales_amt = None
            try:
                sales_amt = float(m["销售额"]) if m.get("销售额") not in (None, "") else None
            except (TypeError, ValueError):
                sales_amt = None
            rows.append({
                "key": dkey, "fee_type": g.get("fee_type") or "", "src": g.get("src") or "", "signed": g.get("signed"),
                "subject": g["subject"], "carrier": g["carrier"], "fee_item": g["fee_item"], "bizline": biz, "doc_no": no,
                "party": (m.get("往来") or "") if (mode == "sales" or g.get("src") == "账单") else (g["dept"] or ""),
                "code": m.get("编码"), "name": m.get("名称"), "baseqty": bq, "baseunit": m.get("基本单位"),
                "fee": fline, "unitfee": round(fline / bq, 4) if bq else None,
                "sales": round(sales_amt, 2) if sales_amt is not None else None,
                "ratio": round(fline / sales_amt, 4) if sales_amt else None,
                "reg_id": g.get("reg_id")})
    # 本页物料名/编码二次过滤（仅在已取物料内，跨页搜索请用单据号）
    if q:
        rows = [r for r in rows if q in (r["doc_no"] or "") or q in (r.get("name") or "") or q in (r.get("code") or "")]
    pages = max(1, (doc_total + size - 1) // size)
    ostat = None
    if mode != "sales":
        def _s(pred):
            gs = [g for g in docfee.values() if pred(g)]
            return {"n": len(gs), "amount": round(sum(g["运费"] for g in gs), 2)}
        ostat = {"bill_ok": _s(lambda g: g.get("src") == "账单" and g.get("signed")),
                 "bill_pending": _s(lambda g: g.get("src") == "账单" and not g.get("signed")),
                 "reg": _s(lambda g: g.get("src") != "账单")}
    return {"ok": True, "period": period, "mode": mode, "count": len(rows), "total": all_total, "other_stat": ostat,
            "doc_count": doc_total, "rows": rows, "page": page, "pages": pages, "size": size,
            "facets": facets, "pending": pending, "nodoc": {"n": nodoc["n"], "amount": round(nodoc["amt"], 2)}}


# ---------- 单据运费·销售出库单（全量）V2.774 ----------
# 用户 2026-10-03：「销售出库这里需要列出来所有的出库单，支持筛选……我需要知道这个月每一笔出库单的运费是多少（除去内部交易的部分）」。
# 以金蝶当月全部销售出库单为底(不是只列有运费的)，把各家账单逐单运费按单号挂上去；没挂到运费的也列出来。
# 内部交易口径同 BP 工作台(kingdee.INTERNAL_CUSTOMERS)：客户＝集团内部主体(孝感卖给深圳两家的那一段)，默认剔除。
#   实证(2026-08)：内部腿的「源单编号」＝对外那张出库单号；账单上填的全是对外单号，没有填内部腿的，剔除不丢运费。
_OS_CACHE = {}      # period -> (ts, {单号: doc})
_OS_TTL = 900
_OS_FIELDS = [("FBillNo", "单号"), ("FDate", "日期"), ("FSaleOrgId.FName", "销售组织"), ("FCustomerID.FName", "客户"),
              ("FMaterialID.FNumber", "编码"), ("FMaterialID.FName", "名称"), ("FBaseUnitQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"),
              ("FAllAmount", "价税合计"), ("FRealQty", "数量"), ("FUnitID.FName", "单位"),
              ("FDocumentStatus", "状态"), ("FBillTypeID.FName", "单据类型"), ("FStockOrgId.FName", "发货组织"), ("FSrcBillNo", "源单"),
              # V2.775 用户要「发货仓库和收货地址」：仓库在分录上(一张单可能几个仓)，收货地址/联系人在单头
              ("FStockID.FName", "仓库"), ("FReceiveAddress", "收货地址"), ("FLinkMan", "联系人")]
# 销售退货单(V2.781，用户「销售退货单是不是也可以加进来」)：字段名和出库单不一样(客户=FRetcustId、基本数量=FBaseunitQty、仓库=FStockId)，2026-10-03 真机实测。
# 进同一张表，销售额、数量记负数(冲减)；RK 开头是对外退货，XSTHD 开头是孝感对深圳的内部镜像(客户=内部主体，照样剔除)。
_RT_FIELDS = [("FBillNo", "单号"), ("FDate", "日期"), ("FSaleOrgId.FName", "销售组织"), ("FRetcustId.FName", "客户"),
              ("FMaterialId.FNumber", "编码"), ("FMaterialId.FName", "名称"), ("FBaseunitQty", "基本数量"), ("FBaseUnitId.FName", "基本单位"),
              ("FAllAmount", "价税合计"), ("FRealQty", "数量"), ("FUnitID.FName", "单位"),
              ("FDocumentStatus", "状态"), ("FBillTypeID.FName", "单据类型"), ("FStockOrgId.FName", "发货组织"), ("FSrcBillNo", "源单"),
              ("FStockId.FName", "仓库"), ("FReceiveAddress", "收货地址"), ("FLinkMan", "联系人")]
_OS_INTERNAL_EXTRA = ("SINKIO LIMITED",)
_BPMAP_CACHE = {}   # "m" -> (ts, combo, mat, err)


def _bp_mapping(fresh=False):
    """BP 工作台「客户物料映射表」→ ({(客户名, 物料码): 行}, {物料码: 首行}, 错误)。只读，缓存 15 分钟；BP 不可达返回空表+原因，不拖垮页面。
    匹配规则照 BP(revenue_ledger.build_combo_lookup / build_material_lookup)：(客户名, 物料码) 精确优先，其次该物料第一条映射；
    同一组合多行取首条。8 月实证：对外出库单 7,419 行物料全部组合命中。"""
    hit = _BPMAP_CACHE.get("m")
    if hit and not fresh and time.time() - hit[0] < _OS_TTL:
        return hit[1], hit[2], hit[3]
    import urllib.request
    combo, mat, err = {}, {}, ""
    try:
        with urllib.request.urlopen(db.BP_API_BASE + "/api/master/mapping", timeout=8) as r:
            rows = json.loads(r.read().decode("utf-8"))
        for m in rows or []:
            if str(m.get("status") or "") == "已停用":
                continue
            mc, cn = str(m.get("materialCode") or "").strip(), str(m.get("customerName") or "").strip()
            if not mc:
                continue
            v = {"brand": m.get("brand") or "", "bu": m.get("businessUnit") or "", "division": m.get("division") or "",
                 "series": m.get("productSeries") or "", "team": m.get("salesTeam") or ""}
            if cn:
                combo.setdefault((cn, mc), v)
            mat.setdefault(mc, v)
    except Exception as e:
        err = "BP 工作台取不到客户物料映射表（%s），品牌暂时显示不了" % str(e)[:80]
    if not err:
        _BPMAP_CACHE["m"] = (time.time(), combo, mat, err)
    return combo, mat, err


def _internal_customers():
    return {o.get("full_name") for o in (db.list_orgs() or []) if o.get("full_name")} | set(_OS_INTERNAL_EXTRA)


_OS_KIND = "kd:sal_outstock"        # period_inputs 里的 kind：金蝶当月销售出库单+退货单(按单聚合好的)


def _month_outstock(period, fresh=False, operator=""):
    """金蝶当月全部销售出库单+退货单 → ({单号: {…}}, 取数时刻)。只读金蝶。
    V2.783 落库(用户问「是每次进来重新取数金蝶吗」)：取一次就存进 period_inputs(平台通用的「取数一次、进页面直接读」那张表)，
    重启不丢、全员共用；以后进页面直接读库，只有点「刷新金蝶」(fresh) 才重取。进程里另留一份免得每次请求都解压。
    原来只放内存 15 分钟，过期或后端重启就得再等 5～10 秒。"""
    y, m = int(period[:4]), int(period[5:7])
    hit = _OS_CACHE.get(period)
    if hit and not fresh:
        return hit[1], hit[0]
    if not fresh:
        try:
            snap = db.get_period_input("kingdee", y, m, _OS_KIND)
        except Exception:
            snap = None
        if snap and isinstance(snap.get("payload"), dict) and snap["payload"].get("docs"):
            ts = float((snap.get("meta") or {}).get("ts") or 0) or time.time()
            _OS_CACHE[period] = (ts, snap["payload"]["docs"])
            return _OS_CACHE[period][1], ts
    nxt = "%04d-%02d-01" % ((y + 1, 1) if m == 12 else (y, m + 1))
    s, conf = kc.login()
    cols, rows = list(_OS_FIELDS), []
    while len(cols) >= 9:               # 后面几列是锦上添花，哪列本账套没有就丢哪列；前 9 列缺了就报错
        try:
            rows = kc._query(s, conf, "SAL_OUTSTOCK", cols, "FDate>='%s-01' and FDate<'%s'" % (period, nxt))
            break
        except Exception:
            if len(cols) == 9:
                raise
            cols = cols[:-1]
    # 销售退货单：取不到不影响出库单(降级成没有退货)
    rcols, rrows = list(_RT_FIELDS), []
    while len(rcols) >= 9:
        try:
            rrows = kc._query(s, conf, "SAL_RETURNSTOCK", rcols, "FDate>='%s-01' and FDate<'%s'" % (period, nxt))
            break
        except Exception:
            if len(rcols) == 9:
                rrows = []
                break
            rcols = rcols[:-1]
    docs = {}
    for r, sign, kind in [(x, 1, "out") for x in rows] + [(x, -1, "ret") for x in rrows]:
        no = str(r.get("单号") or "")
        if not no:
            continue
        d = docs.setdefault(no, {"kind": kind, "no": no, "date": str(r.get("日期") or "")[:10], "org": _short_subject(r.get("销售组织") or ""),
                                 "stock_org": _short_subject(r.get("发货组织") or ""), "customer": r.get("客户") or "",
                                 "status": r.get("状态") or "", "btype": r.get("单据类型") or "", "src": r.get("源单") or "",
                                 "addr": str(r.get("收货地址") or "").strip(), "linkman": str(r.get("联系人") or "").strip(),
                                 "stocks": [], "amount": 0.0, "kg": 0.0, "units": {}, "lines": []})
        wh = str(r.get("仓库") or "").strip()
        if wh and wh not in d["stocks"]:
            d["stocks"].append(wh)
        try:
            bq = sign * float(r.get("基本数量") or 0)
        except (TypeError, ValueError):
            bq = 0.0
        try:
            am = sign * float(r.get("价税合计") or 0)
        except (TypeError, ValueError):
            am = 0.0
        u = str(r.get("基本单位") or "")
        d["amount"] += am
        if "千克" in u or "kg" in u.lower():
            d["kg"] += bq
        d["units"][u] = d["units"].get(u, 0.0) + bq
        d["lines"].append({"code": r.get("编码"), "name": r.get("名称"), "baseqty": bq, "baseunit": u, "amount": round(am, 2), "stock": wh})
    for d in docs.values():
        d["amount"] = round(d["amount"], 2)
        d["kg"] = round(d["kg"], 3)
        d["qty_txt"] = "、".join("%s %s" % (("%.3f" % v).rstrip("0").rstrip("."), k) for k, v in d.pop("units").items() if v)
    _OS_CACHE[period] = (time.time(), docs)
    try:
        db.set_period_input("kingdee", y, m, _OS_KIND, {"docs": docs},
                            meta={"ts": _OS_CACHE[period][0], "n_docs": len(docs), "n_out": sum(1 for d in docs.values() if d.get("kind") != "ret"),
                                  "n_ret": sum(1 for d in docs.values() if d.get("kind") == "ret"), "desc": "金蝶销售出库单+退货单(单据运费·销售出库单页)"},
                            operator=operator or "系统")
    except Exception:
        pass                      # 存不进库不影响本次显示，下次进来再取
    return docs, _OS_CACHE[period][0]


# 费用分三类(V2.778，用户 2026-10-03「运费和装卸费要分开，不然不知道哪些是不含运费的」)：
#   tr 运费＝运输本身及其附加(运费/快递费/提货费/送货费/保费/回单费/压车等其他费用)；
#   ld 装卸费＝装货/卸货/装卸/搬运/上楼；ot 其他＝仓储、冷藏、处置、操作费、箱子(包材)、退货服务费。
#   先看账单行的分项(sub_fees 的每个数字项)，分项加起来不够金额的差额、以及没有分项的行，按该行费用名(fee_item)归类。
_LD_KW = ("装货", "卸货", "装卸", "搬运", "上楼")
_OT_KW = ("仓储", "冷藏", "处置", "操作费", "箱子", "包材", "物料", "退货服务")
_SUB_SKIP = ("标准", "核价差", "公式", "箱型")


def _fee_bucket(name):
    s = str(name or "")
    if any(k in s for k in _LD_KW):
        return "ld"
    if any(k in s for k in _OT_KW):
        return "ot"
    return "tr"


def _line_buckets(amount, fee_item, sub_fees):
    """一行账单 → {tr, ld, ot}，三者之和＝该行金额。"""
    out = {"tr": 0.0, "ld": 0.0, "ot": 0.0}
    amt = float(amount or 0)
    try:
        sf = json.loads(sub_fees) if sub_fees else {}
    except Exception:
        sf = {}
    used = 0.0
    for k, v in (sf or {}).items():
        if not isinstance(v, (int, float)) or isinstance(v, bool) or k in _SUB_SKIP or k.startswith("含税金额"):
            continue
        out[_fee_bucket(k)] += float(v)
        used += float(v)
    if abs(amt - used) > 0.005:
        out[_fee_bucket(fee_item)] += amt - used
    return out


def _outstock_freight(period):
    """各家账单逐单费用按单号归集 → {单号: {ok, pending, tr, ld, ot, parts:{(承运商, 账期, 已复核): {tr, ld, ot}}}}。
    已登记复核的承运商×账期 + 登记制 算「已复核」，其余算「待复核」。一格多单号的平摊。账单账期不限(跨月发货也挂得上)。"""
    with db._engine.connect() as c:
        signed = {(r[0], r[1]) for r in c.execute(select(SG.c.carrier, SG.c.period).where(SG.c.status == "signed")).all()}
        bl = c.execute(select(BL.c.carrier, BL.c.period, BL.c.doc_no, BL.c.amount, BL.c.review_mode, BL.c.fee_item, BL.c.sub_fees).where(
            (BL.c.grain == "detail") & (BL.c.doc_no != "") & (BL.c.doc_no != "无单据") & BL.c.doc_no.isnot(None))).all()
    out = {}
    for car, per, doc, amt, mode, fee_item, sub in bl:
        nos = [p for p in str(doc).split("+") if p]
        if not nos:
            continue
        ok = mode == "register" or (car, per) in signed
        bk = _line_buckets(amt, fee_item, sub)
        n = len(nos)
        for no in nos:
            o = out.setdefault(no, {"ok": 0.0, "pending": 0.0, "tr": 0.0, "ld": 0.0, "ot": 0.0, "parts": {}})
            o["ok" if ok else "pending"] += float(amt or 0) / n
            p = o["parts"].setdefault((car, per, ok), {"tr": 0.0, "ld": 0.0, "ot": 0.0})
            for k in ("tr", "ld", "ot"):
                o[k] += bk[k] / n
                p[k] += bk[k] / n
    return out


_ORPH_CACHE = {}    # (period, 单号集合指纹) -> (ts, {单号: (去向 kind, 月份, 往来)})


def _orphan_where(period, nos):
    """账单上写了销售出库单号、但本月销售出库单里没有的 → 去金蝶按单号查它到底在哪(V2.782，用户问「会不会留存，很可能出现在下个月的出库单上」)：
    别的月份的销售出库单(kind=out，带月份；费用按单号挂，切到那个月会自动出现) / 其他出库单(kind=mis，如迅鸽 XQLCK 开头的媒介领用)
    / 销售退货单(kind=ret) / 都查不到(kind=none，多半单号填错)。只读，缓存 15 分钟。"""
    key = (period, hash(tuple(sorted(nos))))
    hit = _ORPH_CACHE.get(key)
    if hit and time.time() - hit[0] < _OS_TTL:
        return hit[1]
    found = {}
    try:
        s, conf = kc.login()
        for form, kind, party in (("SAL_OUTSTOCK", "out", "FCustomerID.FName"), ("STK_MisDelivery", "mis", "FDeptId.FName"),
                                  ("SAL_RETURNSTOCK", "ret", "FRetcustId.FName")):
            left = [n for n in nos if n not in found]
            for i in range(0, len(left), 200):
                inl = ",".join("'%s'" % n.replace("'", "") for n in left[i:i + 200])
                try:
                    rr = kc._query(s, conf, form, [("FBillNo", "单号"), ("FDate", "日期"), (party, "往来")], "FBillNo in (%s)" % inl)
                except Exception:
                    rr = []
                for r in rr:
                    found.setdefault(str(r.get("单号")), (kind, str(r.get("日期") or "")[:7], r.get("往来") or ""))
    except Exception:
        return {}
    _ORPH_CACHE[key] = (time.time(), found)
    return found


@router.get("/api/logistics-review/outstock-freight")
async def review_outstock_freight(request: Request, period: str = "", internal: int = 0, q: str = "", org: str = "", btype: str = "",
                                  carrier: str = "", state: str = "", sort: str = "date", page: int = 1, size: int = 100, fresh: int = 0,
                                  brand: str = "", stock: str = "", dir: str = ""):
    """销售出库单全量 × 物流费用。internal=1 连内部交易一起列。
    state: has 有运费 / none 没有运费(含只有装卸费的) / ldonly 只有装卸费 / nofee 什么费用都没有 / pending 有待复核 / ok 都已复核。
    费用分 运费 fee_tr / 装卸费 fee_ld / 其他 fee_ot，fee＝三者合计；费比＝合计÷销售额。sort 任一列 + dir。"""
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    if not period or len(period) != 7:
        return JSONResponse({"ok": False, "msg": "缺账期"}, status_code=400)
    from starlette.concurrency import run_in_threadpool
    try:
        _u = _perm(request) or {}
        docs, ts = await run_in_threadpool(_month_outstock, period, bool(fresh), _uname(_u) if _u else "")
        if fresh:
            db.audit(_uname(_u) if _u else "", "单据运费-刷新金蝶", period, "重取销售出库单+退货单 %d 张（只读金蝶）" % len(docs))
    except Exception as e:
        return JSONResponse({"ok": False, "msg": "金蝶取销售出库单失败：%s" % str(e)[:160]}, status_code=502)
    fr = _outstock_freight(period)
    inner = _internal_customers()
    combo, matmap, bp_err = await run_in_threadpool(_bp_mapping, bool(fresh))

    def bp_of(cust, code):
        return combo.get((str(cust or "").strip(), str(code or "").strip())) or matmap.get(str(code or "").strip()) or {}
    rows = []
    n_internal = 0
    for no, d in docs.items():
        is_in = d["customer"] in inner
        if is_in:
            n_internal += 1
            if not internal:
                continue
        f = fr.get(no) or {"ok": 0.0, "pending": 0.0, "tr": 0.0, "ld": 0.0, "ot": 0.0, "parts": {}}
        fee = round(f["ok"] + f["pending"], 2)
        bps = [bp_of(d["customer"], ln.get("code")) for ln in d["lines"]]       # 品牌照 BP 的客户物料映射：一张单几个物料可能几个品牌
        brands = list(dict.fromkeys(b.get("brand") for b in bps if b.get("brand")))
        is_ret = d.get("kind") == "ret"
        rows.append({"kind": d.get("kind") or "out", "brand": "、".join(brands), "brands": brands,
                     "bu": "、".join(dict.fromkeys(b.get("bu") for b in bps if b.get("bu"))),
                     "stock": "、".join(d.get("stocks") or []), "stocks": d.get("stocks") or [],
                     "addr": d.get("addr") or "", "linkman": d.get("linkman") or "",
                     "no": no, "date": d["date"], "org": d["org"], "stock_org": d["stock_org"], "customer": d["customer"],
                     "btype": d["btype"], "status": d["status"], "internal": is_in, "amount": d["amount"], "kg": d["kg"], "qty_txt": d["qty_txt"],
                     "n_lines": len(d["lines"]), "fee_ok": round(f["ok"], 2), "fee_pending": round(f["pending"], 2), "fee": fee,
                     "fee_tr": round(f["tr"], 2), "fee_ld": round(f["ld"], 2), "fee_ot": round(f["ot"], 2),
                     "ratio": round(fee / d["amount"], 4) if d["amount"] and fee and not is_ret else None,      # 退货单不算费比
                     "carriers": [{"carrier": k[0], "period": k[1], "signed": k[2], "amount": round(v["tr"] + v["ld"] + v["ot"], 2),
                                   "tr": round(v["tr"], 2), "ld": round(v["ld"], 2), "ot": round(v["ot"], 2)} for k, v in sorted(f["parts"].items())]})
    # 账单写了销售出库单号、但本月出库单里没有的(发货在别的月 / 单号填错 / 其实是其他出库单)
    orphan = {"n": 0, "amount": 0.0}
    with db._engine.connect() as c:
        cur = c.execute(select(BL.c.doc_no, BL.c.amount, BL.c.carrier).where((BL.c.period == period) & (BL.c.grain == "detail") &
                                                                             (BL.c.doc_no != "") & BL.c.doc_no.isnot(None))).all()
    oitems = {}
    for doc, amt, car in cur:
        nos = [p for p in str(doc).split("+") if p]
        for no in nos:
            if "".join(ch for ch in no if ch.isalpha()) in ("XSCKD", "XQLCK") and no not in docs:
                orphan["n"] += 1
                orphan["amount"] += float(amt or 0) / len(nos)
                it = oitems.setdefault(no, {"no": no, "carrier": car, "amount": 0.0})
                it["amount"] += float(amt or 0) / len(nos)
    # 这些单到底在哪：别的月份的出库单 / 其他出库单 / 退货单 / 金蝶查不到
    where = await run_in_threadpool(_orphan_where, period, list(oitems)) if oitems else {}
    KIND_CN = {"out": "%s 的销售出库单", "mis": "其他出库单（不是销售出库）", "ret": "销售退货单", "none": "金蝶里查不到这个单号"}
    ogrp = {}
    for no, it in oitems.items():
        w = where.get(no)
        kind, mon, party = (w[0], w[1], w[2]) if w else ("none", "", "")
        it.update(kind=kind, month=mon, party=party, amount=round(it["amount"], 2),
                  where=(KIND_CN["out"] % mon) if kind == "out" else KIND_CN[kind])
        g = ogrp.setdefault((kind, mon if kind == "out" else ""), {"kind": kind, "month": mon if kind == "out" else "", "n": 0, "amount": 0.0, "label": it["where"]})
        g["n"] += 1
        g["amount"] += it["amount"]
    order = {"out": 0, "ret": 1, "mis": 2, "none": 3}
    orphan["groups"] = [dict(g, amount=round(g["amount"], 2)) for g in sorted(ogrp.values(), key=lambda g: (order[g["kind"]], g["month"]))]
    orphan["items"] = sorted(oitems.values(), key=lambda x: (order[x["kind"]], x["month"], -x["amount"]))[:400]
    base = rows                       # 口径内全部(剔内部后)：统计卡片按筛选前后各给一份
    bcnt = {}
    for r in base:
        for b in r["brands"] or ["（未映射品牌）"]:
            bcnt[b] = bcnt.get(b, 0) + 1
    facets = {"orgs": sorted({r["org"] for r in base if r["org"]}), "btypes": sorted({r["btype"] for r in base if r["btype"]}),
              "carriers": sorted({x["carrier"] for r in base for x in r["carriers"]}),
              "brands": [[k, v] for k, v in sorted(bcnt.items(), key=lambda kv: (-kv[1], kv[0]))],
              "stocks": sorted({s for r in base for s in r["stocks"]})}
    if brand:
        rows = [r for r in rows if (brand in r["brands"]) or (brand == "（未映射品牌）" and not r["brands"])]
    if stock:
        rows = [r for r in rows if stock in r["stocks"]]
    if org:
        rows = [r for r in rows if r["org"] == org]
    if btype:
        rows = [r for r in rows if r["btype"] == btype]
    if carrier:
        rows = [r for r in rows if any(x["carrier"] == carrier for x in r["carriers"])]
    # 有/没有运费只对出库单说(退货单本来就不该有运费)
    if state == "has":
        rows = [r for r in rows if r["kind"] == "out" and r["fee_tr"]]
    elif state == "none":
        rows = [r for r in rows if r["kind"] == "out" and not r["fee_tr"]]
    elif state == "ldonly":
        rows = [r for r in rows if r["kind"] == "out" and not r["fee_tr"] and r["fee_ld"]]
    elif state == "nofee":
        rows = [r for r in rows if r["kind"] == "out" and not r["fee"]]
    elif state == "pending":
        rows = [r for r in rows if r["fee_pending"]]
    elif state == "ok":
        rows = [r for r in rows if r["fee"] and not r["fee_pending"]]
    if q:
        ql = q.strip().lower()
        rows = [r for r in rows if ql in r["no"].lower() or ql in (r["customer"] or "").lower() or ql in (r["addr"] or "").lower()
                or ql in (r["brand"] or "").lower() or ql in (r["linkman"] or "").lower()]
    # 排序(V2.777 点表头排)：sort=列，dir=asc/desc；不传 dir 时金额类默认从大到小、其余从小到大。空值(没有费比/没有品牌…)永远排最后。
    NUM = {"amount": "amount", "fee": "fee", "fee_pending": "fee_pending", "ratio": "ratio", "kg": "kg",
           "fee_tr": "fee_tr", "fee_ld": "fee_ld", "fee_ot": "fee_ot"}
    TXT = {"date": "date", "no": "no", "org": "org", "brand": "brand", "customer": "customer", "stock": "stock", "addr": "addr",
           "btype": "btype", "carrier": None}
    col = sort if (sort in NUM or sort in TXT) else "date"
    desc = (dir == "desc") if dir in ("asc", "desc") else (col in NUM)
    if col in NUM:
        val = lambda r: r.get(NUM[col])
    elif col == "carrier":
        val = lambda r: "、".join(x["carrier"] for x in r["carriers"])
    else:
        val = lambda r: r.get(TXT[col]) or ""
    has = [r for r in rows if val(r) not in (None, "")]
    non = [r for r in rows if val(r) in (None, "")]
    has.sort(key=lambda r: r["no"])                       # 先按单号排稳，再按所选列排(稳定排序，同值的按单号)
    has.sort(key=val, reverse=desc)
    non.sort(key=lambda r: r["no"])
    rows = has + non

    def stat(rs):
        fee, amt = sum(r["fee"] for r in rs), sum(r["amount"] for r in rs)
        outs = [r for r in rs if r["kind"] == "out"]
        rets = [r for r in rs if r["kind"] == "ret"]
        amt_has = sum(r["amount"] for r in outs if r["fee_tr"]) + sum(r["amount"] for r in rets)
        return {"n": len(rs), "n_out": len(outs), "n_ret": len(rets), "amount_ret": round(sum(r["amount"] for r in rets), 2),
                "fee_ret": round(sum(r["fee"] for r in rets), 2),
                "n_has": sum(1 for r in outs if r["fee_tr"]), "n_none": sum(1 for r in outs if not r["fee_tr"]),
                "n_ldonly": sum(1 for r in outs if not r["fee_tr"] and r["fee_ld"]), "n_nofee": sum(1 for r in outs if not r["fee"]),
                "fee_tr": round(sum(r["fee_tr"] for r in rs), 2), "fee_ld": round(sum(r["fee_ld"] for r in rs), 2),
                "fee_ot": round(sum(r["fee_ot"] for r in rs), 2),
                "fee_ok": round(sum(r["fee_ok"] for r in rs), 2), "fee_pending": round(sum(r["fee_pending"] for r in rs), 2),
                "fee": round(fee, 2), "amount": round(amt, 2), "ratio": round(fee / amt, 4) if amt > 0 else None,
                "amount_none": round(amt - amt_has, 2)}
    page = max(1, int(page))
    size = max(20, min(500, int(size)))
    pages = max(1, (len(rows) + size - 1) // size)
    view = rows[(page - 1) * size: page * size]
    for r in view:
        r["lines"] = [dict(ln, brand=bp_of(r["customer"], ln.get("code")).get("brand") or "",
                           series=bp_of(r["customer"], ln.get("code")).get("series") or "") for ln in docs[r["no"]]["lines"]]
    return {"ok": True, "period": period, "fetched_at": datetime.fromtimestamp(ts).strftime("%m-%d %H:%M"),
            "stale_days": int((time.time() - ts) // 86400), "internal": bool(internal),
            "n_internal": n_internal, "all": stat(base), "cur": stat(rows), "facets": facets, "bp_err": bp_err,
            "orphan": {"n": orphan["n"], "amount": round(orphan["amount"], 2), "groups": orphan.get("groups") or [], "items": orphan.get("items") or []},
            "rows": view, "page": page, "pages": pages, "size": size, "sort": col, "dir": "desc" if desc else "asc"}


# 单据运费·销售出库 tab（旧·按单据汇总，保留兼容）：
@router.get("/api/logistics-review/doc-sales")
def review_doc_sales(request: Request, period: str = "", q: str = "", page: int = 1, size: int = 50):
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    with db._engine.connect() as c:
        rows = [dict(r) for r in c.execute(select(BL).where(
            (BL.c.period == period) & (BL.c.grain == "detail") & (BL.c.review_mode == "audit") &
            (BL.c.fee == "销售出库费用"))).mappings().all()]
    # 按单据汇总
    docs = {}
    for r in rows:
        key = r.get("doc_no") or "无单号"
        g = docs.setdefault(key, {"doc_no": key, "carrier": r.get("carrier"), "subject": r.get("subject"),
                                  "annot": r.get("annot"), "amount": 0.0, "charge_wt": r.get("charge_wt"),
                                  "kd_qty": r.get("kd_qty"), "n": 0})
        g["amount"] += r.get("amount") or 0
        g["n"] += 1
    out = sorted(docs.values(), key=lambda x: -x["amount"])
    for g in out:
        g["amount"] = round(g["amount"], 2)
    if q:
        out = [g for g in out if q in (g["doc_no"] or "") or q in (g["carrier"] or "")]
    page = max(1, int(page))
    total_amt = round(sum(g["amount"] for g in out), 2)
    return {"ok": True, "period": period, "count": len(out), "total": total_amt,
            "rows": out[(page - 1) * size: page * size], "page": page, "size": size, "detail_total": len(out)}


@router.post("/api/logistics-review/register/kingdee-check")
def review_register_kingdee_check(request: Request, period: str = ""):
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    with db._engine.connect() as c:
        rows = [dict(r) for r in c.execute(select(BL.c.id, BL.c.doc_no).where(
            (BL.c.review_mode == "register") & (BL.c.period == period))).mappings().all()]
    # 按前缀分组单号 → 查对应金蝶单据存在性（只读）
    form_docs = {}
    for r in rows:
        for no in [p for p in (r["doc_no"] or "").split("+") if p]:
            pre = "".join(ch for ch in no if ch.isalpha())
            for form in _FORM_BY_PREFIX.get(pre, ["SAL_OUTSTOCK"]):
                form_docs.setdefault(form, set()).add(no)
    exists = set()
    kg = {}   # 单号 → 金蝶基本单位数量之和（kg/吨/件，看物料基本单位）
    try:
        s, conf = kc.login()
        for form, docs in form_docs.items():
            qf = _QTYFIELD_BY_FORM.get(form, "FBaseQty")
            docs = list(docs)
            for i in range(0, len(docs), 200):
                chunk = docs[i:i + 200]
                inlist = ",".join("'%s'" % d.replace("'", "") for d in chunk)
                try:
                    rr_list = kc._query(s, conf, form, [("FBillNo", "单号"), (qf, "基本数量")], "FBillNo in (%s)" % inlist)
                except Exception:
                    try:
                        rr_list = kc._query(s, conf, form, [("FBillNo", "单号")], "FBillNo in (%s)" % inlist)
                    except Exception:
                        rr_list = []
                for rr in rr_list:
                    no = rr.get("单号")
                    if not no:
                        continue
                    exists.add(str(no))
                    try:
                        kg[str(no)] = kg.get(str(no), 0.0) + float(rr.get("基本数量") or 0)
                    except (TypeError, ValueError):
                        pass
    except Exception:
        return JSONResponse({"ok": False, "msg": "金蝶取数失败，稍后重试"}, status_code=502)
    real = miss = 0
    with db._engine.begin() as c:
        for r in rows:
            nos = [p for p in (r["doc_no"] or "").split("+") if p]
            st = "ok" if (nos and all(n in exists for n in nos)) else "miss"
            kd = round(sum(kg.get(n, 0.0) for n in nos), 2) if nos else None
            if st == "ok":
                real += 1
            else:
                miss += 1
            c.execute(update(BL).where(BL.c.id == r["id"]).values(qty_state=st, kd_kg=kd))
    db.audit(u["name"], "物流复核-登记轻核单号", period, "单号真实 %d / 查无 %d（只读）" % (real, miss))
    return {"ok": True, "real": real, "miss": miss}


# ---------- 复核结果（费用项汇总 + 逐单）----------
def _accr_lines(carrier, period, carrier_full=None):
    """取该承运商本月金蝶计提，业务线/费用项目从核算维度(产品分类FF100010、费用项目FFLEX9)读，不抠摘要。
    计提含税＝费用借方(6*/5*不含税)＋暂估进项税(2221.01.07)。返回([{subject,bizline,fee,label,amt}], 含税合计)。"""
    if not period or "-" not in period:
        return [], 0.0
    try:
        y, m = period.split("-")[:2]
        s, conf = kc.login()
        # 计提摘要用的是承运商真名——简称(如"迅鸽")多是别名，金蝶摘要写全称("武汉顺鸽科技有限公司")，
        # 故简称+全称一起 OR 匹配，别名对不上就靠全称兜底。
        names = [carrier] + ([carrier_full] if (carrier_full and carrier_full != carrier) else [])
        name_cond = " or ".join("FEXPLANATION like '%%计提%%%s%%'" % n for n in names)
        cf = "FYear=%d and FPeriod=%d and FDEBIT>0 and (%s)" % (int(y), int(m), name_cond)
        rows = kc._query(s, conf, "GL_VOUCHER",
                         [("FACCOUNTBOOKID.FName", "账簿"), ("FDEBIT", "借"), ("FAccountID.FNumber", "科目"),
                          ("FVOUCHERGROUPID.FName", "凭证字"), ("FVOUCHERGROUPNO", "凭证序号"),
                          ("FDetailID.FF100010.FDataValue", "产品分类"), ("FDetailID.FF100006.FDataValue", "产品项目"),
                          ("FDetailID.FFLEX9.FName", "费用项目")],
                         cf + " and (FAccountID.FNumber like '6%%' or FAccountID.FNumber like '5%%')")
        tx = kc._query(s, conf, "GL_VOUCHER", [("FDEBIT", "借")],
                       cf + " and FAccountID.FNumber like '2221.01.07%%'")
    except Exception:
        return [], 0.0
    orgs = db.list_orgs() or []
    b2s = {o.get("full_name"): o.get("short_name") for o in orgs if o.get("full_name")}
    agg = {}
    vnos = {}
    for r in rows:
        amt = float(r.get("借") or 0)
        if not amt:
            continue
        subj = b2s.get(str(r.get("账簿") or ""), str(r.get("账簿") or ""))
        # 业务线＝产品分类(主)；产品项目仅作补充(专标山姆/kikiherb)
        cls = str(r.get("产品分类") or "").strip()
        proj = str(r.get("产品项目") or "").strip()
        biz = cls or proj or "（无业务线）"
        if proj and proj not in biz and any(k in proj for k in ("山姆", "kiki", "KIKI", "Kiki")):
            biz = "%s·%s" % (biz, proj) if cls else proj
        fee = str(r.get("费用项目") or "").strip() or "运费"
        agg[(subj, biz, fee)] = agg.get((subj, biz, fee), 0.0) + amt
        zi = str(r.get("凭证字") or "记").strip()
        no = str(r.get("凭证序号") or "").strip()
        if no:
            vnos.setdefault((subj, biz, fee), set()).add("%s-%s" % (zi, no))
    out = [{"subject": k[0], "bizline": k[1], "fee": k[2], "fee_norm": _fee_norm(k[2]),
            "label": "%s · %s" % (k[1], k[2]), "amt": round(v, 2),
            "vno": "、".join(sorted(vnos.get(k, [])))}
           for k, v in sorted(agg.items())]
    tax = round(sum(float(x.get("借") or 0) for x in tx), 2)
    if tax:
        out.append({"subject": "—", "bizline": "进项税", "fee": "暂估进项税", "fee_norm": "暂估进项税",
                    "label": "暂估进项税", "amt": tax, "vno": ""})
    return out, round(sum(x["amt"] for x in out), 2)


def _carrier_full(carrier):
    sup = db.list_logi_suppliers() or []
    return next((x.get("full") for x in sup if x.get("short") == carrier), None)


def _build_recon(request, carrier, period):
    """复核结论：金蝶2241计提 vs 账单，按【主体×费用类型】归口对账，业务线是计提侧明细。
    计提侧费用不含税(6*/5*)按毛率等比毛成含税(＋暂估进项税)，与账单(含税)同口径，故合计对平、逐组差异才有意义。
    返回 {groups:[{subject,fee_type,lines:[{bizline,vno,amt}],accr_amt,bill_amt,sales,ratio,diff}], accr_total,bill_total,tax}。"""
    lines, atot = _accr_lines(carrier, period, _carrier_full(carrier))
    tax = next((a["amt"] for a in lines if a.get("subject") == "—"), 0.0)
    feelines = [a for a in lines if a.get("subject") != "—"]
    feetot = sum(a["amt"] for a in feelines) or 0.0
    gf = (atot / feetot) if feetot else 1.0            # 含税毛率：费用行×gf=含税，Σ=计提含税合计
    # 账单金额(含税)：便宜、无需金蝶。**优先取费用项汇总行(grain=accrual，迅鸽月结清单那种，主体干净费用项全)，
    # 没有汇总行才按逐单明细(detail)汇总**——与顶栏 total_bill 同口径；逐单只是支撑，像迅鸽的操作费/物料费没有逐单。
    with db._engine.connect() as c:
        allrows = [dict(r) for r in c.execute(select(
            BL.c.grain, BL.c.subject, BL.c.doc_no, BL.c.fee_item, BL.c.annot, BL.c.amount, BL.c.subj_ovr, BL.c.fee_ovr).where(
            (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain.in_(("accrual", "detail"))))).mappings().all()]
    acc_rows = [r for r in allrows if r.get("grain") == "accrual" and (r.get("amount") or 0)]
    brows = acc_rows if acc_rows else [r for r in allrows if r.get("grain") == "detail"]
    bill_src = "accrual" if acc_rows else "detail"
    bill_by = {}
    for r in brows:
        k = (_eff_subject(r), _eff_fee(r))
        bill_by[k] = bill_by.get(k, 0.0) + float(r.get("amount") or 0)
    # 销售额(费比)：需金蝶物料，仅逐单不多(≤400)时取，避免迅鸽这类几千单拖垮(按 detail 单数判，别被5条汇总行骗)
    sales_by = {}
    n_detail = sum(1 for r in allrows if r.get("grain") == "detail")
    sales_ready = n_detail <= 400
    if sales_ready:
        res = review_result(request, carrier=carrier, period=period, group="all", page=1, size=1000000, q="")
        if not isinstance(res, JSONResponse):
            for d0 in res.get("detail", []):
                if d0.get("sales") is None:
                    continue
                k = (str(d0.get("subject") or "").strip(), _fee_norm(d0.get("fee_item")))
                sales_by[k] = sales_by.get(k, 0.0) + float(d0.get("sales") or 0)
    # 计提按(主体,费用norm)归组，业务线明细(金额毛成含税)
    accr_by = {}
    for a in feelines:
        k = (str(a["subject"]).strip(), a.get("fee_norm") or _fee_norm(a.get("fee")))
        g = accr_by.setdefault(k, {"total": 0.0, "lines": []})
        amt_tax = round((a["amt"] or 0) * gf, 2)
        g["total"] += amt_tax
        g["lines"].append({"bizline": a.get("bizline") or "（无业务线）", "vno": a.get("vno") or "", "amt": amt_tax})
    keys = sorted(set(list(accr_by.keys()) + list(bill_by.keys())))
    groups = []
    for k in keys:
        subj, fnorm = k
        ag = accr_by.get(k, {"total": 0.0, "lines": []})
        bamt = round(bill_by.get(k, 0.0), 2)
        accr_amt = round(ag["total"], 2)
        sales = round(sales_by[k], 2) if k in sales_by else None
        ratio = round(bamt / sales, 4) if (sales and _is_outbound_fee(fnorm)) else None
        groups.append({"subject": subj, "fee_type": fnorm,
                       "lines": ag["lines"] or [{"bizline": "（无计提）", "vno": "", "amt": None}],
                       "accr_amt": accr_amt, "bill_amt": bamt, "sales": sales, "ratio": ratio,
                       "diff": round(accr_amt - bamt, 2)})
    return {"ok": True, "carrier": carrier, "period": period, "groups": groups, "accr_total": round(atot, 2),
            "bill_total": round(sum(g["bill_amt"] for g in groups), 2),
            "tax": round(tax, 2), "gross_factor": round(gf, 4), "sales_ready": sales_ready, "bill_src": bill_src}


# ---------- 第二页·逐笔计提复核（V2.690 起的主线；上面的 _build_recon 按组汇总版保留作对照，验证稳定后删）----------
_BIZ_ALIAS = {"其他零售": "零售", "零售其他": "零售", "山姆零售": "零售", "kikiherb": "Kiki Herb", "KikiHerb": "Kiki Herb", "KIKIHERB": "Kiki Herb"}


def _bill_biz(r):
    """账单侧产品线：bizline 列(账单登记/手改)优先，否则从费用标注 annot 的业务线段解析(销售单-小料→小料)，
    研发费用TOB/TOC→TO B/TO C；同义归一(其他零售/山姆零售→零售)。取不到返回空串→该笔只能按组配。"""
    a = str(r.get("annot") or "").strip()
    b = str(r.get("bizline") or "").strip() or _bizline_of(a)
    if not b:
        u = a.upper().replace(" ", "")
        if u.endswith("TOB"):
            b = "TO B"
        elif u.endswith("TOC"):
            b = "TO C"
    return _BIZ_ALIAS.get(b, b)


def _accr_entries(carrier, period, carrier_full=None):
    """逐笔计提(按跨月核销调整后)：本期的计提分录，加上别的月份、被本期账单选定核销的(标 xm=计提月)，
    去掉本期被别的月份账单选走的(列在 xmoved 里提示)。"""
    got = _accr_entries_raw(carrier, period, carrier_full)
    try:
        xs = [x for x in _xmonth_picks() if x["carrier"] == carrier or (carrier_full and x["payee"] == carrier_full)]
    except Exception:
        xs = []
    if not xs:
        return got
    ents, moved = list(got.get("entries") or []), []
    out_k = {("记-" + x["vno"], x["subject"]): x["bill"] for x in xs if x["accr"] == period and x["bill"] != period}
    if out_k:
        keep = []
        for e in ents:
            to = out_k.get((e["vno"], e["subject"]))
            if to:
                m0 = next((m for m in moved if m["vno"] == e["vno"] and m["subject"] == e["subject"]), None)
                if not m0:
                    m0 = {"vno": e["vno"], "subject": e["subject"], "to": to, "amt": 0.0}
                    moved.append(m0)
                m0["amt"] = round(m0["amt"] + (e.get("amt") or 0), 2)
            else:
                keep.append(e)
        ents = keep
    for x in xs:
        if x["bill"] == period and x["accr"] != period:
            for e in (_accr_entries_raw(carrier, x["accr"], carrier_full).get("entries") or []):
                if e["vno"] == "记-" + x["vno"] and e["subject"] == x["subject"]:
                    ents.append(dict(e, xm=x["accr"], key=e["key"] + "|" + x["accr"]))
    ents.sort(key=lambda e: (e["subject"], e["fee_norm"], e["biz"], e["proj"], e["vno"]))
    return dict(got, entries=ents, xmoved=moved)


def _accr_entries_raw(carrier, period, carrier_full=None):
    """逐笔计提：金蝶费用借方(6*/5*)每条分录一行——主体/费用项目/产品线(产品分类)/产品类型(产品项目)/部门/凭证号。
    税率按凭证：同凭证 2221.01.07 借方 ÷ 该凭证费用借方合计；含税＝费用×(1+税率)。
    分录没有稳定ID，行键=凭证号+科目+维度组合。只有 2221.01.06(税额调整)的凭证不计入，另列 adj 提示。
    上期计提的红冲/更正/核销(本月做的)不计入本月，按凭证汇总净额另列 prior 提示。"""
    if not period or "-" not in period:
        return {"entries": [], "adj": [], "prior": []}
    try:
        y, m = period.split("-")[:2]
        s, conf = kc.login()
        names = [carrier] + ([carrier_full] if (carrier_full and carrier_full != carrier) else [])
        cond = " or ".join("FEXPLANATION like '%%计提%%%s%%'" % n for n in names)
        cf = "FYear=%d and FPeriod=%d and FDEBIT<>0 and (%s)" % (int(y), int(m), cond)   # <>0：红冲是负数借方，也要取到才能认出来
        rows = kc._query(s, conf, "GL_VOUCHER",
                         [("FACCOUNTBOOKID.FName", "账簿"), ("FDEBIT", "借"), ("FAccountID.FNumber", "科目"), ("FAccountID.FName", "科目名"),
                          ("FVOUCHERGROUPID.FName", "凭证字"), ("FVOUCHERGROUPNO", "凭证序号"), ("FEXPLANATION", "摘要"),
                          ("FDetailID.FF100010.FDataValue", "产品分类"), ("FDetailID.FF100006.FDataValue", "产品项目"),
                          ("FDetailID.FFLEX5.FName", "部门"), ("FDetailID.FFLEX9.FName", "费用项目"),
                          # 编码(计提更正单按编码找分录)：账簿/费用项目/部门/产品分类/产品项目/供应商(税额行上挂)
                          ("FACCOUNTBOOKID.FNumber", "账簿码"), ("FDetailID.FFLEX9.FNumber", "费用项目码"),
                          ("FDetailID.FFLEX5.FNumber", "部门码"), ("FDetailID.FF100010.FNumber", "产品分类码"),
                          ("FDetailID.FF100006.FNumber", "产品项目码"),
                          ("FDetailID.FFLEX4.FNumber", "供应商码"), ("FDetailID.FFLEX4.FName", "供应商")], cf)
    except Exception:
        return {"entries": [], "adj": [], "prior": [], "suppliers": []}
    orgs = db.list_orgs() or []
    b2s = {o.get("full_name"): o.get("short_name") for o in orgs if o.get("full_name")}
    fee_by_v, tax_by_v, ents, adj, prior = {}, {}, [], [], {}
    sups = {}
    for r in rows:
        if r.get("供应商码"):
            sups[str(r["供应商码"]).strip()] = str(r.get("供应商") or "").strip()
        amt = float(r.get("借") or 0)
        if not amt:
            continue
        acct = str(r.get("科目") or "")
        vno = "%s-%s" % (str(r.get("凭证字") or "记").strip(), str(r.get("凭证序号") or "").strip())
        if not _accr_is_current(r.get("摘要"), period):
            if acct.startswith("6") or acct.startswith("5"):
                subj = b2s.get(str(r.get("账簿") or ""), str(r.get("账簿") or ""))
                p = prior.setdefault((vno, subj), {"vno": vno, "subject": subj, "net": 0.0, "expl": ""})
                p["net"] += amt
                if not p["expl"] or "更正" in str(r.get("摘要") or ""):
                    p["expl"] = re.sub(r"^(红冲|更正|冲回|冲销|核销)[\d/#、]*", "", str(r.get("摘要") or ""))[:40]
            continue
        if amt < 0:
            continue
        if acct.startswith("2221.01.07"):
            tax_by_v[vno] = tax_by_v.get(vno, 0.0) + amt
            continue
        if not (acct.startswith("6") or acct.startswith("5")):
            adj.append({"vno": vno, "acct": acct, "amt": round(amt, 2)})
            continue
        subj = b2s.get(str(r.get("账簿") or ""), str(r.get("账簿") or ""))
        fee = str(r.get("费用项目") or "").strip() or "运费"
        fee_by_v[vno] = fee_by_v.get(vno, 0.0) + amt
        ents.append({"subject": subj, "fee": fee, "fee_norm": _fee_norm(fee),
                     "biz": str(r.get("产品分类") or "").strip() or "（无业务线）",
                     "proj": str(r.get("产品项目") or "").strip(), "dept": str(r.get("部门") or "").strip(),
                     "vno": vno, "acct": acct, "acct_name": str(r.get("科目名") or "").strip(), "amt_net": round(amt, 2),
                     "book_code": str(r.get("账簿码") or "").strip(), "fee_code": str(r.get("费用项目码") or "").strip(),
                     "dept_code": str(r.get("部门码") or "").strip(), "biz_code": str(r.get("产品分类码") or "").strip(),
                     "proj_code": str(r.get("产品项目码") or "").strip()})
    # 同一凭证的进项税按各费用分录金额精确分摊，舍入余数给最后一笔——保证每张凭证 Σ含税 = 费用+税 分毫不差
    by_v = {}
    for e in ents:
        by_v.setdefault(e["vno"], []).append(e)
    for vno, es in by_v.items():
        f = fee_by_v.get(vno, 0.0)
        t = round(tax_by_v.get(vno, 0.0), 2)
        rate = round(t / f, 4) if f else 0.0
        acc = 0.0
        for i, e in enumerate(es):
            e["tax_rate"] = rate
            if i < len(es) - 1:
                e["tax"] = round(e["amt_net"] * t / f, 2) if f else 0.0
                acc += e["tax"]
            else:
                e["tax"] = round(t - acc, 2)
            e["amt"] = round(e["amt_net"] + e["tax"], 2)
            e["key"] = "|".join([e["vno"], e["acct"], e["subject"], e["fee"], e["biz"], e["proj"], e["dept"]])
    ents.sort(key=lambda e: (e["subject"], e["fee_norm"], e["biz"], e["proj"], e["vno"]))
    plist = [{**p, "net": round(p["net"], 2)} for p in sorted(prior.values(), key=lambda p: (p["subject"], p["vno"]))]
    return {"entries": ents, "adj": adj, "prior": plist,
            "suppliers": [{"code": k, "name": v} for k, v in sorted(sups.items())]}


def _build_lines(request, carrier, period):
    """逐笔计提复核（第二页主表）：以计提分录为主线每笔一行，**按主体分组**(组头=主体小计，用户 2026-09-30 定)；
    组内账单先按 费用类型×产品线 配到笔(level=biz)，配不到的退回按 费用类型 挂该段首笔(level=group)；账单有计提无的单独成行。
    账单侧优先费用项汇总行(accrual)，无则逐单(detail)。差异＝计提含税−账单。附：差异解释/上期更正/供应商复核要点/登记状态。"""
    got = _accr_entries(carrier, period, _carrier_full(carrier))
    ents, adj, prior = got["entries"], got["adj"], got.get("prior") or []
    with db._engine.connect() as c:
        allrows = [dict(r) for r in c.execute(select(
            BL.c.grain, BL.c.subject, BL.c.doc_no, BL.c.fee_item, BL.c.annot, BL.c.bizline, BL.c.amount, BL.c.subj_ovr, BL.c.fee_ovr).where(
            (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain.in_(("accrual", "detail"))))).mappings().all()]
        notes = {r[0]: (r[1] or "") for r in c.execute(select(LN.c.line_key, LN.c.note).where(
            (LN.c.carrier == carrier) & (LN.c.period == period))).all()}
        pts = c.execute(select(CP.c.points).where(CP.c.carrier == carrier)).scalar() or ""
        sg = c.execute(select(SG).where((SG.c.carrier == carrier) & (SG.c.period == period) & (SG.c.status == "signed"))).mappings().first()
    acc_rows = [r for r in allrows if r.get("grain") == "accrual" and (r.get("amount") or 0)]
    brows = acc_rows if acc_rows else [r for r in allrows if r.get("grain") == "detail"]
    bill_src = "accrual" if acc_rows else "detail"
    bill3 = {}
    for r in brows:
        k = (_eff_subject(r), _eff_fee(r), _bill_biz(r))
        bill3[k] = bill3.get(k, 0.0) + float(r.get("amount") or 0)
    # 可逐单：每个账单归口里，有多少金额是带金蝶单号的逐单明细撑着的(无单据/只有月结汇总行的只能按汇总核)
    det3 = {}
    for r in allrows:
        if r.get("grain") != "detail":
            continue
        no = (r.get("doc_no") or "").strip()
        if not no or no == "无单据":
            continue
        d = det3.setdefault((_eff_subject(r), _eff_fee(r), _bill_biz(r)), {"docs": set(), "amt": 0.0})
        d["docs"].add(no)
        d["amt"] += float(r.get("amount") or 0)

    def _od(keys, bill):
        # fbiz：产品线集合用 | 连，空产品线记 "-"，第②步按同一口径筛单据
        n = len(set().union(*[det3[k]["docs"] for k in keys if k in det3]))
        amt = round(sum(det3.get(k, {}).get("amt", 0.0) for k in keys), 2)
        ratio = max(0.0, min(1.0, amt / bill)) if bill and bill > 0 else 0.0
        bizs = sorted({k[2] for k in keys})
        return {"n": n, "amt": amt, "ratio": round(ratio, 4), "fsub": keys[0][0] if keys else "",
                "ffee": keys[0][1] if keys else "", "fbiz": "|".join(b or "-" for b in bizs),
                "blabel": "、".join(b or "无产品线" for b in bizs)}
    # V2.859(用户看丰源第①步：「这里应该是918账单数0」「深圳星期九，计提数0，账单时918」)：付款做账判定「记错主体」的计提
    #   (发票、请款单是另一个主体的)，不拿它去配本主体的账单——单列一行：账单 0、差异＝计提数，标「记错主体 → 那个主体」；
    #   那个主体账单上对应的那一行(账单有、计提无)标上「计提在××记-N」。两行都由主体更正解决，不算有差异没解释。
    xh = _xsubj_hints(carrier, period)
    xents = [e for e in ents if (e.get("subject"), e.get("vno")) in xh]
    groups = {}
    for e in ents:
        if (e.get("subject"), e.get("vno")) in xh:
            continue
        groups.setdefault((e["subject"], e["fee_norm"]), []).append(e)
    # 页面/复核表都带编码(用户 2026-09-30)：主体→账簿编码(本期凭证优先，主体档案兜底)；账单侧产品线→产品分类编码(本期凭证实证优先，BIZLINE_CODE 兜底)
    s2book = {o.get("short_name"): o.get("book_code") for o in (db.list_orgs() or []) if o.get("short_name") and o.get("book_code")}
    biz2code = {}
    try:
        from kernels.logistics_accrual import BIZLINE_CODE
        biz2code.update({k.lower(): v for k, v in BIZLINE_CODE.items()})
    except Exception:
        pass
    for e in ents:
        if e.get("book_code"):
            s2book[e["subject"]] = e["book_code"]
        if e.get("biz_code") and e.get("biz"):
            biz2code[e["biz"].lower()] = e["biz_code"]
    bill2 = {}
    for (s, f, b), v in bill3.items():
        bill2[(s, f)] = bill2.get((s, f), 0.0) + v
    rows, atot, btot = [], 0.0, 0.0
    smk = _subj_marks(period, carrier)
    allkeys = set(list(groups.keys()) + list(bill2.keys()))
    for s in sorted({k[0] for k in allkeys} | {e["subject"] for e in xents}):
      srows, sa, sb = [], 0.0, 0.0
      for gk in sorted(k for k in allkeys if k[0] == s):
        s, f = gk
        es = groups.get(gk, [])
        gb = round(bill2.get(gk, 0.0), 2)
        ga = round(sum(e["amt"] for e in es), 2)
        subs = {}
        for e in es:
            subs.setdefault(e["biz"], []).append(e)
        # 能按产品线配上的子组先排，配不上的排后面连成一段，组级账单才能跨行合并
        ordered = sorted(subs.items(), key=lambda kv: 0 if (s, f, kv[0]) in bill3 else 1)
        grows, matched, used = [], 0.0, set()
        for biz, ses in ordered:
            bb = bill3.get((s, f, biz))
            if bb is not None:
                used.add(biz); matched += bb
            for i, e in enumerate(ses):
                row = {**e, "fee_type": f, "kind": "accr", "level": "biz" if bb is not None else "",
                       "bill": None, "bill_span": 0, "diff": None, "note": notes.get(e["key"], ""),
                       "anc": ses[0]["key"] if bb is not None else ""}   # anc=同一账单金额下的锚点行键(一起标更正用)
                if bb is not None and i == 0:
                    row["bill"] = round(bb, 2); row["bill_span"] = len(ses)
                    row["diff"] = round(sum(x["amt"] for x in ses) - bb, 2)
                grows.append(row)
        rest = round(gb - matched, 2)
        rest_biz = "、".join(b for (ss, ff, b) in bill3 if ss == s and ff == f and b and b not in used)
        un = [x for x in grows if not x["level"]]
        if un:
            un[0]["bill"] = rest; un[0]["bill_span"] = len(un)
            un[0]["diff"] = round(sum(x["amt"] for x in un) - rest, 2)
            if rest_biz:
                un[0]["bill_biz"] = rest_biz
            for x in un:
                x["level"] = "group"
                x["anc"] = un[0]["key"]
        elif abs(rest) >= 0.01:
            key = "bill|%s|%s" % (s, f)
            grows.append({"subject": s, "book_code": s2book.get(s, ""), "biz_code": biz2code.get((rest_biz or "").lower(), ""),
                          "fee": "", "fee_norm": f, "fee_type": f, "biz": rest_biz or "（账单无产品线）",
                          "proj": "", "dept": "", "vno": "", "acct": "", "amt_net": None, "tax_rate": None, "tax": None, "amt": None,
                          "key": key, "kind": "bill_only", "level": "group", "bill": rest, "bill_span": 1,
                          "diff": round(-rest, 2), "note": notes.get(key, "")})
        rkeys = [(s, f, b) for (ss, ff, b) in bill3 if ss == s and ff == f and b not in used]
        for x in grows:
            if x.get("bill") is None:
                continue
            x["od"] = _od([(s, f, x["biz"])] if x.get("level") == "biz" else rkeys, x["bill"])
        if grows:
            grows[0]["ffirst"] = True          # 主体组内每段费用类型的首行(前端画细分隔)
        srows.extend(grows)
        sa += ga; sb += gb
      for e in xents:
          if e["subject"] != s:
              continue
          srows.append({**e, "fee_type": e["fee_norm"], "kind": "accr", "level": "xsubj", "bill": 0.0, "bill_span": 1, "diff": round(e["amt"], 2),
                        "note": notes.get(e["key"], ""), "anc": "", "ffirst": True, "xsubj": xh[(e["subject"], e["vno"])]})
          sa += e["amt"]
      if srows:
          srows[0]["gfirst"] = True
      rows.extend(srows)
      rows.append({"kind": "gtotal", "subject": s, "book_code": s2book.get(s, ""), "fee_type": "", "amt": round(sa, 2), "bill": round(sb, 2),
                   "mark": smk.get((carrier, s)),
                   "diff": round(sa - sb, 2), "key": "gt|%s" % s})
      atot += sa; btot += sb
    for e in xents:      # 那个主体账单上对应的那一行
        h = xh[(e["subject"], e["vno"])]
        b = next((r for r in rows if r.get("kind") == "bill_only" and r.get("subject") == h["to"] and not r.get("xsubj_in")
                  and abs(float(r.get("bill") or 0) - float(e["amt"])) < 0.01), None)
        if b:
            b["xsubj_in"] = {"from": e["subject"], "vno": e["vno"], "amt": round(e["amt"], 2)}
            h["bill_at"] = h["to"]
    n_unexpl = sum(1 for r in rows if r.get("kind") != "gtotal" and r.get("diff") is not None
                   and abs(r["diff"]) >= 0.01 and not (r.get("note") or "").strip() and not r.get("xsubj") and not r.get("xsubj_in"))
    return _attach_xsubj(_attach_fixes({"ok": True, "carrier": carrier, "period": period, "bill_src": bill_src, "rows": rows,
            "accr_total": round(atot, 2), "bill_total": round(btot, 2), "diff_total": round(atot - btot, 2),
            "adj": adj, "prior": prior, "prior_total": round(sum(p["net"] for p in prior), 2),
            "xmoved": got.get("xmoved") or [],
            "xin": [{"vno": e["vno"], "month": e["xm"], "subject": e["subject"], "amt": e.get("amt")} for e in ents if e.get("xm")],
            "od_total": (lambda a: {"amt": a, "ratio": round(max(0.0, min(1.0, a / btot)), 4) if btot > 0 else 0.0})(
                round(sum(det3[k]["amt"] for k in bill3 if k in det3), 2)),
            "points": pts, "signed": (dict(sg) if sg else None), "n_unexplained": n_unexpl,
            "suppliers": got.get("suppliers") or []}, carrier, period), carrier, period)


def _xsubj_hints(carrier, period):
    """付款做账对「计提记错主体」的判断 → {(原主体简称, "记-N"): {to, req_amount, bid, posted, reversed}}。只读，要读金蝶(有缓存)；取不到返回 {}。
    主体对不对是从发票和请款单看出来的：发票开给另一个主体、由它付款，而计提记在了这个主体。"""
    hints = {}
    try:
        from routers import logistics_voucher as LVR
        PRT = store.payreq
        with db._engine.connect() as c:
            reqs = [dict(r) for r in c.execute(select(PRT.c.inst_id, PRT.c.subject, PRT.c.amount, PRT.c.business_id, PRT.c.excluded,
                                                    PRT.c.dt_status, PRT.c.dt_result).where((PRT.c.carrier == carrier) & (PRT.c.period == period))).mappings().all()]
        for r in reqs:
            if r.get("excluded") or str(r.get("dt_status") or "").upper() == "TERMINATED" or str(r.get("dt_result") or "").lower() == "refuse":
                continue
            try:
                d, code = LVR._preview_data(r["inst_id"])
            except Exception:
                continue
            if code != 200:
                continue
            posted = bool((d.get("req") or {}).get("posted"))
            for x in d.get("xbook") or []:
                rv = x.get("reversed") or None
                hints[(x.get("short"), "记-%s" % x.get("vno"))] = {
                    "to": r.get("subject"), "req_amount": round(float(r.get("amount") or 0), 2), "bid": r.get("business_id") or "", "posted": posted,
                    "reversed": ("%s月 记-%s" % (rv.get("month"), rv.get("vno"))) if rv else ""}
    except Exception:
        return {}
    return hints


def _attach_xsubj(L, carrier, period):
    """把标了「记错主体」的行汇总成一份给页面表下的说明用。"""
    L["xsubj"] = [dict(r["xsubj"], subject=r.get("subject"), vno=r.get("vno"), amt=r.get("amt"))
                  for r in L.get("rows", []) if r.get("kind") == "accr" and r.get("xsubj")]
    return L


_FIX_KEYS = ("to_acct", "to_fee", "to_dept", "to_biz", "to_proj", "to_amt_tax", "to_rate", "to_amt", "memo", "split_amt")
_SNAP_KEYS = ("subject", "book_code", "vno", "acct", "acct_name", "fee", "fee_code", "fee_type", "dept", "dept_code",
              "biz", "biz_code", "proj", "proj_code", "amt_net", "tax_rate", "amt")


def _next_period(p):
    """YYYY-MM 的下一个月；格式不对原样返回。"""
    try:
        y, m = int(p[:4]), int(p[5:7])
        return "%04d-%02d" % (y + (m == 12), 1 if m == 12 else m + 1)
    except Exception:
        return p or ""


def _pair_split_fixes(L):
    """拆分更正和「账单有、计提无」配对(V2.858)。
    用户 2026-10-06：「这笔业务实质是计提的时候全都归集在了零售，但是我们在付款复核发现有问题，他应该是电商，所以计提数肯定是0」——
    这种情形页面上是两行：零售那笔计提多出 X(+X)，电商那行账单有、计提无(−X)。登记了「记-N 其中 X 改到电商」以后，这两行的差异就有了着落：
    计提那行标 fix_cover(差异正好＝拆出去的金额)，账单那行标 fix_from(从哪张计提拆过来)，都不再算「有差异没解释」。对账数字不变(金蝶还没改)。"""
    rows = L.get("rows", [])
    for r in rows:
        r.pop("fix_from", None); r.pop("fix_cover", None)
    for a in rows:
        fx = a.get("fix") if a.get("kind") == "accr" else None
        if not fx or not fx.get("split_amt"):
            continue
        try:
            sp = float(fx["split_amt"])
        except ValueError:
            continue
        to_biz = str(fx.get("to_biz") or "").split(" ", 1)[-1].strip()
        cand = [b for b in rows if b.get("kind") == "bill_only" and b.get("subject") == a.get("subject") and not b.get("fix_from")
                and abs(float(b.get("bill") or 0) - sp) < 0.01]
        cand.sort(key=lambda b: (b.get("fee_type") != a.get("fee_type"), bool(to_biz) and to_biz not in str(b.get("biz") or "")))
        if cand:
            cand[0]["fix_from"] = {"vno": a.get("vno"), "key": a.get("key"), "amt": sp}
        if a.get("diff") is not None and abs(float(a["diff"]) - sp) < 0.01:
            a["fix_cover"] = True
    L["n_unexplained"] = sum(1 for r in rows if r.get("kind") != "gtotal" and r.get("diff") is not None and abs(r["diff"]) >= 0.01
                             and not (r.get("note") or "").strip() and not r.get("fix_from") and not r.get("fix_cover")
                             and not r.get("xsubj") and not r.get("xsubj_in"))


def _attach_fixes(L, carrier, period):
    """把已登记的计提更正挂到逐笔行(row.fix)，另回 fixes(按登记快照，金蝶改好后原行不在了也照样列出)。"""
    with db._engine.connect() as c:
        fx = [dict(r) for r in c.execute(select(FX).where((FX.c.carrier == carrier) & (FX.c.period == period))
                                         .order_by(FX.c.id)).mappings().all()]
    by_key, out = {}, []
    for f in fx:
        try:
            snap = json.loads(f.get("snap_json") or "{}")
        except Exception:
            snap = {}
        item = {"key": f["line_key"], "snap": snap, "by": f.get("updated_by") or "", "at": f.get("updated_at") or "",
                "adj_period": f.get("adj_period") or _next_period(period), **{k: (f.get(k) or "") for k in _FIX_KEYS}}
        for k, ck in (("to_acct", "acct"), ("to_fee", "fee_code"), ("to_dept", "dept_code"), ("to_biz", "biz_code"), ("to_proj", "proj_code")):
            if item.get(k) and item[k].split(" ", 1)[0] == str(snap.get(ck) or "").strip():
                item[k] = ""            # 登记时选的和原记账一样的维度不算更正(V2.860 之前存进去的也这样显示)
        by_key[f["line_key"]] = item
        out.append(item)
    live = set()
    for r in L.get("rows", []):
        if r.get("kind") == "accr":
            r["fix"] = by_key.get(r["key"])
            live.add(r["key"])
    _pair_split_fixes(L)
    for it in out:
        it["live"] = it["key"] in live
    out.sort(key=lambda it: (it["snap"].get("subject", ""), it["snap"].get("vno", "")))
    L["fixes"] = out
    return L


def _subfees(v):
    """账单行费用分项 JSON → dict；坏数据/空返回 None。"""
    try:
        d = json.loads(v) if v else None
        return d if isinstance(d, dict) and d else None
    except Exception:
        return None


def _car_norm(s):
    """车型归一：'9.6米冷藏' / '9.6米冷藏车' / '9.6m' → '9.6米'(账单与报价写法不一)。"""
    m = re.search(r"(\d+(?:\.\d+)?)\s*(米|m|M)", str(s or ""))
    return (m.group(1) + "米") if m else str(s or "").strip()


def _fmt_amt(v):
    return ("{:,.0f}" if abs(v - round(v)) < 0.005 else "{:,.2f}").format(v)


def _trip_check(r, tu, n, fee, pmap):
    """包天包趟一行按报价核：运费 = 数量 × 单价(车型+天/趟)；加班费报价未列 → 待确认。
    没车型/不是天趟/承运商没配价目 → None(只能免核量)。"""
    car = str(r.get("carrier_sub") or "").strip()
    if not car or tu not in ("天", "趟") or not pmap:
        return None
    try:
        sf = json.loads(r.get("sub_fees") or "{}") or {}
    except Exception:
        sf = {}
    run = float(sf["运费"]) if sf.get("运费") is not None else float(fee)
    ot = float(sf.get("加班") or 0)
    n = n or 1
    carn = car.replace("冷藏车", "").replace("冷藏", "")
    label = "包天包趟·%s×%g%s" % (carn, n, tu)
    price = pmap.get((_car_norm(car), tu))
    base = {"car": car, "unit": tu, "n": n, "price": price, "run": round(run, 2), "ot": round(ot, 2), "label": label}
    if price is None:
        return {**base, "std": None, "verdict": "diff", "msg": "报价里没有 %s 按%s的价" % (carn, tu)}
    std = round(price * n, 2)
    msgs = []
    if abs(run - std) >= 0.01:
        msgs.append("运费 %s ≠ 报价 %s×%g=%s" % (_fmt_amt(run), _fmt_amt(price), n, _fmt_amt(std)))
    if ot:
        msgs.append("加班 %s 报价未列价，需确认" % _fmt_amt(ot))
    return {**base, "std": std, "verdict": "diff" if msgs else "ok",
            "msg": "；".join(msgs) if msgs else "按报价 %s×%g=%s" % (_fmt_amt(price), n, _fmt_amt(std))}


def _wt_range(carrier):
    """该承运商的毛重比允许范围 (下限, 上限)；没配返回 None=默认差 2% 以内。存供应商档案 logistics_suppliers.wt_lo/wt_hi。"""
    sup = next((x for x in (db.list_logi_suppliers() or []) if x.get("short") == carrier), None) or {}
    return (float(sup["wt_lo"]), float(sup["wt_hi"])) if sup.get("wt_lo") and sup.get("wt_hi") else None


def _box_docs(rsub, carrier):
    """统一物料模板·单据视图：账单每张单据一条 doc（materials=金蝶物料明细，运费按货品kg摊、剔包材）。
    核量三口径：weight=账单重量vs金蝶kg / qty=账单件vs金蝶件(剔包装) / box=有账单重量按重量，否则账单件vs金蝶箱(规格箱规)。
    doc.state：miss 金蝶查无 / qtydiff 数量不符·待核 / info 打托·包车免核·无单据调整(仅提示) / ok 一致。"""
    _rev = "weight" if carrier in _WEIGHT_CARRIERS else ("qty" if carrier in _QTY_CARRIERS else "box")
    by_form = {}
    for r in rsub:
        # V2.864：一格里写了两张单号的(天鹰「CGRK179615/CGRK179614」一车卸两张入库单)，金蝶量要把几张都加上；原来只取第一张，账单 11 吨对金蝶 8.275 吨
        for dx in [x for x in (r.get("doc_no") or "").split("+") if x and x != "无单据"]:
            pre = "".join(ch for ch in dx if ch.isalpha())
            for form in _FORM_BY_PREFIX.get(pre, ["SAL_OUTSTOCK"])[:1]:
                by_form.setdefault(form, set()).add(dx)
    mats = {}
    if by_form:
        try:
            s2, conf2 = kc.login()
            mats = _fetch_doc_materials(s2, conf2, by_form)
        except Exception:
            mats = {}
    # 短驳(包天包趟)价目：承运商配置 spec_json.shuttle_prices=[{car,unit(天/趟),price}]，如极鲜达《仓储收费标准(星期零)》短驳运输
    _sp = _load_spec(carrier) or {}
    pmap = {(_car_norm(x.get("car")), x.get("unit")): float(x.get("price") or 0) for x in (_sp.get("shuttle_prices") or [])}
    # 毛重比允许范围(基础设置·供应商列表，V2.763)：配了就按 账单重量÷金蝶净重 落在范围内算一致；没配走默认(差 2% 以内)
    wt_rng = _wt_range(carrier)
    docs = []
    for r in rsub:
        d0 = (r.get("doc_no") or "").split("+")[0]
        if d0 == "无单据":          # 迅鸽退货/仓储/卸货行单号列写的字面量"无单据"，与空单号同样按无单据处理
            d0 = ""
        lines = [m for dx in (r.get("doc_no") or "").split("+") if dx and dx != "无单据" for m in (mats.get(dx) or [])] if d0 else []
        billcnt = float(r.get("qty") or 0)
        fee = float(r.get("amount") or 0)
        biz = r.get("bizline") or _bizline_of(r.get("annot"))
        is_trip, trip = False, None
        try:
            chg_wt = float(r.get("charge_wt")) if r.get("charge_wt") not in (None, "") else None
        except (TypeError, ValueError):
            chg_wt = None
        use_weight = (_rev == "weight") or (_rev == "box" and chg_wt)
        if use_weight:
            # 有账单重量 → 按重量核：金蝶量=千克计量物料基本数量之和
            per = []
            for m in lines:
                u = str(m.get("基本单位") or "")
                try:
                    bw = float(m.get("基本数量") or 0)
                except (TypeError, ValueError):
                    bw = 0.0
                per.append(bw if ("千克" in u or "kg" in u.lower()) else 0.0)
            kd_sum = round(sum(per), 2)
            wbase = chg_wt if chg_wt else kd_sum
            bill_amt, bill_unit, kd_unit, mode_cn = wbase, "千克", "千克", "按重量"
            if wt_rng and kd_sum:
                cnt_state = "ok" if wt_rng[0] - 1e-9 <= wbase / kd_sum <= wt_rng[1] + 1e-9 else "qtydiff"
            else:
                cnt_state = "ok" if (kd_sum and abs(wbase - kd_sum) <= max(1.0, 0.02 * kd_sum)) else "qtydiff"
            conv = round(wbase / kd_sum, 3) if kd_sum else None   # 按重量：换算系数=账单重量÷金蝶重量(毛重比)
            mkq = lambda m: (float(m.get("基本数量") or 0) if ("千克" in str(m.get("基本单位") or "")) else None)
            mku = lambda m: m.get("基本单位")
        elif _rev == "qty":
            # 快递 → 按件数核：金蝶件数=货品数量件之和(剔包装)，比账单件数
            per = []
            for m in lines:
                if any(k in str(m.get("名称") or "") for k in _PACK_KW):
                    per.append(0.0); continue
                try:
                    per.append(float(m.get("数量件") or 0))
                except (TypeError, ValueError):
                    per.append(0.0)
            kd_sum = round(sum(per), 2)
            bill_amt, bill_unit, kd_unit, mode_cn = billcnt, (r.get("unit") or "件"), "件", "按件数"
            cnt_state = "miss" if not kd_sum else ("ok" if abs(billcnt - kd_sum) <= max(1.0, 0.02 * kd_sum) else "qtydiff")
            if str(r.get("unit") or "").strip() in lr._NOQTY_UNITS:
                mode_cn, cnt_state = "按%s计费·免核量" % str(r.get("unit")).strip(), "na"   # 搬运费按方、存储费按板：核不了件数
            conv = round(kd_sum / billcnt, 3) if billcnt else None
            mkq = lambda m: (float(m.get("数量件")) if m.get("数量件") not in (None, "") else None)
            mku = lambda m: (m.get("计价单位") or m.get("基本单位"))
        else:
            # 无账单重量 → 按件数/箱核：金蝶箱数=数量件÷规格箱规
            per = []
            for m in lines:
                br = _box_div(m)                  # 袋/箱、计价单位箱、20kg/箱·10L/件(按千克/升折)
                try:
                    qcnt = float(m.get("数量件") or 0)
                except (TypeError, ValueError):
                    qcnt = 0.0
                per.append((qcnt / br) if (br and qcnt) else 0.0)
            kd_sum = round(sum(per), 2)
            bill_amt, bill_unit, kd_unit = billcnt, (r.get("unit") or "件"), "箱"
            ratio_tuo = (kd_sum / billcnt) if billcnt else 0
            # 调拨(分布式调出)按包天包趟计费：一张单可能跑几车、账单一车一行 → 核不了量；有车型/天趟/价目时按报价核价(用户 2026-09-30)
            tu = str(r.get("unit") or "").replace("元/", "").strip()
            is_trip = tu in ("天", "趟") or ((not billcnt) and (d0.upper().startswith("FBDR") or "调拨" in str(r.get("annot") or "")))
            trip = _trip_check(r, tu, billcnt, fee, pmap) if is_trip else None
            if is_trip and trip and trip["verdict"] == "ok":
                mode_cn, cnt_state = trip["label"], "ok"
            elif is_trip and trip:
                mode_cn, cnt_state = trip["label"], "price"
            elif kd_sum and is_trip:
                mode_cn, cnt_state = "包天包趟·免核量", "na"
            elif str(r.get("unit") or "").strip() in lr._NOQTY_UNITS:
                mode_cn, cnt_state = "按%s计费·免核量" % str(r.get("unit")).strip(), "na"   # 恒茂出库装卸按吨、冷藏按板
            if is_trip and tu:
                bill_unit = tu                   # 账单量显示 1 天 / 2 趟
            elif kd_sum and abs(billcnt - kd_sum) <= max(1.0, 0.02 * kd_sum):
                mode_cn, cnt_state = "整车按箱", "ok"
            elif kd_sum and billcnt and 3 <= ratio_tuo <= 60:
                mode_cn, cnt_state = "打托(托规%s)" % round(ratio_tuo, 1), "na"
            elif kd_sum and billcnt and ratio_tuo > 60:
                mode_cn, cnt_state = "整车包车·免核", "na"   # 件数为名义值(如1)、无重量、箱数远超→议价包车不核
            elif not kd_sum:
                mode_cn, cnt_state = "无箱规待核", "qtydiff"
            else:
                mode_cn, cnt_state = "待核", "qtydiff"
            conv = round(kd_sum / billcnt, 3) if billcnt else None   # 按件数：换算系数=金蝶箱数÷账单件(整车按箱≈1、打托=托规)
            mkq = lambda m: (float(m.get("数量件")) if m.get("数量件") not in (None, "") else None)
            mku = lambda m: (m.get("计价单位") or m.get("基本单位"))
        if not d0:
            mode_cn, cnt_state = "无单据·账单调整", "na"   # 如托盘丢失扣款：只登记不核量
        base = {"subject": _eff_subject(r), "carrier": carrier, "fee_item": _eff_fee(r),
                "bizline": biz, "doc_no": d0, "bill_amt": round(bill_amt, 2), "bill_unit": bill_unit,
                "kd_sum": kd_sum, "kd_unit": kd_unit, "mode_cn": mode_cn, "conv": conv, "qty_state": cnt_state,
                "note": r.get("note") or "", "bbiz": _bill_biz(r),
                # 人工归类(V2.788)：账单原主体、主体有没有被人工改、人工定的费用类型、原因——页面据此标「特批」、可恢复账单原值
                "subj_bill": _short_subject(str(r.get("subject") or "").strip()),
                "subj_ovr": bool(str(r.get("subj_ovr") or "").strip()) and _eff_subject(r) != _short_subject(str(r.get("subject") or "").strip()),
                "fee_ovr": _fee_ovr(r), "ovr_reason": r.get("ovr_reason") or "",
                "lid": r.get("id"),    # 账单行ID：同一单号账单上可能有多行(按车次收费)，页面勾选/展开按行认
                "sub_fees": _subfees(r.get("sub_fees"))}   # 费用构成(快递费/操作费/箱子+箱型、运费/加班…)，页面展示
        mrows = []
        kgbase = 0.0
        if not lines:
            mrows.append({**base, "party": "", "code": "", "name": "（金蝶无此单据物料）" if d0 else "（无单据）",
                          "base_qty": None, "base_unit": "", "kd": None, "fee": round(fee, 2),
                          "unit_fee": None, "sales": None, "ratio": None})
        else:
            # 运费分摊与单位运费一律按货品基本重量(kg)，剔除包材（包材不摊、单位运费留空）
            kgs = []
            for m in lines:
                ispack = any(k in str(m.get("名称") or "") for k in _PACK_KW)
                u = str(m.get("基本单位") or "")
                try:
                    kg = float(m.get("基本数量") or 0)
                except (TypeError, ValueError):
                    kg = 0.0
                kgs.append(0.0 if ispack else (kg if ("千克" in u or "kg" in u.lower()) else 0.0))
            kgbase = sum(kgs)
            nnp = sum(1 for x in kgs if x)  # 非包材(有kg)物料数，用于kgbase=0时兜底均摊
            for i, m in enumerate(lines):
                kd = round(per[i], 2)
                ispack = any(k in str(m.get("名称") or "") for k in _PACK_KW)
                kg = kgs[i]
                if kgbase:
                    share = kg / kgbase
                elif not ispack and nnp == 0:
                    share = 1.0 / len(lines)   # 无kg基数(异常)时均摊
                else:
                    share = 0.0
                fline = round(fee * share, 2)
                try:
                    sales = float(m.get("销售额")) if m.get("销售额") not in (None, "") else None
                except (TypeError, ValueError):
                    sales = None
                try:
                    bq = mkq(m)
                except (TypeError, ValueError):
                    bq = None
                try:
                    base_kg = float(m.get("基本数量")) if m.get("基本数量") not in (None, "") else None
                except (TypeError, ValueError):
                    base_kg = None
                mrows.append({**base, "party": m.get("往来") or "", "code": m.get("编码"), "name": m.get("名称"),
                              "base_qty": bq, "base_unit": mku(m),
                              "base_kg": base_kg, "kg_unit": m.get("基本单位"), "is_pack": ispack,
                              "kd": kd or None, "spec": m.get("规格"),
                              "fee": fline, "unit_fee": round(fline / kg, 2) if kg else None,
                              "sales": round(sales, 2) if sales is not None else None,
                              "ratio": round(fline / sales, 4) if sales else None})
        if not d0:
            state = "info"
        elif not lines or cnt_state == "miss":
            state = "miss"
        elif cnt_state == "qtydiff":
            state = "qtydiff"
        elif cnt_state == "price":
            state = "price"            # 包天包趟核价不符/加班未列价
        elif cnt_state == "na":
            state = "info"
        else:
            state = "ok"
        svals = [m["sales"] for m in mrows if m.get("sales") is not None]
        ssum = round(sum(svals), 2) if svals else None
        _pts = list(dict.fromkeys(m.get("party") for m in mrows if m.get("party")))
        f_sys, f_fill, f_chk = _fee_check(r, d0, _pts)
        docs.append({**base, "state": state, "fee_sys": f_sys, "fee_fill": f_fill, "fee_chk": f_chk, "doc_fee": round(fee, 2), "doc_kg": round(kgbase, 2) if kgbase else None,
                     "unit_fee": round(fee / kgbase, 2) if kgbase else None, "sales": ssum,
                     "ratio": round(fee / ssum, 4) if ssum else None,
                     "parties": list(dict.fromkeys(m.get("party") for m in mrows if m.get("party"))),
                     "n_mat": len(lines), "q_diff": round(bill_amt - kd_sum, 2) if (kd_sum and not is_trip) else None,
                     "trip": trip,
                     "materials": mrows})
    # 账单公式核价(恒茂冷藏费=天数×计费数量×单价、处置费=吨重×19.4)：差超 1 分标「核价不符」，一致的注明算法
    for x in docs:
        sf = x.get("sub_fees") or {}
        if "标准" in sf:
            tgt = next((k for k in sf if k not in ("标准", "公式", "核价差") and isinstance(sf[k], (int, float))), "金额")
            if "核价差" in sf:
                x["state"], x["mode_cn"] = "price", "核价不符"
                x["trip"] = {"verdict": "diff", "msg": "%s %s ≠ %s %s" % (tgt, _fmt_amt(sf[tgt]), sf.get("公式") or "标准", _fmt_amt(sf["标准"]))}
            else:
                x["trip"] = {"verdict": "ok", "msg": "%s = %s ✓" % (sf.get("公式") or "按公式", _fmt_amt(sf["标准"]))}
    # 同一单号账单上有多行(包天包趟一车一行 / 恒茂一张调拨单拆几个批次)：每行注明本单共几行、合计运费；单位运费/费比按本单合计运费算；
    # 按件数/箱核的，拿本单几行账单量之和比金蝶(不再单行比整单)
    grp = {}
    for x in docs:
        if x["doc_no"]:
            grp.setdefault(x["doc_no"], []).append(x)
    for xs in grp.values():
        if len(xs) > 1:
            tot = round(sum(x["doc_fee"] for x in xs), 2)
            for x in xs:
                x["trips"], x["doc_fee_all"] = len(xs), tot
                if x.get("doc_kg"):
                    x["unit_fee"] = round(tot / x["doc_kg"], 2)
                if x.get("sales"):
                    x["ratio"] = round(tot / x["sales"], 4)
            # 一张单拆几行批次(恒茂入库)：单行比整单会被误判成打托/包车，先把几行账单量加起来再比金蝶
            if all(x.get("kd_sum") for x in xs) and not any(str(x.get("mode_cn", "")).startswith("包天") or x.get("state") == "price" for x in xs)                     and len({x.get("bill_unit") for x in xs}) == 1 and all(x.get("bill_amt") for x in xs):
                bsum = round(sum(float(x.get("bill_amt") or 0) for x in xs), 2)
                kd = xs[0]["kd_sum"]
                st = "ok" if abs(bsum - kd) <= max(1.0, 0.02 * kd) else "qtydiff"
                for x in xs:
                    x["state"], x["q_diff"], x["doc_bill_all"] = st, round(bsum - kd, 2), bsum
                    x["mode_cn"] = "本单 %d 行合计 %s%s" % (len(xs), _fmt_amt(bsum), x.get("bill_unit") or "")
                    x["conv"] = round(kd / bsum, 3) if bsum else None
    return docs


def _code_maps(request, carrier, period):
    """逐单页的主体/费用类型/产品线带金蝶编码，和第①步逐笔计提同一套(用户 2026-10-01「产品线应该和外面对齐」)：
    取第①步的结果(30 分钟缓存，没有就现算一次)——主体→账簿编码、产品线→产品分类编码、(主体,费用类型)→费用项目编码+金蝶名称。"""
    import time as _t
    k = ("lines", carrier, period)
    c = _ACCR_CACHE.get(k)
    L = c[0] if (c and _t.time() - c[2] < 1800) else None
    if L is None:
        try:
            L = _build_lines(request, carrier, period)
            _ACCR_CACHE[k] = (L, None, _t.time())
        except Exception:
            L = {}
    book = {o.get("short_name"): o.get("book_code") for o in (db.list_orgs() or []) if o.get("short_name") and o.get("book_code")}
    biz = {}
    try:
        from kernels.logistics_accrual import BIZLINE_CODE
        biz.update(BIZLINE_CODE)
    except Exception:
        pass
    fee = {}
    for r in (L or {}).get("rows") or []:
        if r.get("book_code") and r.get("subject"):
            book[r["subject"]] = r["book_code"]
        if r.get("biz_code") and r.get("biz") and not str(r["biz"]).startswith("（"):
            biz[r["biz"]] = r["biz_code"]
        if r.get("kind") == "accr" and r.get("fee_code") and r.get("fee_type"):
            fee.setdefault("%s|%s" % (r.get("subject"), r["fee_type"]), [r["fee_code"], r.get("fee") or r["fee_type"]])
            fee.setdefault(r["fee_type"], [r["fee_code"], r.get("fee") or r["fee_type"]])
    return {"book": book, "biz": biz, "fee": fee}


_FULL_DOCS = 1500      # 单据数在这以内的承运商：第②步整家逐张现算(取金蝶物料)再筛选/计数；超过的只算当前页、计数用中间表


@router.get("/api/logistics-review/result")
def review_result(request: Request, carrier: str = "迅鸽", period: str = "",
                  group: str = "ex", page: int = 1, size: int = 50, q: str = "",
                  fsub: str = "", ffee: str = "", fbiz: str = ""):
    # fsub/ffee/fbiz：从第①步「可逐单」点进来时只看这一组(主体/费用类型/产品线，口径同逐笔复核的账单归口)
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    card, ncard = _load_card(carrier)
    accr_lines, accr_total = [], 0.0   # 计提对账走单独异步接口 /accrual，不拖慢主表
    with db._engine.connect() as c:
        rows = [dict(r) for r in c.execute(select(BL).where(
            (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain == "detail"))).mappings().all()]
        accr = [dict(r) for r in c.execute(select(BL).where(
            (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain == "accrual"))).mappings().all()]
    kdmap = {(r.get("doc_no") or "").split("+")[0]: r["kd_qty"] for r in rows if r.get("kd_qty") is not None}
    lr.review_details(rows, card, kdmap)
    summary = lr.summarize(rows)
    counts = lr.verdict_counts(rows)
    total_bill = round(sum((a.get("amount") or 0) for a in accr) or sum((r.get("amount") or 0) for r in rows), 2)

    qs = (q or "").strip()
    q_on = bool(qs)
    n_docs = len({(r.get("doc_no") or "").split("+")[0] for r in rows})   # 大小按单据数算(取金蝶物料的量)，不按账单行数(恒茂 598 行只 43 张单)
    q_small = n_docs <= _FULL_DOCS   # 小承运商搜索时全量取物料，支持按客户/物料名搜(大承运商仍按单号/省预筛省金蝶)

    bucket_on = bool(fsub or ffee or fbiz)
    bset = {("" if t == "-" else t) for t in fbiz.split("|")} if fbiz else None

    def inb(sub, fee, bb):
        return (not fsub or sub == fsub) and (not ffee or fee == ffee) and (bset is None or bb in bset)

    rows_all = rows            # 单据视图缓存按整家承运商建，按组筛在缓存之后

    def _nd(r):
        """没有金蝶单号的账单行(仓储费、设备/托盘搬运费、延迟扣款…)：没法逐单核，第①步按汇总核(可逐单一栏写「否·汇总核」)。
        V2.829(用户 2026-10-05「这几个否的，为什么我还能再逐单看到」)：第②步不再把它们列成单据(原来挂在「免核·无单据·账单调整」里)，只在表上方留一行数。"""
        d = (r.get("doc_no") or "").split("+")[0]
        return not d or d == "无单据"
    # 第②步三个下拉筛选(主体/费用类型/产品线)的选项：取整家本月全部账单行计张数(一行=表里一张，与列表条数一致；口径同 fsub/ffee/fbiz 筛选)
    _fc = {"subject": {}, "fee": {}, "biz": {}}
    for r in rows_all:
        if _nd(r):
            continue
        for fk, fv in (("subject", _eff_subject(r)), ("fee", _eff_fee(r)), ("biz", _bill_biz(r))):
            _fc[fk][fv or ""] = _fc[fk].get(fv or "", 0) + 1
    facets = {k: sorted(v.items(), key=lambda kv: (-kv[1], kv[0])) for k, v in _fc.items()}
    if bucket_on:
        rows = [r for r in rows if inb(_eff_subject(r), _eff_fee(r), _bill_biz(r))]
        counts = lr.verdict_counts(rows)

    conf = _doc_ok(carrier, period)
    nodoc = [{"subject": _eff_subject(r), "fee_item": _eff_fee(r), "bizline": _bill_biz(r), "amount": round(float(r.get("amount") or 0), 2),
              "note": r.get("note") or ""} for r in rows if _nd(r)]

    def d0_of(r):
        return (r.get("doc_no") or "").split("+")[0]

    def keep(r):
        if q_on:
            # 搜索忽略分组(异常/通过都能搜到)；小承运商放行全部到 view 层全字段筛，大承运商按单号/省预筛
            if q_small:
                return True
            return qs in (r.get("doc_no") or "") or qs in (r.get("prov") or "")
        if group == "all":
            return True
        if group == "done":
            return d0_of(r) in conf
        if d0_of(r) in conf:           # 已确认的单据不再算待核/异常/一致
            return False
        if group == "ex":
            return r["qty_state"] in ("miss", "qtydiff")     # 与逐单视图、分组计数同口径(核量)
        if group in ("miss", "qtydiff"):
            return r["qty_state"] == group
        if group == "info":
            return r["qty_state"] == "na"      # 无单据/按方板计费·免核
        if group == "price":
            return False                       # 包天包趟核价只在逐单视图(小承运商)
        if group in ("gap", "free", "over"):
            return r["price_state"] == group
        if group in ("pass", "ok"):
            return r["qty_state"] == "ok"
        return True

    by_weight = carrier in _WEIGHT_CARRIERS
    filt = [r for r in rows if keep(r) and not _nd(r)]
    page = max(1, int(page))
    sl = filt if q_on else filt[(page - 1) * size: page * size]   # 搜索不分页(结果集小)，全量取物料后按全字段筛
    # 统一物料模板：按重量(顺丰/天鹰)/按件数箱(丰源)/快递件数(迅鸽)共用同一分支
    _rev = "weight" if carrier in _WEIGHT_CARRIERS else ("qty" if carrier in _QTY_CARRIERS else "box")
    by_box = (carrier in _BOX_CARRIERS) or (carrier in _WEIGHT_CARRIERS) or (carrier in _QTY_CARRIERS)
    if by_box:
        import time as _t
        # 小承运商(≤300单)：全量出单据视图、缓存10分钟，按真实核对结果筛/计数/翻页，不用每次翻页都拉金蝶；
        # 大承运商(迅鸽几千单)：只按当前页取金蝶物料，计数沿用中间表核量态。
        # V2.863(用户看天鹰 8 月「数量不符 161」，点进去每张都写着一致)：天鹰 306 张单刚好过了原来 300 张的线，页签上的数退回用中间表的粗口径
        #   (账单「箱数」直接比金蝶「千克」，2,300 箱 对 23,000 千克 → 全判不符)，而表里每张是按计费重量现算的(真不符只有 56 张)。
        #   线提到 1500：天鹰全量现算实测 3.8 秒、缓存 10 分钟；只有迅鸽(五千多张)还走大承运商那条路，它按件数核、中间表口径本来就对。
        full = len({(r.get("doc_no") or "").split("+")[0] for r in rows_all}) <= _FULL_DOCS
        if full:
            ck = ("view", carrier, period)
            cc = _ACCR_CACHE.get(ck)
            if cc and _t.time() - cc[2] < 600:
                pool = cc[0]
            else:
                pool = _box_docs(rows_all, carrier)
                _ACCR_CACHE[ck] = (pool, None, _t.time())
            for x in pool:                    # 缓存里的单据每次重挂确认态
                x["confirmed"] = conf.get(x["doc_no"]) if x["doc_no"] else None
            pool = [x for x in pool if x["doc_no"]]      # 无单据的不列(见 _nd)；缓存里那份不动
            ex_all = sum(1 for x in pool if not x.get("confirmed") and x["state"] in ("miss", "qtydiff", "price"))   # 步骤条用整家总数，不随筛选变
            if bucket_on:
                pool = [x for x in pool if inb(x["subject"], x["fee_item"], x.get("bbiz", ""))]
            dc = {"miss": 0, "qtydiff": 0, "price": 0, "info": 0, "ok": 0, "done": 0}
            for x in pool:
                k = "done" if x.get("confirmed") else x["state"]
                dc[k] = dc.get(k, 0) + 1
            dc["all"] = len(pool)
            dc["feediff"] = sum(1 for x in pool if x.get("fee_chk") == "diff")    # 费用类型 物流部填≠系统判(V2.735)
        else:
            ex_all = None
            pool = _box_docs(sl, carrier)
            for x in pool:
                x["confirmed"] = conf.get(x["doc_no"]) if x["doc_no"] else None
            dc = {"miss": counts.get("miss", 0), "qtydiff": counts.get("qtydiff", 0),
                  "info": sum(1 for r in rows if r.get("qty_state") == "na" and not _nd(r)),
                  "ok": sum(1 for r in rows if r.get("qty_state") == "ok"), "all": sum(1 for r in rows if not _nd(r)), "done": 0}
            seen = set()
            for r in rows:                    # 大承运商按中间表行扣掉已确认的(行级近似，单号去重计已确认)
                d0 = d0_of(r)
                if d0 and d0 in conf:
                    if d0 not in seen:
                        seen.add(d0); dc["done"] += 1
                    if r.get("qty_state") in ("miss", "qtydiff"):
                        dc[r["qty_state"]] = max(0, dc[r["qty_state"]] - 1)
                    elif r.get("qty_state") == "ok":
                        dc["ok"] = max(0, dc["ok"] - 1)
                    elif r.get("qty_state") == "na":
                        dc["info"] = max(0, dc["info"] - 1)
        dc["ex"] = dc["miss"] + dc["qtydiff"] + dc.get("price", 0)

        def dhit(x):
            if any(qs in str(x.get(f) or "") for f in ("doc_no", "subject", "fee_item", "bizline", "note", "mode_cn")):
                return True
            if any(qs in str(p) for p in (x.get("parties") or [])):
                return True
            return any(qs in str(m.get(f) or "") for m in x["materials"] for f in ("name", "code", "party"))
        if q_on:
            pool = [x for x in pool if dhit(x)]
        elif full and group == "done":
            pool = [x for x in pool if x.get("confirmed")]
        elif full and group == "feediff":
            pool = [x for x in pool if x.get("fee_chk") == "diff"]
        elif full and group != "all":
            want = {"ex": ("miss", "qtydiff", "price"), "pass": ("ok",)}.get(group, (group,))
            pool = [x for x in pool if not x.get("confirmed") and x["state"] in want]
        if full:
            dtot, docs = len(pool), pool[(page - 1) * size: page * size]
        else:
            dtot, docs = (len(pool) if q_on else len(filt)), pool
        view = [m for x in docs for m in x["materials"]]
        return {"ok": True, "carrier": carrier, "period": period, "price_card_rows": ncard, "codes": _code_maps(request, carrier, period),
                "total_bill": total_bill, "summary": summary, "accrual": accr, "counts": counts,
                "accr_lines": accr_lines, "accr_total": accr_total, "doc_counts": dc,
                "by_box": True, "material": True, "detail_total": dtot, "docs": docs, "detail": view,
                "page": page, "size": size, "facets": facets, "ex_all": ex_all, "wt_range": _wt_range(carrier), "nodoc": nodoc}
    if by_weight:
        # 物料级：每单拆金蝶物料，运费/账单重量按金蝶基本单位重量摊；换算系数＝账单计费重量÷金蝶重量(毛重比)
        by_form = {}
        for r in sl:
            d0 = (r.get("doc_no") or "").split("+")[0]
            if not d0:
                continue
            pre = "".join(ch for ch in d0 if ch.isalpha())
            for form in _FORM_BY_PREFIX.get(pre, ["SAL_OUTSTOCK"])[:1]:
                by_form.setdefault(form, set()).add(d0)
        mats = {}
        if by_form:
            try:
                s2, conf2 = kc.login()
                mats = _fetch_doc_materials(s2, conf2, by_form)
            except Exception:
                mats = {}
        view = []
        for r in sl:
            d0 = (r.get("doc_no") or "").split("+")[0]
            lines = mats.get(d0) or []
            chg = float(r.get("charge_wt") or 0)
            fee = float(r.get("amount") or 0)
            biz = r.get("bizline") or _bizline_of(r.get("annot"))   # 优先用中间表已存业务线
            kgs = []
            for m in lines:
                u = str(m.get("基本单位") or "")
                try:
                    bw = float(m.get("基本数量") or 0)
                except (TypeError, ValueError):
                    bw = 0.0
                kgs.append(bw if ("千克" in u or "kg" in u.lower()) else 0.0)
            kgsum = sum(kgs)
            conv = round(chg / kgsum, 3) if kgsum else None
            base = {"subject": r.get("subject"), "carrier": carrier, "fee_item": r.get("fee_item"),
                    "bizline": biz, "doc_no": d0, "bill_qty_doc": round(chg, 2), "conv": conv,
                    "qty_state": r.get("qty_state")}
            if not lines:
                view.append({**base, "party": "", "code": "", "name": "（金蝶无此单据物料）",
                             "base_wt": None, "base_unit": "", "fee": round(fee, 2), "unit_fee": None,
                             "bill_qty": round(chg, 2), "bill_unit": "千克", "sales": None, "ratio": None})
                continue
            for i, m in enumerate(lines):
                bw = kgs[i]
                share = (bw / kgsum) if kgsum else (1.0 / len(lines))
                fline = round(fee * share, 2)
                bqty = round(chg * share, 2)
                try:
                    sales = float(m.get("销售额")) if m.get("销售额") not in (None, "") else None
                except (TypeError, ValueError):
                    sales = None
                view.append({**base, "party": m.get("往来") or "", "code": m.get("编码"), "name": m.get("名称"),
                             "base_wt": bw or None, "base_unit": m.get("基本单位"),
                             "fee": fline, "unit_fee": round(fline / bw, 4) if bw else None,
                             "bill_qty": bqty, "bill_unit": "千克",
                             "sales": round(sales, 2) if sales is not None else None,
                             "ratio": round(fline / sales, 4) if sales else None})
        return {"ok": True, "carrier": carrier, "period": period, "price_card_rows": ncard, "codes": _code_maps(request, carrier, period),
                "total_bill": total_bill, "summary": summary, "accrual": accr, "counts": counts, "accr_lines": accr_lines, "accr_total": accr_total,
                "by_weight": True, "material": True, "detail_total": len(filt), "detail": view,
                "page": page, "size": size}
    view = [{k: r.get(k) for k in ("doc_no", "carrier_sub", "prov", "charge_wt", "qty", "kd_qty",
             "amount", "base_amount", "std_amount", "price_diff", "price_state", "qty_diff",
             "qty_state", "tier", "verdict", "fee_item")} for r in sl]
    return {"ok": True, "carrier": carrier, "period": period, "price_card_rows": ncard, "codes": _code_maps(request, carrier, period),
            "total_bill": total_bill, "summary": summary, "accrual": accr, "counts": counts, "accr_lines": accr_lines, "accr_total": accr_total,
            "by_weight": by_weight, "detail_total": len(filt), "detail": view, "page": page, "size": size}


_PSTATE_CN = {"ok": "通过", "over": "多收", "under": "账单少收", "free": "账单未收·我方有利",
              "gap": "价卡缺·待确认", "na": "待补价卡"}
_QSTATE_CN = {"ok": "一致", "qtydiff": "不符", "miss": "金蝶查无", "na": "—"}
_VERDICT_CN = {"pass": "两轴通过", "price": "核价多收", "gap": "核价待补", "free": "账单未收",
               "qty": "核量存疑", "registered": "已登记", "doc_miss": "单号查无"}


_ACCR_CACHE = {}   # (carrier,period) / ("lines",carrier,period) -> (result, None, ts)


def _bust(carrier, period):
    """归类/备注/登记/重解析/接金蝶 → 该承运商该月的 逐笔、结论、逐单视图 缓存作废。"""
    for k in ((carrier, period), ("lines", carrier, period), ("view", carrier, period)):
        _ACCR_CACHE.pop(k, None)


def _signed(carrier, period):
    with db._engine.connect() as c:
        return c.execute(select(SG).where((SG.c.carrier == carrier) & (SG.c.period == period) & (SG.c.status == "signed"))).mappings().first()


def _locked(carrier, period):
    """已登记复核的月份不许再改归类/备注：返回 409 响应，否则 None。"""
    sg = _signed(carrier, period)
    if sg:
        return JSONResponse({"ok": False, "msg": "本月已登记复核（%s %s），撤销登记后才能修改" % (
            sg.get("reviewer") or "", sg.get("signed_at") or "")}, status_code=409)
    return None


def _uname(u):
    try:
        return str((u or {}).get("name") or "")[:50]
    except Exception:
        return ""


@router.get("/api/logistics-review/lines")
@router.get("/api/logistics-review/accrual")
def review_lines(request: Request, carrier: str = "", period: str = ""):
    """第二页·逐笔计提复核（老地址 /accrual 也指到这里）。金蝶查询慢，30分钟缓存；改归类/备注/登记会作废缓存。"""
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    import time as _t
    k = ("lines", carrier, period)
    c = _ACCR_CACHE.get(k)
    if c and _t.time() - c[2] < 1800:
        return {**c[0], "cached": True}
    res = _build_lines(request, carrier, period)
    _ACCR_CACHE[k] = (res, None, _t.time())
    return res


@router.post("/api/logistics-review/doc-note")
async def review_doc_note(request: Request):
    """复核台逐单备注：更新该承运商/账期/单据的 note（人工备注）。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, period, doc_no = (b.get("carrier") or "").strip(), (b.get("period") or "").strip(), (b.get("doc_no") or "").strip()
    note = (b.get("note") or "").strip()
    if not carrier or not doc_no:
        return JSONResponse({"ok": False, "msg": "缺承运商/单据号"}, status_code=400)
    lk = _locked(carrier, period)
    if lk:
        return lk
    with db._engine.begin() as c:
        c.execute(update(BL).where((BL.c.carrier == carrier) & (BL.c.period == period) &
                  (func.substr(BL.c.doc_no, 1, len(doc_no)) == doc_no)).values(note=note))
    _ACCR_CACHE.pop(("view", carrier, period), None)   # 逐单视图带备注，作废；逐笔不受影响
    return {"ok": True}


@router.post("/api/logistics-review/doc-classify")
async def review_doc_classify(request: Request):
    """复核台逐单手动改归类：改该单据的账单侧【主体/费用类型】，复核结论按新归类重算(缓存作废)。
    用于把账单错配的行拨到对的计提组(如账单调拨费对不上计提、主体串号)，让逐组差异对平。
    V2.788(用户「特批调整到星期零的要怎么处理…就是主体和费用类型」)：加 reason(原因，留痕) 和 restore(恢复账单原值)；
    费用类型只收规范类(原来填别的会被悄悄当成没填)；改动进操作日志。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier = (b.get("carrier") or "").strip()
    period = (b.get("period") or "").strip()
    doc_no = (b.get("doc_no") or "").strip()
    if not carrier or not doc_no:
        return JSONResponse({"ok": False, "msg": "缺承运商/单据号"}, status_code=400)
    vals = {}
    if b.get("restore"):
        vals = {"subj_ovr": None, "fee_ovr": None, "ovr_reason": None}
    if b.get("subject") is not None:
        vals["subj_ovr"] = (b.get("subject") or "").strip() or None   # 存覆盖列，不动原 subject；空=清覆盖
    if b.get("fee_item") is not None:
        fv = (b.get("fee_item") or "").strip()
        if fv and _fee_norm(fv) not in _CANON_FEES:
            return JSONResponse({"ok": False, "msg": "费用类型只能选：出库运费 / 入库运费 / 退货运费 / 仓储费 / 研发外购 / 搬运费"}, status_code=400)
        vals["fee_ovr"] = fv or None
    if b.get("reason") is not None:
        vals["ovr_reason"] = (b.get("reason") or "").strip()[:200] or None
    if not vals:
        return {"ok": True}
    lk = _locked(carrier, period)
    if lk:
        return lk
    with db._engine.begin() as c:
        c.execute(update(BL).where((BL.c.carrier == carrier) & (BL.c.period == period) &
                  (func.substr(BL.c.doc_no, 1, len(doc_no)) == doc_no)).values(**vals))
    _bust(carrier, period)   # 归类变了 → 逐笔/结论缓存作废，重算
    CN = {"subj_ovr": "主体", "fee_ovr": "费用类型", "ovr_reason": "原因"}
    db.audit(_uname(u), "物流复核-改归类", "%s %s %s" % (carrier, period, doc_no),
             "恢复账单原值" if b.get("restore") else "；".join("%s→%s" % (CN[k], v or "（清空）") for k, v in vals.items()))
    return {"ok": True}


@router.post("/api/logistics-review/subject-mark")
async def review_subject_mark(request: Request):
    """按主体标复核结论：status=ok 通过 / question 有疑问(note 写疑问) / 空=撤销。已登记复核的月份不可改。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, period, subj = (b.get("carrier") or "").strip(), (b.get("period") or "").strip(), (b.get("subject") or "").strip()
    st, note = (b.get("status") or "").strip(), (b.get("note") or "").strip()
    if not carrier or not period or not subj or st not in ("", "ok", "question"):
        return JSONResponse({"ok": False, "msg": "参数不对"}, status_code=400)
    lk = _locked(carrier, period)
    if lk:
        return lk
    with db._engine.begin() as c:
        c.execute(delete(SJ).where((SJ.c.carrier == carrier) & (SJ.c.period == period) & (SJ.c.subject == subj)))
        if st:
            c.execute(insert(SJ).values(carrier=carrier, period=period, subject=subj, status=st, note=note,
                                        updated_by=_uname(u), updated_at=_now()))
    k = ("lines", carrier, period)
    if k in _ACCR_CACHE:                     # 只改组头标记，不重读金蝶
        for r in _ACCR_CACHE[k][0].get("rows", []):
            if r.get("kind") == "gtotal" and r.get("subject") == subj:
                r["mark"] = {"status": st, "note": note, "by": _uname(u), "at": _now()} if st else None
    return {"ok": True}


@router.post("/api/logistics-review/doc-confirm")
async def review_doc_confirm(request: Request):
    """逐单「已确认」：批量打标/取消(on=false)。只认有金蝶单号的单据；已登记复核的月份不可改。不动账单、不影响对账。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, period = (b.get("carrier") or "").strip(), (b.get("period") or "").strip()
    nos = list(dict.fromkeys(str(x).strip() for x in (b.get("doc_nos") or []) if str(x).strip()))
    on = b.get("on", True) is not False
    if not carrier or not period or not nos:
        return JSONResponse({"ok": False, "msg": "缺承运商/账期/单号"}, status_code=400)
    lk = _locked(carrier, period)
    if lk:
        return lk
    have = _doc_ok(carrier, period)
    with db._engine.begin() as c:
        if on:
            todo = [n for n in nos if n not in have]
            for n in todo:
                c.execute(insert(DK).values(carrier=carrier, period=period, doc_no=n, confirmed_by=_uname(u), confirmed_at=_now()))
        else:
            todo = [n for n in nos if n in have]
            if todo:
                c.execute(delete(DK).where((DK.c.carrier == carrier) & (DK.c.period == period) & (DK.c.doc_no.in_(todo))))
    return {"ok": True, "n": len(todo), "on": on}


@router.post("/api/logistics-review/line-note")
async def review_line_note(request: Request):
    """第二页逐笔备注＝差异解释（有差异时写为什么差）。键=凭证号+维度组合。已登记月份不可改。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, period = (b.get("carrier") or "").strip(), (b.get("period") or "").strip()
    key, note = (b.get("line_key") or "").strip(), (b.get("note") or "").strip()
    if not carrier or not period or not key:
        return JSONResponse({"ok": False, "msg": "缺承运商/账期/行键"}, status_code=400)
    lk = _locked(carrier, period)
    if lk:
        return lk
    with db._engine.begin() as c:
        ex = c.execute(select(LN.c.id).where((LN.c.carrier == carrier) & (LN.c.period == period) & (LN.c.line_key == key))).scalar()
        vals = dict(note=note, updated_by=_uname(u), updated_at=_now())
        if ex:
            c.execute(update(LN).where(LN.c.id == ex).values(**vals))
        else:
            c.execute(insert(LN).values(carrier=carrier, period=period, line_key=key, **vals))
    _bust(carrier, period)
    return {"ok": True}


class _FixErr(Exception):
    pass


@router.post("/api/logistics-review/line-fix")
async def review_line_fix(request: Request):
    """计提更正：对一笔或同一账单下几笔计提登记"应改为"(科目/费用项目/部门/产品分类/产品项目/金额+说明)。不改金蝶、不影响本页对账，
    只进导出的《计提更正单》打印交专人去改。五项全空=撤掉更正。已登记月份不可改。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, period = (b.get("carrier") or "").strip(), (b.get("period") or "").strip()
    keys = [str(k).strip() for k in (b.get("line_keys") or []) if str(k).strip()]
    if not carrier or not period or not keys:
        return JSONResponse({"ok": False, "msg": "缺承运商/账期/行键"}, status_code=400)
    lk = _locked(carrier, period)
    if lk:
        return lk
    vals = {k: str(b.get(k) or "").strip() for k in _FIX_KEYS}
    for k, lb in (("to_amt", "不含税金额"), ("to_amt_tax", "金额(含税)"), ("to_rate", "税率")):
        if vals[k]:             # 金额可带千分位，存两位小数；税率可写 9 / 9% / 0.09，存小数
            try:
                v = float(vals[k].replace(",", "").replace("，", "").replace("%", ""))
            except ValueError:
                return JSONResponse({"ok": False, "msg": "应改为%s不是数字：%s" % (lb, vals[k])}, status_code=400)
            vals[k] = ("%.4f" % (v / 100 if v >= 1 else v)) if k == "to_rate" else "%.2f" % v
    # 只改其中一部分(V2.855，用户「我要改到电商」——记-389 4,117.46 里只有调拨的 1,450 该是电商)：填含税金额；
    #   必须同时选了要改成的维度，且不能和「改金额/税率」混用(那是整笔的事)；金额要大于 0、小于这一笔。
    if vals["split_amt"]:
        try:
            sv = float(vals["split_amt"].replace(",", "").replace("，", ""))
        except ValueError:
            return JSONResponse({"ok": False, "msg": "「只改其中」的金额不是数字：%s" % vals["split_amt"]}, status_code=400)
        if sv <= 0:
            return JSONResponse({"ok": False, "msg": "「只改其中」的金额要大于 0"}, status_code=400)
        if not any(vals[k] for k in ("to_acct", "to_fee", "to_dept", "to_biz", "to_proj")):
            return JSONResponse({"ok": False, "msg": "填了「只改其中」的金额，还要选这一部分改成什么（科目/费用项目/部门/产品分类/产品项目至少一项）"}, status_code=400)
        if any(vals[k] for k in ("to_amt", "to_amt_tax", "to_rate")):
            return JSONResponse({"ok": False, "msg": "「只改其中一部分」不能和改金额、改税率一起填"}, status_code=400)
        if len(keys) > 1:
            return JSONResponse({"ok": False, "msg": "「只改其中一部分」一次只能对一笔计提登记"}, status_code=400)
        vals["split_amt"] = "%.2f" % sv
    # 调账月份：不填默认归属月份的下个月(复核多在次月，原月份一般已结账)；只有它不算"有更正"
    adj = str(b.get("adj_period") or "").strip()[:7] or _next_period(period)
    import time as _t
    ck = ("lines", carrier, period)
    cc = _ACCR_CACHE.get(ck)
    L = cc[0] if cc else _build_lines(request, carrier, period)
    rows = {r["key"]: r for r in L.get("rows", []) if r.get("kind") == "accr"}
    try:
      with db._engine.begin() as c:
        for key in keys:
            ex = c.execute(select(FX.c.id).where((FX.c.carrier == carrier) & (FX.c.period == period) & (FX.c.line_key == key))).scalar()
            if not any(vals.values()):
                if ex:
                    c.execute(delete(FX).where(FX.c.id == ex))
                continue
            r = rows.get(key)
            if vals["split_amt"] and r and r.get("amt") is not None and float(vals["split_amt"]) >= float(r["amt"]) - 0.004:
                raise _FixErr("「只改其中」的金额 %s 不小于这一笔的金额 %.2f：要整笔改就把它留空" % (vals["split_amt"], float(r["amt"])))
            rec = dict(vals, adj_period=adj, updated_by=_uname(u), updated_at=_now())
            if r:   # 金额/税率填的和原记账一样 → 存空(更正单印"不变")
                for k, ok, tol in (("to_amt", "amt_net", 0.005), ("to_amt_tax", "amt", 0.005), ("to_rate", "tax_rate", 0.00005)):
                    if rec.get(k) and r.get(ok) is not None and abs(float(rec[k]) - float(r[ok])) < tol:
                        rec[k] = ""
            if r:   # 维度选的和原记账一样 → 存空(更正单、待更正标记里就不会把没变的也列一遍)
                for k, ck in (("to_acct", "acct"), ("to_fee", "fee_code"), ("to_dept", "dept_code"), ("to_biz", "biz_code"), ("to_proj", "proj_code")):
                    if rec.get(k) and rec[k].split(" ", 1)[0] == str(r.get(ck) or "").strip():
                        rec[k] = ""
                if rec.get("split_amt") and not any(rec.get(k) for k in ("to_acct", "to_fee", "to_dept", "to_biz", "to_proj")):
                    raise _FixErr("「只改其中」那一部分要改成的维度和原记账一样，等于没改：请选一个不同的科目 / 费用项目 / 部门 / 产品分类 / 产品项目")
            if r:
                rec["snap_json"] = json.dumps({k: r.get(k) for k in _SNAP_KEYS}, ensure_ascii=False)
            if ex:
                c.execute(update(FX).where(FX.c.id == ex).values(**rec))
            elif r:
                c.execute(insert(FX).values(carrier=carrier, period=period, line_key=key, **rec))
    except _FixErr as e:
        return JSONResponse({"ok": False, "msg": str(e)}, status_code=400)
    # 首页待办区：登记了更正 → 给计提更正处理人记一笔；这家这个月的更正全撤了 → 待办撤回。失败只留痕，不拦。
    todo_scenes.fix_touch(carrier, period, _uname(u))
    # 只重挂更正，不作废逐笔缓存(不用重读金蝶)
    _attach_fixes(L, carrier, period)
    if not cc:
        _ACCR_CACHE[ck] = (L, None, _t.time())
    return {"ok": True, "fixes": L["fixes"]}


_DIMOPT_CACHE = {}


@router.get("/api/logistics-review/dim-options")
def review_dim_options(request: Request):
    """计提更正「应改为」下拉：金蝶主数据 编码+名称(科目/费用项目/部门/产品分类/产品项目)。改账按编码找，缓存12小时。"""
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    import time as _t
    cc = _DIMOPT_CACHE.get("v")
    if cc and _t.time() - cc[1] < 43200:
        return cc[0]
    out = {"ok": True, "acct": [], "fee": [], "dept": [], "biz": [], "proj": []}
    try:
        s, conf = kc.login()
    except Exception as e:
        return {**out, "ok": False, "msg": "金蝶登录失败：%s" % e}

    def q(form, fields, flt):
        try:
            return kc._query(s, conf, form, fields, flt)
        except Exception:
            return []

    def pairs(rows, keep=None):
        seen = {}
        for r in rows:
            code, name = str(r.get("c") or "").strip(), str(r.get("n") or "").strip()
            if code and code not in seen and (keep is None or keep(r)):
                seen[code] = name
        return [{"code": k, "name": v} for k, v in sorted(seen.items())]
    CN = [("FNumber", "c"), ("FName", "n")]
    out["acct"] = pairs(q("BD_Account", CN, "FNumber in ('5101','5301','6401','6402','6601','6602','6603','6711')"))
    out["fee"] = pairs(q("BD_Expense", CN, "FNumber like 'FYXM%'"))
    out["dept"] = pairs(q("BD_Department", CN, "FNumber like '%'"))
    ad = q("BOS_ASSISTANTDATA_DETAIL", [("FNumber", "c"), ("FDataValue", "n"), ("FId.FNumber", "t")], "FId.FNumber in ('CPFL','CPXM')")
    out["biz"] = pairs(ad, lambda r: r.get("t") == "CPFL")
    out["proj"] = pairs(ad, lambda r: r.get("t") == "CPXM")
    if all(out[k] for k in ("acct", "fee", "dept", "biz", "proj")):
        _DIMOPT_CACHE["v"] = (out, _t.time())
    return out


@router.post("/api/logistics-review/wt-range")
async def review_wt_range(request: Request):
    """毛重比允许范围(一家一档，不分月)：账单重量÷金蝶净重 落在 [lo, hi] 内算重量一致；两个都空=恢复默认(差 2% 以内)。
    快递/快运计费重量含包装、抛重，一般配 1~2(用户 2026-10-03 定放进基础设置，V2.763)。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier = (b.get("carrier") or "").strip()
    sup = next((x for x in (db.list_logi_suppliers() or []) if x.get("short") == carrier), None)
    if not sup:
        return JSONResponse({"ok": False, "msg": "供应商档案里没有「%s」，先建档再配毛重比" % carrier}, status_code=400)
    try:
        db.save_logi_supplier({"short": sup["short"], "full": sup.get("full"), "kd_code": sup.get("kd_code"), "channel": sup.get("channel"),
                               "note": sup.get("note"), "wt_lo": b.get("lo"), "wt_hi": b.get("hi")}, _uname(u))
    except ValueError as e:
        return JSONResponse({"ok": False, "msg": str(e)}, status_code=400)
    db.audit(_uname(u), "物流复核-毛重比范围", carrier, "%s ~ %s" % (b.get("lo") or "默认", b.get("hi") or "默认"))
    for k in [k for k in list(_ACCR_CACHE) if isinstance(k, tuple) and carrier in k]:
        _ACCR_CACHE.pop(k, None)          # 各月逐单视图缓存一并作废，按新范围重判
    return {"ok": True, "wt_range": _wt_range(carrier)}


@router.post("/api/logistics-review/carrier-points")
async def review_carrier_points(request: Request):
    """供应商复核要点（一家一段，挂逐笔复核页顶部：顺丰按重量、丰源按件数箱…）。不分月，不受锁月限制。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, points = (b.get("carrier") or "").strip(), (b.get("points") or "").strip()
    if not carrier:
        return JSONResponse({"ok": False, "msg": "缺承运商"}, status_code=400)
    with db._engine.begin() as c:
        ex = c.execute(select(CP.c.id).where(CP.c.carrier == carrier)).scalar()
        vals = dict(points=points, updated_by=_uname(u), updated_at=_now())
        if ex:
            c.execute(update(CP).where(CP.c.id == ex).values(**vals))
        else:
            c.execute(insert(CP).values(carrier=carrier, **vals))
    for k in [k for k in _ACCR_CACHE if k[0] == "lines" and k[1] == carrier]:
        _ACCR_CACHE.pop(k, None)
    return {"ok": True}


@router.post("/api/logistics-review/sign")
async def review_sign_month(request: Request):
    """确认通过→登记已复核：整月一家一次，存复核人/时间/结论快照；登记后当月归类与备注锁定，总表显示已复核。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, period, note = (b.get("carrier") or "").strip(), (b.get("period") or "").strip(), (b.get("note") or "").strip()
    if not carrier or not period:
        return JSONResponse({"ok": False, "msg": "缺承运商/账期"}, status_code=400)
    if _signed(carrier, period):
        return JSONResponse({"ok": False, "msg": "本月已登记，无需重复"}, status_code=409)
    L = _build_lines(request, carrier, period)
    snap = {k: L.get(k) for k in ("accr_total", "bill_total", "diff_total", "n_unexplained", "bill_src")}
    if carrier in _register_carriers():      # 登记制：快照里记请款合计(没有账单合计)
        PRT = store.payreq
        with db._engine.connect() as c:
            rq = c.execute(select(PRT.c.amount, PRT.c.excluded, PRT.c.dt_status, PRT.c.dt_result).where(
                (PRT.c.carrier == carrier) & (PRT.c.period == period))).all()
        snap.update(mode="register", req_total=round(sum(float(a or 0) for a, ex, st, rs in rq if not ex and st != "TERMINATED" and rs != "refuse"), 2))
    rec = dict(carrier=carrier, period=period, status="signed", reviewer=_uname(u), signed_at=_now(), note=note,
               snap_json=json.dumps(snap, ensure_ascii=False))
    with db._engine.begin() as c:
        c.execute(insert(SG).values(**rec))
    _bust(carrier, period)
    return {"ok": True, "signed": rec}


@router.post("/api/logistics-review/unsign")
async def review_unsign_month(request: Request):
    """撤销复核登记（要改归类/备注时先撤销）。"""
    u = _perm(request)
    if not u:
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    b = await request.json()
    carrier, period = (b.get("carrier") or "").strip(), (b.get("period") or "").strip()
    with db._engine.begin() as c:
        c.execute(delete(SG).where((SG.c.carrier == carrier) & (SG.c.period == period)))
    _bust(carrier, period)
    return {"ok": True}


def _fix_to_txt(fx):
    """更正内容一句话：科目 · 费用项目 · 部门 · 产品分类 · 产品项目 · 金额 · 税率 · 不含税（空=不变，维度值为「编码 名称」）；原因。"""
    parts = [("科目 " + fx["to_acct"]) if fx.get("to_acct") else "", ("费用项目 " + fx["to_fee"]) if fx.get("to_fee") else "",
             ("部门 " + fx["to_dept"]) if fx.get("to_dept") else "", ("产品分类 " + fx["to_biz"]) if fx.get("to_biz") else "",
             ("产品项目 " + fx["to_proj"]) if fx.get("to_proj") else "",
             ("金额 {:,.2f}".format(float(fx["to_amt_tax"]))) if fx.get("to_amt_tax") else "",
             ("税率 {:g}%".format(round(float(fx["to_rate"]) * 100, 2))) if fx.get("to_rate") else "",
             ("不含税 {:,.2f}".format(float(fx["to_amt"]))) if fx.get("to_amt") else ""]
    s = " · ".join(p for p in parts if p) or "（见原因）"
    if fx.get("split_amt"):
        s = "其中 {:,.2f}（含税）改为 ".format(float(fx["split_amt"])) + s + "，其余不动"
    return s + (("；原因：" + fx["memo"]) if fx.get("memo") else "")


def _cn(code, name):
    """维度显示＝「编码 名称」(改账按编码找)；缺哪个写哪个，都缺返回空串。"""
    code, name = str(code or "").strip(), str(name or "").strip()
    return (code + " " + name).strip() if code != name else code


def _dw(t):
    """估打印宽度：汉字按2、其余按1(openpyxl 列宽单位≈1个英文字符)。"""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in str(t or ""))


def _merge_fast(ws, r1, c1, r2, c2):
    """同 ws.merge_cells，但跳过 openpyxl「这一块是不是已经在合并区里」的逐块比对——合并几万块时那一步是平方级的：
    迅鸽 2026-08 复核明细 7,725 行、每张单十几列要合并，线上实测整个导出 501 秒，全耗在这里(V2.843 查出)。调用方保证各块互不重叠。"""
    from openpyxl.worksheet.merge import MergedCellRange
    from openpyxl.utils import get_column_letter
    mcr = MergedCellRange(ws, "%s%d:%s%d" % (get_column_letter(c1), r1, get_column_letter(c2), r2))
    rs = ws.merged_cells.ranges
    (rs.add if hasattr(rs, "add") else rs.append)(mcr)
    ws._clean_merge_range(mcr)


def _raw_json_exists(carrier, period):
    """早先手工备的原账单 json(raw_bills/承运商_账期.json，带右侧复核列)在不在——在就照旧用它，不再另附原件。"""
    import os as _os
    return _os.path.exists(_os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))), "..", "raw_bills", "%s_%s.json" % (carrier, period)))


def _fix_sheet(wb, carrier, period, fixes, suppliers=None, carrier_full=""):
    """《计提更正单》（第二页，可直接打印交专人）。抬头两行居中：①计提更正单 ②供应商编码/名称(取计提凭证上挂的供应商维度)。
    每笔三行(用户 2026-09-30 定)：原记账 / 应改为(不变的写"不变") / 原因(横跨维度与金额列)；
    维度都印「编码 名称」；金额拆 金额(含税)/税率/不含税金额 三列，放产品项目后面；
    序号/主体/凭证号/费用归属月份/调账月份/更正人合并三行；表下合计(金额有改时分原记账/应改为两行)、填写说明、签字栏。"""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.properties import PageSetupProperties
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet("计提更正单", 1)
    thin = Side(style="thin", color="B8C4CC")
    thick = Side(style="medium", color="7A8791")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    BDE = Border(left=thin, right=thin, top=thin, bottom=thick)     # 每笔最后一行下边线加粗，分隔各笔
    C = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    R = Alignment(horizontal="right", vertical="center")
    # 列：序号 主体 凭证号 费用归属月份 调账月份 | 行标 | 科目 费用项目 部门 产品分类 产品项目 | 金额 税率 不含税金额 | 更正人
    WID = [5, 15, 9, 9, 9, 7, 16, 23, 18, 13, 13, 12, 7, 12, 12]
    N = len(WID)
    SHARED = (1, 2, 3, 4, 5, 15)             # 三行合并
    LBL, DIM0, AMT, RATE, NET = 6, 7, 12, 13, 14
    last = get_column_letter(N)
    ws.merge_cells("A1:%s1" % last)
    ws.cell(1, 1, "计提更正单").font = Font(bold=True, size=16, color="1B2733")
    ws.cell(1, 1).alignment = C
    ws.row_dimensions[1].height = 30
    sups = suppliers or []
    sname = "、".join(dict.fromkeys(x.get("name") for x in sups if x.get("name"))) or carrier_full or carrier
    scode = "、".join(x.get("code") for x in sups if x.get("code")) or "（未取到）"
    ws.merge_cells("A2:%s2" % last)
    c2 = ws.cell(2, 1, "供应商编码：%s　　　供应商名称：%s" % (scode, sname))
    c2.alignment = C; c2.font = Font(bold=True, size=11, color="1B2733")
    ws.row_dimensions[2].height = 24
    HR = 3
    heads = ["序号", "主体", "凭证号", "费用归属\n月份", "调账\n月份", "", "科目", "费用项目", "部门", "产品分类", "产品项目",
             "金额\n(含税)", "税率", "不含税\n金额", "更正人/日期"]
    for j, h in enumerate(heads, 1):
        cl = ws.cell(HR, j, h)
        cl.font = Font(bold=True, color="FFFFFF"); cl.alignment = C; cl.border = BD
        cl.fill = PatternFill("solid", fgColor="5E6B78")
    ws.row_dimensions[HR].height = 32
    GRAY = Font(color="9AA5AE")
    HOT = Font(bold=True, color="8A5A00")
    OLD_FILL = PatternFill("solid", fgColor="EEF2F4")
    NEW_FILL = PatternFill("solid", fgColor="FBF0DA")
    HOT_FILL = PatternFill("solid", fgColor="FDF6E8")
    WHY_FILL = PatternFill("solid", fgColor="FFFFFF")

    def nlines(t, width):
        return max(1, -(-_dw(t) // max(1, width - 3)))     # 留余量：Excel 实际折行比估算早

    def fnum(v):
        try:
            return float(v) if v not in (None, "") else None
        except (TypeError, ValueError):
            return None

    def rate_fmt(r):
        return "0%" if abs(r * 100 - round(r * 100)) < 0.005 else "0.00%"
    t_old = [0.0, 0.0]      # 含税, 不含税
    t_new = [0.0, 0.0]
    for i, fx in enumerate(fixes, 1):
        s = fx.get("snap") or {}
        r1 = HR + 3 * i - 2
        r2, r3 = r1 + 1, r1 + 2
        adj = fx.get("adj_period") or _next_period(period)
        shared = {1: i, 2: _cn(s.get("book_code"), s.get("subject")), 3: s.get("vno"), 4: period, 5: adj, 15: None}
        for j in SHARED:
            ws.merge_cells(start_row=r1, start_column=j, end_row=r3, end_column=j)
            cl = ws.cell(r1, j, shared[j])
            cl.alignment = C if j in (1, 3, 4, 5) else LFT
            if j == 5 and adj != period:
                cl.font = HOT                    # 跨月调账醒目
        for rr_, lb, fill, font in ((r1, "原记账", OLD_FILL, Font(color="5E6B78")), (r2, "应改为", NEW_FILL, HOT),
                                    (r3, "原因", WHY_FILL, Font(bold=True, color="5E6B78"))):
            cl = ws.cell(rr_, LBL, lb); cl.fill = fill; cl.alignment = C; cl.font = font
        # 维度：原记账 / 应改为
        old = [_cn(s.get("acct"), s.get("acct_name")), _cn(s.get("fee_code"), s.get("fee")), _cn(s.get("dept_code"), s.get("dept")),
               _cn(s.get("biz_code"), "" if str(s.get("biz") or "").startswith("（") else s.get("biz")),   # （无业务线）=空
               _cn(s.get("proj_code"), s.get("proj"))]
        new = [fx.get("to_acct"), fx.get("to_fee"), fx.get("to_dept"), fx.get("to_biz"), fx.get("to_proj")]
        for k in range(5):
            c1 = ws.cell(r1, DIM0 + k, old[k] or "空"); c1.alignment = LFT; c1.fill = OLD_FILL
            c2_ = ws.cell(r2, DIM0 + k, new[k] or "不变"); c2_.alignment = LFT
            if new[k]:
                c2_.font = HOT; c2_.fill = HOT_FILL
            else:
                c2_.font = GRAY
        # 金额：含税 / 税率 / 不含税
        o_t, o_r, o_n = fnum(s.get("amt")), fnum(s.get("tax_rate")), fnum(s.get("amt_net"))
        n_t, n_r, n_n = fnum(fx.get("to_amt_tax")), fnum(fx.get("to_rate")), fnum(fx.get("to_amt"))
        for col, ov, nv, kind in ((AMT, o_t, n_t, "m"), (RATE, o_r, n_r, "r"), (NET, o_n, n_n, "m")):
            c1 = ws.cell(r1, col, ov); c1.alignment = R; c1.fill = OLD_FILL
            if ov is not None:
                c1.number_format = "#,##0.00" if kind == "m" else rate_fmt(ov)
            c2_ = ws.cell(r2, col, nv if nv is not None else "不变"); c2_.alignment = R
            if nv is not None:
                c2_.number_format = "#,##0.00" if kind == "m" else rate_fmt(nv)
                c2_.font = HOT; c2_.fill = HOT_FILL
            else:
                c2_.font = GRAY
        t_old[0] += o_t or 0; t_old[1] += o_n or 0
        t_new[0] += n_t if n_t is not None else (o_t or 0)
        t_new[1] += n_n if n_n is not None else (o_n or 0)
        # 原因：横跨维度与金额列
        ws.merge_cells(start_row=r3, start_column=DIM0, end_row=r3, end_column=NET)
        why = ws.cell(r3, DIM0, (("只改其中 {:,.2f}（含税），其余不动。".format(float(fx["split_amt"]))) if fx.get("split_amt") else "") + (fx.get("memo") or ""))
        why.alignment = LFT
        for j in range(1, N + 1):            # 合并之后再上边框，每笔最后一行下沿粗线
            ws.cell(r1, j).border = BD
            ws.cell(r2, j).border = BD
            ws.cell(r3, j).border = BDE
        # 行高按折行估：原记账/应改为按最长的格；原因按合并宽度；主体(三行合并)不够时摊到三行
        l1 = max(nlines(old[k] or "空", WID[DIM0 + k - 1]) for k in range(5))
        l2 = max(nlines(new[k] or "不变", WID[DIM0 + k - 1]) for k in range(5))
        l3 = nlines(fx.get("memo") or "", sum(WID[DIM0 - 1:NET]))
        extra = max(0, nlines(shared[2], WID[1]) - l1 - l2 - l3)
        for rr_, ln in ((r1, l1), (r2, l2), (r3, l3)):
            ws.row_dimensions[rr_].height = 15 * (ln + extra / 3) + 7
    rr = HR + 3 * len(fixes) + 1
    t_old = [round(v, 2) for v in t_old]
    t_new = [round(v, 2) for v in t_new]
    changed = any(abs(a - b) >= 0.005 for a, b in zip(t_old, t_new))
    tots = [("原记账合计" if changed else "合计", t_old)] + ([("应改为合计", t_new)] if changed else [])   # 金额有改才出第二行
    TOT_FILL = [PatternFill("solid", fgColor="E7ECEF"), PatternFill("solid", fgColor="FBF0DA")]
    for k, (lb, v) in enumerate(tots):
        # 合计行撑满整行(用户 2026-09-30)：标签合并到金额列前，金额两列，其余格也上边框和底色
        ws.merge_cells(start_row=rr + k, start_column=1, end_row=rr + k, end_column=AMT - 1)
        cl = ws.cell(rr + k, 1, lb); cl.font = Font(bold=True, color="8A5A00" if k else "1B2733"); cl.alignment = R
        for col, val in ((AMT, v[0]), (NET, v[1])):
            c_ = ws.cell(rr + k, col, val); c_.number_format = "#,##0.00"; c_.alignment = R
            c_.font = Font(bold=True, color="8A5A00" if k else "1B2733")
        for j in range(1, N + 1):
            ws.cell(rr + k, j).border = BD
            ws.cell(rr + k, j).fill = TOT_FILL[min(k, 1)]
        ws.row_dimensions[rr + k].height = 22
    rn = rr + len(tots)
    ws.merge_cells(start_row=rn, start_column=1, end_row=rn, end_column=N)
    nt = ws.cell(rn, 1, "填写说明：以上计提分录需要更正（登记更正不影响复核台对账）。每笔三行：原记账、应改为、原因；应改为写「不变」的不用动；各项均为「编码 名称」，请按编码修改。"
                        "金额/税率变动的，按应改为的含税金额、税率、不含税金额调整进项税与费用。调账月份＝费用归属月份的，直接修改原凭证；晚于费用归属月份的，在调账月份做调整凭证。改完在右侧签字。")
    nt.alignment = LFT; nt.font = Font(color="5E6B78", size=10)
    ws.row_dimensions[rn].height = 32
    by = sorted({fx.get("by") for fx in fixes if fx.get("by")})
    ws.merge_cells(start_row=rn + 2, start_column=1, end_row=rn + 2, end_column=N)
    ws.cell(rn + 2, 1, "登记人：%s        更正人：______________        更正日期：______________        复核人：______________"
            % ("、".join(by) or "______________")).alignment = LFT
    ws.row_dimensions[rn + 2].height = 26
    for j, w in enumerate(WID, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:3"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    return ws


@router.get("/api/logistics-review/export")
def review_export(request: Request, carrier: str = "迅鸽", period: str = ""):
    """导出该承运商本月复核结果 xlsx：费用项汇总 + 逐单/物料级复核明细（按重量承运商=物料级17列）。"""
    if not _perm(request):
        return JSONResponse({"ok": False, "msg": "无权限"}, status_code=403)
    res = review_result(request, carrier=carrier, period=period, group="all", page=1, size=1000000, q="")
    if isinstance(res, JSONResponse):
        return res
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.properties import PageSetupProperties
    from io import BytesIO
    from urllib.parse import quote
    HFONT = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="1F6E8C")
    wb = Workbook()
    # sheet1 复核表（可直接打印）：逐笔计提为主线 主体/费用类型/产品线/产品类型/部门/凭证号/计提含税/税率/账单/差异/差异解释
    L = _build_lines(request, carrier, period)
    ws = wb.active; ws.title = "复核表"
    acc_t = L.get("accr_total") or 0
    bill_t = L.get("bill_total") or 0
    dif_t = round(acc_t - bill_t, 2)
    NCOL = 11
    thin = Side(style="thin", color="B8C4CC")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    right = Alignment(horizontal="right", vertical="center")
    RED = Font(color="B23B2E", bold=True); GREEN = Font(color="2E7D57", bold=True)
    sg = L.get("signed") or {}
    ws.merge_cells("A1:K1")
    t = ws.cell(1, 1, "物流账单复核表  ·  %s  ·  %s" % (carrier, period))
    t.font = Font(bold=True, size=15, color="1B2733"); t.alignment = center
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:K2")
    st = ("已复核  复核人 %s  %s" % (sg.get("reviewer") or "", sg.get("signed_at") or "")) if sg else "待复核"
    ws.cell(2, 1, "复核状态：%s    |    计提合计 %s    |    账单合计 %s    |    差异 %s%s" % (
        st, f"{acc_t:,.2f}", f"{bill_t:,.2f}", f"{dif_t:,.2f}", "（对平）" if abs(dif_t) < 0.01 else "")
    ).alignment = Alignment(horizontal="left", vertical="center")
    ws.cell(2, 1).font = Font(bold=True, color="1F6E8C"); ws.row_dimensions[2].height = 20
    ws.merge_cells("A3:K3")
    ws.cell(3, 1, "复核要点：%s" % (L.get("points") or "（未填写）")).alignment = left
    ws.row_dimensions[3].height = 34
    ws.append([None] * NCOL)
    ws.append(["主体", "费用类型", "产品线", "产品类型", "部门", "凭证号", "计提金额(含税)", "税率", "账单金额", "差异", "备注·差异解释"])
    HROW = 5
    for hc0 in ws[HROW]:
        hc0.font = HFONT; hc0.fill = HFILL; hc0.alignment = center; hc0.border = BORDER
    gt_rows, fix_rows, spans = [], [], []
    for r in L.get("rows", []):
        if r.get("kind") == "gtotal":
            sj = _cn(r.get("book_code"), r.get("subject"))
            ws.append([("%s · %s 小计" % (sj, r.get("fee_type"))) if r.get("fee_type") else ("%s 小计" % sj), None, None, None, None, None,
                       r.get("amt"), None, r.get("bill"), r.get("diff"), None])
            gt_rows.append(ws.max_row)
            continue
        biz = r.get("biz") or ""
        if not biz.startswith("（"):
            biz = _cn(r.get("biz_code"), biz)          # 维度都带编码：CPFL013 Kiki Herb
        if r.get("bill_biz"):
            biz = "%s（账单:%s）" % (biz, r["bill_biz"])
        nt = r.get("note") or ""
        if r.get("fix"):
            nt = ("%s\n" % nt if nt else "") + "【待更正】改为 " + _fix_to_txt(r["fix"])
        if r.get("xsubj"):
            nt = ("%s\n" % nt if nt else "") + "【记错主体】发票、请款单是%s的：付款做账时本主体红冲、%s补提" % (r["xsubj"]["to"], r["xsubj"]["to"])
        if r.get("xsubj_in"):
            nt = ("%s\n" % nt if nt else "") + "【记错主体】计提记在%s %s：付款做账时补提到本主体" % (r["xsubj_in"]["from"], r["xsubj_in"]["vno"])
        fee = _cn(r.get("fee_code"), r.get("fee")) if r.get("fee_code") else r.get("fee_type")   # 费用类型=金蝶费用项目(编码 名称)，账单有计提无的用复核归类
        ws.append([_cn(r.get("book_code"), r.get("subject")), fee, biz,
                   _cn(r.get("proj_code"), r.get("proj")) or None, _cn(r.get("dept_code"), r.get("dept")) or None, r.get("vno") or None,
                   r.get("amt"), r.get("tax_rate"), r.get("bill"), r.get("diff"), nt or None])
        if r.get("fix"):
            fix_rows.append(ws.max_row)
        if (r.get("bill_span") or 0) > 1:
            spans.append((ws.max_row, int(r["bill_span"])))
    ws.append(["合计", None, None, None, None, None, acc_t, None, bill_t, dif_t, None])
    LAST = ws.max_row
    for rr in range(HROW + 1, LAST + 1):
        is_total = (rr == LAST); is_gt = rr in gt_rows
        for cc in range(1, NCOL + 1):
            cell = ws.cell(rr, cc); cell.border = BORDER
            if cc in (7, 9, 10):
                cell.number_format = "#,##0.00"; cell.alignment = right
            elif cc == 8:
                cell.number_format = "0.0%"; cell.alignment = right
            else:
                cell.alignment = left
            if is_total:
                cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="E1EEF3")
            elif is_gt:
                cell.font = Font(bold=True); cell.fill = PatternFill("solid", fgColor="EDF2F5")
        dv = ws.cell(rr, 10).value
        if isinstance(dv, (int, float)):
            ws.cell(rr, 10).font = GREEN if abs(dv) < 0.01 else RED
        if rr in fix_rows:
            ws.cell(rr, 11).fill = PatternFill("solid", fgColor="FBF0DA")
            ws.cell(rr, 11).font = Font(color="8A5A00")
    # V2.843(用户看易风达导出「为什么小料出库运费没有数据」)：几笔计提合起来对同一个账单数的(记-553 6,255 + 记-566 68,900 对账单 75,155)，
    #   账单金额、差异两格跨这几行合并，和页面一致；并在表下注明。
    vmid = Alignment(horizontal="right", vertical="center")
    for r1, n in spans:
        r2 = min(r1 + n - 1, LAST - 1)
        if r2 <= r1:
            continue
        for cc in (9, 10):
            ws.merge_cells(start_row=r1, start_column=cc, end_row=r2, end_column=cc)
            ws.cell(r1, cc).alignment = vmid
    fixes = L.get("fixes") or []
    raw_files = [] if _raw_json_exists(carrier, period) else _raw_bill_files(carrier, period)
    foot = []
    if spans:
        foot.append(("账单金额、差异跨几行合并的：那几笔计提合起来对同一个账单数（账单上分不到每一笔计提）。", "5E6B78"))
    if fixes:
        foot.append(("另有 %d 笔计提需要更正（登记更正不影响本表对账），明细见《计提更正单》，打印交专人在金蝶修改。" % len(fixes), "8A5A00"))
    else:
        foot.append(("复核台第①步没有登记维度/金额的更正，所以这份导出里没有《计提更正单》页。要登记：第①步在那一笔上点「改维度」，再导出。", "5E6B78"))
    # 发票对暂估(税率开错、税额尾差、记错主体)的更正不在本表(本表比含税金额，看不出来)：在「付款做账」里随付款凭证红冲更正，更正单也在那边打印。
    #   用户 2026-10-06 看易风达导出问「计提更正单是在哪里出」「不是有差异吗」——记-565 计提按 6%、发票 9%，已在付款凭证里更正，导出里却一个字没有。
    try:
        from routers import logistics_payreq as _PQ
        blocks = _PQ.inv_blocks(carrier, period)
    except Exception:
        blocks = None
    if blocks is None:
        foot.append(("发票对暂估：这次没取到（读金蝶没成功），请看复核台第③步「发票与暂估」。", "B23B2E"))
    else:
        nfix = 0
        for b in blocks:
            if b.get("kind") in (None, "", "hx"):
                continue
            nfix += 1
            pv = b.get("posted") or {}
            txt = "发票对暂估 · %s 请款 %s：%s。" % (b.get("subject") or "", f"{float(b.get('amount') or 0):,.2f}", b.get("kind_text") or b.get("kind_cn") or "")
            if b.get("kind") in ("tail", "redo", "subj"):
                txt += ("已在付款凭证 记-%s 里红冲更正；" % pv.get("vno") if pv.get("vno") else "还没做付款凭证，做的时候随付款凭证红冲更正；") + \
                       "这种《计提更正单》在「付款做账」这张请款单的预览里打印。"
            foot.append((txt, "8A5A00"))
        if blocks and not nfix:
            foot.append(("发票对暂估：%d 张请款单的发票和计提暂估都一致，只核销，不用更正。" % len(blocks), "5E6B78"))
    if not _raw_json_exists(carrier, period):
        foot.append(("原账单：见后面「原账单-…」页（%d 份文件，原样附上）。" % len(raw_files), "5E6B78") if raw_files else
                    ("原账单：系统里没有这家这月的账单原件（这版之前导入的账单没留原件，钉钉请款单里也没找到已导入的那份）；重新导入一次，再导出就会附上。", "B23B2E"))
    for i, (txt, color) in enumerate(foot):
        r0 = LAST + 2 + i
        ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=NCOL)
        c0 = ws.cell(r0, 1, txt)
        c0.font = Font(bold=(color == "8A5A00"), color=color); c0.alignment = left
        ws.row_dimensions[r0].height = 32 if len(txt) > 70 else 18
    for i, w in enumerate([16, 26, 20, 17, 20, 10, 14, 7, 14, 13, 30], 1):     # 主体/产品线/产品类型/部门带编码后加宽
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A6"
    # 打印设置：横向(11列)、按宽缩放到1页、居中、重复表头、窄边距
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:5"
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if fixes:
        _fix_sheet(wb, carrier, period, fixes, L.get("suppliers") or [], _carrier_full(carrier))
    # sheet2 复核明细
    ws2 = wb.create_sheet("复核明细")
    from openpyxl.utils import get_column_letter
    if res.get("material"):
        # 接上原账单：每单费用分项(sub_fees)+运输方式(carrier_sub)+账单计入金额+账单计费重量，doc级仅首个物料行填
        with db._engine.connect() as c:
            braw = {r[0]: (r[1], r[2], r[3], r[4]) for r in c.execute(select(
                BL.c.doc_no, BL.c.carrier_sub, BL.c.sub_fees, BL.c.amount, BL.c.charge_wt).where(
                (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain == "detail"))).all()}
        feekeys = []
        for _no, (cs, sf, amt, cw) in braw.items():
            if sf:
                try:
                    for k in json.loads(sf):
                        if k not in feekeys:
                            feekeys.append(k)     # 含「箱型」文字列(迅鸽)，照样作一列
                except Exception:
                    pass
        _ORD = {"运费": 0, "快递费": 0, "操作费": 1, "箱子": 2, "箱型": 3, "加班": 4}
        feekeys.sort(key=lambda k: (_ORD.get(k, 9), k))
        # 分组：ERP/金蝶数据、账单原始数据、复核数据分开配色；★=重点关注列。(组名, 组色, 列头浅色, doc级仅首行填, [(列头, 取值键)])
        groups = [
            ("归属·计提", "5E6B78", "E7ECEF", False, [("费用主体", "subject"), ("承运商", "carrier"),
                ("费用类型", "fee_item"), ("业务线", "bizline"), ("单据号", "doc_no")]),
            ("ERP·金蝶数据", "2E7D57", "DCEFE4", False, [("客户/仓库", "party"), ("物料编码", "code"),
                ("物料名称", "name"), ("基本单位数量", "base_kg"), ("基本单位", "kg_unit"),
                ("金蝶数量", "base_qty"), ("数量单位", "base_unit"), ("金蝶核对量", "kd"), ("销售额", "sales")]),
            ("账单·%s原账单" % carrier, "B06A12", "FBF0DA", True, [("运输方式", "_cs")] +
                [(k, "_fee:" + k) for k in feekeys] + [("账单计入金额", "_amt"), ("账单计费量", "_cw")]),
            ("复核数据", "B23B2E", "F8DDD8", False, [("账单量", "bill_amt"), ("账单单位", "bill_unit"),
                ("★计费方式", "mode_cn"), ("★换算系数", "conv"), ("运费(分摊)", "fee"),
                ("单位运费(元/kg)", "unit_fee"), ("★费比", "ratio"), ("备注", "note"), ("已确认", "_ok")]),
        ]
        okmap = _doc_ok(carrier, period)   # 逐单已确认：确认人+时间
        col = 1
        col_of_key = {}
        for name, gc, hc, _dl, cols in groups:
            span = len(cols)
            ws2.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + span - 1)
            gcell = ws2.cell(row=1, column=col, value=name)
            gcell.font = HFONT; gcell.fill = PatternFill("solid", fgColor=gc)
            gcell.alignment = Alignment(horizontal="center")
            for j, (h, k) in enumerate(cols):
                col_of_key[k] = col + j
                hc2 = ws2.cell(row=2, column=col + j, value=h)
                hc2.font = Font(bold=True, color="B23B2E" if h.startswith("★") else "1B2733")
                hc2.fill = PatternFill("solid", fgColor=hc)
            col += span
        # 单据级列(跨该单所有物料行合并单元格)；物料级列(客户/物料/数量/运费分摊/单位运费/费比/销售额)不合并
        _DOCLVL = {"subject", "carrier", "fee_item", "bizline", "doc_no", "_cs", "_amt", "_cw",
                   "bill_amt", "bill_unit", "mode_cn", "conv", "note", "_ok"}
        def _is_doclvl(k):
            return k in _DOCLVL or k.startswith("_fee:")
        rownum = 3
        prev = None
        doc_start = None
        ranges = []
        for r in res.get("detail", []):
            d0 = r.get("doc_no"); firstdoc = (d0 != prev)
            if firstdoc:
                if doc_start is not None:
                    ranges.append((doc_start, rownum - 1))
                doc_start = rownum
            cs, sf, amt, cw = braw.get(d0, (None, None, None, None))
            sfd = {}
            if sf:
                try:
                    sfd = json.loads(sf)
                except Exception:
                    sfd = {}
            col = 1
            for name, gc, hc, doclvl, cols in groups:
                for (_h, k) in cols:
                    if _is_doclvl(k) and not firstdoc:
                        val = None                       # 单据级列仅首行写值，其余留空待合并
                    elif k == "_cs":
                        val = cs
                    elif k == "_amt":
                        val = amt
                    elif k == "_cw":
                        val = cw
                    elif k == "_ok":
                        ok_ = okmap.get(d0) if d0 else None
                        val = ("✓ %s %s" % (ok_["by"], ok_["at"])).strip() if ok_ else None
                    elif k.startswith("_fee:"):
                        val = sfd.get(k[5:])
                    elif k == "ratio":
                        val = round(r["ratio"], 4) if r.get("ratio") is not None else None
                    else:
                        val = r.get(k)
                    ws2.cell(row=rownum, column=col, value=val)
                    col += 1
            prev = d0
            rownum += 1
        if doc_start is not None:
            ranges.append((doc_start, rownum - 1))
        midv = Alignment(vertical="center")
        for s, e in ranges:
            if e <= s:
                continue
            for k, ci in col_of_key.items():
                if _is_doclvl(k):
                    _merge_fast(ws2, s, ci, e, ci)
                    ws2.cell(row=s, column=ci).alignment = midv
        ws2.freeze_panes = "F3"
        widths = ([12, 14, 12, 10, 15] + [16, 12, 22, 11, 8, 9, 8, 11, 10] + [14] + [10] * len(feekeys) + [13, 13] +
                  [10, 8, 14, 9, 10, 12, 8, 16])
        for i, w in enumerate(widths, 1):
            ws2.column_dimensions[get_column_letter(i)].width = w
    else:
        cols = ["金蝶单号", "快递/子类", "省", "计费重量", "账单数量", "金蝶数量", "账单金额", "标准费",
                "核价差", "核价", "核量", "归一态", "计价档"]
        ws2.append(cols)
        for r in res.get("detail", []):
            ws2.append([r.get("doc_no"), r.get("carrier_sub"), r.get("prov"), r.get("charge_wt"), r.get("qty"),
                        r.get("kd_qty"), r.get("amount"), r.get("std_amount"), r.get("price_diff"),
                        _PSTATE_CN.get(r.get("price_state"), r.get("price_state")),
                        _QSTATE_CN.get(r.get("qty_state"), r.get("qty_state")),
                        _VERDICT_CN.get(r.get("verdict"), r.get("verdict")), r.get("tier")])
        for c in ws2[1]:
            c.font = HFONT; c.fill = HFILL
        ws2.freeze_panes = "A2"
        for i, w in enumerate([15, 16, 8, 10, 10, 10, 11, 10, 9, 9, 8, 10, 14], 1):
            ws2.column_dimensions[get_column_letter(i)].width = w
    # sheet3+ 原账单（可追溯）：若已存原始账单，逐 sheet 原样附上，右侧接复核列（金蝶重量/换算系数/核量结论），按金蝶单号匹配
    import os as _os
    raw_fp = _os.path.join(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__))), "..", "raw_bills",
                           "%s_%s.json" % (carrier, period))
    if _os.path.exists(raw_fp):
        try:
            raw = json.load(open(raw_fp, encoding="utf-8"))
        except Exception:
            raw = None
        if raw:
            with db._engine.connect() as c:
                rv = {(r[0] or "").split("+")[0].strip(): (r[1], r[2], r[3]) for r in c.execute(select(
                    BL.c.doc_no, BL.c.kd_qty, BL.c.charge_wt, BL.c.qty_state).where(
                    (BL.c.carrier == carrier) & (BL.c.period == period) & (BL.c.grain == "detail"))).all()}
            docmode = {}   # 单据→计费方式(按重量/整车按箱/打托…)，取自复核结果
            for dd in res.get("detail", []):
                dn = dd.get("doc_no")
                if dn and dn not in docmode:
                    docmode[dn] = dd.get("mode_cn")
            addcols = ["计费类别", "ERP重量", "账单数量", "差异", "结论"]   # 单据运费复核数据，只接右侧、单独配色
            AFILL = PatternFill("solid", fgColor="B06A12")      # 追加列表头：琥珀底
            ALIGHT = PatternFill("solid", fgColor="FBF0DA")     # 追加列数据：淡琥珀底
            acenter = Alignment(horizontal="center", vertical="center")
            for sh in raw.get("sheets", []):
                title = ("原账单-" + sh.get("name", ""))[:31]
                ws3 = wb.create_sheet(title)
                hdr = list(sh.get("header", []))
                base_n = len(hdr)
                ws3.append(hdr + addcols)
                for j in range(len(addcols)):   # 仅追加列表头配色，原始表头不动(保留原格式)
                    cc = ws3.cell(1, base_n + 1 + j)
                    cc.font = HFONT; cc.fill = AFILL; cc.alignment = acenter
                ki = sh.get("kidx")
                rn = 1
                for row in sh.get("rows", []):
                    rn += 1
                    # 一格多单号(换行/区间)按首单号匹配，与中间表 doc_no 用"+"连的首段一致
                    no = re.split(r"[\n\-+]", str(row[ki]))[0].strip() if (ki is not None and ki < len(row) and row[ki] not in (None, "")) else ""
                    kq, cw, qs = rv.get(no, (None, None, None))
                    conv = round(cw / kq, 3) if (cw and kq) else None
                    diff = round((cw or 0) - (kq or 0), 2) if (no and (cw is not None or kq is not None)) else None
                    concl = "" if not no else ("一致" if qs == "ok" else "%s多报" % carrier if (conv and conv > 1)
                                               else "%s少报" % carrier if (conv and conv < 1) else "待核")
                    ws3.append(list(row) + [docmode.get(no), kq, cw, diff, concl])
                    for j in range(len(addcols)):   # 仅追加列数据配淡底，原始列不动
                        ws3.cell(rn, base_n + 1 + j).fill = ALIGHT
                ws3.freeze_panes = "A2"
    if raw_files:
        _attach_raw_sheets(wb, raw_files)
    bio = BytesIO(); wb.save(bio)
    fn = "%s_%s_复核结果.xlsx" % (carrier, period)
    return Response(content=bio.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename*=UTF-8''%s" % quote(fn)})
